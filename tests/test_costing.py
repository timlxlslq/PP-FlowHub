import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from traveler_assistant.core import Config
from traveler_assistant.costing import _display_cost_lines, calculate_order_cost, export_order_cost
from traveler_assistant.database import connect_database
from traveler_assistant.inventory import Product, _replace_product_database


class CostingTests(unittest.TestCase):
    # 创建隔离的成本测试配置，并导入带价格的商品样本。
    # self：当前测试用例或测试替身实例。
    # root：隔离测试目录根路径。
    def _config(self, root: Path) -> Config:
        config = Config(
            state_dir=root / "state",
            order_root=root / "orders",
            backup_root=root / "backups",
        )
        config.prepare_storage()
        _replace_product_database(config.workflow_database, [
            Product("Panel", "M0004", "Plywood", "18mm", "启用", unit="pcs", cost_price=10.0),
            Product("Panel", "M1010", "Ivory Oak", "19.1mm", "启用", unit="pcs", cost_price=25.0),
            Product("Panel", "M1011", "Unknown", "19.1mm", "启用", unit="pcs", cost_price=None),
            Product("Edge band", "M2010", "Ivory Oak Edge Banding", "", "启用", unit="m", cost_price=2.0),
        ])
        return config

    # 验证订单总成本采用订单材料汇总及未显示舍入的原始数量。
    # self：当前测试用例或测试替身实例。
    def test_order_total_uses_order_material_summary_and_raw_quantities(self):
        with tempfile.TemporaryDirectory() as temporary:
            config = self._config(Path(temporary))
            connection = connect_database(config.workflow_database)
            connection.execute(
                """insert into material_items(
                    order_id,product_code,quantity,source_type,source_path,
                    source_fingerprint,updated_at
                ) values(?,?,?,?,?,?,?)""",
                ("PP9999", "M0004", 2.5, "aihouse", "", "", "now"),
            )
            connection.execute(
                """insert into material_items(
                    order_id,product_code,quantity,source_type,source_path,
                    source_fingerprint,updated_at
                ) values(?,?,?,?,?,?,?)""",
                ("PP9999", "M1010", 1, "aihouse", "", "", "now"),
            )
            connection.commit()
            connection.close()

            report = calculate_order_cost(config, "PP9999")
            self.assertEqual(report["status"], "已完成")
            self.assertAlmostEqual(report["total_cost"], 50.0)
            self.assertEqual(report["factory_totals"][0]["factory_order"], "材料汇总")
            self.assertAlmostEqual(report["factory_totals"][0]["total"], 50.0)
            self.assertTrue(all(row["factory_order"] == "材料汇总" for row in report["lines"]))

            exported = export_order_cost(config, "PP9999")
            self.assertTrue(Path(exported["export_path"]).is_file())
            workbook = load_workbook(exported["export_path"], data_only=False, read_only=True)
            self.assertEqual(workbook.sheetnames, ["成本汇总", "价格与来源", "成本明细", "缺失项目"])
            summary_values = [cell for row in workbook["成本汇总"].iter_rows(values_only=True) for cell in row]
            self.assertIn("材料汇总", summary_values)
            self.assertNotIn("待分配", summary_values)
            workbook.close()

    # 验证缺失成本价被标记为缺失，不当作零成本。
    # self：当前测试用例或测试替身实例。
    def test_missing_cost_price_is_not_treated_as_zero(self):
        with tempfile.TemporaryDirectory() as temporary:
            config = self._config(Path(temporary))
            connection = connect_database(config.workflow_database)
            connection.execute(
                """insert into material_items(
                    order_id,product_code,quantity,source_type,source_path,
                    source_fingerprint,updated_at
                ) values(?,?,?,?,?,?,?)""",
                ("PP9998", "M1011", 1, "aihouse", "", "", "now"),
            )
            connection.commit()
            connection.close()
            report = calculate_order_cost(config, "PP9998")
            self.assertIsNone(report["total_cost"])
            self.assertEqual(report["status"], "待补充")
            self.assertTrue(report["missing_items"])

    # 验证成本明细按材料业务顺序显示，并合并相同五金。
    # self：当前测试用例或测试替身实例。
    def test_display_lines_use_material_business_order_and_aggregate_hardware(self):
        # 构造成本明细，按数量和单价计算金额。
        # category：成本项目类别。
        # code：商品 SKU 编码。
        # name：测试样本名称。
        # quantity：测试项目数量。
        # factory：明细所属工厂单或材料汇总标签。
        # unit：商品计量单位。
        # price：单件成本价。
        def line(category, code, name, quantity, *, factory="材料汇总", unit="pcs", price=1.0):
            return {
                "category": category,
                "factory_order": factory,
                "room_name": "",
                "name": name,
                "spec": "",
                "quantity": quantity,
                "unit": unit,
                "product_code": code,
                "cost_price": price,
                "amount": quantity * price,
                "source": "商品编号",
                "missing": "",
            }

        raw_lines = [
            line("板材", "M0049", "Cashmere SM", 7),
            line("封边条", "M0050", "Cashmere SM Edge Banding", 20, unit="m"),
            line("板材", "M0002", "1/4 Finished,UV1S", 4),
            line("板材", "M0004", "3/4 Finished,UV2S", 11),
            line("五金", "M1001", "Unihopper Hinge", 4, factory="F100", unit="Piece", price=0.47),
            line("板材", "M0003", "5/8 Finished,UV2S", 3),
            line("五金", "M1001", "Unihopper Hinge", 50, factory="F200", unit="Piece", price=0.47),
        ]

        display_lines = _display_cost_lines(raw_lines)
        self.assertEqual(
            [item["product_code"] for item in display_lines],
            ["M0004", "M0003", "M0002", "M0049", "M0050", "M1001"],
        )
        hardware = display_lines[-1]
        self.assertEqual(hardware["factory_order"], "五金汇总")
        self.assertAlmostEqual(hardware["quantity"], 54)
        self.assertAlmostEqual(hardware["amount"], 25.38)


if __name__ == "__main__":
    unittest.main()

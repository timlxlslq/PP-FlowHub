"""Legacy Chinese AICNC reports share identity and quantity safety contracts."""
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from traveler_assistant.core import RuleError, parse_fittings_groups
from traveler_assistant.fittings import is_fittings_report
from traveler_assistant.order_index import _report_files
from traveler_assistant.order_workflow import _choose_fittings


class LegacyFittingsTests(unittest.TestCase):
    def test_discovery_and_parsing_across_report_versions(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for legacy in (False, True):
                with self.subTest(legacy=legacy):
                    folder = root / str(legacy)
                    folder.mkdir()
                    path = folder / ('五金料单F123.xlsx' if legacy else 'FittingslistF123.xlsx')
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet.title = 'Page1'
                    sheet['A3'] = '订 单 号' if legacy else 'Order No.'
                    sheet['C3'] = 'F123'
                    for col, label in ({3: '名称', 5: '编号', 6: '尺寸', 11: '数量'} if legacy else {3: 'Name', 5: 'Code', 6: 'Size', 11: 'Quantity'}).items():
                        sheet.cell(8, col, label)
                    for row, name in ((9, 'Left Rail'), (10, 'Right Rail')):
                        sheet.cell(row, 3, name)
                        sheet.cell(row, 5, 'H-Rail')
                        sheet.cell(row, 11, '9')
                    workbook.save(path)
                    self.assertTrue(is_fittings_report(path))
                    self.assertFalse(is_fittings_report(path.with_name('~$' + path.name)))
                    self.assertEqual(_report_files(folder), [(path, 'fittings')])
                    groups = parse_fittings_groups(path)
                    self.assertEqual(groups[0][0], 'F123')
                    self.assertEqual([(i.name, i.quantity) for i in groups[0][1]], [('Left Rail', 9)])
                    self.assertEqual(_choose_fittings(folder)[0]['F123'], groups[0][1])
                    sheet['K10'] = '8'
                    workbook.save(path)
                    with self.assertRaises(RuleError) as error:
                        parse_fittings_groups(path)
                    self.assertEqual(error.exception.code, 'paired_rail_quantity_mismatch')
                    sheet['C3'] = 'NOT-A-FACTORY'
                    workbook.save(path)
                    with self.assertRaises(RuleError) as error:
                        parse_fittings_groups(path)
                    self.assertEqual(error.exception.code, 'fittings_identity')
                    sheet['C3'] = 'F123'
                    sheet['K8'] = 'Unknown'
                    workbook.save(path)
                    with self.assertRaises(RuleError) as error:
                        parse_fittings_groups(path)
                    self.assertEqual(error.exception.code, 'fittings_schema')

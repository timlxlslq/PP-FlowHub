import json
import io
import os
import sqlite3
import subprocess
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from traveler_assistant.core import Config, RuleError
from traveler_assistant.database import connect_database
from traveler_assistant.inventory import (
    InventoryMappings,
    InventoryOperationJournal,
    InventoryPreview,
    InventorySyncStore,
    OutboundItem,
    Product,
    ProductDatabase,
    ProductCatalog,
    _catalog_change_summary,
    close_inventory_chrome,
    _keychain_password,
    _jdy_error_detail,
    _find_existing_inventory_page,
    _resolve_jdy_runtime,
    _is_inventory_authenticated_url,
    _is_inventory_domain_url,
    _is_inventory_service_workbench_url,
    _persist_completed_outbound_results,
    _persist_single_outbound_result,
    _assert_single_server_material_source,
    build_database_preview,
    database_document_items,
    build_preview,
    check_stock,
    database_outbound_fingerprint,
    import_catalog,
    list_traveler_names,
    mark_no_hardware_outbound,
    mark_customer_supplied_outbound,
    open_inventory_chrome,
    order_stock_requirements,
    parse_traveler,
    list_inventory_mappings,
    remove_manual_mapping,
    reconcile_folder_status,
    run_jdy,
    resolve_inventory_items,
    outbound_scope_decisions,
    set_ignored_mapping,
    save_manual_mapping,
    set_outbound_scope,
    stock_requirements,
    bootstrap_product_database,
    TravelerItem,
    TravelerData,
    update_catalog_online,
    update_manual_mapping,
)


class _FakeNodeInput:
    """金蝶流式 Node 子进程测试使用的最小可写标准输入替身。"""

    # 初始化模拟输入流的内容和关闭状态。
    # self：当前测试用例或测试替身实例。
    def __init__(self):
        self.value = ""
        self.closed = False

    # 将输入文本追加到模拟标准输入，并返回写入字符数。
    # self：当前测试用例或测试替身实例。
    # value：要写入模拟输入流的文本。
    def write(self, value):
        self.value += value
        return len(value)

    # 标记模拟输入流已关闭。
    # self：当前测试用例或测试替身实例。
    def close(self):
        self.closed = True


class _FakeNodeProcess:
    """模拟 Popen 进程接口：标准输出提供最终 JSON，标准错误提供实时进度。"""

    # 初始化模拟子进程的输入输出、退出状态及超时行为。
    # self：当前测试用例或测试替身实例。
    # result：模拟子进程的退出码及标准输出和错误输出。
    # timeout：模拟等待超时值或异常。
    def __init__(self, result=None, timeout=None):
        result = result or SimpleNamespace(returncode=0, stdout="", stderr="")
        self.stdin = _FakeNodeInput()
        self.stdout = io.StringIO(result.stdout)
        self.stderr = io.StringIO(result.stderr)
        self.returncode = result.returncode
        self._timeout = timeout
        self.wait_timeout = None
        self.killed = False

    # 记录等待时限，按配置抛出超时或返回模拟退出码。
    # self：当前测试用例或测试替身实例。
    # timeout：模拟等待超时值或异常。
    def wait(self, timeout=None):
        self.wait_timeout = timeout
        if self._timeout is not None and timeout is not None:
            raise self._timeout
        return self.returncode

    # 记录模拟进程已被终止。
    # self：当前测试用例或测试替身实例。
    def kill(self):
        self.killed = True


from traveler_assistant.order_index import (
    OrderIndexStore,
    _load_outbound_records,
    _has_factory_hardware_outbound_record,
    _refresh_outbound_status,
    assert_factory_orders_outbound_allowed,
    reconcile_outbound_statuses,
)


# 按材料与五金项目生成包含用料表和领料单的 Traveler。
# path：测试文件的读写路径。
# items：来源名称与数量的项目列表。
def make_traveler(path: Path, items):
    workbook = Workbook()
    main = workbook.active
    main.title = "WorkOrderTraveler"
    main.append(["Name/工厂单名称", "PP0099"])
    usage = workbook.create_sheet("Usage List")
    usage.append(["Job:", "PP0099"])
    usage.append([
        "Room/section", "", "3/4 Plywood", "5/8 Plywood", "1/4 Plywood",
        "3/4 Finish Panel", "1/4 Finish Panel", "Edge Banding (m)", "Color",
    ])
    plywood = {"18mm--Plywood": 0, "14.5mm--Plywood": 0, "5.4mm--Plywood": 0}
    colored_materials = {}
    hardware = []
    for name, quantity in items:
        if name in plywood:
            plywood[name] = quantity
        elif "--" in name and (
            name.lower().startswith(("8mm--", "9mm--", "19.1mm--", "edge banding--"))
        ):
            color = name.split("--", 1)[1]
            values = colored_materials.setdefault(color, {"panel": 0, "back": 0, "edge": 0})
            if name.lower().startswith("19.1mm--"):
                values["panel"] = quantity
            elif name.lower().startswith(("8mm--", "9mm--")):
                values["back"] = quantity
            else:
                values["edge"] = quantity
        else:
            hardware.append((name, quantity))
    colors = list(colored_materials) or ["TEST COLOR"]
    for index, color in enumerate(colors):
        values = colored_materials.get(color, {"panel": 0, "back": 0, "edge": 0})
        usage.append([
            "Test" if index == 0 else "", "",
            plywood["18mm--Plywood"] if index == 0 else 0,
            plywood["14.5mm--Plywood"] if index == 0 else 0,
            plywood["5.4mm--Plywood"] if index == 0 else 0,
            values["panel"], values["back"], values["edge"], color,
        ])
    usage.cell(13, 1).value = None
    usage.append(["Total Qty:"])
    usage.append(["Color Table"])
    usage.append(["Color:"])
    usage.append(["Sheets (3/4):"])
    usage.append(["Sheets (1/4):"])
    usage.append(["Edge Banding (m):"])
    picking = workbook.create_sheet("Picking List")
    picking.append(["Picking List领料单"])
    picking.append(["Name/工厂单名称", "PP0099-KITCHEN"])
    picking.append(["Panel板材"])
    picking.append(["No.", "Name名字", "QTY数量"])
    for number, (name, quantity) in enumerate(hardware, 1):
        picking.append([number, name, quantity])
    workbook.save(path)


# 生成含库存单位的最小商品目录工作簿。
# path：测试文件的读写路径。
def make_catalog(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品类别", "*商品编号", "商品名称", "规格型号", "状态", "计量单位"])
    rows = [
        ("Plywood", "M0004", "3/4 Finished UV2S", "18mm", "启用", "张"),
        ("Hardware", "M1001", "Unihopper Hinge", "", "启用", "件"),
        ("Hardware", "M1068", "Push Open A", "", "启用", "件"),
        ("Hardware", "M1069", "Push Open B", "", "启用", "件"),
        ("Edge band", "M0020", "Woodline 4 Edge Banding", "22mm", "启用", "m"),
    ]
    for row in rows:
        sheet.append(row)
    workbook.save(path)


# 生成同时包含已知成本价和缺失成本价的商品目录。
# path：测试文件的读写路径。
def make_priced_catalog(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商品类别", "*商品编号", "商品名称", "规格型号", "状态", "计量单位", "预计采购价"])
    sheet.append(["Plywood", "M0004", "3/4 Finished UV2S", "18mm", "启用", "张", 31.76])
    sheet.append(["Panel", "M0005", "No price panel", "19.1mm", "启用", "张", None])
    workbook.save(path)


# 写入数据库测试样本所需的有限商品目录。
# connection：隔离测试数据库连接。
def seed_sku_products(connection) -> None:
    rows = [
        ("Plywood", "M0004", "3/4 Finished UV2S", "18mm", "启用", "张", "plywood", "", "18"),
        ("Edge band", "M0020", "Woodline 4 Edge Banding", "22mm", "启用", "m", "edge", "Woodline 4", ""),
        ("Panel", "M-CUSTOM", "Customer Panel", "19.1mm", "启用", "张", "panel", "Customer Panel", "19.1"),
        ("Panel", "M-BLANCO", "Blanco HG", "19.1mm", "启用", "张", "panel", "Blanco HG", "19.1"),
        ("Hardware", "M1001", "Unihopper Hinge", "", "启用", "件", "", "", ""),
        ("Hardware", "M1003", "L-Rail", "", "启用", "件", "", "", ""),
        ("Hardware", "M1013", "Adjustable shelf holder", "", "启用", "件", "", "", ""),
        ("Hardware", "M-LED", "LED", "", "启用", "件", "", "", ""),
    ]
    connection.executemany(
        """insert or ignore into products(
               category, code, name, spec, status, unit,
               normalized_code, normalized_name, normalized_spec,
               normalized_category, normalized_remark,
               material_kind, material_color, material_thickness, catalog_present
           ) values(?,?,?,?,?,?,replace(lower(?),'-',''),replace(lower(?),' ',''),
                    replace(lower(?),'.',''),lower(?),'',?,?,?,1)""",
        [
            (
                category, code, name, spec, status, unit,
                code, name, spec, category, kind, color, thickness,
            )
            for category, code, name, spec, status, unit, kind, color, thickness in rows
        ],
    )


class InventoryTests(unittest.TestCase):
    # 验证存在多个基础 Server 材料来源时阻止数据库出库。
    # self：当前测试用例或测试替身实例。
    def test_database_outbound_blocks_multiple_base_server_material_sources(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            connection = connect_database(config.workflow_database)
            seed_sku_products(connection)
            for path in (
                root / "server" / "Optimized Orders" / "PP0072" / "pp0072 materials.xlsx",
                root / "fixtures" / "Optimized Orders" / "PP0072" / "pp0072 materials.xlsx",
            ):
                connection.execute(
                    """insert into material_items(
                        order_id, product_code, quantity, source_type, source_path,
                        source_fingerprint, updated_at
                    ) values(?,?,?,?,?,?,?)""",
                    ("PP0072", "M0004", 22, "aihouse", str(path), "fixture-fingerprint", "now"),
                )
            connection.commit()
            connection.close()

            with self.assertRaisesRegex(RuleError, "多个 Server 材料来源"):
                _assert_single_server_material_source(config, "PP0072")

    # 验证没有五金的出货只更新状态，不创建库存单据。
    # self：当前测试用例或测试替身实例。
    def test_shipment_without_hardware_marks_status_without_inventory_document(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state", order_root=root / "generated")
            config.prepare_storage()
            (config.state_dir / "inventory").mkdir(parents=True, exist_ok=True)
            make_catalog(config.state_dir / "inventory" / "current-products.xlsx")
            store = OrderIndexStore(config.workflow_database)
            store.connection.execute(
                "insert into orders(order_id, order_type, updated_at) values(?,?,?)",
                ("PP0100", "owned", "now"),
            )
            store.connection.execute(
                "insert into factory_orders(factory_order,order_id,factory_name,updated_at,stage) select v.factory_order,v.order_id,v.factory_name,v.updated_at,case when v.outbound_status='已出库' then '已出货' when v.optimized then '已优化' else '已拆单' end from (select ? as factory_order,? as order_id,? as factory_name,? as optimized,? as outbound_status,? as updated_at) v",
                ("F1000", "PP0100", "PP0100-NO-HARDWARE", 1, "未出库", "now"),
            )
            store.commit()
            store.close()

            preview = build_database_preview(config, "PP0100", ["F1000"], shipment_only=True)
            self.assertTrue(preview.ready)
            self.assertTrue(preview.no_outbound_required)
            self.assertEqual(preview.outbound_items, [])

            result = mark_no_hardware_outbound(config, "PP0100", ["F1000"])
            self.assertEqual(result["outbound_mode"], "no_hardware")
            row = sqlite3.connect(config.workflow_database).execute(
                "select (case when stage='已出货' then '已出库' else '未出库' end) as outbound_status, outbound_document, outbound_mode from factory_orders where factory_order='F1000'"
            ).fetchone()
            self.assertEqual(row, ("已出库", "", "no_hardware"))

    # 验证客供订单出库仅记录数据库状态，事实变化后重新要求处理。
    # self：当前测试用例或测试替身实例。
    def test_customer_supplied_outbound_marks_database_only_and_reopens_on_fact_change(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.connection.execute(
                "insert into orders(order_id, order_type, updated_at) values(?,?,?)",
                ("CS001", "cutToSize", "now"),
            )
            store.connection.execute(
                "insert into factory_orders(factory_order,order_id,factory_name,updated_at,stage) select v.factory_order,v.order_id,v.factory_name,v.updated_at,case when v.outbound_status='已出库' then '已出货' when v.optimized then '已优化' else '已拆单' end from (select ? as factory_order,? as order_id,? as factory_name,? as optimized,? as outbound_status,? as updated_at) v",
                ("F1001", "CS001", "CS001-KITCHEN", 1, "未出库", "now"),
            )
            store.connection.execute(
                "insert into material_items(order_id, product_code, quantity, source_type, updated_at) values(?,?,?,?,?)",
                ("CS001", "M-CUSTOM", 4, "aihouse", "now"),
            )
            store.commit()
            store.close()

            set_outbound_scope(
                config,
                "CS001",
                "material",
                "customer_supplied",
                reason="客户提供材料，本公司只负责加工",
            )
            fingerprint = database_outbound_fingerprint(config, "CS001", "F1001")
            result = mark_customer_supplied_outbound(config, "CS001", ["F1001"])
            self.assertEqual(result["outbound_mode"], "customer_supplied")
            self.assertEqual(result["outbound_status"], "已出库")

            connection = sqlite3.connect(config.workflow_database)
            row = connection.execute(
                "select (case when stage='已出货' then '已出库' else '未出库' end) as outbound_status, outbound_document, outbound_mode, outbound_fingerprint from factory_orders where factory_order='F1001'"
            ).fetchone()
            self.assertEqual(row, ("已出库", "", "customer_supplied", fingerprint))
            self.assertEqual(
                _refresh_outbound_status(
                    config,
                    {"order_id": "CS001", "factory_order": "F1001", "factory_name": "CS001-KITCHEN"},
                    [],
                ),
                ("已出库", ""),
            )
            connection.execute(
                "update material_items set quantity=5 where order_id='CS001'"
            )
            connection.commit()
            connection.close()
            self.assertNotEqual(
                database_outbound_fingerprint(config, "CS001", "F1001"),
                fingerprint,
            )
            self.assertEqual(
                _refresh_outbound_status(
                    config,
                    {"order_id": "CS001", "factory_order": "F1001", "factory_name": "CS001-KITCHEN"},
                    [],
                ),
                ("需要更新", ""),
            )

    # 验证五金出库范围必须具有实际正数量的五金事实。
    # self：当前测试用例或测试替身实例。
    def test_hardware_scope_requires_actual_positive_hardware_facts(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            store.connection.execute(
                "insert into orders(order_id, order_type, updated_at) values(?,?,?)",
                ("CS004", "cutToSize", "now"),
            )
            store.connection.execute(
                "insert into factory_orders(factory_order,order_id,factory_name,updated_at,stage) select v.factory_order,v.order_id,v.factory_name,v.updated_at,case when v.outbound_status='已出库' then '已出货' when v.optimized then '已优化' else '已拆单' end from (select ? as factory_order,? as order_id,? as factory_name,? as optimized,? as outbound_status,? as updated_at) v",
                ("F4004", "CS004", "CS004-KITCHEN", 1, "未出库", "now"),
            )
            store.commit()
            store.close()

            scope = outbound_scope_decisions(config, "CS004")
            self.assertEqual(scope["hardware"], {})
            with self.assertRaisesRegex(RuleError, "没有.*可出库五金数据"):
                set_outbound_scope(config, "CS004", "hardware", "required", factory_order="F4004")

    # 验证读取出库范围时采用最新且相关的已保存决定。
    # self：当前测试用例或测试替身实例。
    def test_outbound_scope_read_returns_latest_relevant_saved_decision(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.connection.execute(
                "insert into orders(order_id, order_type, updated_at) values(?,?,?)",
                ("CS005", "cutToSize", "now"),
            )
            store.connection.execute(
                "insert into factory_orders(factory_order,order_id,factory_name,updated_at,stage) select v.factory_order,v.order_id,v.factory_name,v.updated_at,case when v.outbound_status='已出库' then '已出货' when v.optimized then '已优化' else '已拆单' end from (select ? as factory_order,? as order_id,? as factory_name,? as optimized,? as outbound_status,? as updated_at) v",
                ("F5005", "CS005", "CS005-KITCHEN", 1, "未出库", "now"),
            )
            store.connection.execute(
                "insert into material_items(order_id, product_code, quantity, source_type, updated_at) values(?,?,?,?,?)",
                ("CS005", "M-BLANCO", 1, "manual", "now"),
            )
            store.connection.execute(
                'insert into hardware_items(order_id,factory_order,scope,product_code,quantity,source_type,updated_at) values(?,?,?,?,?,?,?)',
                ('CS005', 'F5005', 'factory_order', 'M1001', 1, 'database', 'now'),
            )
            store.commit()
            store.close()

            set_outbound_scope(config, "CS005", "material", "customer_supplied", reason="客户提供板材")
            material_scope = outbound_scope_decisions(config, "CS005", ["F5005"])
            self.assertEqual(material_scope["last_decision"]["scope_type"], "material")
            self.assertEqual(material_scope["material"]["requirement"], "customer_supplied")
            self.assertEqual(material_scope["last_decision"]["reason"], "客户提供板材")

            set_outbound_scope(config, "CS005", "hardware", "required", factory_order="F5005")
            hardware_scope = outbound_scope_decisions(config, "CS005", ["F5005"])
            self.assertEqual(hardware_scope["last_decision"]["scope_type"], "hardware")
            self.assertEqual(hardware_scope["last_decision"]["factory_order"], "F5005")

    # 验证代切客供材料保留为事实，但不作为库存出库项目。
    # self：当前测试用例或测试替身实例。
    def test_cut_to_size_customer_supplied_material_stays_fact_but_is_not_outbound(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state", order_root=root / "generated")
            config.prepare_storage()
            (config.state_dir / "inventory").mkdir(parents=True, exist_ok=True)
            make_catalog(config.state_dir / "inventory" / "current-products.xlsx")
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.connection.execute(
                "insert into orders(order_id, order_type, source_folder, updated_at) values(?,?,?,?)",
                ("CS002", "cutToSize", str(root / "source" / "CS002"), "now"),
            )
            store.connection.execute(
                "insert into factory_orders(factory_order,order_id,factory_name,updated_at,stage) select v.factory_order,v.order_id,v.factory_name,v.updated_at,case when v.outbound_status='已出库' then '已出货' when v.optimized then '已优化' else '已拆单' end from (select ? as factory_order,? as order_id,? as factory_name,? as optimized,? as outbound_status,? as updated_at) v",
                ("F2002", "CS002", "CS002-Hardware", 1, "未出库", "now"),
            )
            store.connection.execute(
                "insert into material_items(order_id, product_code, quantity, source_type, updated_at) values(?,?,?,?,?)",
                ("CS002", "M0004", 5, "aihouse", "now"),
            )
            store.connection.execute(
                "insert into material_items(order_id, product_code, quantity, source_type, updated_at) values(?,?,?,?,?)",
                ("CS002", "M0020", 100, "aihouse", "now"),
            )
            store.connection.execute(
                'insert into hardware_items(order_id,factory_order,scope,product_code,quantity,source_type,updated_at) values(?,?,?,?,?,?,?)',
                ('CS002', 'F2002', 'factory_order', 'M1001', 2, 'aicnc', 'now'),
            )
            store.commit()
            store.close()
            set_outbound_scope(config, "CS002", "material", "customer_supplied", reason="客户提供板材和封边")
            with patch("traveler_assistant.inventory.bootstrap_product_database", return_value=config.state_dir / "inventory" / "current-products.xlsx"):
                preview = build_database_preview(config, "CS002", ["F2002"])
            self.assertTrue(preview.ready)
            self.assertEqual({item.product_code for item in preview.outbound_items}, {"M1001"})
            self.assertEqual(preview.outbound_items[0].document_remark, "CS002-Hardware")
            self.assertEqual(preview.scope_decisions[0]["requirement"], "customer_supplied")
            facts = sqlite3.connect(config.workflow_database).execute(
                "select quantity from material_items where order_id='CS002' order by product_code"
            ).fetchall()
            self.assertEqual([row[0] for row in facts], [5.0, 100.0])

    # 验证工厂单名称的订单前缀不匹配时阻止出库。
    # self：当前测试用例或测试替身实例。
    def test_database_outbound_blocks_mismatched_factory_name_order_prefix(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.connection.execute(
                "insert into orders(order_id, order_type, updated_at) values(?,?,?)",
                ("PP0072", "owned", "now"),
            )
            store.connection.execute(
                "insert into factory_orders(factory_order,order_id,factory_name,sales_order_name,updated_at,stage) select v.factory_order,v.order_id,v.factory_name,v.sales_order_name,v.updated_at,case when v.outbound_status='已出库' then '已出货' when v.optimized then '已优化' else '已拆单' end from (select ? as factory_order,? as order_id,? as factory_name,? as sales_order_name,? as optimized,? as outbound_status,? as updated_at) v",
                ("F2608190230", "PP0072", "P0072-BED 2", "PP0072", 1, "未出库", "now"),
            )
            store.connection.execute(
                'insert into hardware_items(order_id,factory_order,scope,product_code,quantity,source_type,updated_at) values(?,?,?,?,?,?,?)',
                ('PP0072', 'F2608190230', 'factory_order', 'M1001', 1, 'aicnc', 'now'),
            )
            store.commit()
            store.close()

            with self.assertRaisesRegex(RuleError, "名称前缀 P0072.*PP0072"):
                database_document_items(config, "PP0072", ["F2608190230"])

    # 验证自有订单不接受客供出库范围决定。
    # self：当前测试用例或测试替身实例。
    def test_customer_supplied_scope_is_rejected_for_owned_order(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            store.connection.execute(
                "insert into orders(order_id, order_type, updated_at) values(?,?,?)",
                ("PP9999", "owned", "now"),
            )
            store.commit()
            store.close()
            with self.assertRaisesRegex(RuleError, "自有订单不支持设置出库范围"):
                set_outbound_scope(config, "PP9999", "material", "customer_supplied", reason="误操作")

    # 验证自有订单即使收到标准化后的客供范围也会拒绝。
    # self：当前测试用例或测试替身实例。
    def test_owned_order_rejects_even_normalized_outbound_scope_decision(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            store.connection.execute(
                "insert into orders(order_id, order_type, updated_at) values(?,?,?)",
                ("PP9998", "owned", "now"),
            )
            store.connection.commit()
            store.close()
            with self.assertRaisesRegex(RuleError, "自有订单不支持设置出库范围"):
                set_outbound_scope(config, "PP9998", "material", "required")

    # 验证余料处理决定允许空订单完成而不打开浏览器。
    # self：当前测试用例或测试替身实例。
    def test_remainder_decision_allows_empty_order_without_opening_browser(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            store.connection.execute(
                "insert into orders(order_id, order_type, updated_at) values(?,?,?)",
                ("CS002", "cutToSize", "now"),
            )
            store.commit()
            store.close()
            set_outbound_scope(config, "CS002", "material", "remainder", reason="本单只用余料生产")
            preview = build_database_preview(config, "CS002", [])
            self.assertTrue(preview.ready)
            self.assertTrue(preview.no_outbound_required)
            self.assertEqual(preview.outbound_items, [])

    # 验证数据库订单出库预览无需 Traveler 文件即可映射商品。
    # self：当前测试用例或测试替身实例。
    def test_database_order_outbound_preview_maps_without_traveler_file(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            state = root / "state"
            config = Config(state_dir=state, order_root=root / "generated")
            config.prepare_storage()
            (state / "inventory").mkdir(parents=True, exist_ok=True)
            make_catalog(state / "inventory" / "current-products.xlsx")
            (state / "inventory" / "mappings.json").write_text(
                '{"manual": {}, "ignored": {}}', encoding="utf-8"
            )
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.connection.execute(
                "insert into orders(order_id, order_type, source_folder, updated_at) values(?,?,?,?)",
                ("PP9999", "owned", str(root / "source" / "PP9999"), "2026-08-15T10:00:00"),
            )
            store.connection.execute(
                "insert into factory_orders(factory_order,order_id,factory_name,updated_at,stage) select v.factory_order,v.order_id,v.factory_name,v.updated_at,case when v.outbound_status='已出库' then '已出货' when v.optimized then '已优化' else '已拆单' end from (select ? as factory_order,? as order_id,? as factory_name,? as optimized,? as outbound_status,? as updated_at) v",
                ("F9999", "PP9999", "PP9999-KITCHEN", 1, "未出库", "2026-08-15T10:00:00"),
            )
            store.connection.execute(
                "insert into material_items(order_id, product_code, quantity, source_type, updated_at) values(?,?,?,?,?)",
                ("PP9999", "M0004", 2, "database", "2026-08-15T10:00:00"),
            )
            store.connection.execute(
                "insert into material_items(order_id, product_code, quantity, source_type, updated_at) values(?,?,?,?,?)",
                ("PP9999", "M0020", 12.5, "database", "2026-08-15T10:00:00"),
            )
            store.connection.execute(
                'insert into hardware_items(order_id,factory_order,scope,product_code,quantity,source_type,updated_at) values(?,?,?,?,?,?,?)',
                ('PP9999', 'F9999', 'factory_order', 'M1001', 3, 'database', '2026-08-15T10:00:00'),
            )
            store.connection.commit()
            store.close()

            with patch(
                "traveler_assistant.inventory.bootstrap_product_database",
                return_value=state / "inventory" / "current-products.xlsx",
            ):
                preview = build_database_preview(config, "PP9999", ["F9999"])

            self.assertTrue(preview.ready)
            self.assertEqual(
                {(item.document_remark, item.product_code) for item in preview.outbound_items},
                {("PP9999", "M0004"), ("PP9999", "M0020"), ("PP9999-KITCHEN", "M1001")},
            )
            self.assertEqual(preview.traveler.path, config.workflow_database.resolve())
            self.assertFalse((config.order_root / "PP9999" / "Work Order Traveler(PP9999).xlsx").exists())

    # 验证订单上下文来源不再以 Traveler 存在为前提。
    # self：当前测试用例或测试替身实例。
    def test_order_context_source_no_longer_requires_traveler_gate(self):
        root = Path(__file__).resolve().parents[1]
        dashboard = (root / "macos" / "OrderDashboardView.swift").read_text(encoding="utf-8")
        swift = (root / "macos" / "TravelerAssistant.swift").read_text(encoding="utf-8")
        self.assertNotIn("model.orderExistingTravelerPath.isEmpty", dashboard)
        self.assertIn("func startDirectProduction", swift)
        self.assertIn("func startDirectOrderShipment", swift)
        self.assertIn('"--production-materials-json"', swift)
        self.assertIn('var arguments = ["outbound", "--order-id", normalizedOrderID, "--shipment-only", "--confirm-save"]', swift)
        self.assertIn('beginDashboardOperation("inventory", label: "库存系统出货")', swift)
        self.assertIn("inventoryInactivityTimeoutSeconds", swift)
        self.assertIn("DispatchSource.makeTimerSource", swift)
        self.assertIn("inventoryFailureNeedsVerification", swift)
        self.assertIn('"get-outbound-scope", "--order-id"', swift)
        self.assertIn("loadOutboundScope", swift)
        self.assertIn('Button("确认并直接出货")', dashboard)
        self.assertIn('Button("确认生产并扣减材料")', dashboard)
        self.assertNotIn("case .inventory", swift)
        self.assertNotIn("selection = .inventory", swift)
        self.assertNotIn("onChange(of: model.inventoryMappingRequestPath", swift)
        self.assertIn("showInventoryMappingWorkspace", swift)
        self.assertIn("inventoryMappingTargetNames", swift)
        self.assertIn("PendingInventoryMappingWorkspace", swift)
        self.assertIn("PendingInventoryIgnoreSheet", swift)
        self.assertIn('Button("忽略")', swift)
        self.assertIn('Button("加入全局忽略")', swift)
        self.assertIn('Button("处理映射")', swift)
        self.assertIn("inventoryMappingSourceFolderPath", swift)
        resume = swift.split("private func resumePendingMappingOperationAfterMapping()", 1)[1].split(
            "func retryPendingMappingPreview()", 1
        )[0]
        self.assertIn('"preview-server-changes"', resume)
        self.assertNotIn('"list-index"', resume)
        self.assertNotIn('"process-server-folder"', resume)
        self.assertIn("OrderShipmentConfirmationSheet(", dashboard)
        detail_card = dashboard.split("struct OrderDashboardDetailCard", 1)[1].split(
            "struct OutboundScopeSheet", 1
        )[0]
        self.assertIn("?.produced != true", detail_card)
        self.assertIn("AppLayout.controlHeight", detail_card)
        self.assertIn('if orderType != "owned"', detail_card)
        self.assertIn('Button { onOpenScope() } label:', detail_card)
        self.assertIn('Text("设置出库范围")', detail_card)
        detail = dashboard.split("struct OrderDashboardDetailPage", 1)[1].split(
            "private func orderInstallationDateFormatter", 1
        )[0]
        self.assertNotIn('Text("订单详情")', detail)
        self.assertIn("VStack(alignment: .leading, spacing: 14)", detail)
        self.assertIn("min(max(0, geometry.size.width - AppLayout.contentPadding * 2), 1120)", detail)
        self.assertIn("ScrollView(.vertical)", detail)
        self.assertEqual(detail.count("ScrollView(.vertical)"), 1)
        self.assertIn("ScrollView(.vertical)", detail)
        self.assertIn("boardAndEdgeSection", detail)
        self.assertIn("hardwareSection", detail)
        self.assertIn("HStack(alignment: .center", detail)
        self.assertIn('Text("订单 (\\(order.orderId))")', detail)
        self.assertNotIn("order.orderType", detail)
        self.assertIn('Button("生成 Traveler")', detail)
        self.assertIn('Button("关闭")', detail)
        self.assertIn(".font(.title2.weight(.semibold))", detail)
        self.assertIn("orderDetailCardMinHeight", detail)
        self.assertIn("value: row.quantity.formatted()", detail)
        self.assertNotIn("model.loadOrderDetailFromDatabase(order)", detail)
        self.assertNotIn("if model.selectedOrderId.caseInsensitiveCompare(order.orderId)", detail)
        self.assertIn('本次生产消耗材料', swift)
        self.assertIn("orderDetailGridColumnCount", detail)
        self.assertIn("inventoryIgnoredMappings", swift)
        self.assertIn("inventoryManualMappings", swift)
        self.assertIn("五金全局忽略列表", swift)
        self.assertIn("struct InventoryManualMappingsSheet", swift)
        self.assertIn("saveInventoryIgnoredMapping", swift)
        self.assertIn("updateInventoryIgnoredMapping", swift)
        self.assertIn("removeInventoryIgnoredMapping", swift)
        self.assertIn("saveSettingsManualMapping", swift)
        self.assertIn("updateSettingsManualMapping", swift)
        self.assertIn("removeSettingsManualMapping", swift)
        self.assertIn("manual_display_names", swift)
        self.assertIn("displayName: String", swift)
        self.assertIn("显示名称（可选）", swift)
        settings = swift.split("struct SettingsView", 1)[1].split(
            "struct OperationLogViewerView", 1
        )[0]
        self.assertNotIn("Traveler 材料名称", settings)
        self.assertNotIn("设置材料映射", settings)
        self.assertIn("库存资料与规则", settings)
        self.assertIn("showIgnoredHardwareList = true", settings)
        self.assertIn("struct InventoryIgnoredMappingsSheet", swift)

    # 验证持久订单数据变化后数据库出库状态随之变化。
    # self：当前测试用例或测试替身实例。
    def test_database_outbound_status_changes_when_persisted_order_data_changes(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            state = root / "state"
            config = Config(state_dir=state, order_root=root / "generated")
            config.prepare_storage()
            (state / "inventory").mkdir(parents=True, exist_ok=True)
            make_catalog(state / "inventory" / "current-products.xlsx")
            (state / "inventory" / "mappings.json").write_text(
                '{"manual": {}, "ignored": {}}', encoding="utf-8"
            )
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.connection.execute(
                "insert into orders(order_id, order_type, source_folder, updated_at) values(?,?,?,?)",
                ("PP9999", "owned", str(root / "source" / "PP9999"), "2026-08-15T10:00:00"),
            )
            store.connection.execute(
                "insert into factory_orders(factory_order,order_id,factory_name,updated_at,stage) select v.factory_order,v.order_id,v.factory_name,v.updated_at,case when v.outbound_status='已出库' then '已出货' when v.optimized then '已优化' else '已拆单' end from (select ? as factory_order,? as order_id,? as factory_name,? as optimized,? as outbound_status,? as updated_at) v",
                ("F9999", "PP9999", "PP9999-KITCHEN", 1, "已出库", "2026-08-15T10:00:00"),
            )
            store.connection.execute(
                "insert into material_items(order_id, product_code, quantity, source_type, updated_at) values(?,?,?,?,?)",
                ("PP9999", "M0004", 2, "database", "2026-08-15T10:00:00"),
            )
            store.connection.execute(
                'insert into hardware_items(order_id,factory_order,scope,product_code,quantity,source_type,updated_at) values(?,?,?,?,?,?,?)',
                ('PP9999', 'F9999', 'factory_order', 'M1001', 3, 'database', '2026-08-15T10:00:00'),
            )
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.inventory.bootstrap_product_database",
                return_value=state / "inventory" / "current-products.xlsx",
            ):
                preview = build_database_preview(config, "PP9999", ["F9999"])

            factory = {
                "order_id": "PP9999",
                "factory_order": "F9999",
                "factory_name": "PP9999-KITCHEN",
            }
            record = {
                "order_id": "PP9999",
                "remark": "PP9999-KITCHEN",
                "status": "已出库",
                "document_number": "QTCK-1",
                "traveler_path": str(config.workflow_database),
                "raw_fingerprint": InventorySyncStore.raw_document_fingerprint(
                    preview.traveler.documents["PP9999-KITCHEN"]
                ),
            }
            self.assertEqual(_refresh_outbound_status(config, factory, [record]), ("已出库", "QTCK-1"))

            connection = sqlite3.connect(config.workflow_database)
            connection.execute(
                "update hardware_items set quantity=4 where order_id=? and factory_order=?",
                ("PP9999", "F9999"),
            )
            connection.commit()
            connection.close()
            self.assertEqual(_refresh_outbound_status(config, factory, [record]), ("需要更新", "QTCK-1"))

    # 验证库存页面识别支持租户工作台子域名。
    # self：当前测试用例或测试替身实例。
    def test_inventory_page_detection_accepts_tenant_workbench_subdomain(self):
        url = "https://vip2-hz.jdy.com/default-new.jsp?dbid=7937191793066#/beginner-guide"

        self.assertTrue(_is_inventory_domain_url(url))
        self.assertTrue(_is_inventory_authenticated_url(url))
        with patch(
            "traveler_assistant.inventory._inventory_cdp_pages",
            return_value=[{"type": "page", "url": url}],
        ):
            self.assertEqual(_find_existing_inventory_page("http://127.0.0.1:9222"), {"type": "page", "url": url})

    # 验证库存服务工作台可作为有效入口页。
    # self：当前测试用例或测试替身实例。
    def test_inventory_service_workbench_is_an_entry_page(self):
        service_url = "https://service.jdy.com/workbench/web/index.html"

        self.assertTrue(_is_inventory_authenticated_url(service_url))
        self.assertTrue(_is_inventory_service_workbench_url(service_url))

    # 验证登录页及非库存网址仍被排除为库存操作页。
    # self：当前测试用例或测试替身实例。
    def test_inventory_page_detection_still_rejects_login_and_non_inventory_urls(self):
        self.assertFalse(_is_inventory_authenticated_url("https://www.jdy.com/login/"))
        self.assertFalse(_is_inventory_authenticated_url("https://www.jdy.com/global/?logout=true"))
        self.assertFalse(_is_inventory_domain_url("https://example.com/default-new.jsp"))

    # 验证金蝶调试连接清理使用 Playwright 的浏览器关闭接口。
    # self：当前测试用例或测试替身实例。
    def test_jdy_cdp_cleanup_uses_playwright_browser_close(self):
        source = (Path(__file__).resolve().parents[1] / "tools" / "jdy_inventory.mjs").read_text(encoding="utf-8")

        self.assertIn("await browser.close()", source)
        self.assertIn("await remoteBrowser.close()", source)
        self.assertIn('request.action === "closeChrome"', source)
        self.assertIn('session.send("Browser.close")', source)
        self.assertNotIn("browser.disconnect()", source)
        self.assertNotIn("remoteBrowser.disconnect()", source)

    # 验证出库导航复用已有列表，并点击可见菜单项。
    # self：当前测试用例或测试替身实例。
    def test_outbound_navigation_reuses_existing_list_and_opens_visible_menu_item(self):
        source = (Path(__file__).resolve().parents[1] / "tools" / "jdy_inventory.mjs").read_text(encoding="utf-8")

        self.assertIn("const currentOtherOutboundListFrame", source)
        self.assertIn("const isOtherOutboundListFrame", source)
        self.assertIn("const isOtherOutboundListURL", source)
        self.assertIn("const isOtherOutboundFormURL", source)
        self.assertIn("storage/other-outbound", source)
        form_check = source.split("const isOtherOutboundFormFrame = async frame => {", 1)[1].split("\n};", 1)[0]
        self.assertIn("const hasSave = await hasVisibleOutboundSaveControl(frame);", form_check)
        self.assertIn('const hasTable = await frame.locator("thead:visible th").count() > 0;', form_check)
        self.assertIn("return hasSave && hasTable && (", form_check)
        self.assertIn("action=initOiList", source)
        self.assertIn("const currentOtherOutboundFormFrame", source)
        self.assertIn("const isServiceWorkbenchURL", source)
        self.assertIn("const ensureInventoryBusinessWorkbench", source)
        self.assertIn("const waitForInventoryActionShell", source)
        self.assertIn("点击服务工作台“进入使用”", source)
        self.assertIn("const waitForVisibleLeftNavigationItem", source)
        self.assertIn("const waitForOtherOutboundListFrame", source)
        self.assertIn("const waitForOtherOutboundFormFrame", source)
        self.assertIn("const clickOtherOutboundHistory", source)
        self.assertIn("const clickOtherOutboundHistoryWithRetry", source)
        self.assertIn("其他出库单中的历史单据", source)
        self.assertIn("准备点击“历史单据”进入记录列表", source)
        self.assertIn("const assertOutboundFormMatchesRequest", source)
        self.assertIn("const outboundMaterialRows", source)
        self.assertNotIn("const saveAvailable", source)
        self.assertIn("按直接编辑表单处理", source)
        self.assertIn("const waitForOutboundSaveControl", source)
        self.assertIn("#edit:visible", source)
        self.assertIn("aria-disabled", source)
        self.assertIn("历史单据回读确认成功", source)
        self.assertIn('td[aria-describedby="grid_invNumber"]', source)
        self.assertIn("snapshot.productCode === item.productCode", source)
        self.assertNotIn("const meaningfulCells = cells.filter", source)
        self.assertIn("await waitForOtherOutboundListFrame(page)", source)
        self.assertIn("isOtherOutboundFormFrame(frame)", source)
        self.assertNotIn('log("已复用当前页面中的空白其他出库单表单")', source)
        self.assertIn(".quick-datepicker-start:visible", source)
        preferred_tab = source.index("if (await currentOtherOutboundListFrame(candidate) ||")
        fallback_tab = source.index("existingPage ||= inventoryPages.find")
        self.assertLess(preferred_tab, fallback_tab)
        self.assertIn("existingPage = candidate;", source[preferred_tab:fallback_tab])
        self.assertIn("break;", source[preferred_tab:fallback_tab])
        self.assertIn('log("已复用当前页面中的其他出库单记录列表")', source)
        self.assertIn('const openOtherOutboundMenuItem', source)
        self.assertIn('visibleTextLocatorsAcrossFrames(page, label)', source)
        self.assertIn("warehouse.hover({ force: true })", source)
        self.assertIn("本机已有出库单", source)
        self.assertIn("已停止，不新建重复出库单", source)

    # 验证出库超时后先保存已完成单据，再允许重试。
    # self：当前测试用例或测试替身实例。
    def test_outbound_timeout_persists_completed_documents_before_retry(self):
        source = (Path(__file__).resolve().parents[1] / "traveler_assistant" / "inventory.py").read_text(encoding="utf-8")
        self.assertIn("_persist_completed_outbound_results", source)
        self.assertIn("后续单据结果需要先查询库存历史再重试", source)
        self.assertIn("partial_external_confirmed", source)
        self.assertIn("_persist_single_outbound_result", source)

    # 验证金蝶出库使用直接编辑入口，保存后核对历史单据。
    # self：当前测试用例或测试替身实例。
    def test_jdy_outbound_uses_direct_edit_and_post_save_history_verification(self):
        source = (Path(__file__).resolve().parents[1] / "tools" / "jdy_inventory.mjs").read_text(encoding="utf-8")
        self.assertNotIn("const saveAvailable", source)
        self.assertIn("按直接编辑表单处理", source)
        self.assertIn("waitForOutboundSaveControl", source)
        self.assertIn("await quantityEditor.fill(String(item.quantity));", source)
        self.assertNotIn("quantityEditor.evaluate((input, value)", source)
        self.assertIn("历史单据回读确认成功", source)
        self.assertIn("findExactOutboundRows", source)
        listener_start = source.index("const loginCheckPromise = page.waitForResponse(")
        submit_start = source.index("loginSubmittedAt = performance.now();", listener_start)
        self.assertLess(listener_start, submit_start)
        self.assertIn('response.request().method() === "POST"', source[listener_start:submit_start])
        self.assertIn('response.url().includes("/commonservice/ajaxChecking.do")', source[listener_start:submit_start])
        self.assertIn("trace.zip", source)

    # 验证金蝶错误详情说明复用页面时的菜单超时原因。
    # self：当前测试用例或测试替身实例。
    def test_jdy_error_detail_explains_reused_page_menu_timeout(self):
        detail = _jdy_error_detail(
            "库存系统自动操作失败：左侧菜单中等待可见“仓库”超过 30 秒；当前页面："
            "https://vip2-hz.jdy.com/default-new.jsp"
        )
        self.assertIn("左侧“仓库”菜单", detail)

    # 验证出库预览隐藏底部写入提示文案。
    # self：当前测试用例或测试替身实例。
    def test_outbound_preview_hides_bottom_write_note(self):
        source = (Path(__file__).resolve().parents[1] / "macos" / "TravelerAssistant.swift").read_text(encoding="utf-8")
        self.assertNotIn("真实写入会创建库存出库单，并保存本机同步记录。", source)

    # 验证出库预览支持整行点击，各区域独立滚动。
    # self：当前测试用例或测试替身实例。
    def test_outbound_preview_rows_are_full_row_clickable_and_scroll_independently(self):
        root = Path(__file__).resolve().parents[1]
        source = (root / "macos" / "TravelerAssistant.swift").read_text(encoding="utf-8")
        dashboard_source = (root / "macos" / "OrderDashboardView.swift").read_text(encoding="utf-8")
        self.assertIn("private func previewRowContent(_ row: InventoryPreviewRow)", source)
        self.assertNotIn("selectedPreviewRows", source)
        self.assertNotIn("togglePreviewRow(row.id)", source)
        self.assertNotIn("点击记录行即可整行选中；按住 Command 可多选", source)
        self.assertIn(".frame(maxWidth: .infinity)\n        .frame(height: AppLayout.operationLogHeight)", source)
        self.assertRegex(
            source,
            r'HStack\(spacing: 8\)\s*\{\s*Button\("取消"\)\s*\{\s*confirmRealSave = false\s*\}',
        )
        self.assertIn("ScrollView(.vertical)", source)
        self.assertIn("inventoryWriteBlocked", source)
        self.assertIn("inventoryWriteCompleted", source)
        self.assertIn('text: model.inventoryWriteCompleted', source)
        self.assertIn('active: confirmRealSave || model.inventoryWriteCompleted', source)
        self.assertIn('self.inventoryWriteCompleted = true', source)
        self.assertIn('finishRunningInventoryStep(', source)
        self.assertIn('.contentShape(Rectangle())', source)
        self.assertNotIn('Text(model.inventorySuccessMessage.isEmpty ? model.inventoryStatus : model.inventorySuccessMessage)', source)
        self.assertNotIn('· 已选 \\(selectedPreviewRows.count)', source)
        self.assertNotIn('已选工厂单：', source)
        self.assertNotIn('Text("\\(model.inventoryPreviewRows.count) 行")', source)
        self.assertIn('if model.inventoryChromeStatus.hasPrefix("❌")', source)
        self.assertNotIn('Text("已选 \\(selectedFactoryIDs.count) 个工厂单")', dashboard_source)

    # 验证订单面板的出库刷新和饰面板颜色布局符合约定。
    # self：当前测试用例或测试替身实例。
    def test_order_dashboard_outbound_refresh_and_panel_color_layout_contract(self):
        root = Path(__file__).resolve().parents[1]
        source = (root / "macos" / "TravelerAssistant.swift").read_text(encoding="utf-8")
        dashboard_source = (root / "macos" / "OrderDashboardView.swift").read_text(encoding="utf-8")
        self.assertIn("func refreshDashboardOrdersAfterOutbound()", source)
        self.assertIn('runOrder(["list-index"], failureStatus: "订单列表刷新失败"', source)
        self.assertIn("pendingDashboardOutboundRefresh", source)
        self.assertIn("HStack(alignment: .center, spacing: AppLayout.actionSpacing)", source)
        self.assertIn(".frame(maxWidth: .infinity, alignment: .center)", source)
        self.assertIn("refreshDashboardOrdersAfterOutbound()", source)
        self.assertIn('.fixedSize(horizontal: true, vertical: false)', dashboard_source)
        self.assertIn('.lineLimit(2)', dashboard_source)
        self.assertIn('.multilineTextAlignment(.leading)', dashboard_source)

    # 验证 Server 五金来源选择突出显示当前选中的操作。
    # self：当前测试用例或测试替身实例。
    def test_server_hardware_choice_highlights_selected_action(self):
        source = (Path(__file__).resolve().parents[1] / "macos" / "OrderDashboardView.swift").read_text(encoding="utf-8")
        choice = source.split("private func cutToSizeHardwareChoice", 1)[1].split("    private func materialChangeRow", 1)[0]
        selected_branch = choice.split("if skipped {", 1)[1].split("} else {", 1)[0]
        unselected_branch = choice.split("} else {", 1)[1].split("            if skipped {", 1)[0]
        self.assertIn('Button("本次不写入五金")', selected_branch)
        self.assertIn(".buttonStyle(.glassProminent)", selected_branch)
        self.assertIn(".tint(AppPalette.warning)", selected_branch)
        self.assertIn('Button("本次写入五金")', unselected_branch)
        self.assertIn(".buttonStyle(.glassProminent)", unselected_branch)
        self.assertIn(".tint(AppPalette.success)", unselected_branch)

    # 验证库存浏览器成功结果不附带隐藏的登录提示。
    # self：当前测试用例或测试替身实例。
    def test_inventory_chrome_success_result_has_no_hidden_login_prompt(self):
        source = (Path(__file__).resolve().parents[1] / "traveler_assistant" / "inventory.py").read_text(encoding="utf-8")
        swift_source = (Path(__file__).resolve().parents[1] / "macos" / "TravelerAssistant.swift").read_text(encoding="utf-8")
        self.assertNotIn("库存专用 Chrome 已打开，请手工登录并完成验证码后再操作", source)
        open_chrome = swift_source.split('func openInventoryChrome() {', 1)[1].split('func updateInventoryCatalog()', 1)[0]
        self.assertNotIn('object["message"] as? String', open_chrome)

    # 验证订单材料预览可映射到库存检查项目。
    # self：当前测试用例或测试替身实例。
    def test_order_preview_materials_can_be_mapped_for_stock_check(self):
        with tempfile.TemporaryDirectory() as directory:
            state = Path(directory) / "state"
            catalog = state / "inventory" / "current-products.xlsx"
            mappings = state / "inventory" / "mappings.json"
            catalog.parent.mkdir(parents=True)
            make_catalog(catalog)
            mappings.write_text('{"manual": {}, "ignored": {}}', encoding="utf-8")
            preview = SimpleNamespace(
                order_id="PP0068",
                materials=[SimpleNamespace(kind="plywood", thickness=18.0, color="", quantity=2)],
                edge_banding={"Woodline 4": 12.5},
            )
            with patch("traveler_assistant.order_workflow.preview_order", return_value=preview):
                order_id, rows = order_stock_requirements(Config(state_dir=state), Path(directory) / "PP0068")
            self.assertEqual(order_id, "PP0068")
            self.assertEqual([row["productCode"] for row in rows], ["M0004", "M0020"])
            self.assertEqual([row["requiredQuantity"] for row in rows], [2, 13])
            self.assertEqual([row["unit"] for row in rows], ["张", "m"])

    # 验证金蝶运行环境支持路径覆盖，依赖缺失时明确拒绝运行。
    # self：当前测试用例或测试替身实例。
    def test_jdy_runtime_uses_portable_overrides_and_rejects_missing_dependencies(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            node = root / "runtime" / "node"
            node.parent.mkdir(parents=True)
            node.touch()
            node.chmod(0o755)
            modules = root / "runtime" / "node_modules"
            (modules / "playwright").mkdir(parents=True)
            (modules / "playwright-core").mkdir()
            with patch.dict(
                "os.environ",
                {"TRAVELER_NODE": str(node), "TRAVELER_NODE_MODULES": str(modules)},
            ):
                self.assertEqual(_resolve_jdy_runtime(root), (node, modules))

            with patch.dict(
                "os.environ",
                {
                    "TRAVELER_NODE": str(root / "missing-node"),
                    "TRAVELER_NODE_MODULES": str(modules),
                },
            ):
                with self.assertRaises(RuleError) as raised:
                    _resolve_jdy_runtime(root)
                self.assertEqual(raised.exception.code, "inventory_runtime")

    # 验证库存需求默认只含材料，并可显式包含五金。
    # self：当前测试用例或测试替身实例。
    def test_stock_requirements_default_to_materials_and_can_include_hardware(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_traveler(traveler, [("18mm--Plywood", 2), ("Hinge", 3)])
            make_catalog(catalog)
            mappings.write_text('{"manual": {"Hinge": "M1001"}, "ignored": {}}', encoding="utf-8")
            preview = build_preview(traveler, catalog, mappings)

            materials = stock_requirements(preview)
            all_items = stock_requirements(preview, include_hardware=True)

            self.assertEqual([item["productCode"] for item in materials], ["M0004"])
            self.assertEqual([item["productCode"] for item in all_items], ["M0004", "M1001"])

    # 验证库存检查比较所需数量与可用数量。
    # self：当前测试用例或测试替身实例。
    def test_stock_check_compares_required_and_available_quantities(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            state = root / "state"
            catalog = state / "inventory" / "current-products.xlsx"
            mappings = state / "inventory" / "mappings.json"
            catalog.parent.mkdir(parents=True)
            make_traveler(traveler, [("18mm--Plywood", 2), ("Hinge", 3)])
            make_catalog(catalog)
            mappings.write_text('{"manual": {"Hinge": "M1001"}, "ignored": {}}', encoding="utf-8")
            config = Config(state_dir=state)
            browser_result = {
                "ok": True,
                "results": [{
                    "productCode": "M0004",
                    "productName": "3/4 Finished UV2S",
                    "availableQuantity": 1,
                }],
            }
            with patch("traveler_assistant.inventory.run_jdy", return_value=browser_result) as run:
                result = check_stock(config, traveler)

            run.assert_called_once()
            self.assertTrue(result["hasShortage"])
            self.assertEqual(result["rows"][0]["shortageQuantity"], 1)
            self.assertFalse(result["rows"][0]["sufficient"])

    # 验证代切文件夹状态可与实际出库记录对账。
    # self：当前测试用例或测试替身实例。
    def test_cut_to_size_folder_status_can_be_reconciled(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folder = root / "orders" / "CS004"
            folder.mkdir(parents=True)
            (folder / "Work Order Traveler(CS004).xlsx").touch()
            config = Config(
                order_root=root / "orders",
                backup_root=root / "backups",
                state_dir=root / "state",
            )
            traveler = SimpleNamespace(order_name="CS004")
            with patch("traveler_assistant.inventory.parse_traveler", return_value=traveler), \
                 patch(
                     "traveler_assistant.inventory.run_jdy",
                     return_value={"matches": ["CS004 QTCK000123"]},
                 ):
                result = reconcile_folder_status(config, "cs004")
            self.assertEqual(result["found"][0]["document_number"], "QTCK000123")

    # 验证在线目录更新完成导出、校验及安装。
    # self：当前测试用例或测试替身实例。
    def test_online_catalog_update_exports_validates_and_installs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")

            # 核对商品导出动作，并在下载路径生成测试目录文件。
            # _config：为兼容真实调用签名保留的测试配置。
            # action：期望模拟的库存浏览器动作。
            # kwargs：转发给被模拟接口的关键字参数。
            def fake_export(_config, action, **kwargs):
                self.assertEqual(action, "exportProducts")
                make_catalog(kwargs["download_path"])
                return {"ok": True}

            with patch("traveler_assistant.inventory.run_jdy", side_effect=fake_export):
                result = update_catalog_online(config)

            installed = root / "state" / "inventory" / "current-products.xlsx"
            database = root / "state" / "workflow.sqlite3"
            self.assertTrue(result["ok"])
            self.assertEqual(result["count"], len(ProductCatalog(installed).products))
            loaded = ProductDatabase(database)
            self.assertEqual(result["count"], loaded.count())
            self.assertEqual(loaded.require_code("M0004").name, "3/4 Finished UV2S")
            loaded.close()
            self.assertFalse(any(installed.parent.glob(".products-download-*.xlsx")))

    # 验证目录导入保存成本价，缺失成本价仍为空值。
    # self：当前测试用例或测试替身实例。
    def test_catalog_import_persists_cost_price_and_keeps_missing_price_null(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            state = root / "state"
            source = root / "priced-products.xlsx"
            make_priced_catalog(source)

            parsed = {product.code: product for product in ProductCatalog(source).products}
            self.assertAlmostEqual(parsed["M0004"].cost_price, 31.76)
            self.assertIsNone(parsed["M0005"].cost_price)

            import_catalog(Config(state_dir=state), source)
            database = state / "workflow.sqlite3"
            connection = sqlite3.connect(database)
            columns = {row[1] for row in connection.execute("pragma table_info(products)")}
            connection.close()
            self.assertIn("cost_price", columns)
            self.assertFalse((state / "inventory" / "inventory.sqlite3").exists())
            loaded = ProductDatabase(database)
            self.assertAlmostEqual(loaded.require_code("M0004").cost_price, 31.76)
            self.assertIsNone(loaded.require_code("M0005").cost_price)
            loaded.close()

    # 验证目录变化摘要区分新增、更新及移除的商品。
    # self：当前测试用例或测试替身实例。
    def test_catalog_change_summary_reports_added_updated_and_removed_products(self):
        previous = [
            Product("A", "M001", "Old name", "", "启用"),
            Product("B", "M002", "Removed", "", "启用"),
        ]
        current = [
            Product("A", "M001", "New name", "", "启用"),
            Product("C", "M003", "Added", "", "启用"),
        ]

        self.assertEqual(
            _catalog_change_summary(previous, current),
            {"added_count": 1, "updated_count": 1, "removed_count": 1},
        )

    # 验证成本价变化会被目录差异摘要识别。
    # self：当前测试用例或测试替身实例。
    def test_catalog_change_summary_detects_cost_price_change(self):
        previous = [Product("Plywood", "M001", "Panel", "18mm", "启用", cost_price=10.0)]
        current = [Product("Plywood", "M001", "Panel", "18mm", "启用", cost_price=12.5)]

        self.assertEqual(
            _catalog_change_summary(previous, current),
            {"added_count": 0, "updated_count": 1, "removed_count": 0},
        )

    # 验证关闭库存浏览器采用专用调试操作。
    # self：当前测试用例或测试替身实例。
    def test_close_inventory_chrome_uses_dedicated_cdp_action(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            browser_result = SimpleNamespace(
                returncode=0,
                stdout='{"ok": true, "closed": true}',
                stderr="",
            )
            with patch("traveler_assistant.inventory._resolve_jdy_runtime", return_value=(Path("/node"), Path("/modules"))), \
                 patch("traveler_assistant.inventory.subprocess.run", return_value=browser_result) as run:
                result = close_inventory_chrome(config)

            request = json.loads(run.call_args.kwargs["input"])
            self.assertTrue(result["closed"])
            self.assertEqual(request["action"], "closeChrome")
            self.assertEqual(request["cdpEndpoint"], "http://127.0.0.1:9222")
            self.assertEqual(run.call_args.kwargs["timeout"], 5)

    # 验证延迟列出 Traveler 时仍报告已有商品目录状态。
    # self：当前测试用例或测试替身实例。
    def test_lazy_traveler_listing_reports_existing_catalog_status(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            state = root / "state"
            source = root / "products.xlsx"
            make_catalog(source)
            config = Config(state_dir=state, order_root=root / "orders")
            import_catalog(config, source)

            result = list_traveler_names(config)

            installed = state / "inventory" / "current-products.xlsx"
            self.assertEqual(result["catalog"]["count"], len(ProductCatalog(installed).products))
            self.assertFalse(result["catalog"]["stale"])

    # 验证目录刷新仅保留最新的 Excel 备份。
    # self：当前测试用例或测试替身实例。
    def test_catalog_refresh_keeps_only_latest_xlsx_backup(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            state = root / "state"
            inventory_dir = state / "inventory"
            inventory_dir.mkdir(parents=True)
            old_backup = inventory_dir / "current-products-2026-08-10.xlsx"
            old_backup.write_bytes(b"old")
            source = root / "export.xlsx"
            make_catalog(source)

            result = import_catalog(Config(state_dir=state), source)

            self.assertEqual(result["count"], 5)
            self.assertTrue((inventory_dir / "current-products.xlsx").is_file())
            self.assertFalse(old_backup.exists())
            self.assertTrue((state / "workflow.sqlite3").is_file())
            self.assertFalse((inventory_dir / "inventory.sqlite3").exists())

    # 验证删除目录 Excel 后仍可从数据库搜索商品。
    # self：当前测试用例或测试替身实例。
    def test_runtime_product_search_uses_database_after_xlsx_is_removed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            state = root / "state"
            catalog = state / "inventory" / "current-products.xlsx"
            catalog.parent.mkdir(parents=True)
            make_catalog(catalog)
            database = bootstrap_product_database(Config(state_dir=state))
            catalog.unlink()

            loaded = ProductDatabase(database)
            self.assertEqual([item.code for item in loaded.find(contains="Unihopper")], ["M1001"])
            loaded.close()

    # 验证浏览器错误保留具体失败原因。
    # self：当前测试用例或测试替身实例。
    def test_browser_error_preserves_specific_reason(self):
        stderr = (
            '{"event":"progress","message":"正在保存"}\n'
            "库存系统自动操作失败：保存需要人工确认：库存不足\n"
        )
        self.assertEqual(_jdy_error_detail(stderr), "保存需要人工确认：库存不足")

    # 验证多行浏览器错误保留首条具体原因。
    # self：当前测试用例或测试替身实例。
    def test_multiline_browser_error_keeps_first_specific_reason(self):
        stderr = (
            "库存系统自动操作失败：locator.hover: Timeout 30000ms exceeded.\n"
            "Call log:\n"
            "；当前页面：https://example.invalid/\n"
        )
        self.assertEqual(_jdy_error_detail(stderr), "locator.hover: Timeout 30000ms exceeded.")

    # 验证登录失败隐藏整页内容转储，并说明重试方式。
    # self：当前测试用例或测试替身实例。
    def test_login_failure_hides_full_page_dump_and_explains_retry(self):
        stderr = (
            "库存系统自动操作失败：库存系统登录未成功，请检查账号密码、验证码或登录限制后重试"
            "；当前页面：https://www.jdy.com/login/；页面提示：很长的整页内容\n"
        )
        detail = _jdy_error_detail(stderr)
        self.assertIn("库存系统登录未成功", detail)
        self.assertIn("再次查询", detail)
        self.assertNotIn("当前页面", detail)
        self.assertNotIn("整页内容", detail)

    # 验证安全验证失败时提示使用可见浏览器恢复登录。
    # self：当前测试用例或测试替身实例。
    def test_security_challenge_error_explains_visible_login_recovery(self):
        stderr = (
            "库存系统自动操作失败：库存系统要求完成验证码或安全验证；请先使用可见浏览器完成登录验证，再重试\n"
        )
        self.assertEqual(
            _jdy_error_detail(stderr),
            "库存系统要求完成验证码或安全验证；请先使用可见浏览器完成登录验证，再重试",
        )

    # 验证浏览器配置锁冲突被报告为可重试错误。
    # self：当前测试用例或测试替身实例。
    def test_profile_lock_failure_is_reported_as_retryable(self):
        stderr = (
            "库存系统自动操作失败：browserType.launchPersistentContext: "
            "Failed to create a ProcessSingleton for your profile directory.\n"
        )
        self.assertEqual(
            _jdy_error_detail(stderr),
            "上一次库存查询浏览器尚未完全退出，请稍候后再次点击查询",
        )

    # 验证运行环境噪声不会遮盖具体导出超时原因。
    # self：当前测试用例或测试替身实例。
    def test_browser_runtime_noise_does_not_hide_specific_export_timeout(self):
        stderr = (
            "库存系统自动操作失败：点击导出后 60 秒内没有检测到文件下载事件\n"
            "Node.js v24.19.0\n"
            "    at processTicksAndRejections (node:internal/process/task_queues:95:5)\n"
        )
        self.assertEqual(
            _jdy_error_detail(stderr),
            "点击导出后 60 秒内没有检测到文件下载事件",
        )

    # 验证金蝶调用传入可连接的浏览器调试端点。
    # self：当前测试用例或测试替身实例。
    def test_run_jdy_passes_attachable_chrome_endpoint(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            browser_result = SimpleNamespace(
                returncode=0,
                stdout='{"ok": true, "url": "https://service.jdy.com/workbench"}',
                stderr="",
            )
            with patch.dict(os.environ, {"TRAVELER_CHROME_CDP_ENDPOINT": "http://127.0.0.1:9333"}), \
                 patch("traveler_assistant.inventory._local_setting", return_value="18108100188"), \
                 patch("traveler_assistant.inventory._keychain_password", return_value="secret"), \
                 patch("traveler_assistant.inventory._resolve_jdy_runtime", return_value=(Path("/node"), Path("/modules"))), \
                 patch(
                     "traveler_assistant.inventory.subprocess.Popen",
                     return_value=_FakeNodeProcess(browser_result),
                 ) as popen:
                result = run_jdy(config, "preflight")

            process = popen.return_value
            request = json.loads(process.stdin.value)
            self.assertTrue(result["ok"])
            self.assertEqual(request["cdpEndpoint"], "http://127.0.0.1:9333")
            self.assertFalse(request["keepBrowserOpen"])
            self.assertEqual(process.wait_timeout, 90)

    # 验证复用已有库存页面时无需读取钥匙串。
    # self：当前测试用例或测试替身实例。
    def test_run_jdy_reuses_existing_inventory_page_without_reading_keychain(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            browser_result = SimpleNamespace(
                returncode=0,
                stdout='{"ok": true, "url": "https://www.jdy.com/workbench/web/index.html"}',
                stderr="",
            )
            with patch("traveler_assistant.inventory._local_setting", return_value="18108100188"), \
                 patch(
                     "traveler_assistant.inventory._find_existing_inventory_page",
                     return_value={"url": "https://www.jdy.com/workbench/web/index.html"},
                 ), \
                 patch(
                     "traveler_assistant.inventory._keychain_password",
                     side_effect=AssertionError("已登录页面不应读取钥匙串"),
                 ), \
                 patch("traveler_assistant.inventory._resolve_jdy_runtime", return_value=(Path("/node"), Path("/modules"))), \
                 patch(
                     "traveler_assistant.inventory.subprocess.Popen",
                     return_value=_FakeNodeProcess(browser_result),
                 ) as popen:
                result = run_jdy(config, "preflight")

            request = json.loads(popen.return_value.stdin.value)
            self.assertTrue(result["ok"])
            self.assertNotIn("password", request)
            self.assertEqual(request["cdpEndpoint"], "http://127.0.0.1:9222")

    # 验证更换重试批次号后仍复用已经确认的生产操作日志。
    # self：当前测试用例或测试替身实例。
    def test_inventory_operation_journal_reuses_confirmed_production_across_retry_batch_numbers(self):
        with tempfile.TemporaryDirectory() as directory:
            database = Path(directory) / "workflow.sqlite3"
            journal = InventoryOperationJournal(database)
            base_payload = {
                "documents": [{"remark": "CS004", "items": [{"productCode": "M0019", "quantity": 5}]}],
                "production_draft": {
                    "batch_number": "MP-FIRST",
                    "order_id": "CS004",
                    "selected_factory_orders": ["F-VANITY"],
                    "materials": [{"key": "panel", "quantity": 1}],
                },
            }
            first = journal.prepare("production", "CS004", ["F-VANITY"], base_payload)
            journal.update(
                first["operation_id"],
                "external_confirmed",
                results=[{"saved": True, "remark": "CS004", "documentNumber": "QTCK-001"}],
            )

            retry_payload = json.loads(json.dumps(base_payload))
            retry_payload["production_draft"]["batch_number"] = "MP-RETRY"
            retry = journal.prepare("production", "CS004", ["F-VANITY"], retry_payload)

            self.assertEqual(retry["operation_id"], first["operation_id"])
            self.assertEqual(retry["status"], "external_confirmed")
            self.assertEqual(
                journal.decoded_payload(retry)["production_draft"]["batch_number"],
                "MP-FIRST",
            )
            self.assertEqual(
                journal.decoded_results(retry)[0]["documentNumber"],
                "QTCK-001",
            )

    # 验证部分单据确认后重试从未完成步骤继续。
    # self：当前测试用例或测试替身实例。
    def test_inventory_operation_journal_resumes_after_partial_document_confirmation(self):
        with tempfile.TemporaryDirectory() as directory:
            database = Path(directory) / "workflow.sqlite3"
            journal = InventoryOperationJournal(database)
            payload = {
                "documents": [
                    {"remark": "PP0072-OFFICE", "items": [{"productCode": "M1001", "quantity": 24}]},
                    {"remark": "PP0072-HALLWAY", "items": [{"productCode": "M1001", "quantity": 16}]},
                ],
                "production_draft": None,
            }
            first = journal.prepare("shipment", "PP0072", ["F2608190229", "F2608190232"], payload)
            completed = {
                "saved": True,
                "remark": "PP0072-OFFICE",
                "documentNumber": "QTCK20260822002",
            }
            journal.update(
                first["operation_id"],
                "partial_external_confirmed",
                results=[completed],
            )

            retry = journal.prepare("shipment", "PP0072", ["F2608190229", "F2608190232"], payload)

            self.assertEqual(retry["status"], "partial_external_confirmed")
            self.assertEqual(journal.decoded_results(retry), [completed])
            self.assertEqual(journal.decoded_payload(retry)["documents"][1]["remark"], "PP0072-HALLWAY")

    # 验证打开库存浏览器采用独立配置目录及调试端口。
    # self：当前测试用例或测试替身实例。
    def test_open_inventory_chrome_launches_dedicated_profile_and_debug_port(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            executable = Path(directory) / "Google Chrome"
            executable.touch()
            with patch("traveler_assistant.inventory._inventory_cdp_pages", return_value=[]), \
                 patch("traveler_assistant.inventory._inventory_chrome_executable", return_value=executable), \
                 patch.dict(os.environ, {"TRAVELER_CHROME_CDP_ENDPOINT": "http://127.0.0.1:9333"}), \
                 patch("traveler_assistant.inventory.subprocess.Popen") as popen:
                result = open_inventory_chrome(config)

            self.assertTrue(result["ok"])
            self.assertTrue(result["launched"])
            arguments = popen.call_args.args[0]
            self.assertIn("--remote-debugging-port=9333", arguments)
            self.assertIn("--user-data-dir=" + result["profileDir"], arguments)
            self.assertEqual(arguments[-1], "https://www.jdy.com/login/")

    # 验证已有登录页时不会再启动第二个浏览器。
    # self：当前测试用例或测试替身实例。
    def test_open_inventory_chrome_does_not_launch_second_browser_on_login_page(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            with patch(
                "traveler_assistant.inventory._inventory_cdp_pages",
                return_value=[{"type": "page", "url": "https://www.jdy.com/login/"}],
            ), patch("traveler_assistant.inventory.subprocess.Popen") as popen:
                result = open_inventory_chrome(config)

            self.assertTrue(result["ok"])
            self.assertTrue(result["waitingForLogin"])
            popen.assert_not_called()

    # 验证材料映射缺失时在打开浏览器前阻止出库。
    # self：当前测试用例或测试替身实例。
    def test_outbound_stops_before_browser_when_material_mapping_is_missing(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state", order_root=root / "orders")
            traveler = root / "orders" / "Work Order Traveler(PP0099).xlsx"
            traveler.parent.mkdir(parents=True)
            make_traveler(traveler, [("Unmapped Hardware", 2)])
            with patch("traveler_assistant.inventory._local_setting", return_value="18108100188"), \
                 patch("traveler_assistant.inventory._keychain_password", return_value="secret"), \
                 patch("traveler_assistant.inventory.subprocess.run") as run:
                with self.assertRaisesRegex(RuleError, "未映射材料：Unmapped Hardware") as raised:
                    run_jdy(config, "outbound", traveler, confirm_save=True)

            self.assertEqual(raised.exception.code, "inventory_mapping_required")
            run.assert_not_called()

    # 验证数据库已出货工厂单禁止再次正常出库。
    # self：当前测试用例或测试替身实例。
    def test_database_shipped_factory_is_hard_blocked(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state", order_root=root / "orders")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("CS005", validation_status="正常")
            store.upsert_factory(
                "F2608120222",
                order_id="CS005",
                factory_name="CS005-KITCHEN",
                sales_order_name="CS005",
                name_source="AIMES",
                ownership_status="已确认",
                optimized=True,
                outbound_status="已出库",
                outbound_document="QTCK20260815001",
            )
            store.commit()
            store.close()

            with self.assertRaisesRegex(RuleError, "已出库") as raised:
                assert_factory_orders_outbound_allowed(config, "CS005", ["F2608120222"])
            self.assertEqual(raised.exception.code, "inventory_already_outbound")

    # 验证已出货工厂单的数据发生变化后可走更新流程。
    # self：当前测试用例或测试替身实例。
    def test_database_shipped_factory_with_changed_data_can_be_updated(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state", order_root=root / "orders")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("CS005", validation_status="正常")
            store.upsert_factory(
                "F2608120222",
                order_id="CS005",
                factory_name="CS005-KITCHEN",
                sales_order_name="CS005",
                name_source="AIMES",
                ownership_status="已确认",
                optimized=True,
                outbound_status="已出库",
                outbound_document="QTCK20260815001",
            )
            store.commit()
            store.close()

            with patch("traveler_assistant.order_index.reconcile_outbound_statuses", return_value=0):
                assert_factory_orders_outbound_allowed(
                    config,
                    "CS005",
                    ["F2608120222"],
                    changed_factory_orders=["F2608120222"],
                )

    # 验证金蝶浏览器超时转成可指导操作的业务错误。
    # self：当前测试用例或测试替身实例。
    def test_run_jdy_converts_browser_timeout_to_actionable_rule_error(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            timeout = subprocess.TimeoutExpired(["node", "jdy_inventory.mjs"], 90, stderr="页面未响应")
            with patch("traveler_assistant.inventory._local_setting", return_value="18108100188"), \
                 patch("traveler_assistant.inventory._keychain_password", return_value="secret"), \
                 patch("traveler_assistant.inventory._resolve_jdy_runtime", return_value=(Path("/node"), Path("/modules"))), \
                 patch(
                     "traveler_assistant.inventory.subprocess.Popen",
                     return_value=_FakeNodeProcess(timeout=timeout),
                 ):
                with self.assertRaises(RuleError) as raised:
                    run_jdy(config, "preflight")

            self.assertEqual(raised.exception.code, "jdy_timeout")
            self.assertIn("超过 90 秒", str(raised.exception))

    # 验证钥匙串读取超时归类为凭据错误。
    # self：当前测试用例或测试替身实例。
    def test_keychain_timeout_is_reported_as_credentials_error(self):
        with tempfile.TemporaryDirectory() as directory:
            helper = Path(directory) / "keychain-read"
            helper.write_text("placeholder")
            config = Config(state_dir=Path(directory) / "state")
            with patch("traveler_assistant.inventory.subprocess.run", side_effect=subprocess.TimeoutExpired([str(helper), "account"], 10)), \
                 patch("traveler_assistant.inventory.Path.is_file", return_value=True):
                with self.assertRaises(RuleError) as raised:
                    # 替换后的辅助脚本路径仍通过模块的
                    # 正常位置解析；本断言仅关注子进程调用。
                    _keychain_password("account")

            self.assertEqual(raised.exception.code, "jdy_credentials")
            self.assertIn("超过 10 秒", str(raised.exception))

    # 验证 Traveler 解析支持动态区域及零数量。
    # self：当前测试用例或测试替身实例。
    def test_parse_dynamic_regions_and_zero(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            make_traveler(path, [("18mm--Plywood", 2), ("Hinge", 0)])
            traveler = parse_traveler(path)
            self.assertEqual(traveler.order_name, "PP0099")
            self.assertEqual(traveler.items[0].quantity, 2)
            self.assertEqual(traveler.zero_items[0].quantity, 0)

    # 验证固定映射生效，并展开按压开启配件需求。
    # self：当前测试用例或测试替身实例。
    def test_fixed_mapping_and_push_open_expansion(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_traveler(traveler, [("18mm--Plywood", 2), ("Push Open", 3), ("Adjustable shelf holder", 4)])
            make_catalog(catalog)
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["商品类别", "*商品编号", "商品名称", "规格型号", "状态"])
            for product in ProductCatalog(catalog).products:
                sheet.append([product.category, product.code, product.name, product.spec, product.status])
            sheet.append(["Hardware", "M1013", "Adjustable Shelf Holder", "", "启用"])
            workbook.save(catalog)
            mappings.write_text('{"manual": {}, "ignored": {}}', encoding="utf-8")
            preview = build_preview(traveler, catalog, mappings)
            self.assertTrue(preview.ready)
            self.assertEqual([item.product_code for item in preview.outbound_items], ["M0004", "M1068", "M1069", "M1013"])
            self.assertEqual([item.quantity for item in preview.outbound_items], [2, 3, 3, 4])

    # 验证选择工厂单后保留订单材料，仅保留选中工厂单的五金。
    # self：当前测试用例或测试替身实例。
    def test_factory_selection_keeps_order_materials_and_selected_hardware_only(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_traveler(traveler, [("18mm--Plywood", 2), ("Hinge", 3)])
            make_catalog(catalog)
            mappings.write_text('{"manual": {"Hinge": "M1001"}, "ignored": {}}', encoding="utf-8")

            preview = build_preview(
                traveler,
                catalog,
                mappings,
                selected_document_remarks=["PP0099-KITCHEN"],
            )

            self.assertTrue(preview.ready)
            self.assertEqual(
                [(item.document_remark, item.product_code) for item in preview.outbound_items],
                [("PP0099", "M0004"), ("PP0099-KITCHEN", "M1001")],
            )
            self.assertEqual(
                set(document["remark"] for document in preview.document_payloads()),
                {"PP0099", "PP0099-KITCHEN"},
            )

    # 验证零数量项目不进入出库明细。
    # self：当前测试用例或测试替身实例。
    def test_zero_quantity_items_are_not_outbound_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_traveler(traveler, [("18mm--Plywood", 2), ("5.4mm--Plywood", 0)])
            make_catalog(catalog)
            mappings.write_text('{"manual": {}, "ignored": {}}', encoding="utf-8")

            preview = build_preview(traveler, catalog, mappings)

            self.assertEqual([item.traveler_name for item in preview.outbound_items], ["18mm--Plywood"])
            self.assertIn("5.4mm--Plywood", [item["name"] for item in preview.payload()["zero_items"]])

    # 验证唯一精确库存名称可映射到指定商品。
    # self：当前测试用例或测试替身实例。
    def test_unique_exact_inventory_name_maps_bls36(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-PANTRY).xlsx"
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_traveler(traveler, [("BLS36", 2)])
            make_catalog(catalog)
            workbook = load_workbook(catalog)
            workbook.active.append(["Hardware", "M1093", "BLS36", "TR-270B", "启用"])
            workbook.save(catalog)
            mappings.write_text('{"manual": {}, "ignored": {}}', encoding="utf-8")
            preview = build_preview(traveler, catalog, mappings)
            self.assertTrue(preview.ready)
            self.assertEqual(preview.outbound_items[0].product_code, "M1093")
            self.assertEqual(preview.outbound_items[0].match_source, "库存商品名称精确匹配")

    # 验证忽略材料必须填写原因，且在预览中可见。
    # self：当前测试用例或测试替身实例。
    def test_ignored_material_requires_reason_and_is_visible(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_traveler(traveler, [("Old Brand", 1)])
            make_catalog(catalog)
            mappings.write_text(json.dumps({"manual": {}, "ignored": {"Old Brand": "不再采购"}}), encoding="utf-8")
            preview = build_preview(traveler, catalog, mappings)
            self.assertTrue(preview.ready)
            self.assertEqual(preview.ignored_items[0]["reason"], "不再采购")

    # 验证封边出库数量按四舍五入规则取整。
    # self：当前测试用例或测试替身实例。
    def test_edge_quantity_is_rounded_half_up_for_inventory(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_traveler(traveler, [("Edge banding--Woodline 4", 37.75)])
            make_catalog(catalog)
            mappings.write_text('{"manual": {}, "ignored": {}}', encoding="utf-8")
            preview = build_preview(traveler, catalog, mappings)
            self.assertEqual(preview.outbound_items[0].quantity, 38)
            self.assertIn("37.75m→38m", preview.outbound_items[0].match_source)

    # 验证人工映射的封边同样按库存规则舍入数量。
    # self：当前测试用例或测试替身实例。
    def test_manual_edge_mapping_also_rounds_quantity_for_inventory(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_traveler(traveler, [("Edge banding--Woodline 4", 94.25)])
            make_catalog(catalog)
            mappings.write_text(
                json.dumps({"manual": {"Edge banding--Woodline 4": "M0020"}, "ignored": {}}),
                encoding="utf-8",
            )

            preview = build_preview(traveler, catalog, mappings)

            self.assertTrue(preview.ready)
            self.assertEqual(preview.outbound_items[0].product_code, "M0020")
            self.assertEqual(preview.outbound_items[0].quantity, 94)
            self.assertIn("人工指定", preview.outbound_items[0].match_source)
            self.assertIn("94.25m→94m", preview.outbound_items[0].match_source)

    # 验证封边优先匹配同色且带 24mm 后缀的 ABS 封边商品。
    # self：当前测试用例或测试替身实例。
    def test_edge_prefers_matching_color_abs_banding_with_24mm_suffix(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            traveler = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_traveler(traveler, [("Edge banding--Ivory Oak", 20.43)])
            make_catalog(catalog)
            workbook = load_workbook(catalog)
            sheet = workbook.active
            sheet.append(["Edge band", "M0145", "Ivory Oak ABS banding 24mm", "1mm*24mm*50M", "启用"])
            sheet.append(["Edge band", "M0146", "Ivory Oak ABS banding 48mm", "1mm*48mm*50M", "启用"])
            sheet.append(["Edge band", "M0147", "Ivory Oak Wood veneer banding 24mm", "0.6mm*24mm*50M", "启用"])
            workbook.save(catalog)
            mappings.write_text('{"manual": {}, "ignored": {}}', encoding="utf-8")
            preview = build_preview(traveler, catalog, mappings)
            self.assertTrue(preview.ready)
            self.assertEqual(preview.outbound_items[0].product_code, "M0145")
            self.assertEqual(preview.outbound_items[0].quantity, 20)
            self.assertIn("ABS banding + 24mm", preview.outbound_items[0].match_source)

    # 验证背板 8mm 与 9mm 规格可双向匹配别名。
    # self：当前测试用例或测试替身实例。
    def test_back_panel_8mm_and_9mm_are_bidirectional_aliases(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            catalog = root / "products.xlsx"
            mappings = root / "mappings.json"
            make_catalog(catalog)
            workbook = load_workbook(catalog)
            workbook.active.append(["Panel", "M1134", "Rosales 3", "8*1220*2745mm", "启用"])
            workbook.save(catalog)
            mappings.write_text('{"manual": {}, "ignored": {}}', encoding="utf-8")
            for thickness in ("8", "9"):
                traveler = root / f"Work Order Traveler(PP0099-{thickness}MM).xlsx"
                make_traveler(traveler, [(f"{thickness}mm--Rosales 3", 2)])
                preview = build_preview(traveler, catalog, mappings)
                self.assertTrue(preview.ready)
                self.assertEqual(preview.outbound_items[0].product_code, "M1134")

    # 验证忽略映射可保存并删除。
    # self：当前测试用例或测试替身实例。
    def test_ignored_mapping_can_be_saved_and_removed(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mappings.json"
            mappings = InventoryMappings(path)
            mappings.save_ignored("Special Material", "不再采购")
            self.assertEqual(
                InventoryMappings(path).ignored_reason("special material"),
                "不再采购",
            )
            InventoryMappings(path).remove_ignored("SPECIAL MATERIAL")
            self.assertIsNone(InventoryMappings(path).ignored_reason("Special Material"))

    # 验证来源编码不能被直接当作库存 SKU。
    # self：当前测试用例或测试替身实例。
    def test_source_codes_are_not_treated_as_inventory_skus(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            make_catalog(config.state_dir / "inventory" / "current-products.xlsx")
            set_ignored_mapping(config, "LED", True, "用户确认全局忽略")

            result = resolve_inventory_items(
                config,
                [
                    (TravelerItem(1, "五金", "LED", 2, "F0099"), "WJ-CBD"),
                    (TravelerItem(2, "五金", "未建档五金", 1, "F0099"), "WJ-UNKNOWN"),
                ],
            )

            self.assertEqual([item["name"] for item in result["ignored"]], ["LED"])
            self.assertEqual([item["name"] for item in result["missing"]], ["未建档五金"])
            self.assertEqual(result["missing"][0]["source_code"], "WJ-UNKNOWN")
            self.assertEqual(result["outbound"], [])

    # 验证低导轨名称可解析到对应低导轨 SKU。
    # self：当前测试用例或测试替身实例。
    def test_lower_rail_names_resolve_to_l_rail_sku(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            catalog_path = config.state_dir / "inventory" / "current-products.xlsx"
            make_catalog(catalog_path)
            workbook = load_workbook(catalog_path)
            workbook.active.append(["Hardware", "M1003", "L-Rail", "", "启用", "件"])
            workbook.save(catalog_path)

            result = resolve_inventory_items(
                config,
                [(TravelerItem(1, "五金", "Lower Left Rail", 1, "F0099"), "")],
            )

            self.assertEqual(result["missing"], [])
            self.assertEqual(result["outbound"][0].product_code, "M1003")

    # 验证忽略五金映射不会删除已确认的 SKU 事实。
    # self：当前测试用例或测试替身实例。
    def test_ignoring_hardware_preserves_confirmed_sku_facts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            connection = connect_database(config.workflow_database)
            seed_sku_products(connection)
            connection.executemany(
                'insert into hardware_items(order_id,factory_order,product_code,quantity,source_type,updated_at) values(?,?,?,?,?,?)',
                [('PP0099', 'F0099', 'M1013', 2, 'aicnc', 'now'), ('PP0099', 'F0099', 'M1001', 4, 'aicnc', 'now'), ('PP0099', 'F0099', 'M-LED', 5, 'aicnc', 'now')],
            )
            connection.commit()
            connection.close()

            result = set_ignored_mapping(config, "Shelf Holder", True, "不再采购")
            self.assertEqual(result["removed_database_rows"], 0)
            connection = sqlite3.connect(config.workflow_database)
            try:
                self.assertEqual(connection.execute("select count(*) from hardware_items").fetchone()[0], 3)
            finally:
                connection.close()

            result = set_ignored_mapping(config, "Hinge", True, "不再采购")
            self.assertEqual(result["removed_database_rows"], 0)
            connection = sqlite3.connect(config.workflow_database)
            try:
                self.assertEqual(connection.execute("select count(*) from hardware_items").fetchone()[0], 3)
            finally:
                connection.close()

            result = set_ignored_mapping(config, "LED", True, "不再采购")
            self.assertEqual(result["removed_database_rows"], 0)
            connection = sqlite3.connect(config.workflow_database)
            try:
                self.assertEqual(connection.execute("select count(*) from hardware_items").fetchone()[0], 3)
            finally:
                connection.close()

    # 验证人工映射保存后替换旧忽略决定。
    # self：当前测试用例或测试替身实例。
    def test_manual_mapping_can_be_saved_and_replaces_ignore(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mappings.json"
            mappings = InventoryMappings(path)
            mappings.save_ignored("Special Material", "旧品牌")
            InventoryMappings(path).save_manual("Special Material", "m1093")
            reloaded = InventoryMappings(path)
            self.assertEqual(reloaded.manual_code("special material"), "M1093")
            self.assertIsNone(reloaded.ignored_reason("Special Material"))

    # 验证人工映射支持查看、更新及删除。
    # self：当前测试用例或测试替身实例。
    def test_manual_mapping_can_be_viewed_updated_and_removed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            make_catalog(config.state_dir / "inventory" / "current-products.xlsx")

            save_manual_mapping(config, "Hinge", "M1001")
            self.assertEqual(list_inventory_mappings(config)["manual"]["Hinge"], "M1001")
            update_manual_mapping(config, "Hinge", "Cabinet Hinge", "M1001")
            self.assertEqual(list_inventory_mappings(config)["manual"], {"Cabinet Hinge": "M1001"})
            remove_manual_mapping(config, "Cabinet Hinge")
            self.assertEqual(list_inventory_mappings(config)["manual"], {})

    # 验证同一 SKU 的不同别名共享五金显示名称。
    # self：当前测试用例或测试替身实例。
    def test_hardware_display_name_is_shared_by_sku_aliases(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            make_catalog(config.state_dir / "inventory" / "current-products.xlsx")

            save_manual_mapping(config, "Hinge", "M1001", "柜门铰链")
            save_manual_mapping(config, "TestFullHinge", "M1001")

            mappings = InventoryMappings(config.workflow_database)
            manual = list_inventory_mappings(config)
            self.assertEqual(manual["manual"]["Hinge"], "M1001")
            self.assertEqual(manual["manual"]["TestFullHinge"], "M1001")
            self.assertEqual(manual["manual_display_names"]["Hinge"], "柜门铰链")
            self.assertEqual(manual["manual_display_names"]["TestFullHinge"], "柜门铰链")
            self.assertEqual(
                mappings.display_name_for_hardware("M1001", "TestFullHinge", "71T950A"),
                "柜门铰链",
            )

    # 验证 Traveler 内容指纹变化会改变同步状态。
    # self：当前测试用例或测试替身实例。
    def test_sync_status_changes_with_traveler_fingerprint(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "Work Order Traveler(PP0099-KITCHEN).xlsx"
            catalog_path = root / "products.xlsx"
            mapping_path = root / "mappings.json"
            make_traveler(path, [("18mm--Plywood", 2)])
            make_catalog(catalog_path)
            mapping_path.write_text('{"manual": {}, "ignored": {}}', encoding="utf-8")
            preview = build_preview(path, catalog_path, mapping_path)
            store = InventorySyncStore(root / "workflow.sqlite3", root / "backups")
            store.save_success(preview, [{
                "remark": "PP0099",
                "saved": True,
                "documentNumber": "CK-001",
            }])
            self.assertEqual(store.status_for(parse_traveler(path))[0], "已出库")
            make_traveler(path, [("18mm--Plywood", 3)])
            reloaded = InventorySyncStore(root / "workflow.sqlite3", root / "backups")
            self.assertEqual(reloaded.status_for(parse_traveler(path))[0], "需要更新")

    # 验证订单材料出库仅关联被选中的拆分工厂单。
    # self：当前测试用例或测试替身实例。
    def test_order_material_outbound_links_only_selected_split_factories(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.upsert_order("CS004", validation_status="正常")
            for factory_order, factory_name in (
                ("F-KITCHEN", "CS004-KITCHEN"),
                ("F-VANITY", "CS004-vanity"),
            ):
                store.upsert_factory(
                    factory_order,
                    order_id="CS004",
                    factory_name=factory_name,
                    sales_order_name="CS004",
                    name_source="AIMES",
                    ownership_status="已确认",
                    optimized=True,
                    outbound_status="未出库",
                )
            store.connection.execute(
                "insert into material_items(order_id, product_code, quantity, source_type, updated_at) values(?,?,?,?,?)",
                ("CS004", "M0004", 2, "database", "now"),
            )
            store.commit()
            store.close()

            source_item = TravelerItem(1, "板材与封边", "18mm--Plywood", 2, "CS004")
            outbound_item = OutboundItem(
                traveler_name="18mm--Plywood",
                product_code="M0004",
                product_name="18mm Plywood",
                quantity=2,
                section="板材与封边",
                match_source="test",
                document_remark="CS004",
                unit="张",
            )

            # 构造仅选择指定工厂单的部分范围出库预览。
            # factory_order：本次预览选中的工厂单编号。
            def preview(factory_order):
                traveler = TravelerData(
                    path=config.workflow_database.resolve(),
                    pp_folder="CS004",
                    order_id="CS004",
                    order_name="CS004",
                    items=[source_item],
                    zero_items=[],
                    documents={"CS004": [source_item]},
                    modified_at="2026-08-20T00:00:00",
                    fingerprint="test-fingerprint",
                )
                return InventoryPreview(
                    traveler=traveler,
                    outbound_items=[outbound_item],
                    selected_factory_orders=(factory_order,),
                    partial_scope=True,
                    source_type="database",
                )

            sync = InventorySyncStore(
                config.workflow_database,
                config.backup_root,
            )
            sync.save_success(preview("F-KITCHEN"), [{
                "remark": "CS004",
                "saved": True,
                "documentNumber": "QTCK-001",
            }])
            connection = sqlite3.connect(config.workflow_database)
            links = connection.execute(
                "select factory_order from outbound_document_factories "
                "where document_number='QTCK-001' order by factory_order"
            ).fetchall()
            connection.close()
            self.assertEqual(links, [("F-KITCHEN",)])

            # 第二个工厂单复用未变化的订单材料
            # 单据；应追加其明确身份关联，
            # 而不另建库存单据。
            sync = InventorySyncStore(
                config.workflow_database,
                config.backup_root,
            )
            sync.save_success(preview("F-VANITY"), [{
                "remark": "CS004",
                "unchanged": True,
                "saved": True,
                "documentNumber": "QTCK-001",
            }])
            connection = sqlite3.connect(config.workflow_database)
            links = connection.execute(
                "select factory_order from outbound_document_factories "
                "where document_number='QTCK-001' order by factory_order"
            ).fetchall()
            self.assertEqual(links, [("F-KITCHEN",), ("F-VANITY",)])
            connection.close()
            self.assertEqual(reconcile_outbound_statuses(config), 2)
            connection = sqlite3.connect(config.workflow_database)
            statuses = connection.execute(
                "select factory_order, (case when stage='已出货' then '已出库' else '未出库' end) as outbound_status, outbound_document from factory_orders where order_id='CS004' order by factory_order"
            ).fetchall()
            connection.close()
            self.assertEqual(
                statuses,
                [
                    ("F-KITCHEN", "已出库", "QTCK-001"),
                    ("F-VANITY", "已出库", "QTCK-001"),
                ],
            )

            connection = sqlite3.connect(config.workflow_database)
            connection.execute("update material_items set quantity=3 where order_id='CS004'")
            connection.commit()
            connection.close()
            self.assertEqual(reconcile_outbound_statuses(config), 2)
            connection = sqlite3.connect(config.workflow_database)
            changed_statuses = connection.execute(
                "select factory_order, (case when stage='已出货' then '已出库' else '未出库' end) as outbound_status, outbound_document from factory_orders where order_id='CS004' order by factory_order"
            ).fetchall()
            connection.close()
            self.assertEqual(
                changed_statuses,
                [
                    ("F-KITCHEN", "未出库", "QTCK-001"),
                    ("F-VANITY", "未出库", "QTCK-001"),
                ],
            )

    # 验证生产材料出库不会把工厂单标记为已出货。
    # self：当前测试用例或测试替身实例。
    def test_production_material_outbound_does_not_mark_factory_orders_shipped(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state", backup_root=root / "backups")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.upsert_order("PP0063-2", validation_status="正常")
            store.upsert_factory(
                "F-PRODUCTION",
                order_id="PP0063-2",
                factory_name="PP0063-2-KITCHEN",
                sales_order_name="PP0063-2",
                name_source="AIMES",
                ownership_status="已确认",
                optimized=True,
                outbound_status="未出库",
            )
            store.commit()
            store.close()

            source_item = TravelerItem(1, "板材与封边", "18mm--Plywood", 2, "PP0063-2")
            outbound_item = OutboundItem(
                traveler_name="18mm--Plywood",
                product_code="M0004",
                product_name="18mm Plywood",
                quantity=2,
                section="板材与封边",
                match_source="test",
                document_remark="PP0063-2",
                unit="张",
            )
            preview = InventoryPreview(
                traveler=TravelerData(
                    path=config.workflow_database.resolve(),
                    pp_folder="PP0063-2",
                    order_id="PP0063-2",
                    order_name="PP0063-2",
                    items=[source_item],
                    zero_items=[],
                    documents={"PP0063-2": [source_item]},
                    modified_at="2026-08-21T00:00:00",
                    fingerprint="test-fingerprint",
                ),
                outbound_items=[outbound_item],
                selected_factory_orders=("F-PRODUCTION",),
                source_type="database",
            )
            production_draft = {
                "batch_number": "MP-PRODUCTION-001",
                "order_id": "PP0063-2",
                "selected_factory_orders": ["F-PRODUCTION"],
                "materials": [{"product_code": "M0004", "material_type": "plywood", "color": "", "thickness": "18", "edge": "", "unit": "张", "quantity": 2}],
            }
            sync = InventorySyncStore(
                config.workflow_database,
                config.backup_root,
            )
            result = {"remark": "PP0063-2", "saved": True, "documentNumber": "QTCK-PRODUCTION"}
            _persist_single_outbound_result(config, preview, result, production_draft=production_draft)
            # 新的存储对象在最终提交前不得重建出货关联。
            sync = InventorySyncStore(config.workflow_database, config.backup_root)
            with sqlite3.connect(config.workflow_database) as connection:
                self.assertEqual(connection.execute("select count(*) from production_records").fetchone()[0], 0)
                self.assertEqual(connection.execute("select count(*) from outbound_document_factories").fetchone()[0], 0)
                self.assertEqual(connection.execute("select document_type from outbound_documents").fetchone()[0], "production_materials")
            sync.save_success(
                preview,
                [{"remark": "PP0063-2", "saved": True, "documentNumber": "QTCK-PRODUCTION"}],
                production_draft=production_draft,
            )

            connection = sqlite3.connect(config.workflow_database)
            links = connection.execute(
                "select factory_order from outbound_document_factories "
                "where document_number='QTCK-PRODUCTION'"
            ).fetchall()
            status = connection.execute(
                "select (case when stage='已出货' then '已出库' else '未出库' end) as outbound_status, outbound_document from factory_orders where factory_order='F-PRODUCTION'"
            ).fetchone()
            batch = connection.execute(
                "select status from production_records"
            ).fetchone()
            connection.close()
            self.assertEqual(links, [])
            self.assertEqual(status, ("未出库", ""))
            self.assertEqual(batch, ("completed",))

    # 验证浏览器生产出库成功后提交材料消耗和生产批次。
    # self：当前测试用例或测试替身实例。
    def test_production_browser_success_commits_materials_and_batch(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state", backup_root=root / "backups")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.upsert_order("PP0063-2", validation_status="正常")
            store.upsert_factory(
                "F-PRODUCTION",
                order_id="PP0063-2",
                factory_name="PP0063-2-KITCHEN",
                sales_order_name="PP0063-2",
                name_source="AIMES",
                ownership_status="已确认",
                optimized=True,
                outbound_status="未出库",
            )
            store.commit()
            store.close()

            source_item = TravelerItem(1, "板材与封边", "18mm--Plywood", 2, "PP0063-2")
            outbound_item = OutboundItem(
                traveler_name="18mm--Plywood",
                product_code="M0004",
                product_name="18mm Plywood",
                quantity=2,
                section="板材与封边",
                match_source="test",
                document_remark="PP0063-2",
                unit="张",
            )
            preview = InventoryPreview(
                traveler=TravelerData(
                    path=config.workflow_database.resolve(),
                    pp_folder="PP0063-2",
                    order_id="PP0063-2",
                    order_name="PP0063-2",
                    items=[source_item],
                    zero_items=[],
                    documents={"PP0063-2": [source_item]},
                    modified_at="2026-08-21T00:00:00",
                    fingerprint="test-fingerprint",
                ),
                outbound_items=[outbound_item],
                selected_factory_orders=("F-PRODUCTION",),
                source_type="database",
            )
            production_draft = {
                "batch_number": "MP-PRODUCTION-001",
                "order_id": "PP0063-2",
                "selected_factory_orders": ["F-PRODUCTION"],
                "materials": [{"key": "M0004", "material_type": "plywood", "color": "", "thickness": "18", "edge": "", "unit": "张", "quantity": 2}],
            }
            browser_result = SimpleNamespace(returncode=0, stdout=json.dumps({
                "remark": "PP0063-2", "saved": True, "documentNumber": "QTCK-PRODUCTION"
            }), stderr="")
            with patch("traveler_assistant.inventory._local_setting", return_value=""), \
                 patch("traveler_assistant.inventory._find_existing_inventory_page", return_value={"url": "https://vip2-hz.jdy.com/default-new.jsp"}), \
                 patch("traveler_assistant.production.cumulative_production_materials", return_value=production_draft["materials"]), \
                 patch("traveler_assistant.inventory.build_database_preview", return_value=preview), \
                 patch("traveler_assistant.inventory._resolve_jdy_runtime", return_value=(Path("/node"), Path("/modules"))), \
                 patch("traveler_assistant.inventory.subprocess.Popen", return_value=_FakeNodeProcess(browser_result)):
                result = run_jdy(config, "outbound", confirm_save=True, order_id="PP0063-2",
                                 selected_factory_orders=["F-PRODUCTION"],
                                 production_request_id=production_draft["batch_number"],
                                 production_materials=production_draft["materials"])
            self.assertTrue(result["productionCompleted"])
            with sqlite3.connect(config.workflow_database) as connection:
                self.assertEqual(connection.execute("select status from inventory_operations").fetchone()[0], "local_committed")

            connection = sqlite3.connect(config.workflow_database)
            links = connection.execute(
                "select factory_order from outbound_document_factories "
                "where document_number='QTCK-PRODUCTION'"
            ).fetchall()
            status = connection.execute(
                "select (case when stage='已出货' then '已出库' else '未出库' end) as outbound_status, outbound_document from factory_orders where factory_order='F-PRODUCTION'"
            ).fetchone()
            batch = connection.execute(
                "select status from production_records"
            ).fetchone()
            self.assertEqual(connection.execute("select product_code, quantity from production_materials").fetchall(), [("M0004", 2)])
            connection.close()
            self.assertEqual(links, [])
            self.assertEqual(status, ("未出库", ""))
            self.assertEqual(batch, ("completed",))

    # 验证拆分生产材料单据更新时清理过期工厂单关联。
    # self：当前测试用例或测试替身实例。
    def test_split_production_material_document_cleans_stale_factory_links(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            seed_sku_products(store.connection)
            store.upsert_order("PP0057", validation_status="正常")
            for factory_order, factory_name in (
                ("F-PRODUCTION-KITCHEN", "PP0057-KITCHEN"),
                ("F-PRODUCTION-LAUNDRY", "PP0057-LAUNDRY"),
            ):
                store.upsert_factory(
                    factory_order,
                    order_id="PP0057",
                    factory_name=factory_name,
                    sales_order_name="PP0057",
                    name_source="AIMES",
                    ownership_status="已确认",
                    optimized=True,
                    outbound_status="未出库",
                )
            store.commit()
            store.close()

            source_item = TravelerItem(1, "板材与封边", "18mm--Plywood", 2, "PP0057")
            outbound_item = OutboundItem(
                traveler_name="18mm--Plywood",
                product_code="M0004",
                product_name="18mm Plywood",
                quantity=2,
                section="板材与封边",
                match_source="test",
                document_remark="PP0057",
                unit="张",
            )
            preview = InventoryPreview(
                traveler=TravelerData(
                    path=config.workflow_database.resolve(),
                    pp_folder="PP0057",
                    order_id="PP0057",
                    order_name="PP0057",
                    items=[source_item],
                    zero_items=[],
                    documents={"PP0057": [source_item]},
                    modified_at="2026-08-25T00:00:00",
                    fingerprint="test-fingerprint",
                ),
                outbound_items=[outbound_item],
                selected_factory_orders=("F-PRODUCTION-KITCHEN", "F-PRODUCTION-LAUNDRY"),
                source_type="database",
            )
            production_draft = {
                "batch_number": "MP-PRODUCTION-SPLIT-001",
                "order_id": "PP0057",
                "selected_factory_orders": ["F-PRODUCTION-KITCHEN", "F-PRODUCTION-LAUNDRY"],
                "materials": [{"product_code": "M0004", "material_type": "plywood", "color": "", "thickness": "18", "edge": "", "unit": "张", "quantity": 2}],
            }
            sync = InventorySyncStore(
                config.workflow_database,
                config.backup_root,
            )
            sync.save_success(
                preview,
                [{"remark": "PP0057", "saved": True, "documentNumber": "QTCK-PRODUCTION-SPLIT"}],
                production_draft=production_draft,
            )

            connection = sqlite3.connect(config.workflow_database)
            connection.execute(
                "insert into outbound_document_factories(document_number, order_id, factory_order, created_at, updated_at) values(?,?,?,?,?)",
                ("QTCK-PRODUCTION-SPLIT", "PP0057", "F-PRODUCTION-KITCHEN", "now", "now"),
            )
            connection.execute(
                "update factory_orders set stage='已出货', outbound_document='QTCK-PRODUCTION-SPLIT', outbound_mode='inventory' where order_id='PP0057'"
            )
            connection.commit()
            connection.close()

            self.assertEqual(reconcile_outbound_statuses(config), 2)
            connection = sqlite3.connect(config.workflow_database)
            links = connection.execute(
                "select factory_order from outbound_document_factories where document_number='QTCK-PRODUCTION-SPLIT'"
            ).fetchall()
            statuses = connection.execute(
                "select factory_order, (case when stage='已出货' then '已出库' else '未出库' end) as outbound_status, outbound_document from factory_orders where order_id='PP0057' order by factory_order"
            ).fetchall()
            kind = connection.execute(
                "select document_type from outbound_documents where document_number='QTCK-PRODUCTION-SPLIT'"
            ).fetchone()
            connection.close()
            self.assertEqual(links, [])
            self.assertEqual(statuses, [
                ("F-PRODUCTION-KITCHEN", "未出库", ""),
                ("F-PRODUCTION-LAUNDRY", "未出库", ""),
            ])
            self.assertEqual(kind, ("production_materials",))

    # 验证之前存在的五金出库块变空时要求人工作废旧单。
    # self：当前测试用例或测试替身实例。
    def test_previous_hardware_block_becoming_empty_requires_manual_void(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "Work Order Traveler(PP0099).xlsx"
            catalog_path = root / "products.xlsx"
            mapping_path = root / "mappings.json"
            make_traveler(path, [("Hinge", 2)])
            make_catalog(catalog_path)
            mapping_path.write_text('{"manual": {"Hinge": "M1001"}, "ignored": {}}', encoding="utf-8")
            preview = build_preview(path, catalog_path, mapping_path)
            store = InventorySyncStore(root / "workflow.sqlite3", root / "backups")
            plans = store.prepare_documents(preview)
            results = [
                {
                    "remark": plan["remark"],
                    "saved": True,
                    "documentNumber": f"CK-{index:03d}",
                }
                for index, plan in enumerate(plans, 1)
            ]
            store.save_success(preview, results)
            make_traveler(path, [])
            empty_preview = build_preview(path, catalog_path, mapping_path)
            with self.assertRaisesRegex(RuleError, "人工删除或作废"):
                InventorySyncStore(root / "workflow.sqlite3", root / "backups").prepare_documents(empty_preview)

    # 验证五金出货判断不采用之前的订单材料出库单。
    # self：当前测试用例或测试替身实例。
    def test_hardware_shipment_ignores_previous_order_material_document(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            sync_path = root / "workflow.sqlite3"
            store = InventorySyncStore(sync_path, root / "backups")
            connection = sqlite3.connect(store.database)
            connection.execute(
                """insert into outbound_documents(
                       document_number, document_type, order_id, factory_order,
                       status, source, raw_fingerprint, mapped_fingerprint, updated_at
                   ) values(?,?,?,?,?,?,?,?,?)""",
                (
                    "QTCK20260822001", "materials", "PP0072", "PP0072",
                    "已出库", "金蝶", "materials-source", "materials-mapped", "now",
                ),
            )
            connection.commit()
            connection.close()

            hardware_item = TravelerItem(1, "五金", "Hinge", 2, "PP0072-OFFICE")
            traveler = TravelerData(
                path=root / "workflow.sqlite3",
                pp_folder="PP0072",
                order_id="PP0072",
                order_name="PP0072",
                items=[hardware_item],
                zero_items=[],
                documents={"PP0072": [], "PP0072-OFFICE": [hardware_item]},
                modified_at="2026-08-21T21:00:00",
                fingerprint="shipment-preview",
            )
            preview = InventoryPreview(
                traveler=traveler,
                outbound_items=[OutboundItem(
                    traveler_name="Hinge",
                    product_code="M1001",
                    product_name="Hinge",
                    quantity=2,
                    section="五金",
                    match_source="test",
                    document_remark="PP0072-OFFICE",
                )],
                selected_document_remarks=("PP0072-OFFICE",),
                source_type="database",
                document_kinds=("hardware",),
            )

            plans = InventorySyncStore(sync_path, root / "backups").prepare_documents(preview)
            self.assertEqual([plan["remark"] for plan in plans], ["PP0072-OFFICE"])

    # 验证出库状态优先采用工厂单五金记录，而非订单材料记录。
    # self：当前测试用例或测试替身实例。
    def test_outbound_status_prefers_factory_hardware_over_order_material_record(self):
        factory = {
            "order_id": "PP0072",
            "factory_order": "F2608190229",
            "factory_name": "PP0072-OFFICE",
            "sales_order_name": "PP0072",
            "outbound_mode": "inventory",
            "outbound_fingerprint": "",
        }
        records = [
            {
                "order_id": "PP0072",
                "remark": "F2608190229",
                "document_remark": "PP0072",
                "kind": "",
                "status": "已出库",
                "document_number": "QTCK20260822001",
            },
            {
                "order_id": "PP0072",
                "remark": "F2608190229",
                "document_remark": "PP0072-OFFICE",
                "kind": "hardware",
                "status": "已出库",
                "document_number": "QTCK20260822002",
            },
        ]

        self.assertEqual(
            _refresh_outbound_status(Config(), factory, records, factory_group=[factory]),
            ("已出库", "QTCK20260822002"),
        )
        self.assertTrue(_has_factory_hardware_outbound_record(factory, records))

    # 验证 SQLite 出库记录依据同步类型执行五金对账。
    # self：当前测试用例或测试替身实例。
    def test_sqlite_outbound_record_uses_sync_kind_for_hardware_reconciliation(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Config(state_dir=Path(directory) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP0072", validation_status="正常")
            store.upsert_factory(
                "F2608190232",
                order_id="PP0072",
                factory_name="PP0072-HALLWAY",
                sales_order_name="PP0072",
                name_source="AIMES",
                ownership_status="已确认",
                optimized=True,
                outbound_status="需要更新",
                outbound_document="QTCK-MATERIAL",
            )
            store.commit()
            store.close()

            connection = sqlite3.connect(config.workflow_database)
            connection.execute(
                """
                insert into outbound_documents(
                    document_number, document_type, order_id, factory_order,
                    status, source, issued_at, raw_fingerprint, mapped_fingerprint, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    "QTCK-HARDWARE",
                    "hardware",
                    "PP0072",
                    "F2608190232",
                    "已出库",
                    "金蝶",
                    "2026-08-21T22:09:24",
                    "hardware-source",
                    "hardware-mapped",
                    "2026-08-21T22:09:24",
                ),
            )
            connection.execute(
                "insert into outbound_document_factories(document_number, order_id, factory_order, created_at, updated_at) values(?,?,?,?,?)",
                (
                    "QTCK-HARDWARE",
                    "PP0072",
                    "F2608190232",
                    "2026-08-21T22:09:24",
                    "2026-08-21T22:09:24",
                ),
            )
            connection.commit()
            connection.close()
            records = _load_outbound_records(config)
            self.assertEqual(records[0]["kind"], "hardware")
            self.assertEqual(reconcile_outbound_statuses(config), 1)
            reopened = sqlite3.connect(config.workflow_database)
            row = reopened.execute(
                "select (case when stage='已出货' then '已出库' else '未出库' end) as outbound_status, outbound_document from factory_orders where factory_order='F2608190232'"
            ).fetchone()
            reopened.close()
            self.assertEqual(row, ("已出库", "QTCK-HARDWARE"))

    # 验证同一工厂单内同名五金会合并数量。
    # self：当前测试用例或测试替身实例。
    def test_same_hardware_name_is_aggregated_within_factory(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Work Order Traveler(PP0099).xlsx"
            make_traveler(path, [("Hinge", 2), ("Hinge", 3)])
            traveler = parse_traveler(path)
            hardware = traveler.documents["PP0099-KITCHEN"]
            self.assertEqual(len(hardware), 1)
            self.assertEqual(hardware[0].quantity, 5)


if __name__ == "__main__":
    unittest.main()

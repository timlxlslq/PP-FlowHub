"""Server piece-to-pair conversion through real reports, SKU mapping and writes."""
import copy
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from traveler_assistant.core import Config, RuleError
from traveler_assistant.hardware_facts import server_hardware_quantity
from traveler_assistant.inventory import InventoryMappings, import_catalog
from traveler_assistant.order_index import (
    OrderIndexStore, preview_server_changes, confirm_server_material_preview_memory,
    sync_order_index,
)
from tests.test_order_workflow import make_materials, make_board, make_fittings


class ServerRailUnitsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.config = Config(state_dir=root / 'state', source_root=root / 'server', order_root=root / 'orders')
        self.config.prepare_storage()
        self.folder = self.config.source_root / 'PP9999'
        report = self.folder / 'Report'
        report.mkdir(parents=True)
        make_materials(self.folder / 'PP9999 materials.xlsx')
        make_board(report / '板材清单.xlsx', 'F100', 'PP9999-KITCHEN')
        self.report = report / 'Fittingslist.xlsx'
        catalog = root / 'products.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.append(['商品类别', '*商品编号', '商品名称', '规格型号', '状态', '计量单位'])
        for code in ['M1094', 'M1095', 'M1096', 'M1097', 'M1142']:
            ws.append(['Hardware/Rail', code, 'Product ' + code, '', '启用', 'Sets' if code != 'M1142' else 'PCS'])
        ws.append(['Hardware/Rail', 'M1002', 'H-Rail', '', '启用', '套'])
        ws.append(['Hardware/Rail', 'M1003', 'L-Rail', '', '启用', '套'])
        wb.save(catalog)
        import_catalog(self.config, catalog)
        mappings = InventoryMappings(self.config.workflow_database)
        for code in ['M1094', 'M1095', 'M1096', 'M1097', 'M1142']:
            mappings.save_manual('Source ' + code, code)
        for name in ['18mm--Plywood', '14.5mm--Plywood', '5.4mm--Plywood',
                     '19.1mm--Basalto SM', '8mm--Basalto SM', 'Edge banding--Basalto SM']:
            mappings.save_ignored(name, 'Unrelated materials in rail test')
        store = OrderIndexStore(self.config.workflow_database)
        store.upsert_order('PP9999', validation_status='正常', source_folder=str(self.folder))
        store.upsert_factory('F100', order_id='PP9999', factory_name='PP9999-KITCHEN',
                             sales_order_name='PP9999', name_source='AIMES', ownership_status='已确认')
        store.commit()
        store.close()

    def write_report(self, quantities):
        make_fittings(self.report, [('F100', 2)])
        wb = load_workbook(self.report)
        ws = wb.active
        for row, (code, quantity) in enumerate(quantities.items(), start=7):
            ws.cell(row, 3, 'Source ' + code)
            ws.cell(row, 5, 'DTC-' + code)
            ws.cell(row, 9, 'Piece')
            ws.cell(row, 11, quantity)
        wb.save(self.report)

    def rows(self):
        with sqlite3.connect(self.config.workflow_database) as connection:
            return connection.execute('select h.product_code,h.quantity,p.unit from hardware_items h join products p on p.code=h.product_code order by h.product_code').fetchall()

    def test_preview_and_repeated_confirmation_convert_raw_counts_only_once(self):
        self.write_report({'M1094': 12, 'M1095': 6, 'M1096': 4, 'M1097': 2, 'M1142': 5})
        payload = preview_server_changes(self.config, [self.folder])['server_write_preview']
        expected = [('M1094', 6, 'Sets'), ('M1095', 3, 'Sets'), ('M1096', 2, 'Sets'), ('M1097', 1, 'Sets'), ('M1142', 5, 'PCS')]
        self.assertEqual(sorted((r['product_code'], r['quantity']) for r in payload['write_records']['hardware_items']), [(code, qty) for code, qty, unit in expected])
        self.assertEqual(payload['hardware_source_items'][0]['items'][0]['quantity'], 12)
        self.assertEqual(self.rows(), [])
        for _ in range(2):
            confirm_server_material_preview_memory(self.config, copy.deepcopy(payload), confirm_write=True)
            self.assertEqual(self.rows(), expected)
        self.assertEqual(payload['hardware_source_items'][0]['items'][0]['quantity'], 12)

    def test_odd_report_blocks_preview_and_preserves_existing_facts(self):
        self.write_report({'M1094': 12})
        payload = preview_server_changes(self.config, [self.folder])['server_write_preview']
        confirm_server_material_preview_memory(self.config, payload, confirm_write=True)
        self.write_report({'M1094': 5})
        selection = preview_server_changes(self.config, [self.folder])['hardware_source_selection']
        candidate = next(row for row in selection['conflicts'][0]['candidates']
                         if row['items'][0]['quantity'] == 5)
        with self.assertRaises(RuleError) as error:
            preview_server_changes(self.config, [self.folder], hardware_source_choices={'F100': candidate['id']})
        self.assertEqual(error.exception.code, 'server_rail_pair_quantity_invalid')
        for text in ['F100', 'M1094', '5', '偶数']:
            self.assertIn(text, str(error.exception))
        self.assertEqual(self.rows(), [('M1094', 6, 'Sets')])

    def test_confirmation_revalidates_raw_counts_before_any_write(self):
        self.write_report({'M1094': 6, 'M1095': 4})
        payload = preview_server_changes(self.config, [self.folder])['server_write_preview']
        payload['hardware_source_items'][0]['items'][1]['quantity'] = 3
        with self.assertRaises(RuleError) as error:
            confirm_server_material_preview_memory(self.config, payload, confirm_write=True)
        self.assertEqual(error.exception.code, 'server_rail_pair_quantity_invalid')
        self.assertEqual(self.rows(), [])

    def test_direct_sync_converts_and_reports_odd_quantity_without_replacing(self):
        self.write_report({'M1097': 10})
        sync_order_index(self.config, refresh_outbound_statuses=False, reconcile_outbound=False)
        self.assertEqual(self.rows(), [('M1097', 5, 'Sets')])
        self.write_report({'M1097': 7})
        # A committed source now requires approval of the new revision first.
        from traveler_assistant.fittings import select_latest_fittings, fittings_candidate
        from traveler_assistant.hardware_source_decisions import load_source_decisions
        from traveler_assistant.report_read_context import report_read_session
        selected, *_ = select_latest_fittings([self.report])
        choice = fittings_candidate(selected['F100'])['id']
        with report_read_session({'F100': choice}) as context:
            context.locked_decisions = load_source_decisions(self.config)
            sync_order_index(self.config, refresh_outbound_statuses=False, reconcile_outbound=False)
        self.assertEqual(self.rows(), [('M1097', 5, 'Sets')])
        store = OrderIndexStore(self.config.workflow_database)
        try:
            self.assertTrue(any('M1097' in row['message'] and '偶数' in row['message'] for row in store.active_issues()))
        finally:
            store.close()

    def test_h_and_l_rail_source_pairs_keep_canonical_quantity(self):
        self.write_report({'M1094': 6})
        wb = load_workbook(self.report)
        for row, (name, code) in enumerate([
            ('Left Rail', 'H-Rail'), ('Right Rail', 'H-Rail'),
            ('Lower Left Rail', 'L-Rail'), ('Lower Right Rail', 'L-Rail'),
        ], start=8):
            wb.active.cell(row, 3, name)
            wb.active.cell(row, 5, code)
            wb.active.cell(row, 9, 'Piece')
            wb.active.cell(row, 11, 3)
        wb.save(self.report)
        payload = preview_server_changes(self.config, [self.folder])['server_write_preview']
        confirm_server_material_preview_memory(self.config, payload, confirm_write=True)
        self.assertEqual(self.rows(), [('M1002', 3, '套'), ('M1003', 3, '套'), ('M1094', 3, 'Sets')])

    def test_legacy_order_preview_persistence_converts_once_and_rolls_back_on_odd(self):
        from traveler_assistant.order_workflow import preview_order, persist_preview
        self.write_report({'M1095': 6})
        preview = preview_order(self.config, self.folder, persist_facts=False)
        for _ in range(2):
            persist_preview(self.config, preview)
            self.assertEqual(self.rows(), [('M1095', 3, 'Sets')])
        self.write_report({'M1095': 5})
        preview = preview_order(self.config, self.folder, persist_facts=False)
        with self.assertRaises(RuleError):
            persist_preview(self.config, preview)
        self.assertEqual(self.rows(), [('M1095', 3, 'Sets')])

    def test_existing_pair_rules_and_other_skus_are_not_halved(self):
        for code in ['M1002', 'M1003', 'M1142']:
            self.assertEqual(server_hardware_quantity(code, 3, '套'), {'quantity': 3, 'unit': '套'})
        for code in ['M1094', 'M1095', 'M1096', 'M1097']:
            for quantity in [1, 3, 2.5, -2, float('inf')]:
                with self.subTest(code=code, quantity=quantity), self.assertRaises(RuleError):
                    server_hardware_quantity(code, quantity, 'Piece', factory_order='F100')

"""Behavioral contract for SKU-backed material and product references.

Every test uses a disposable workflow database.  The tests intentionally
exercise application entry points or SQLite's real foreign-key enforcement;
they do not mock persistence or run against the user's workflow database.
"""

from __future__ import annotations

import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from traveler_assistant.core import Config, RuleError
from traveler_assistant.database import connect_database, ensure_schema
from traveler_assistant.inventory import (
    InventoryMappings,
    confirm_product_material_attributes,
    database_stock_requirements,
    import_catalog,
    save_manual_mapping,
    update_manual_mapping,
)
from traveler_assistant.order_details import order_detail
from traveler_assistant.order_index import OrderIndexStore
from traveler_assistant.order_index import confirm_server_material_preview_memory
from traveler_assistant.order_workflow import (
    MaterialItem,
    OrderPreview,
    persist_preview,
    preview_order,
)
from traveler_assistant.production import production_preview


class MaterialSkuContractTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name).resolve()
        self.config = Config(
            state_dir=self.root / "state",
            source_root=self.root / "server",
            order_root=self.root / "travelers",
        )
        self.config.prepare_storage()
        store = OrderIndexStore(self.config.workflow_database)
        try:
            store.upsert_order(
                "CS901",
                order_type="cutToSize",
                source_folder=str(self.config.source_root / "CS901"),
                validation_status="正常",
            )
            store.commit()
        finally:
            store.close()

    def _write_catalog(self, name: str, rows: list[tuple[str, ...]]) -> Path:
        path = self.root / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "商品类别", "*商品编号", "商品名称", "规格型号", "状态",
            "品牌", "计量单位", "备注",
        ])
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        return path

    def _import_standard_catalog(self) -> None:
        import_catalog(
            self.config,
            self._write_catalog(
                "products.xlsx",
                [
                    ("Panel", "M-PANEL-A", "Contract Oak", "19.1*1220*2745mm", "启用", "Brand A", "SHT", "original"),
                    ("Panel", "M-PANEL-B", "Contract Oak Alternate", "19.1*1220*2745mm", "启用", "Brand B", "SHT", "alternate"),
                    ("Panel", "M-DISABLED", "Disabled Oak", "19.1*1220*2745mm", "停用", "Brand D", "SHT", "disabled"),
                    ("Edge band", "M-EDGE-A", "Contract Oak Edge Banding", "22mm*1mm*225m", "启用", "Brand A", "M", "original"),
                    ("Edge band", "M-EDGE-B", "Alternate Edge Banding", "22mm*1mm*225m", "启用", "Brand B", "M", "alternate"),
                ],
            ),
        )

    def _write_material_source(
        self,
        order_id: str = "CS901",
        *,
        color: str = "Contract Oak",
        panel_quantity: float = 2,
        edge_quantity: float = 6,
    ) -> Path:
        folder = self.config.source_root / order_id
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{order_id} materials.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"], sheet["B1"] = "Job:", order_id
        headers = [
            "Room/section", "", "3/4 Plywood", "5/8 Plywood", "1/4 Plywood",
            "3/4 Finish Panel", "1/4 Finish Panel", "Edge Banding (m)", "Color",
        ]
        for column, value in enumerate(headers, 1):
            sheet.cell(2, column).value = value
        values = ["Kitchen", "", 0, 0, 0, panel_quantity, 0, edge_quantity, color]
        for column, value in enumerate(values, 1):
            sheet.cell(3, column).value = value
        totals = ["Total Qty:", "", 0, 0, 0, panel_quantity, 0, edge_quantity, color]
        for column, value in enumerate(totals, 1):
            sheet.cell(14, column).value = value
        sheet["A15"] = "Color Table"
        sheet["A16"] = "Color:"
        sheet["A17"] = "Sheets (3/4):"
        sheet["A18"] = "Sheets (1/4):"
        sheet["A19"] = "Edge Banding (m):"
        sheet["C16"] = color
        sheet["C17"], sheet["C18"], sheet["C19"] = panel_quantity, 0, edge_quantity
        workbook.save(path)
        return folder

    def _confirmed_material_snapshot(self, order_id: str = "CS901") -> list[tuple]:
        with connect_database(self.config.workflow_database) as connection:
            return connection.execute(
                """select order_id, product_code, quantity, source_type, source_path,
                          source_fingerprint, updated_at
                   from material_items where order_id=? order by product_code""",
                (order_id,),
            ).fetchall()

    def _confirm_standard_materials(self) -> OrderPreview:
        self._import_standard_catalog()
        save_manual_mapping(self.config, "19.1mm--Contract Oak", "M-PANEL-A")
        save_manual_mapping(self.config, "Edge banding--Contract Oak", "M-EDGE-A")
        folder = self._write_material_source()
        preview = preview_order(
            self.config,
            folder,
            include_hardware=False,
            persist_facts=False,
        )
        persist_preview(self.config, preview)
        return preview

    def test_parse_confirm_reopen_and_repeat_keep_saved_skus_and_quantities(self) -> None:
        preview = self._confirm_standard_materials()

        reopened = Config(
            state_dir=self.config.state_dir,
            source_root=self.config.source_root,
            order_root=self.config.order_root,
        )
        reopened.prepare_storage()
        materials = {
            row["product_code"]: row
            for row in order_detail(reopened, "CS901")["materials"]
        }
        self.assertEqual(set(materials), {"M-PANEL-A", "M-EDGE-A"})
        self.assertEqual(materials["M-PANEL-A"]["quantity"], 2)
        self.assertEqual(materials["M-PANEL-A"]["material_type"], "panel")
        self.assertEqual(materials["M-PANEL-A"]["color"], "Contract Oak")
        self.assertEqual(materials["M-PANEL-A"]["thickness"], "19.1")
        self.assertEqual(materials["M-PANEL-A"]["unit"], "SHT")
        self.assertEqual(materials["M-PANEL-A"]["brand"], "Brand A")
        self.assertEqual(materials["M-EDGE-A"]["quantity"], 6)
        self.assertEqual(materials["M-EDGE-A"]["material_type"], "edge")
        self.assertEqual(materials["M-EDGE-A"]["unit"], "M")

        first = self._confirmed_material_snapshot()
        persist_preview(reopened, preview)
        second = self._confirmed_material_snapshot()
        self.assertEqual([(row[0], row[1], row[2:6]) for row in first],
                         [(row[0], row[1], row[2:6]) for row in second])
        self.assertEqual(len(second), 2)

    def test_mapping_changes_do_not_rebind_confirmed_materials(self) -> None:
        self._confirm_standard_materials()

        update_manual_mapping(
            self.config,
            "19.1mm--Contract Oak",
            "19.1mm--Contract Oak",
            "M-PANEL-B",
        )
        update_manual_mapping(
            self.config,
            "Edge banding--Contract Oak",
            "Edge banding--Contract Oak",
            "M-EDGE-B",
        )

        detail_codes = {
            row["product_code"]
            for row in order_detail(self.config, "CS901")["materials"]
        }
        self.assertEqual(detail_codes, {"M-PANEL-A", "M-EDGE-A"})
        self.assertEqual(
            {row[1] for row in self._confirmed_material_snapshot()},
            {"M-PANEL-A", "M-EDGE-A"},
        )
        _, requirements = database_stock_requirements(self.config, "CS901")
        quantities = {
            row["productCode"]: row["requiredQuantity"]
            for row in requirements
        }
        self.assertEqual(quantities, {"M-PANEL-A": 2, "M-EDGE-A": 6})

    def test_historical_production_and_server_allocations_keep_product_attributes_bound(self) -> None:
        self._import_standard_catalog()
        with connect_database(self.config.workflow_database) as connection:
            confirmed = {
                "M-PANEL-A": ("Contract Oak", "19.1"),
                "M-PANEL-B": ("Contract Oak Alternate", "19.1"),
            }
            for code, (color, thickness) in confirmed.items():
                confirm_product_material_attributes(
                    connection, code, "panel", color, thickness
                )
                connection.execute(
                    """insert into material_items(
                           order_id, product_code, quantity, source_type,
                           source_path, source_fingerprint, updated_at
                       ) values('CS901',?,1,'test',?,'fixture','now')""",
                    (code, f"/{code}.xlsx"),
                )

            connection.execute(
                """insert into production_records(
                       production_time, source, status,
                       created_at, updated_at
                   ) values('','manual','completed','now','now')"""
            )
            batch_id = connection.execute("select last_insert_rowid()").fetchone()[0]
            connection.execute(
                """insert into production_materials(
                       batch_id, order_id, product_code, quantity
                   ) values(?,?,?,?)""",
                (batch_id, "CS901", "M-PANEL-A", 1),
            )
            connection.execute(
                """insert into server_material_allocations(
                       source_path, source_material_key, product_code,
                       source_quantity, order_id, allocated_quantity,
                       source_fingerprint, created_at, updated_at
                   ) values('/server/materials.xlsx','panel-b','M-PANEL-B',1,
                            'CS901',1,'fixture','now','now')"""
            )
            connection.execute("delete from material_items where order_id='CS901'")
            self.assertEqual(
                connection.execute(
                    "select count(*) from material_items where order_id='CS901'"
                ).fetchone()[0],
                0,
            )

            for code in confirmed:
                with self.subTest(product_code=code):
                    with self.assertRaisesRegex(RuleError, "不能静默重绑"):
                        confirm_product_material_attributes(
                            connection, code, "panel", "Rebound Color", "18"
                        )

            attributes = {
                row[0]: (row[1], row[2], row[3])
                for row in connection.execute(
                    """select code, material_kind, material_color,
                              material_thickness
                       from products where code in ('M-PANEL-A','M-PANEL-B')"""
                )
            }
            self.assertEqual(attributes, {
                "M-PANEL-A": ("panel", "Contract Oak", "19.1"),
                "M-PANEL-B": ("panel", "Contract Oak Alternate", "19.1"),
            })

    def test_long_catalog_names_and_imperial_specs_project_source_business_attributes(self) -> None:
        import_catalog(
            self.config,
            self._write_catalog(
                "mixed-spec-products.xlsx",
                [
                    (
                        "Panel", "M1042",
                        "Penelope FA44 Finished Decorative Panel - Fingerprint Resistant",
                        '3/4 Board | 81-1/2W x 109L', "启用", "Long Name Brand", "SHT", "",
                    ),
                    (
                        "Edge band", "M1043",
                        "PVC Edge Banding for Penelope FA44 Finished Decorative Panel",
                        "22mm x 1mm x 225m", "启用", "Long Name Brand", "M", "",
                    ),
                ],
            ),
        )
        save_manual_mapping(self.config, "19.1mm--Penelope FA44", "M1042")
        save_manual_mapping(self.config, "Edge banding--Penelope FA44", "M1043")
        folder = self._write_material_source(color="Penelope FA44")
        preview = preview_order(
            self.config, folder, include_hardware=False, persist_facts=False
        )

        persist_preview(self.config, preview)

        materials = {
            row["product_code"]: row
            for row in order_detail(self.config, "CS901")["materials"]
        }
        self.assertEqual(materials["M1042"]["material_type"], "panel")
        self.assertEqual(materials["M1042"]["color"], "Penelope FA44")
        self.assertEqual(materials["M1042"]["thickness"], "19.1")
        self.assertEqual(materials["M1043"]["material_type"], "edge")
        self.assertEqual(materials["M1043"]["color"], "Penelope FA44")

    def test_server_memory_confirmation_copies_verified_long_name_product_attributes(self) -> None:
        catalog = self._write_catalog(
            "server-long-name-products.xlsx",
            [
                (
                    "Panel", "M1042",
                    "Penelope FA44 Finished Decorative Panel - Fingerprint Resistant",
                    '3/4 Board | 81-1/2W x 109L', "启用", "Long Name Brand", "SHT", "",
                ),
            ],
        )
        import_catalog(self.config, catalog)

        preview_config = Config(state_dir=self.root / "server-preview-state")
        preview_config.prepare_storage()
        import_catalog(preview_config, catalog)
        with connect_database(preview_config.workflow_database) as preview_connection:
            confirm_product_material_attributes(
                preview_connection, "M1042", "panel", "Penelope FA44", "19.1"
            )
            preview_attributes = preview_connection.execute(
                """select material_kind, material_color, material_thickness
                   from products where code='M1042'"""
            ).fetchone()
        self.assertEqual(preview_attributes, ("panel", "Penelope FA44", "19.1"))

        source_path = "/server/CS901/CS901 materials.xlsx"
        payload = {
            "has_business_changes": True,
            "validation_recomputed": True,
            "source_folders": [],
            "orders": [
                {
                    "order_id": "CS901",
                    "validation_status": "正常",
                    "material_changes": [{"product_code": "M1042"}],
                    "factories": [],
                },
            ],
            "materials": [
                {
                    "source_order_id": "CS901",
                    "source_path": source_path,
                    "product_code": "M1042",
                    "source_quantity": 2,
                    "material_type": preview_attributes[0],
                    "color": preview_attributes[1],
                    "thickness": preview_attributes[2],
                },
            ],
            "write_records": {
                "orders": [
                    {
                        "order_id": "CS901",
                        "order_type": "cutToSize",
                        "source_folder": "/server/CS901",
                        "validation_status": "正常",
                        "updated_at": "now",
                    },
                ],
                "material_items": [
                    {
                        "order_id": "CS901",
                        "product_code": "M1042",
                        "quantity": 2,
                        "source_type": "aihouse",
                        "source_path": source_path,
                        "source_fingerprint": "server-fixture",
                        "updated_at": "now",
                    },
                ],
                "factory_orders": [],
                "hardware_items": [],
                "optimization_artifacts": [],
                "source_files": [],
                "batch_evidence": [],
                "production_batches": [],
                "server_scan_xml_state": [],
            },
        }

        result = confirm_server_material_preview_memory(
            self.config, payload, confirm_write=True
        )

        self.assertTrue(result["server_material_write_confirmed"])
        with connect_database(self.config.workflow_database) as connection:
            production_attributes = connection.execute(
                """select material_kind, material_color, material_thickness
                   from products where code='M1042'"""
            ).fetchone()
            material_rows = connection.execute(
                """select product_code, quantity, source_type, source_path
                   from material_items where order_id='CS901'"""
            ).fetchall()
        self.assertEqual(production_attributes, preview_attributes)
        self.assertEqual(
            material_rows,
            [("M1042", 2.0, "aihouse", source_path)],
        )
        detail = order_detail(self.config, "CS901")["materials"]
        self.assertEqual(len(detail), 1)
        self.assertEqual(detail[0]["product_code"], "M1042")
        self.assertEqual(detail[0]["material_type"], "panel")
        self.assertEqual(detail[0]["color"], "Penelope FA44")
        self.assertEqual(detail[0]["thickness"], "19.1")

    def test_unmatched_and_disabled_materials_leave_old_facts_unchanged(self) -> None:
        good_preview = self._confirm_standard_materials()
        before = self._confirmed_material_snapshot()

        unmatched = OrderPreview(
            order_id=good_preview.order_id,
            folder=good_preview.folder,
            materials_path=good_preview.materials_path,
            materials_sheet_name=good_preview.materials_sheet_name,
            materials=[MaterialItem("panel", 19.1, "No Such Product", 9)],
            edge_banding={},
            factories=[],
            warnings=[],
            include_hardware=False,
        )
        with self.assertRaises(RuleError):
            persist_preview(self.config, unmatched)
        self.assertEqual(self._confirmed_material_snapshot(), before)

        InventoryMappings(self.config.workflow_database).save_manual(
            "19.1mm--Disabled Oak", "M-DISABLED"
        )
        disabled = OrderPreview(
            order_id=good_preview.order_id,
            folder=good_preview.folder,
            materials_path=good_preview.materials_path,
            materials_sheet_name=good_preview.materials_sheet_name,
            materials=[MaterialItem("panel", 19.1, "Disabled Oak", 9)],
            edge_banding={},
            factories=[],
            warnings=[],
            include_hardware=False,
        )
        with self.assertRaises(RuleError):
            persist_preview(self.config, disabled)
        self.assertEqual(self._confirmed_material_snapshot(), before)

    def test_catalog_refresh_preserves_referenced_missing_products(self) -> None:
        confirmed_preview = self._confirm_standard_materials()
        confirmed_facts = self._confirmed_material_snapshot()
        replacement = self._write_catalog(
            "replacement-products.xlsx",
            [
                ("Panel", "M-PANEL-B", "Replacement Oak", "18*1220*2440mm", "启用", "Brand Z", "SHT", "new"),
            ],
        )

        import_catalog(self.config, replacement)

        with connect_database(self.config.workflow_database) as connection:
            products = {
                row[0]: row[1]
                for row in connection.execute(
                    "select code, catalog_present from products order by code"
                )
            }
            self.assertEqual(products["M-PANEL-A"], 0)
            self.assertEqual(products["M-EDGE-A"], 0)
            self.assertEqual(products["M-PANEL-B"], 1)
            self.assertEqual(connection.execute("pragma foreign_key_check").fetchall(), [])
        self.assertEqual(
            {row["product_code"] for row in order_detail(self.config, "CS901")["materials"]},
            {"M-PANEL-A", "M-EDGE-A"},
        )
        with self.assertRaises(RuleError):
            persist_preview(self.config, confirmed_preview)
        self.assertEqual(self._confirmed_material_snapshot(), confirmed_facts)
        with self.assertRaises(RuleError):
            database_stock_requirements(self.config, "CS901")
        self.assertEqual(self._confirmed_material_snapshot(), confirmed_facts)

    def test_every_sku_fact_table_declares_and_enforces_product_foreign_key(self) -> None:
        self._import_standard_catalog()
        tables = {
            "material_items",
            "production_materials",
            "server_material_allocations",
            "hardware_items",
            "inventory_resolution_rules",
        }
        with connect_database(self.config.workflow_database) as connection:
            self.assertEqual(connection.execute("pragma foreign_keys").fetchone()[0], 1)
            self.assertEqual(
                {row[1] for row in connection.execute("pragma table_info(material_items)")},
                {
                    "id", "order_id", "product_code", "quantity", "source_type",
                    "source_path", "source_fingerprint", "updated_at",
                },
            )
            self.assertTrue(
                {"material_kind", "material_color", "material_thickness", "catalog_present"}
                <= {row[1] for row in connection.execute("pragma table_info(products)")}
            )
            required_quantity_column = {
                "production_materials": "quantity",
                "server_material_allocations": "allocated_quantity",
            }
            for table, quantity_column in required_quantity_column.items():
                columns = {row[1] for row in connection.execute(f"pragma table_info({table})")}
                self.assertTrue({"product_code", quantity_column} <= columns, table)
                self.assertTrue(
                    {"material_type", "color", "thickness", "edge", "unit"}.isdisjoint(columns),
                    table,
                )
            for table in tables:
                foreign_keys = connection.execute(
                    f"pragma foreign_key_list({table})"
                ).fetchall()
                product_key = [
                    row for row in foreign_keys
                    if row[2] == "products" and row[3] == "product_code" and row[4] == "code"
                ]
                self.assertEqual(len(product_key), 1, table)
                self.assertEqual(product_key[0][6].upper(), "RESTRICT", table)

            connection.execute(
                """insert into production_records(
                       production_time, source, status,
                       created_at, updated_at
                   ) values('','manual','completed','now','now')"""
            )
            batch_id = connection.execute("select last_insert_rowid()").fetchone()[0]
            invalid_statements = [
                (
                    """insert into material_items(
                           order_id, product_code, quantity, source_type, source_path,
                           source_fingerprint, updated_at
                       ) values('CS901','M-MISSING',1,'test','','','now')""",
                    (),
                ),
                (
                    """insert into production_materials(
                           batch_id, order_id, product_code, quantity
                       ) values(?,?,?,?)""",
                    (batch_id, "CS901", "M-MISSING", 1),
                ),
                (
                    """insert into server_material_allocations(
                           source_path, source_material_key, product_code,
                           source_quantity, order_id, allocated_quantity,
                           source_fingerprint, created_at, updated_at
                       ) values('source','row-1','M-MISSING',1,'CS901',1,'','now','now')""",
                    (),
                ),
                (
                    """insert into hardware_items(
                           order_id, factory_order, product_code,
                           quantity, updated_at
                       ) values('CS901','F901','M-MISSING',1,'now')""",
                    (),
                ),
                (
                    """insert into inventory_resolution_rules(
                           rule_type, source_name, normalized_name, product_code,
                           created_at, updated_at
                       ) values('mapping','Missing','MISSING','M-MISSING','now','now')""",
                    (),
                ),
            ]
            for statement, parameters in invalid_statements:
                with self.subTest(statement=statement.split("(", 1)[0].strip()):
                    with self.assertRaises(sqlite3.IntegrityError):
                        connection.execute(statement, parameters)

            connection.execute(
                """insert into inventory_resolution_rules(
                       rule_type, source_name, normalized_name, product_code,
                       reason, created_at, updated_at
                   ) values('ignore','Ignored','IGNORED',null,'expected','now','now')"""
            )
            connection.execute(
                """insert into material_items(
                       order_id, product_code, quantity, source_type, source_path,
                       source_fingerprint, updated_at
                   ) values('CS901','M-PANEL-A',1,'test','','','now')"""
            )
            with self.assertRaises(sqlite3.IntegrityError):
                connection.execute("delete from products where code='M-PANEL-A'")
            self.assertEqual(connection.execute("pragma foreign_key_check").fetchall(), [])

        store = OrderIndexStore(self.config.workflow_database)
        try:
            self.assertEqual(store.connection.execute("pragma foreign_keys").fetchone()[0], 1)
        finally:
            store.close()

    def test_legacy_migration_preserves_consumption_and_is_idempotent(self) -> None:
        database = self.root / "legacy.sqlite3"
        self._create_legacy_database(database)

        ensure_schema(database)
        first = self._migration_snapshot(database)
        ensure_schema(database)
        self.assertEqual(self._migration_snapshot(database), first)

        config = Config(state_dir=self.root / "legacy-state")
        config.storage_prepared = True
        config.workflow_database.parent.mkdir(parents=True, exist_ok=True)
        database.replace(config.workflow_database)
        store = OrderIndexStore(config.workflow_database)
        try:
            store.connection.execute(
                "insert into orders(order_id, order_type, updated_at) values('CS990','cutToSize','now')"
            )
            store.connection.execute(
                "insert into factory_orders(factory_order,order_id,factory_name,updated_at,stage) select v.factory_order,v.order_id,v.factory_name,v.updated_at,case when v.outbound_status='已出库' then '已出货' when v.optimized then '已优化' else '已拆单' end from (select 'F990' as factory_order,'CS990' as order_id,'CS990-KITCHEN' as factory_name,1 as optimized,'未查询' as outbound_status,'now' as updated_at) v"
            )
            store.commit()
        finally:
            store.close()

        result = production_preview(config, "CS990", ["F990"])
        self.assertEqual(len(result["materials"]), 1)
        self.assertEqual(result["materials"][0]["product_code"], "M-LEGACY")
        self.assertEqual(result["materials"][0]["total_quantity"], 10)
        self.assertEqual(result["materials"][0]["remaining_quantity"], 6)

    def test_legacy_migration_failure_rolls_back_and_can_retry(self) -> None:
        database = self.root / "legacy-failure.sqlite3"
        self._create_legacy_database(database, include_unresolved=True)
        before = self._legacy_material_snapshot(database)
        before_schema = self._legacy_target_schema(database)

        with self.assertRaisesRegex(ValueError, "无法唯一迁移"):
            ensure_schema(database)
        self.assertEqual(self._legacy_material_snapshot(database), before)
        self.assertEqual(self._legacy_target_schema(database), before_schema)
        with sqlite3.connect(database) as connection:
            columns = {row[1] for row in connection.execute("pragma table_info(material_items)")}
            self.assertIn("material_type", columns)
            self.assertNotIn("product_code", columns)
            connection.execute(
                """insert into products(
                       category, code, name, spec, status, brand, remark, unit,
                       cost_price, normalized_code, normalized_name, normalized_spec,
                       normalized_category, normalized_remark
                   ) values('Panel','M-UNKNOWN','Unknown','19.1mm','启用','','','SHT',null,
                            'MUNKNOWN','UNKNOWN','191MM','PANEL','')"""
            )

        ensure_schema(database)
        migrated = self._migration_snapshot(database)
        ensure_schema(database)
        self.assertEqual(self._migration_snapshot(database), migrated)
        self.assertEqual(
            {row[1] for row in migrated["materials"]},
            {"M-LEGACY", "M-UNKNOWN"},
        )

    def test_legacy_migration_upgrades_an_empty_production_material_table(self) -> None:
        database = self.root / "legacy-empty-production.sqlite3"
        self._create_legacy_database(database, include_production=False)

        ensure_schema(database)

        with connect_database(database) as connection:
            columns = {
                row[1]
                for row in connection.execute(
                    "pragma table_info(production_materials)"
                )
            }
            self.assertEqual(columns, {"batch_id", "order_id", "product_code", "quantity"})
            self.assertEqual(
                connection.execute(
                    "select count(*) from production_materials"
                ).fetchone()[0],
                0,
            )

    @staticmethod
    def _create_legacy_database(
        path: Path,
        *,
        include_unresolved: bool = False,
        include_production: bool = True,
    ) -> None:
        connection = sqlite3.connect(path)
        connection.executescript(
            """
            create table products(
                category text not null default '', code text primary key,
                name text not null default '', spec text not null default '',
                status text not null default '', brand text not null default '',
                remark text not null default '', unit text not null default '',
                cost_price real, normalized_code text not null,
                normalized_name text not null, normalized_spec text not null,
                normalized_category text not null, normalized_remark text not null
            );
            insert into products values(
                'Panel','M-LEGACY','Legacy Oak','19.1mm','启用','','','SHT',null,
                'MLEGACY','LEGACYOAK','191MM','PANEL',''
            );
            create table material_items(
                id integer primary key, order_id text not null,
                material_type text not null, color text not null,
                thickness text not null, quantity real not null,
                unit text not null, edge text not null,
                source_type text not null, source_path text not null,
                source_fingerprint text not null, updated_at text not null
            );
            insert into material_items values(
                1,'CS990','panel','Legacy Oak','19.1',10,'pcs','','aihouse',
                '/legacy/material.xlsx','legacy-fingerprint','2026-01-01'
            );
            create table manual_production_batches(
                batch_id integer primary key, batch_number text not null unique,
                order_id text not null, production_time text not null default '',
                source text not null default 'manual', status text not null default 'prepared',
                created_at text not null, updated_at text not null
            );
            insert into manual_production_batches values(
                1,'LEGACY-BATCH','CS990','','manual','completed','2026-01-02','2026-01-02'
            );
            create table manual_production_batch_materials(
                batch_id integer not null, order_id text not null,
                material_type text not null, color text not null,
                thickness text not null, edge text not null, unit text not null,
                quantity real not null
            );
            """
        )
        if include_production:
            connection.execute(
                """insert into manual_production_batch_materials values(
                       1,'CS990','panel','Legacy Oak','19.1','','pcs',4)"""
            )
        if include_unresolved:
            connection.execute(
                """insert into material_items values(
                       2,'CS991','panel','Unknown','19.1',3,'pcs','','aihouse',
                       '/legacy/unknown.xlsx','unknown-fingerprint','2026-01-01')"""
            )
        connection.commit()
        connection.close()

    @staticmethod
    def _legacy_material_snapshot(path: Path) -> list[tuple]:
        with sqlite3.connect(path) as connection:
            return connection.execute(
                "select * from material_items order by id"
            ).fetchall()

    @staticmethod
    def _legacy_target_schema(path: Path) -> list[tuple]:
        with sqlite3.connect(path) as connection:
            return connection.execute(
                """select type, name, sql from sqlite_master
                   where type='table'
                     and (name in ('material_items', 'manual_production_batch_materials')
                          or name like 'material_items_legacy%'
                          or name like 'manual_production_batch_materials_legacy%')
                   order by type, name"""
            ).fetchall()

    @staticmethod
    def _migration_snapshot(path: Path) -> dict[str, list[tuple]]:
        with connect_database(path) as connection:
            return {
                "materials": connection.execute(
                    """select order_id, product_code, quantity, source_type,
                              source_path, source_fingerprint, updated_at
                       from material_items order by order_id, product_code"""
                ).fetchall(),
                "production": connection.execute(
                    """select batch_id, order_id, product_code, quantity
                       from production_materials
                       order by batch_id, product_code"""
                ).fetchall(),
                "foreign_key_check": connection.execute(
                    "pragma foreign_key_check"
                ).fetchall(),
            }


if __name__ == "__main__":
    unittest.main()

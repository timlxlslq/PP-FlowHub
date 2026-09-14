import os
import sqlite3
import tempfile
import unittest
from datetime import datetime
from types import SimpleNamespace
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from traveler_assistant.core import Config, RuleError
from traveler_assistant.database import ensure_schema
from traveler_assistant.order_index import (
    OrderIndexStore,
    _aimes_order_fingerprint,
    _aimes_row_issue,
    _business_aimes_message,
    _business_report_message,
    _business_validation_message,
    _clear_stale_mapping_validation_status,
    _complete_aimes_stage_durations,
    _server_change_message,
    _summarize_server_read_items,
    _is_mixed_order_folder,
    _effective_factory_candidate,
    _exact_resolve_unowned_factories,
    _factory_order_before_initial_date,
    _merge_candidate,
    _load_outbound_records,
    _mark_initial_orders_shipped,
    _merge_database_factory_candidates,
    _orders_requiring_server_scan,
    _optimization_artifacts,
    _memory_factory_selection,
    install_shared_workflow_connection,
    clear_shared_workflow_connection,
    _partition_aimes_rows,
    _record_generated_material_baseline,
    _reconcile_temporary_order_projections,
    _report_files,
    _replace_server_material_facts,
    _reconcile_authoritative_server_material_sources,
    reconcile_outbound_statuses,
    _server_snapshot,
    _server_optimization_monitor_files,
    _server_snapshot_folder,
    _server_folder_rename_pairs,
    _server_folders_for_sync,
    _server_data_change_message,
    _server_preview_payload,
    _source_path_in_dashboard_scope,
    _valid_aimes_order_id,
    _visible_aimes_row,
    assign_aimes_factory_order,
    scan_server_changes,
    ignore_aimes_factories,
    list_order_index,
    save_order_annotations,
    process_server_folder,
    process_server_changes,
    preview_server_changes,
    allocate_server_material,
    confirm_server_material_allocations,
    confirm_server_material_preview,
    confirm_server_material_preview_memory,
    confirm_server_preview,
    confirm_server_preview_memory,
    mark_temporary_folder_manual,
    record_temporary_outbound,
    restore_aimes_factories,
    restore_aimes_order_assignment,
    sync_aimes_index,
    sync_order_index,
)
from traveler_assistant.order_workflow import MaterialItem
from traveler_assistant.inventory import InventoryMappings


class OrderIndexTests(unittest.TestCase):
    @staticmethod
    def _set_permanent_server_policy(store, order_id, folder):
        store.upsert_order(order_id, source_folder=str(folder))
        store.save_server_scan_policy(
            order_id,
            policy="permanent",
            aimes_fingerprint=_aimes_order_fingerprint(store, order_id),
            updated_at="2026-08-31T10:00:00",
        )

    def test_desktop_cs004_and_pp0072_are_offline_memory_preview_fixtures(self):
        """The desktop copies exercise parsing only, not current Server state."""
        desktop = Path.home() / "Desktop"
        folders = [desktop / "cs004", desktop / "pp0072"]
        if not all(folder.is_dir() for folder in folders):
            self.skipTest("Desktop CS004/PP0072 fixtures are not available")
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=desktop,
                order_root=root / "orders",
            )
            config.prepare_storage()
            preview = preview_server_changes(config, folders)["server_write_preview"]
            self.assertNotIn("token", preview)
            self.assertEqual(
                {row["order_id"] for row in preview["orders"]},
                {"CS004", "PP0072"},
            )
            self.assertIn("write_records", preview)
            self.assertFalse((config.state_dir / "server-previews").exists())

    def test_server_preview_summarizes_changes_and_excludes_shipped_factory_orders(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            folder = root / "server" / "CUT TO SIZE" / "cs004"
            folder.mkdir(parents=True)
            source_path = str(folder / "cs004 material.xlsx")
            current = OrderIndexStore(config.workflow_database)
            current.upsert_order("CS004", validation_status="正常")
            current.upsert_factory(
                "F-KITCHEN", order_id="CS004", factory_name="CS004-KITCHEN",
                sales_order_name="CS004", name_source="AIMES", ownership_status="已确认",
                optimized=True, outbound_status="已出库", outbound_document="QTCK-001",
            )
            current.connection.execute(
                """insert into material_items(
                    order_id, material_type, color, thickness, quantity, unit, edge,
                    source_type, source_path, source_fingerprint, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)""",
                ("CS004", "panel", "Woodline 4", "19.1", 4, "pcs", "",
                 "aihouse", source_path, "old", "2026-08-20T10:00:00"),
            )
            current.connection.execute(
                """insert into material_items(
                    order_id, material_type, color, thickness, quantity, unit, edge,
                    source_type, source_path, source_fingerprint, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)""",
                ("CS004", "edge", "Woodline 4", "", 100, "m", "",
                 "aihouse", source_path, "old", "2026-08-20T10:00:00"),
            )
            current.connection.execute(
                """insert into hardware_items(
                    order_id, factory_order, product_code, name, spec, quantity,
                    unit, source_type, source_path, active, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)""",
                ("CS004", "F-KITCHEN", "H1", "Hinge", "", 2, "pcs",
                 "aicnc", "", 1, "2026-08-20T10:00:00"),
            )
            current.commit()
            current.close()

            preview_path = root / "preview.sqlite3"
            preview = OrderIndexStore(preview_path)
            preview.upsert_order("CS004", validation_status="正常", source_folder=str(folder))
            preview.upsert_factory(
                "F-KITCHEN", order_id="CS004", factory_name="CS004-KITCHEN",
                sales_order_name="CS004", name_source="AIMES", ownership_status="已确认",
                optimized=True, outbound_status="已出库", outbound_document="QTCK-001",
            )
            preview.upsert_factory(
                "F-VANITY", order_id="CS004", factory_name="CS004-vanity",
                sales_order_name="CS004", name_source="AIMES", ownership_status="已确认",
                optimized=True, outbound_status="未查询",
            )
            preview.connection.executemany(
                """insert into material_items(
                    order_id, material_type, color, thickness, quantity, unit, edge,
                    source_type, source_path, source_fingerprint, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)""",
                [
                    ("CS004", "panel", "Woodline 4", "19.1", 5, "pcs", "",
                     "aihouse", source_path, "new", "2026-08-20T10:30:00"),
                    ("CS004", "edge", "Woodline 4", "", 120, "m", "",
                     "aihouse", source_path, "new", "2026-08-20T10:30:00"),
                ],
            )
            preview.connection.execute(
                """insert into hardware_items(
                    order_id, factory_order, product_code, name, spec, quantity,
                    unit, source_type, source_path, active, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)""",
                ("CS004", "F-VANITY", "H2", "Drawer slide", "", 4, "pcs",
                 "aicnc", "", 1, "2026-08-20T10:30:00"),
            )
            preview.commit()
            preview.close()

            payload = _server_preview_payload(
                config, preview_path, "token", [folder], include_hardware=True
            )
            order = payload["orders"][0]
            self.assertEqual(
                [item["factory_order"] for item in order["factories"]], ["F-VANITY"]
            )
            self.assertEqual(
                [(item["factory_order"], item["outbound_document"])
                 for item in order["excluded_factories"]],
                [("F-KITCHEN", "QTCK-001")],
            )
            material_changes = {
                (item["material_type"], item["old_quantity"], item["new_quantity"])
                for item in order["material_changes"]
            }
            self.assertEqual(
                material_changes,
                {("panel", 4.0, 5.0), ("edge", 100.0, 120.0)},
            )
            self.assertEqual(
                [(item["factory_order"], item["name"], item["new_quantity"])
                 for item in order["hardware_changes"]],
                [("F-VANITY", "Drawer slide", 4.0)],
            )

    def test_invalid_preview_folder_reports_path_without_copying_database(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source", order_root=root / "orders")
            config.prepare_storage()
            folder = config.source_root / "PP0062-KITCHEN_20260908145832"
            folder.mkdir(parents=True)
            (folder / "nesting_result.xml").write_text("<result/>")
            with patch("traveler_assistant.order_index._report_files", wraps=_report_files) as reports, patch("traveler_assistant.order_index.sqlite3.connect") as connect:
                with self.assertRaises(RuleError) as raised:
                    preview_server_changes(config, [folder])
            self.assertEqual(raised.exception.code, "server_folder_invalid")
            self.assertIn(str(folder), str(raised.exception))
            self.assertIn("上一级订单目录", str(raised.exception))
            reports.assert_called_once_with(folder.resolve())
            connect.assert_not_called()

    def test_server_preview_requires_factory_confirmation_before_production_write(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            report = folder / "Report"
            report.mkdir(parents=True)
            from tests.test_order_workflow import make_board_material_report, make_fittings, make_materials
            make_materials(folder / "PP9999 materials.xlsx")
            make_board_material_report(report / "pp-板材清单.xlsx", factory="F100", name="PP9999-KITCHEN")
            make_fittings(report / "Fittingslist.xlsx", [("F100", 2)])
            stale = OrderIndexStore(config.workflow_database)
            stale.upsert_order(
                "PP9999",
                validation_status="数据异常",
                source_folder=str(folder),
            )
            stale.upsert_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                name_source="AIMES",
                ownership_status="已确认",
            )
            stale.connection.execute(
                "update orders set validation_message=? where order_id=?",
                ("旧的订单校验错误", "PP9999"),
            )
            stale.commit()
            stale.close()

            from traveler_assistant.order_workflow import parse_fittings_groups
            from traveler_assistant.order_index import _server_preview_payload
            with patch("traveler_assistant.order_workflow.parse_fittings_groups", wraps=parse_fittings_groups) as parse, patch("traveler_assistant.order_index._server_preview_payload", wraps=_server_preview_payload) as build, patch(
                "traveler_assistant.order_index.sync_order_index",
                wraps=sync_order_index,
            ) as sync, patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                return_value={"missing": []},
            ):
                preview = preview_server_changes(config, [folder])
            parse.assert_called_once_with((report / "Fittingslist.xlsx").resolve())
            build.assert_called_once()
            self.assertTrue(preview["server_write_preview"]["hardware_source_items"])
            sync.assert_called_once()
            self.assertFalse(sync.call_args.kwargs["full_refresh"])
            self.assertTrue(sync.call_args.kwargs["validate_selected_orders"])
            self.assertFalse(sync.call_args.kwargs["refresh_outbound_statuses"])
            self.assertFalse(sync.call_args.kwargs["reconcile_outbound"])
            payload = preview["server_write_preview"]
            self.assertEqual(payload["orders"][0]["order_id"], "PP9999")
            self.assertEqual(payload["orders"][0]["validation_status"], "正常")
            self.assertEqual(payload["orders"][0]["validation_message"], "")
            self.assertEqual(payload["orders"][0]["factories"][0]["factory_order"], "F100")
            self.assertIn("materials", payload)
            self.assertEqual(
                [item["material_type"] for item in payload["orders"][0]["materials"]],
                ["plywood", "plywood", "plywood", "panel", "panel", "edge"],
            )
            store = OrderIndexStore(config.workflow_database)
            try:
                self.assertEqual(store.connection.execute("select count(*) from orders").fetchone()[0], 1)
                self.assertEqual(
                    store.connection.execute("select count(*) from material_items").fetchone()[0],
                    0,
                )
            finally:
                store.close()

    def test_selected_server_folder_validation_does_not_touch_other_aimes_orders(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            config.prepare_storage()
            selected_folder = config.source_root / "PP9999"
            report = selected_folder / "Report"
            report.mkdir(parents=True)
            from tests.test_order_workflow import make_board_material_report, make_fittings, make_materials
            make_materials(selected_folder / "PP9999 materials.xlsx")
            make_board_material_report(report / "pp-板材清单.xlsx", factory="F100", name="PP9999-KITCHEN")
            make_fittings(report / "Fittingslist.xlsx", [("F100", 2)])

            other_folder = config.source_root / "PP8888"
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order(
                "PP8888",
                validation_status="数据异常",
                source_folder=str(other_folder),
            )
            store.connection.execute(
                "update orders set validation_message=? where order_id=?",
                ("保留的历史校验结果", "PP8888"),
            )
            store.upsert_factory(
                "F200",
                order_id="PP8888",
                factory_name="PP8888-KITCHEN",
                sales_order_name="PP8888",
                name_source="AIMES",
                ownership_status="已确认",
            )
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                return_value={"missing": []},
            ):
                preview_server_changes(config, [selected_folder])

            store = OrderIndexStore(config.workflow_database)
            try:
                self.assertEqual(
                    store.connection.execute(
                        "select validation_status, validation_message from orders where order_id=?",
                        ("PP8888",),
                    ).fetchone(),
                    ("数据异常", "保留的历史校验结果"),
                )
            finally:
                store.close()

    def test_server_preview_preserves_recut_material_and_requires_hardware_choice(self):
        """A recut board adds material without replacing the base hardware report."""
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            base_report = folder / "Report"
            recut_report = folder / "PP9999-recut" / "Report"

            from tests.test_order_workflow import (
                make_board_material_report,
                make_fittings,
                make_materials,
            )

            folder.mkdir(parents=True)
            make_materials(folder / "PP9999 materials.xlsx")
            make_board_material_report(
                base_report / "pp-板材清单.xlsx",
                factory="F100",
                name="PP9999-KITCHEN",
                plywood_qty=2,
            )
            make_board_material_report(
                recut_report / "pp-板材清单.xlsx",
                factory="F100",
                name="PP9999-KITCHEN",
                plywood_qty=1,
            )
            make_fittings(base_report / "Fittingslist.xlsx", [("F100", 2)])
            make_fittings(recut_report / "Fittingslist.xlsx", [("F100", 7)])

            current = OrderIndexStore(config.workflow_database)
            current.upsert_order(
                "PP9999",
                validation_status="数据异常",
                source_folder=str(folder),
            )
            current.upsert_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                name_source="AIMES",
                ownership_status="已确认",
                optimized=True,
            )
            current.connection.execute(
                """insert into material_items(
                    order_id, material_type, color, thickness, quantity, unit, edge,
                    source_type, source_path, source_fingerprint, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    "PP9999", "plywood", "", "18.0", 2, "pcs", "",
                    "aihouse", str((folder / "PP9999 materials.xlsx").resolve()), "old", "old",
                ),
            )
            current.connection.execute(
                """insert into hardware_items(
                    order_id, factory_order, product_code, name, spec, quantity,
                    unit, source_type, source_path, active, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    "PP9999", "F100", "71T950A", "Hinge", "", 2,
                    "pcs", "aicnc", str(folder.resolve()), 1, "old",
                ),
            )
            current.commit()
            self.assertEqual(
                current.connection.execute(
                    "select quantity, source_path from material_items where order_id='PP9999'"
                ).fetchone(),
                (2.0, str((folder / "PP9999 materials.xlsx").resolve())),
            )
            current.close()

            with patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                return_value={"missing": [], "ignored": [], "outbound": []},
            ):
                selection = preview_server_changes(config, [folder])["hardware_source_selection"]
                candidates = selection["conflicts"][0]["candidates"]
                chosen = next(item for item in candidates if item["path"] == str((base_report / "Fittingslist.xlsx").resolve()))
                payload = preview_server_changes(config, [folder], hardware_source_choices={"F100": chosen["id"]})["server_write_preview"]

            order = next(item for item in payload["orders"] if item["order_id"] == "PP9999")
            plywood_changes = [
                item for item in order["material_changes"]
                if item["material_type"] == "plywood"
                and item["thickness"] == "18.0"
            ]
            self.assertTrue(
                any(
                    item["old_quantity"] == 2.0 and item["new_quantity"] == 3.0
                    for item in plywood_changes
                ),
                plywood_changes,
            )
            self.assertEqual(order["hardware_changes"], [])
            self.assertTrue(
                any(
                    str(item["source_path"]).endswith(
                        "PP9999-recut/Report/pp-板材清单.xlsx"
                    )
                    and item["material_type"] == "plywood"
                    and item["thickness"] == "18.0"
                    and item["quantity"] == 1.0
                    for item in payload["materials"]
                )
            )
            self.assertEqual(
                {
                    str(item["source_path"])
                    for item in payload["hardware_source_items"]
                },
                {str((base_report / "Fittingslist.xlsx").resolve())},
            )

    def test_server_confirmation_writes_materials_and_factory_hardware_after_mapping(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            report = folder / "Report"
            report.mkdir(parents=True)
            from tests.test_order_workflow import make_board_material_report, make_fittings, make_materials
            make_materials(folder / "PP9999 materials.xlsx")
            make_board_material_report(report / "pp-板材清单.xlsx", factory="F100", name="PP9999-KITCHEN")
            make_fittings(report / "Fittingslist.xlsx", [("F100", 2)])

            current = OrderIndexStore(config.workflow_database)
            current.upsert_order(
                "PP9999",
                validation_status="正常",
                source_folder=str(folder),
            )
            current.upsert_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                name_source="AIMES",
                ownership_status="已确认",
            )
            current.commit()
            current.close()

            fitting = SimpleNamespace(
                code="WJ-UNMAPPED", name="Unmapped Hinge", size="Full", unit="Piece", quantity=2
            )

            def resolve_items(current_config, pairs):
                mappings = InventoryMappings(current_config.workflow_database)
                if not any(item.name == "Unmapped Hinge" for item, _ in pairs):
                    return {"missing": [], "ignored": [], "outbound": []}
                if mappings.manual_code("Unmapped Hinge"):
                    return {"missing": [], "ignored": [], "outbound": []}
                return {
                    "missing": [{
                        "name": "Unmapped Hinge",
                        "source_code": "WJ-UNMAPPED",
                        "quantity": 2,
                        "message": "找不到有效库存 SKU",
                    }],
                    "ignored": [],
                    "outbound": [],
                }

            with patch(
                "traveler_assistant.order_workflow.parse_fittings_groups",
                return_value=[("F100", [fitting])],
            ), patch(
                "traveler_assistant.order_workflow.resolve_inventory_items",
                side_effect=resolve_items,
            ), patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                side_effect=resolve_items,
            ):
                preview = preview_server_changes(config, [folder])

            payload = preview["server_write_preview"]
            self.assertEqual(
                [item["name"] for item in payload["hardware_mapping_requirements"]],
                ["Unmapped Hinge"],
            )
            invalid_payload = dict(payload)
            invalid_payload["orders"] = [
                dict(payload["orders"][0], validation_status="数据异常", validation_message="旧的校验错误")
            ]
            with self.assertRaises(RuleError) as blocked_validation:
                confirm_server_material_preview_memory(
                    config, invalid_payload, confirm_write=True
                )
            self.assertEqual(blocked_validation.exception.code, "order_validation")
            with patch(
                "traveler_assistant.order_workflow.parse_fittings_groups",
                return_value=[("F100", [fitting])],
            ), patch(
                "traveler_assistant.order_workflow.resolve_inventory_items",
                side_effect=resolve_items,
            ), patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                side_effect=resolve_items,
            ), self.assertRaises(RuleError) as blocked:
                confirm_server_material_preview_memory(config, payload, confirm_write=True)
            self.assertEqual(blocked.exception.code, "hardware_mapping_required")

            InventoryMappings(config.workflow_database).save_manual("Unmapped Hinge", "M1001")
            with patch(
                "traveler_assistant.order_workflow.parse_fittings_groups",
                return_value=[("F100", [fitting])],
            ), patch(
                "traveler_assistant.order_workflow.resolve_inventory_items",
                side_effect=resolve_items,
            ), patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                side_effect=resolve_items,
            ):
                confirmed = confirm_server_material_preview_memory(
                    config, payload, confirm_write=True
                )

            self.assertEqual(confirmed["confirmation_mode"], "order_materials_and_factory_hardware")
            self.assertEqual(confirmed["factory_orders"], ["F100"])
            store = OrderIndexStore(config.workflow_database)
            try:
                self.assertEqual(
                    store.connection.execute(
                        "select count(*) from material_items where order_id='PP9999'"
                    ).fetchone()[0],
                    6,
                )
                hardware = store.connection.execute(
                    "select order_id, factory_order, product_code, name, quantity from hardware_items"
                ).fetchall()
                self.assertEqual(hardware, [("PP9999", "F100", "WJ-UNMAPPED", "Unmapped Hinge", 2.0)])
                self.assertNotIn("factory_order", {
                    row[1] for row in store.connection.execute("pragma table_info(material_items)").fetchall()
                })
            finally:
                store.close()

            with self.assertRaises(RuleError):
                confirm_server_preview_memory(config, payload, "PP9999", "F100")
            confirm_server_preview_memory(
                config,
                payload,
                "PP9999",
                "F100",
                confirm_write=True,
            )
            store = OrderIndexStore(config.workflow_database)
            self.assertEqual(
                store.connection.execute("select order_id from orders").fetchall(),
                [("PP9999",)],
            )
            self.assertEqual(
                store.connection.execute("select factory_order, order_id from factory_orders").fetchall(),
                [("F100", "PP9999")],
            )
            material_rows = store.connection.execute(
                "select order_id, material_type, quantity from material_items"
            ).fetchall()
            self.assertEqual(
                sorted(material_rows),
                sorted([
                    ("PP9999", "edge", 12.5),
                    ("PP9999", "panel", 4.0),
                    ("PP9999", "panel", 1.0),
                    ("PP9999", "plywood", 2.0),
                    ("PP9999", "plywood", 1.0),
                    ("PP9999", "plywood", 3.0),
                ]),
            )
            store.close()

    def test_cut_to_size_server_confirmation_can_skip_hardware_for_entire_order(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            config.prepare_storage()
            folder = config.source_root / "CS999"
            current = OrderIndexStore(config.workflow_database)
            current.upsert_order(
                "CS999",
                validation_status="正常",
                source_folder=str(folder),
            )
            current.commit()
            current.close()
            report = folder / "Report"
            report.mkdir(parents=True)
            from tests.test_order_workflow import make_board_material_report, make_fittings, make_materials
            make_materials(folder / "CS999 materials.xlsx")
            make_board_material_report(report / "cs-板材清单.xlsx", factory="F200", name="CS999-VANITY")
            make_fittings(report / "Fittingslist.xlsx", [("F200", 2)])

            fitting = SimpleNamespace(
                code="WJ-CUSTOMER", name="Customer Hinge", size="Full", unit="Piece", quantity=2
            )

            resolution_calls = 0

            def unresolved(current_config, pairs):
                nonlocal resolution_calls
                resolution_calls += 1
                if resolution_calls == 1:
                    # Let the normal read-only index build complete; the
                    # preview refresh below is the SKU gate under test.
                    return {"missing": [], "ignored": [], "outbound": []}
                return {
                    "missing": [{
                        "name": "Customer Hinge",
                        "source_code": "WJ-CUSTOMER",
                        "quantity": 2,
                        "message": "找不到有效库存 SKU",
                    }],
                    "ignored": [],
                    "outbound": [],
                }

            with patch(
                "traveler_assistant.order_workflow.parse_fittings_groups",
                return_value=[("F200", [fitting])],
            ), patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                side_effect=unresolved,
            ):
                preview = preview_server_changes(config, [folder])
                payload = preview["server_write_preview"]
                self.assertEqual(payload["orders"][0]["order_type"], "cutToSize")
                self.assertEqual(
                    [item["name"] for item in payload["hardware_mapping_requirements"]],
                    ["Customer Hinge"],
                )
                confirmed = confirm_server_material_preview_memory(
                    config,
                    payload,
                    confirm_write=True,
                    skip_hardware_order_ids=["CS999"],
                )

            self.assertEqual(confirmed["hardware_skipped_orders"], ["CS999"])
            self.assertEqual(confirmed["factory_orders"], ["F200"])
            store = OrderIndexStore(config.workflow_database)
            try:
                self.assertEqual(
                    store.connection.execute(
                        "select order_id from factory_orders where factory_order='F200'"
                    ).fetchone()[0],
                    "CS999",
                )
                self.assertEqual(
                    store.connection.execute(
                        "select count(*) from hardware_items where factory_order='F200'"
                    ).fetchone()[0],
                    0,
                )
                self.assertEqual(
                    store.connection.execute(
                        "select count(*) from material_items where order_id='CS999'"
                    ).fetchone()[0],
                    6,
                )
            finally:
                store.close()
    def test_server_scan_blocks_material_preview_until_source_file_is_fixed(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            folder.mkdir(parents=True)
            material_path = folder / "PP9999 materials.xlsx"
            material_path.write_bytes(b"not an Excel workbook")

            first_scan = scan_server_changes(config)
            scan_timing = first_scan["operation_timing"]
            self.assertAlmostEqual(
                scan_timing["total_seconds"],
                sum(stage["duration_seconds"] for stage in scan_timing["stages"]),
                places=3,
            )
            self.assertEqual(
                [stage["stage"] for stage in scan_timing["stages"]],
                ["server_metadata", "optimization_evidence", "scan_finalize"],
            )
            self.assertFalse(any(
                issue["kind"] == "material_validation"
                for issue in first_scan["current_issues"]
            ))
            with self.assertRaises(RuleError) as raised:
                preview_server_changes(config, [folder])
            self.assertEqual(raised.exception.code, "material_validation")

            from tests.test_order_workflow import make_materials
            make_materials(material_path)
            with patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                return_value={"missing": []},
            ):
                second_scan = scan_server_changes(config)
                self.assertFalse(any(
                    issue["kind"] == "material_validation"
                    for issue in second_scan["current_issues"]
                ))
                preview = preview_server_changes(config, [folder])

            preview_timing = preview["operation_timing"]
            self.assertAlmostEqual(
                preview_timing["total_seconds"],
                sum(stage["duration_seconds"] for stage in preview_timing["stages"]),
                places=3,
            )
            self.assertEqual(
                [stage["stage"] for stage in preview_timing["stages"]],
                ["source_selection", "preview_database", "server_parse", "preview_build"],
            )

            self.assertEqual(
                preview["server_write_preview"]["orders"][0]["order_id"],
                "PP9999",
            )

    def test_server_preview_validates_material_before_room_allocation(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            config.prepare_storage()
            folder = config.source_root / "CS004"
            folder.mkdir(parents=True)
            material_path = folder / "CS004 material.xlsx"
            from tests.test_order_workflow import make_materials
            make_materials(material_path, order_id="CS004")
            workbook = load_workbook(material_path)
            workbook.active["I3"] = None
            workbook.save(material_path)

            with patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                return_value={"missing": []},
            ):
                with self.assertRaises(RuleError) as raised:
                    preview_server_changes(config, [folder])

            self.assertEqual(raised.exception.code, "material_validation")
            self.assertIn("Color 为空", str(raised.exception))
            self.assertNotIn("请手工修正 Room/section", str(raised.exception))
            self.assertIn("重新预览", str(raised.exception))

    def test_server_material_allocation_splits_one_source_row_between_orders(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            config.prepare_storage()
            token = "a" * 32
            preview_path = config.state_dir / "server-previews" / token / "workflow.sqlite3"
            ensure_schema(preview_path)
            preview = OrderIndexStore(preview_path)
            preview.upsert_order("PP9999", source_folder="/server/mixed")
            preview.upsert_order("PP8888", source_folder="/server/mixed")
            preview.upsert_factory(
                "F100", order_id="PP9999", factory_name="PP9999 KITCHEN",
                sales_order_name="PP9999", name_source="AIMES",
            )
            preview.upsert_factory(
                "F200", order_id="PP8888", factory_name="PP8888 KITCHEN",
                sales_order_name="PP8888", name_source="AIMES",
            )
            preview.connection.execute(
                """
                insert into material_items(
                    order_id, material_type, color, thickness, quantity, unit,
                    edge, source_type, source_path, source_fingerprint, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    "PP9999", "panel", "Frappe 3", "19.1", 13, "pcs", "",
                    "aihouse", "/server/mixed/material.xlsx", "fingerprint", "now",
                ),
            )
            preview.connection.execute(
                "insert into server_material_preview_scopes(source_folder) values(?)",
                ("/server/mixed",),
            )
            preview.commit()
            material_id = preview.connection.execute(
                "select id from material_items"
            ).fetchone()[0]
            preview.close()

            first = allocate_server_material(config, token, material_id, "PP9999", 8)
            second = allocate_server_material(config, token, material_id, "PP8888", 5)
            self.assertEqual(first["material"]["remaining_quantity"], 5)
            self.assertEqual(second["material"]["remaining_quantity"], 0)

            confirmed = confirm_server_material_allocations(
                config, token, confirm_write=True
            )
            self.assertEqual(confirmed["orders"], ["PP8888", "PP9999"])
            store = OrderIndexStore(config.workflow_database)
            rows = store.connection.execute(
                """
                select order_id, quantity
                from material_items
                order by order_id
                """
            ).fetchall()
            allocations = store.connection.execute(
                """
                select order_id, allocated_quantity
                from server_material_allocations
                order by order_id
                """
            ).fetchall()
            store.close()

        self.assertEqual(rows, [("PP8888", 5.0), ("PP9999", 8.0)])
        self.assertEqual(allocations, [("PP8888", 5.0), ("PP9999", 8.0)])

    def test_server_material_allocation_ignores_stale_sqlite_row_ids(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            config.prepare_storage()
            token = "b" * 32
            preview_path = config.state_dir / "server-previews" / token / "workflow.sqlite3"
            ensure_schema(preview_path)
            preview = OrderIndexStore(preview_path)
            preview.upsert_order("PP9999", source_folder="/server/one")
            source_path = "/server/one/material.xlsx"
            preview.connection.execute(
                """
                insert into material_items(
                    order_id, material_type, color, thickness, quantity, unit,
                    edge, source_type, source_path, source_fingerprint, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)
                """,
                ("PP9999", "panel", "Ivory Oak", "19.1", 3, "pcs", "",
                 "aihouse", source_path, "fingerprint", "now"),
            )
            ivory_id = preview.connection.execute(
                "select id from material_items where color='Ivory Oak'"
            ).fetchone()[0]
            preview.connection.execute(
                """
                insert into material_items(
                    order_id, material_type, color, thickness, quantity, unit,
                    edge, source_type, source_path, source_fingerprint, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)
                """,
                ("PP9999", "plywood", "", "5.4", 6, "pcs", "",
                 "aihouse", source_path, "fingerprint", "now"),
            )
            # This row has the old panel id but the fields of the plywood
            # fact.  The former implementation counted it against Ivory Oak.
            preview.connection.execute(
                """
                insert into server_material_allocations(
                    source_material_id, source_path, source_material_key,
                    material_type, color, thickness, unit, edge,
                    source_quantity, order_id, allocated_quantity,
                    source_fingerprint, created_at, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (ivory_id, source_path, str(ivory_id), "plywood", "", "5.4", "pcs", "",
                 6, "PP9999", 6, "fingerprint", "now", "now"),
            )
            preview.connection.execute(
                "insert into server_material_preview_scopes(source_folder) values(?)",
                ("/server/one",),
            )
            preview.commit()
            preview.close()

            confirmed = confirm_server_material_allocations(
                config, token, confirm_write=True
            )
            self.assertEqual(confirmed["orders"], ["PP9999"])
            store = OrderIndexStore(config.workflow_database)
            rows = store.connection.execute(
                """
                select material_type, color, quantity
                from material_items
                where order_id='PP9999'
                order by material_type, color
                """
            ).fetchall()
            store.close()

        self.assertEqual(rows, [("panel", "Ivory Oak", 3.0), ("plywood", "", 6.0)])

    def test_resolved_mapping_clears_stale_order_validation_error(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order(
                "PP9999",
                order_type="temporary",
                validation_status="数据异常",
                source_folder=str(root / "source" / "PP9999"),
            )
            store.connection.execute(
                "update orders set validation_message = ? where order_id = ?",
                ("订单存在未完成商品 SKU 处理：LED。", "PP9999"),
            )
            cleared = _clear_stale_mapping_validation_status(
                store,
                {"PP9999"},
                set(),
                "2026-08-19T10:00:00",
            )
            row = store.connection.execute(
                "select validation_status, validation_message from orders where order_id = ?",
                ("PP9999",),
            ).fetchone()
            store.close()

        self.assertEqual(cleared, 1)
        self.assertEqual(row, ("正常", ""))

    def test_unresolved_mapping_keeps_order_validation_error(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order(
                "PP9999",
                order_type="temporary",
                validation_status="数据异常",
                source_folder=str(root / "source" / "PP9999"),
            )
            store.connection.execute(
                "update orders set validation_message = ? where order_id = ?",
                ("订单存在未完成商品 SKU 处理：LED。", "PP9999"),
            )
            cleared = _clear_stale_mapping_validation_status(
                store,
                {"PP9999"},
                {"hardware_mapping:PP9999:/server/Fittingslist.xlsx"},
                "2026-08-19T10:00:00",
            )
            row = store.connection.execute(
                "select validation_status, validation_message from orders where order_id = ?",
                ("PP9999",),
            ).fetchone()
            store.close()

        self.assertEqual(cleared, 0)
        self.assertEqual(row, ("数据异常", "订单存在未完成商品 SKU 处理：LED。"))

    def test_fully_shipped_temporary_order_is_excluded_from_unfinished_stage(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order(
                "PP9999",
                order_type="temporary",
                validation_status="正常",
                stage="待人工处理",
                source_folder=str(root / "source" / "temporary-production"),
            )
            store.upsert_factory(
                "F999",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                name_source="AIMES",
                ownership_status="已确认",
                outbound_status="已出库",
                outbound_document="QTCK-999",
            )
            store.commit()

            row = store.summaries()[0]
            persisted_stage = store.connection.execute(
                "select stage from orders where order_id = ?",
                ("PP9999",),
            ).fetchone()[0]
            store.close()

        self.assertEqual(row["stage"], "已出货")
        self.assertEqual(persisted_stage, "已出货")

    def test_temporary_projection_is_removed_without_overwriting_formal_order(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            formal_folder = config.source_root / "PP9999"
            formal_folder.mkdir(parents=True)
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order(
                "PP9999",
                order_type="temporary",
                source_folder=str(root / "source" / "temporary-recut"),
            )
            store.upsert_factory(
                "F999",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                name_source="AIMES",
                source_folder=str(formal_folder),
            )
            store.upsert_order(
                "B12",
                order_type="temporary",
                source_folder=str(root / "source" / "B12 repair"),
            )

            self.assertEqual(_reconcile_temporary_order_projections(store), 2)
            formal = store.connection.execute(
                "select order_type, source_folder from orders where order_id = 'PP9999'"
            ).fetchone()
            orphan = store.connection.execute(
                "select 1 from orders where order_id = 'B12'"
            ).fetchone()
            store.close()

        self.assertEqual(formal, ("owned", str(formal_folder)))
        self.assertIsNone(orphan)

    def test_aimes_stage_durations_exclude_aggregate_and_account_for_backend_overhead(self):
        stages = _complete_aimes_stage_durations([
            {"stage": "login", "label": "登录 AIMES", "duration_seconds": 1.2},
            {"stage": "attempt", "label": "获取 AIMES 数据成功，总计用时", "duration_seconds": 9.9},
        ], 2.0)

        self.assertEqual(
            [item["label"] for item in stages],
            ["登录 AIMES", "后台准备与收尾"],
        )
        self.assertAlmostEqual(
            sum(float(item["duration_seconds"]) for item in stages),
            2.0,
            places=6,
        )

    def test_order_annotations_store_single_actual_installation_start_date(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP9999", source_folder="/server/PP9999")
            store.upsert_factory(
                "F999",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                name_source="AIMES",
                ownership_status="已确认",
            )
            store.commit()
            store.close()

            saved = save_order_annotations(
                config,
                "PP9999",
                user_note="客户要求安装前确认台面颜色",
                planned_days=[
                    {"date": "2026-07-08", "installer": "安装组 A"},
                ],
                actual_days=[
                    {"date": "2026-07-08", "installer": "安装组 A"},
                ],
            )

            self.assertTrue(saved["saved"])
            self.assertEqual(
                [item["order_id"] for item in saved["orders"]],
                ["PP9999"],
            )
            row = saved["order"]
            self.assertEqual(row["user_note"], "客户要求安装前确认台面颜色")
            self.assertEqual(row["installation"]["planned"]["start_date"], "2026-07-08")
            self.assertEqual(row["installation"]["planned"]["end_date"], "2026-07-08")
            self.assertEqual(row["installation"]["planned"]["day_count"], 1)
            self.assertEqual(row["installation"]["actual"]["days"], [
                {"date": "2026-07-08", "installer": "安装组 A"},
            ])
            self.assertEqual(row["installation"]["actual"]["start_date"], "2026-07-08")
            self.assertEqual(row["installation"]["actual"]["end_date"], "2026-07-08")
            self.assertEqual(row["installation"]["actual"]["day_count"], 1)

            reopened = OrderIndexStore(config.workflow_database)
            reopened.upsert_order("PP9999", server_seen="2026-08-18T12:00:00")
            persisted = next(item for item in reopened.summaries() if item["order_id"] == "PP9999")
            reopened.close()
            self.assertEqual(persisted["user_note"], "客户要求安装前确认台面颜色")
            self.assertEqual(persisted["installation"]["actual"]["day_count"], 1)

    def test_order_annotations_reject_multiple_actual_installation_dates(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP9998", source_folder="/server/PP9998")
            store.commit()
            with self.assertRaisesRegex(ValueError, "实际安装日期只能填写一个开始日期"):
                store.save_order_annotations(
                    "PP9998",
                    user_note="",
                    planned_days=[],
                    actual_days=[
                        {"date": "2026-07-08", "installer": "安装组 A"},
                        {"date": "2026-07-11", "installer": "安装组 A"},
                    ],
                )
            store.close()

    def test_order_index_collapses_historical_actual_installation_dates(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP9997", source_folder="/server/PP9997")
            # Simulate a legacy database from before the single-start-date
            # unique index was introduced.
            store.connection.execute("drop index idx_order_installation_actual_start")
            store.connection.executemany(
                """
                insert into order_installation_days(
                    order_id, date_type, install_date, installer, updated_at
                ) values(?,?,?,?,?)
                """,
                [
                    ("PP9997", "actual", "2026-07-11", "安装组 B", "2026-07-11T08:00:00"),
                    ("PP9997", "actual", "2026-07-08", "安装组 A", "2026-07-08T08:00:00"),
                ],
            )
            store.connection.commit()
            store.close()

            reopened = OrderIndexStore(config.workflow_database)
            rows = reopened.connection.execute(
                """
                select install_date, installer
                from order_installation_days
                where order_id = ? and date_type = 'actual'
                """,
                ("PP9997",),
            ).fetchall()
            reopened.close()
            self.assertEqual(rows, [("2026-07-08", "安装组 A")])

    def test_order_annotations_allow_missing_installer_but_reject_duplicate_dates(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP9999")
            store.commit()
            saved = store.save_order_annotations(
                "PP9999",
                user_note="",
                planned_days=[{"date": "2026-07-08", "installer": ""}],
                actual_days=[],
            )
            self.assertTrue(saved["saved"])
            self.assertEqual(
                saved["order"]["installation"]["planned"]["days"],
                [{"date": "2026-07-08", "installer": ""}],
            )
            with self.assertRaisesRegex(ValueError, "计划安装日期只能填写一个开始日期"):
                store.save_order_annotations(
                    "PP9999",
                    user_note="",
                    planned_days=[
                        {"date": "2026-07-08", "installer": ""},
                        {"date": "2026-07-11", "installer": ""},
                    ],
                    actual_days=[],
                )
            with self.assertRaisesRegex(ValueError, "实际安装日期只能填写一个开始日期"):
                store.save_order_annotations(
                    "PP9999",
                    user_note="",
                    planned_days=[
                        {"date": "2026-07-08", "installer": "安装组"},
                    ],
                    actual_days=[
                        {"date": "2026-07-08", "installer": "安装组"},
                        {"date": "2026-07-08", "installer": "安装组"},
                    ],
                )
            store.close()

    def test_server_folder_rename_requires_unique_identical_report_signature(self):
        old = "/server/PP0099"
        new = "/server/PP0099-renamed"
        previous = {
            old: {"source_folder": old, "kind": "folder", "order_id": "PP0099", "modified_at": 10, "size": 1},
            f"{old}/material.xlsx": {"source_folder": old, "kind": "material", "order_id": "PP0099", "modified_at": 20, "size": 200},
        }
        current = {
            new: {"source_folder": new, "kind": "folder", "order_id": "PP0099", "modified_at": 99, "size": 1},
            f"{new}/material.xlsx": {"source_folder": new, "kind": "material", "order_id": "PP0099", "modified_at": 20, "size": 200},
        }
        self.assertEqual(_server_folder_rename_pairs(previous, current), [(old, new)])
        current[f"{new}/extra.xlsx"] = {"source_folder": new, "kind": "material", "order_id": "PP0099", "modified_at": 20, "size": 200}
        self.assertEqual(_server_folder_rename_pairs(previous, current), [])
    def test_server_material_replacement_collapses_old_source_path_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            old_path = root / "Downloads" / "PP0065 materials.xlsx"
            current_path = root / "server" / "PP0065 materials.xlsx"
            current_path.parent.mkdir(parents=True)
            current_path.write_bytes(b"same workbook content")
            store.connection.execute(
                """insert into material_items(
                    order_id, material_type, color, thickness, quantity, unit,
                    source_type, source_path, updated_at
                ) values(?,?,?,?,?,?,?,?,?)""",
                ("PP0065", "panel", "Rosales 3", "19.1", 17, "pcs",
                 "aihouse", str(old_path), "before"),
            )
            _replace_server_material_facts(
                store,
                "PP0065",
                current_path,
                [MaterialItem("panel", 19.1, "Rosales 3", 17)],
                {"Rosales 3": 239.56},
                InventoryMappings(config.workflow_database),
                "after",
            )
            rows = store.connection.execute(
                "select source_path, source_fingerprint, count(*) from material_items where order_id='PP0065' group by source_path, source_fingerprint"
            ).fetchall()
            store.close()

            self.assertEqual(len(rows), 1)
            self.assertEqual({row[0] for row in rows}, {str(current_path)})
            self.assertEqual(rows[0][2], 2)
            self.assertTrue(rows[0][1])

    def test_server_material_scope_retires_rows_from_previous_server_root(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            current_folder = root / "server" / "Optimized Orders" / "PP0072"
            old_folder = root / "fixtures" / "Optimized Orders" / "PP0072"
            current_path = current_folder / "PP0072 materials.xlsx"
            old_path = old_folder / "PP0072 materials.xlsx"
            for path in (current_path, old_path):
                path.parent.mkdir(parents=True, exist_ok=True)
            for path in (current_path, old_path):
                store.connection.execute(
                    """insert into material_items(
                        order_id, material_type, color, thickness, quantity, unit,
                        source_type, source_path, updated_at
                    ) values(?,?,?,?,?,?,?,?,?)""",
                    ("PP0072", "plywood", "", "18.0", 22, "pcs", "aihouse", str(path), "now"),
                )
                store.connection.execute(
                    """insert into source_files(
                        path, source_folder, kind, order_id, modified_at, size, last_seen
                    ) values(?,?,?,?,?,?,?)""",
                    (str(path), str(path.parent), "material", "PP0072", 1, 1, "now"),
                )
            removed = _reconcile_authoritative_server_material_sources(
                store, {"PP0072": {str(current_folder)}}
            )
            rows = store.connection.execute(
                "select source_path, quantity from material_items where order_id='PP0072' order by source_path"
            ).fetchall()
            source_rows = store.connection.execute(
                "select source_folder from source_files where order_id='PP0072' order by source_folder"
            ).fetchall()
            store.close()

            self.assertEqual(removed, {"PP0072"})
            self.assertEqual(rows, [(str(current_path), 22.0)])
            self.assertEqual(source_rows, [(str(current_folder),)])

    def test_prepared_sync_can_resolve_material_mappings_before_fittings_import_path(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            material = config.source_root / "PP9999" / "PP9999 materials.xlsx"
            material.parent.mkdir(parents=True)
            material.write_bytes(b"placeholder")
            config.prepare_storage()

            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=[]), \
                 patch(
                     "traveler_assistant.order_workflow.parse_order_materials",
                     return_value=(
                         "Sheet1",
                         [MaterialItem("panel", 19.1, "Test Oak", 1)],
                         {"Test Oak": 1},
                     ),
                 ), \
                 patch(
                     "traveler_assistant.inventory.resolve_inventory_items",
                     return_value={"missing": []},
                 ):
                result = sync_order_index(config)

            self.assertEqual(result["orders"], [])

    def test_server_read_trace_is_grouped_by_folder_and_file_kind(self):
        items = [
            ("/Volumes/server/CUT TO SIZE/cs001", "/Volumes/server/CUT TO SIZE/cs001", "folder"),
            ("/Volumes/server/CUT TO SIZE/cs001/CS001 materials.xlsx", "/Volumes/server/CUT TO SIZE/cs001", "material"),
            ("/Volumes/server/CUT TO SIZE/cs001/Report/pp-板材清单-new.xlsx", "/Volumes/server/CUT TO SIZE/cs001", "board"),
            ("/Volumes/server/CUT TO SIZE/cs001/Report/Fittingslist.xlsx", "/Volumes/server/CUT TO SIZE/cs001", "fittings"),
            ("/Volumes/server/Optimized Orders/PP0035", "/Volumes/server/Optimized Orders/PP0035", "folder"),
            ("/Volumes/server/Optimized Orders/PP0035/PP0035 materials.xlsx", "/Volumes/server/Optimized Orders/PP0035", "material"),
        ]

        summary = _summarize_server_read_items(items)

        self.assertIn("读取 2 个文件夹及 4 个相关文件", summary)
        self.assertIn("cs001 文件夹（下属：1 个 material 文件、1 个板材清单、1 个五金清单）", summary)
        self.assertIn("PP0035 文件夹（下属：1 个 material 文件）", summary)
        self.assertNotIn("/Volumes/server/CUT TO SIZE", summary)

    def test_server_read_trace_limits_folder_examples(self):
        items = [
            (f"/server/CS{i:03d}", f"/server/CS{i:03d}", "folder")
            for i in range(1, 8)
        ]

        summary = _summarize_server_read_items(items)

        self.assertIn("读取 7 个文件夹及 0 个相关文件", summary)
        self.assertIn("另有 1 个文件夹已省略", summary)
        self.assertNotIn("CS007 文件夹", summary)

    def test_server_scan_covers_owned_and_cut_to_size_roots(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "server" / "Optimized Orders")
            owned = config.source_root / "PP9999"
            cut_to_size = config.source_root.parent / "CUT TO SIZE" / "CS999"
            owned.mkdir(parents=True)
            cut_to_size.mkdir(parents=True)
            (owned / "PP9999 materials.xlsx").write_bytes(b"material")
            (cut_to_size / "CS999 materials.xlsx").write_bytes(b"material")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100", order_id="PP9999", factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999", split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            store.upsert_aimes_factory(
                "F101", order_id="CS999", factory_name="CS999-KITCHEN",
                sales_order_name="CS999", split_time="2026-08-10T08:31:00",
                seen_at="2026-08-10T08:31:00",
            )
            store.commit()
            root_path, snapshot = _server_snapshot(config, store)
            store.close()

            self.assertEqual(root_path, config.source_root)
            folder_paths = {
                path for path, item in snapshot.items() if item["kind"] == "folder"
            }
            self.assertIn(str(owned), folder_paths)
            self.assertIn(str(cut_to_size), folder_paths)
            store = OrderIndexStore(config.workflow_database)
            _, folders = _server_folders_for_sync(config, None, store=store)
            store.close()
            self.assertEqual({str(item) for item in folders}, {str(owned), str(cut_to_size)})

    def test_successful_cut_to_size_preview_does_not_claim_optimization_without_aicnc_evidence(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "server" / "Optimized Orders")
            folder = config.source_root.parent / "CUT TO SIZE" / "CS005"
            report = folder / "Report" / "板材清单.xlsx"
            report.parent.mkdir(parents=True)
            report.write_bytes(b"placeholder")
            aimes_row = {
                "factory_order": "F100",
                "factory_name": "CS005-KITCHEN",
                "sales_order_name": "CS005",
                "split_time": "2026-08-10T08:30:00",
            }
            preview = SimpleNamespace(factories=[])
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=[aimes_row]), \
                 patch("traveler_assistant.order_workflow.parse_board_identity", return_value=("F100", "CS005-KITCHEN")), \
                 patch("traveler_assistant.order_workflow.preview_order", return_value=preview):
                result = sync_order_index(config)

            self.assertEqual(result["orders"][0]["order_id"], "CS005")
            self.assertEqual(result["orders"][0]["optimized_count"], 0)
            self.assertEqual(result["orders"][0]["stage"], "已拆单待优化")

    def test_cut_to_size_optimization_artifact_marks_status_when_material_is_absent(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "server" / "Optimized Orders")
            folder = config.source_root.parent / "CUT TO SIZE" / "CS005"
            report = folder / "Report" / "pp-板材清单.xlsx"
            artifact = folder / "New Nesting" / "Optimize file" / "layout file" / "nesting_result.xml"
            report.parent.mkdir(parents=True)
            artifact.parent.mkdir(parents=True)
            report.write_bytes(b"placeholder")
            artifact.write_text('<Nesting><BoardControl OrderID="F100" /></Nesting>', encoding="utf-8")
            aimes_row = {
                "factory_order": "F100",
                "factory_name": "CS005-KITCHEN",
                "sales_order_name": "CS005",
                "split_time": "2026-08-10T08:30:00",
            }
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=[aimes_row]), \
                 patch("traveler_assistant.order_workflow.parse_board_identity", return_value=("F100", "CS005-KITCHEN")), \
                 patch("traveler_assistant.order_workflow.preview_order", side_effect=RuleError("missing_material", "缺少 material")):
                result = sync_order_index(config)

            order = result["orders"][0]
            self.assertEqual(order["stage"], "已优化")
            self.assertEqual(order["optimized_count"], 1)
            self.assertEqual(order["validation_status"], "待校验")
            self.assertEqual(order["material_status"], "待校验")
            self.assertTrue(order["optimization_completed_at"])
            self.assertEqual(order["factories"][0]["optimization_source_path"], str(artifact))

    def test_order_is_optimized_only_after_every_active_factory_has_aicnc_evidence(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "server" / "Optimized Orders")
            folder = config.source_root / "PP0099"
            first = folder / "Kitchen" / "Optimize file" / "layout file" / "nesting_result.xml"
            first.parent.mkdir(parents=True)
            first.write_text('<Nesting><BoardControl OrderID="F100" /></Nesting>', encoding="utf-8")
            rows = [
                {"factory_order": "F100", "factory_name": "PP0099-KITCHEN", "sales_order_name": "PP0099", "split_time": "2026-08-10T08:30:00"},
                {"factory_order": "F200", "factory_name": "PP0099-VANITY", "sales_order_name": "PP0099", "split_time": "2026-08-11T08:30:00"},
            ]
            preview = SimpleNamespace(factories=[])
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=rows), \
                 patch("traveler_assistant.order_workflow.preview_order", return_value=preview):
                partial = sync_order_index(config)
            self.assertEqual(partial["orders"][0]["stage"], "部分优化")
            self.assertEqual(partial["orders"][0]["optimized_count"], 1)

            second = folder / "Vanity" / "Optimize file" / "layout file" / "nesting_result.xml"
            second.parent.mkdir(parents=True)
            second.write_text('<Nesting><BoardControl OrderID="F200" /></Nesting>', encoding="utf-8")
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=rows), \
                 patch("traveler_assistant.order_workflow.preview_order", return_value=preview):
                complete = sync_order_index(config)
            self.assertEqual(complete["orders"][0]["stage"], "已优化")
            self.assertEqual(complete["orders"][0]["optimized_count"], 2)
            self.assertTrue(complete["orders"][0]["optimization_completed_at"])

            original_mtime = first.stat().st_mtime
            first.write_text(
                '<Nesting><BoardControl OrderID="F100" /><BoardControl OrderID="F100" /></Nesting>',
                encoding="utf-8",
            )
            os.utime(first, (original_mtime + 60, original_mtime + 60))
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=rows), \
                 patch("traveler_assistant.order_workflow.preview_order", return_value=preview):
                refreshed = sync_order_index(config)
            f100 = next(item for item in refreshed["orders"][0]["factories"] if item["factory_order"] == "F100")
            self.assertNotEqual(
                f100["optimization_first_completed_at"],
                f100["optimization_latest_completed_at"],
            )
            evidence_connection = sqlite3.connect(config.workflow_database)
            evidence = evidence_connection.execute(
                "select count(*) from optimization_artifacts where factory_order='F100'"
            ).fetchone()[0]
            evidence_connection.close()
            self.assertEqual(evidence, 2)

            rows.append({"factory_order": "F300", "factory_name": "PP0099-OFFICE", "sales_order_name": "PP0099", "split_time": "2026-08-12T08:30:00"})
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=rows), \
                 patch("traveler_assistant.order_workflow.preview_order", return_value=preview):
                expanded = sync_order_index(config)
            self.assertEqual(expanded["orders"][0]["stage"], "部分优化")
            self.assertEqual(expanded["orders"][0]["optimized_count"], 2)

    def test_optimization_marker_scan_uses_known_paths_without_recursive_file_walk(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / "PP0099"
            kitchen_root = folder / "Kitchen" / "Optimize file"
            kitchen_root.mkdir(parents=True)
            (kitchen_root / "Optimize file.xml").write_text("<Optimize />", encoding="utf-8")
            (kitchen_root / "layout file").mkdir()
            nesting = kitchen_root / "layout file" / "nesting_result.xml"
            nesting.write_text("<Nesting />", encoding="utf-8")

            unexpected = folder / "Kitchen" / "output" / "Optimize file" / "nesting_result.xml"
            unexpected.parent.mkdir(parents=True)
            unexpected.write_text("<Unexpected />", encoding="utf-8")

            with patch.object(Path, "rglob", side_effect=AssertionError("不应递归遍历文件")):
                artifacts = _optimization_artifacts(folder)
                monitor_files = _server_optimization_monitor_files(folder)

            self.assertEqual(artifacts, [kitchen_root / "Optimize file.xml", nesting])
            self.assertEqual(
                [path for path, _ in monitor_files],
                [kitchen_root / "Optimize file.xml", nesting],
            )

    def test_folder_timing_total_equals_final_file_timing_sum(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / "PP0099"
            optimize_root = folder / "Kitchen" / "Optimize file"
            optimize_root.mkdir(parents=True)
            (optimize_root / "Optimize file.xml").write_text("<Optimize />", encoding="utf-8")
            (optimize_root / "layout file").mkdir()
            (optimize_root / "layout file" / "nesting_result.xml").write_text(
                "<Nesting />", encoding="utf-8"
            )

            _, timing = _server_snapshot_folder(folder)

            file_total = round(sum(item["duration_seconds"] for item in timing["files"]), 6)
            self.assertEqual(timing["duration_seconds"], file_total)
            self.assertGreaterEqual(timing["wall_duration_seconds"], timing["duration_seconds"])

    def test_visible_server_scan_refreshes_aicnc_evidence_and_returns_updated_orders(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "server" / "Optimized Orders")
            config.prepare_storage()
            folder = config.source_root / "PP0099"
            artifact = folder / "Kitchen" / "Optimize file" / "layout file" / "nesting_result.xml"
            artifact.parent.mkdir(parents=True)
            artifact.write_text('<Nesting><BoardControl OrderID="F100" /></Nesting>', encoding="utf-8")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP0099", source_folder=str(folder), validation_status="正常")
            store.upsert_factory(
                "F100",
                order_id="PP0099",
                factory_name="PP0099-KITCHEN",
                sales_order_name="PP0099",
                split_time="2026-08-10T08:30:00",
                name_source="AIMES",
                ownership_status="已确认",
            )
            store.commit()
            store.close()

            result = scan_server_changes(config)

            order = next(item for item in result["orders"] if item["order_id"] == "PP0099")
            self.assertEqual(order["stage"], "已优化")
            self.assertEqual(order["optimized_count"], 1)
            self.assertEqual(result["server"]["scan_stats"]["optimization_artifact_refresh_count"], 1)
            self.assertEqual(
                [stage["stage"] for stage in result["operation_timing"]["stages"]],
                ["server_metadata", "optimization_evidence", "scan_finalize"],
            )

    def test_exact_standard_order_folder_wins_over_mixed_factory_report_folder(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "server" / "Optimized Orders")
            exact_folder = config.source_root / "PP0035"
            mixed_folder = config.source_root / "PP0035-2"
            (mixed_folder / "Report").mkdir(parents=True)
            (mixed_folder / "Report" / "板材清单.xlsx").write_bytes(b"board")
            exact_folder.mkdir(parents=True)
            rows = [
                {
                    "factory_order": "F100",
                    "factory_name": "PP0035-OFFICE",
                    "sales_order_name": "PP0035",
                    "split_time": "2026-08-10T08:30:00",
                },
                {
                    "factory_order": "F101",
                    "factory_name": "PP0035-2-MASTER",
                    "sales_order_name": "PP0035-2",
                    "split_time": "2026-08-10T08:31:00",
                },
            ]
            preview = SimpleNamespace(factories=[])
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=rows), \
                 patch("traveler_assistant.order_workflow.parse_board_identity", return_value=("F100", "PP0035-OFFICE")), \
                 patch("traveler_assistant.order_workflow.preview_order", return_value=preview):
                sync_order_index(config)

            store = OrderIndexStore(config.workflow_database)
            source_folders = {
                row[0]: row[1]
                for row in store.connection.execute(
                    "select order_id, source_folder from orders where order_id in ('PP0035', 'PP0035-2')"
                ).fetchall()
            }
            store.close()
            self.assertEqual(source_folders["PP0035"], str(exact_folder))
            self.assertEqual(source_folders["PP0035-2"], str(mixed_folder))

    def test_cut_to_size_fittings_are_not_persisted_as_hardware(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "server" / "Optimized Orders")
            folder = config.source_root.parent / "CUT TO SIZE" / "CS999"
            folder.mkdir(parents=True)
            (folder / "CS999 materials.xlsx").write_bytes(b"material")
            (folder / "Fittingslist.xlsx").write_bytes(b"fittings")
            config.prepare_storage()
            fitting = SimpleNamespace(code="M0001", name="Test hardware", size="", unit="Piece", quantity=1)
            preview = SimpleNamespace(factories=[])
            row = {
                "factory_order": "F999",
                "factory_name": "CS999-KITCHEN",
                "sales_order_name": "CS999",
                "split_time": "2026-08-10T08:30:00",
            }
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=[row]), \
                 patch("traveler_assistant.order_workflow.parse_order_materials", return_value=("Sheet1", [], {})), \
                 patch("traveler_assistant.order_workflow.parse_fittings_groups", return_value=[("F999", [fitting])]), \
                 patch("traveler_assistant.order_workflow.preview_order", return_value=preview):
                sync_order_index(config)

            connection = __import__("sqlite3").connect(config.workflow_database)
            count = connection.execute(
                "select count(*) from hardware_items where order_id='CS999'"
            ).fetchone()[0]
            connection.close()
            self.assertEqual(count, 0)

    def test_incremental_sync_reuses_unchanged_server_report_and_rechecks_changes(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            folder = config.source_root / "PP9999"
            report = folder / "Report" / "板材清单.xlsx"
            report.parent.mkdir(parents=True)
            report.write_bytes(b"placeholder")
            aimes_row = {
                "factory_order": "F100",
                "factory_name": "PP9999-KITCHEN",
                "sales_order_name": "PP9999",
                "split_time": "2026-08-10T08:30:00",
            }
            preview = SimpleNamespace(factories=[])
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=[aimes_row]), \
                 patch("traveler_assistant.order_workflow.parse_board_identity", return_value=("F100", "PP9999-KITCHEN")) as parse, \
                 patch("traveler_assistant.order_workflow.preview_order", return_value=preview) as preview_order:
                first = sync_order_index(config)
                first_parse_count = parse.call_count
                second = sync_order_index(config)

                self.assertGreaterEqual(first_parse_count, 1)
                self.assertEqual(parse.call_count, first_parse_count)
                self.assertEqual(preview_order.call_count, 1)
                self.assertEqual(first["index_stats"]["parsed_report_count"], 1)
                self.assertEqual(second["index_stats"]["reused_report_count"], 1)
                self.assertEqual(second["index_stats"]["parsed_report_count"], 0)
                self.assertEqual(second["index_stats"]["validated_order_count"], 0)
                self.assertEqual(second["orders"][0]["optimized_count"], 0)

                report.write_bytes(b"changed")
                third = sync_order_index(config)

            self.assertGreater(parse.call_count, first_parse_count)
            self.assertEqual(preview_order.call_count, 2)
            self.assertEqual(third["index_stats"]["parsed_report_count"], 1)
            self.assertEqual(third["index_stats"]["validated_order_count"], 1)

    def test_server_sync_preserves_hardware_when_report_is_unchanged_or_mapping_fails(self):
        """A scan must not erase the last valid hardware projection."""
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            fittings = folder / "Report" / "Fittingslist.xlsx"
            fittings.parent.mkdir(parents=True)
            from tests.test_order_workflow import make_fittings

            make_fittings(fittings, [("F100", 2)])
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP9999", validation_status="正常", source_folder=str(folder))
            store.upsert_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                name_source="AIMES",
                ownership_status="已确认",
                has_hardware=True,
            )
            store.upsert_source_file(
                fittings,
                source_folder=folder,
                kind="fittings",
                order_id="PP9999",
                factory_order="F100",
                changed_at="2026-08-31T10:00:00",
            )
            store.connection.execute(
                """insert into hardware_items(
                    order_id, factory_order, product_code, source_code, name, spec,
                    quantity, unit, source_type, source_path, active, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    "PP9999", "F100", "M1001", "71T950A", "TestFullHinge", "",
                    2, "Piece", "aicnc", str(fittings), 1, "2026-08-31T10:00:00",
                ),
            )
            store.commit()
            store.close()

            unchanged = sync_order_index(
                config,
                refresh_outbound_statuses=False,
                reconcile_outbound=False,
            )
            self.assertGreaterEqual(unchanged["index_stats"]["reused_report_count"], 1)

            store = OrderIndexStore(config.workflow_database)
            try:
                self.assertEqual(
                    store.connection.execute(
                        "select quantity from hardware_items where order_id='PP9999'"
                    ).fetchone()[0],
                    2.0,
                )
            finally:
                store.close()

            make_fittings(fittings, [("F100", 5)])
            unresolved = {
                "missing": [{
                    "name": "TestFullHinge",
                    "source_code": "71T950A",
                    "quantity": 5,
                    "message": "找不到有效库存 SKU",
                }],
                "ignored": [],
                "outbound": [],
                "accepted": [],
            }
            with patch(
                "traveler_assistant.inventory.resolve_inventory_items",
                return_value=unresolved,
            ):
                failed = sync_order_index(
                    config,
                    refresh_outbound_statuses=False,
                    reconcile_outbound=False,
                )

            self.assertGreaterEqual(failed["index_stats"]["parsed_report_count"], 1)
            store = OrderIndexStore(config.workflow_database)
            try:
                self.assertEqual(
                    store.connection.execute(
                        "select quantity from hardware_items where order_id='PP9999'"
                    ).fetchone()[0],
                    2.0,
                )
            finally:
                store.close()

    def test_server_sync_replaces_hardware_by_factory_order_across_source_paths(self):
        """A re-optimization path must replace, not add to, factory hardware."""
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            first = folder / "Report" / "Fittingslist-first.xlsx"
            second = folder / "second-optimization" / "Report" / "Fittingslist-second.xlsx"
            third = folder / "third-optimization" / "Report" / "Fittingslist-third.xlsx"
            from tests.test_order_workflow import make_fittings

            first.parent.mkdir(parents=True)
            make_fittings(first, [("F100", 2)])
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP9999", validation_status="正常", source_folder=str(folder))
            store.upsert_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                name_source="AIMES",
                ownership_status="已确认",
                has_hardware=True,
            )
            store.commit()
            store.close()

            sync_order_index(
                config,
                refresh_outbound_statuses=False,
                reconcile_outbound=False,
            )
            second.parent.mkdir(parents=True)
            make_fittings(second, [("F100", 2)])

            sync_order_index(
                config,
                refresh_outbound_statuses=False,
                reconcile_outbound=False,
            )

            store = OrderIndexStore(config.workflow_database)
            try:
                unchanged_rows = store.connection.execute(
                    """
                    select quantity, source_path
                    from hardware_items
                    where factory_order='F100' and source_type='aicnc' and active=1
                    order by id
                    """
                ).fetchall()
            finally:
                store.close()

            self.assertEqual(unchanged_rows, [(2.0, str(first))])

            third.parent.mkdir(parents=True)
            make_fittings(third, [("F100", 5)])
            sync_order_index(
                config,
                refresh_outbound_statuses=False,
                reconcile_outbound=False,
            )

            store = OrderIndexStore(config.workflow_database)
            try:
                rows = store.connection.execute(
                    """
                    select quantity, source_path
                    from hardware_items
                    where factory_order='F100' and source_type='aicnc' and active=1
                    order by id
                    """
                ).fetchall()
            finally:
                store.close()

            self.assertEqual(rows, [(2.0, str(first))])
            from traveler_assistant.fittings import select_latest_fittings
            from traveler_assistant.report_read_context import report_read_session
            with self.assertRaises(RuleError) as conflict:
                select_latest_fittings([first, second, third])
            option = next(item for item in conflict.exception.context["conflicts"][0]["candidates"] if item["path"] == str(third.resolve()))
            with report_read_session({"F100": option["id"]}):
                sync_order_index(config, full_refresh=True, refresh_outbound_statuses=False, reconcile_outbound=False)
            store = OrderIndexStore(config.workflow_database)
            try:
                rows = store.connection.execute("select quantity, source_path from hardware_items where factory_order='F100' and source_type='aicnc' and active=1").fetchall()
                self.assertEqual(rows, [(5.0, str(third))])
            finally:
                store.close()

    def test_server_snapshot_uses_bounded_read_only_workers_and_preserves_records(self):
        class TrackingExecutor:
            def __init__(self, max_workers):
                self.max_workers = max_workers

            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc_value, traceback):
                return False

            def map(self, function, folders):
                return [function(folder) for folder in folders]

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            for name in ("PP0035", "PP0046"):
                folder = config.source_root / name
                folder.mkdir(parents=True)
                (folder / f"{name} materials.xlsx").write_bytes(b"material")

            executors = []

            def make_executor(max_workers):
                executor = TrackingExecutor(max_workers)
                executors.append(executor)
                return executor

            with patch(
                "traveler_assistant.order_index._orders_requiring_server_scan",
                return_value={"PP0035", "PP0046"},
            ), patch(
                "traveler_assistant.order_index.ThreadPoolExecutor",
                side_effect=make_executor,
            ):
                store = OrderIndexStore(config.workflow_database)
                root_path, snapshot = _server_snapshot(config, store)
                store.close()

            self.assertEqual(root_path, config.source_root)
            self.assertEqual(len(executors), 1)
            self.assertEqual(executors[0].max_workers, 2)
            self.assertEqual(
                sorted(Path(path).name for path, item in snapshot.items() if item["kind"] == "folder"),
                ["PP0035", "PP0046"],
            )
            self.assertEqual(
                sorted(Path(path).name for path, item in snapshot.items() if item["kind"] != "folder"),
                [],
            )

    def test_sync_index_reuses_scan_snapshot_and_reports_phase_durations(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            folder = config.source_root / "PP9999"
            report = folder / "Report" / "板材清单.xlsx"
            report.parent.mkdir(parents=True)
            report.write_bytes(b"placeholder")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            store.commit()
            store.close()

            scan = scan_server_changes(config)
            snapshot_path = Path(scan["server"]["snapshot_path"])
            preview = SimpleNamespace(factories=[])
            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=[]), \
                 patch("traveler_assistant.order_index._server_folders_for_sync", side_effect=AssertionError("snapshot should avoid folder discovery")), \
                 patch("traveler_assistant.order_workflow.parse_board_identity", return_value=("F100", "PP9999-KITCHEN")), \
                 patch("traveler_assistant.order_workflow.preview_order", return_value=preview):
                result = sync_order_index(config, server_snapshot_path=snapshot_path)

            self.assertTrue(result["index_stats"]["server_snapshot_reused"])
            self.assertEqual(result["index_stats"]["server_snapshot_entry_count"], 1)
            self.assertIn("server_metadata_and_report_sync", result["index_stats"]["phase_durations"])
            self.assertIn("索引阶段耗时：", result["operation_trace"]["sync"][-1])

    def test_prebaseline_temporary_folder_is_excluded_and_stale_pending_cleared(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                server_scan_baseline_at="2099-01-01T00:00:00-08:00",
            )
            folder = config.source_root / "pp001"
            folder.mkdir(parents=True)
            store = OrderIndexStore(config.workflow_database)
            store.upsert_source_file(
                folder,
                source_folder=folder,
                kind="folder",
                changed_at="2026-08-12T09:00:00",
            )
            store.upsert_active_issue(
                issue_key=f"temporary_processing:{folder}",
                kind="temporary_processing",
                path=str(folder),
                message="旧的临时文件夹问题",
            )
            store.commit()
            store.close()

            result = scan_server_changes(config)

            self.assertEqual(result["server"]["changes"], [])
            self.assertEqual(result["current_issues"], [])
            reopened = OrderIndexStore(config.workflow_database)
            self.assertEqual(
                reopened.connection.execute("select count(*) from source_files").fetchone()[0],
                0,
            )
            self.assertEqual(reopened.active_issues(), [])
            reopened.close()

    def test_named_mixed_folder_is_not_marked_as_temporary(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            folder = config.source_root / "CS003 PP0047"
            folder.mkdir(parents=True)

            self.assertTrue(_is_mixed_order_folder(folder))
            changes = scan_server_changes(config)["server"]["changes"]
            folder_change = next(item for item in changes if item["path"] == str(folder))
            self.assertFalse(folder_change["manual_only"])
            self.assertTrue(folder_change["mixed_order"])
            self.assertEqual(folder_change["order_id"], "CS003、PP0047")
            self.assertIn("混单文件夹", folder_change["message"])

    def test_fully_shipped_mixed_folder_is_watched_then_reopened_by_aimes(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            folder = config.source_root / "CS003 PP0047"
            folder.mkdir(parents=True)
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100", order_id="CS003", factory_name="CS003 KITCHEN",
                sales_order_name="CS003", split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            store.upsert_aimes_factory(
                "F101", order_id="PP0047", factory_name="PP0047 KITCHEN",
                sales_order_name="PP0047", split_time="2026-08-10T08:31:00",
                seen_at="2026-08-10T08:31:00",
            )
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[
                    {"order_id": "CS003", "remark": "CS003 KITCHEN", "status": "已出库", "document_number": "OUT-CS003"},
                    {"order_id": "PP0047", "remark": "PP0047 KITCHEN", "status": "已出库", "document_number": "OUT-PP0047"},
                ],
            ):
                watched = scan_server_changes(config)["server"]
                self.assertTrue(watched["changed"])
                self.assertEqual(watched["scan_stats"]["order_folder_count"], 1)
                self.assertTrue(any(
                    item["change_type"] == "missing_report" for item in watched["changes"]
                ))

                reopened = OrderIndexStore(config.workflow_database)
                reopened.upsert_aimes_factory(
                    "F102", order_id="PP0047", factory_name="PP0047 VANITY",
                    sales_order_name="PP0047", split_time="2026-08-17T08:31:00",
                    seen_at="2026-08-17T08:31:00",
                )
                reopened.commit()
                reopened.close()
                reopened_scan = scan_server_changes(config)["server"]

        self.assertTrue(reopened_scan["changed"])
        self.assertEqual(reopened_scan["scan_stats"]["order_folder_count"], 1)
        self.assertEqual(reopened_scan["changes"][0]["path"], str(folder))

    def test_shipped_server_order_becomes_permanent_after_seven_day_watch(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            folder = config.source_root / "PP9999"
            folder.mkdir(parents=True)
            (folder / "PP9999 materials.xlsx").write_bytes(b"material")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100", order_id="PP9999", factory_name="PP9999 KITCHEN",
                sales_order_name="PP9999", split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            store.commit()
            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[{
                    "order_id": "PP9999", "remark": "PP9999 KITCHEN",
                    "status": "已出库", "document_number": "QTCK001",
                }],
            ):
                _, first = _server_snapshot(config, store)
                self.assertIn(str(folder), first)
                self.assertEqual(store.server_scan_policy("PP9999")["policy"], "watching")
                store.connection.execute(
                    "update orders set server_scan_watch_until = ? where order_id = ?",
                    ("2026-08-01T00:00:00", "PP9999"),
                )
                store.commit()
                _, expired = _server_snapshot(config, store)
                self.assertNotIn(str(folder), expired)
                self.assertEqual(store.server_scan_policy("PP9999")["policy"], "permanent")
            store.close()

    def test_initial_date_orders_are_marked_shipped_without_fabricating_documents(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            old_folder = config.source_root / "PP0010"
            old_folder.mkdir(parents=True)
            old_mtime = datetime.fromisoformat("2026-07-21T12:00:00").timestamp()
            os.utime(old_folder, (old_mtime, old_mtime))
            factory_folder = config.source_root / "PP0011"
            factory_folder.mkdir(parents=True)
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F2605120103", order_id="PP0011", factory_name="Kitchen",
                sales_order_name="PP0011", split_time="2026-05-12T02:59:55",
                seen_at="2026-08-31T10:00:00",
            )
            store.upsert_order("PP0011", source_folder=str(factory_folder))
            store.connection.execute(
                "update factory_orders set outbound_status = '未查询', outbound_document = '' where factory_order = ?",
                ("F2605120103",),
            )
            store.connection.execute(
                "update orders set validation_status = '数据异常', validation_message = '旧的缺少 material 提醒' where order_id = ?",
                ("PP0011",),
            )
            store.commit()

            updated = _mark_initial_orders_shipped(config, store)

            shipped_order = store.connection.execute(
                "select stage, validation_status, validation_message from orders where order_id = ?", ("PP0011",)
            ).fetchone()
            shipped_factory = store.connection.execute(
                "select outbound_status, outbound_document, outbound_mode from factory_orders where factory_order = ?",
                ("F2605120103",),
            ).fetchone()
            store.close()

        self.assertEqual(updated, 1)
        self.assertEqual(shipped_order, ("已出货", "待校验", ""))
        self.assertEqual(shipped_factory, ("已出库", "", "historical_initial_date"))

    def test_mark_temporary_folder_manual_starts_three_day_xml_watch(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            folder = config.source_root / "manual-temporary"
            xml_root = folder / "New Nesting" / "Optimize file"
            xml_root.mkdir(parents=True)
            (xml_root / "Optimize file.xml").write_text("<Optimize />", encoding="utf-8")
            (xml_root / "layout file").mkdir()
            (xml_root / "layout file" / "nesting_result.xml").write_text(
                '<Nesting><BoardControl OrderID="F100" /></Nesting>', encoding="utf-8"
            )
            (folder / "manual materials.xlsx").write_bytes(b"material")

            result = mark_temporary_folder_manual(config, folder)

            self.assertTrue(result["ok"])
            self.assertEqual(result["outbound_status"], "已出库")
            self.assertEqual(result["server_scan_policy"], "watching")
            store = OrderIndexStore(config.workflow_database)
            temporary = store.temporary_order(str(folder))
            xml_count = store.connection.execute(
                "select count(*) from server_scan_xml_state where source_folder = ?",
                (str(folder),),
            ).fetchone()[0]
            self.assertEqual(temporary["processing_status"], "已人工处理")
            self.assertEqual(temporary["outbound_status"], "已出库")
            self.assertEqual(temporary["server_scan_policy"], "watching")
            self.assertEqual(xml_count, 2)
            store.close()
            self.assertEqual(scan_server_changes(config)["server"]["changes"], [])

    def test_removed_server_folder_ignore_table_is_cleaned_on_open(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            config.state_dir.mkdir(parents=True)
            connection = sqlite3.connect(config.workflow_database)
            connection.execute(
                "create table ignored_server_folders(path text primary key, ignored_at text)"
            )
            connection.execute(
                "insert into ignored_server_folders(path, ignored_at) values(?, ?)",
                ("/old/server/folder", "2026-08-01T00:00:00"),
            )
            connection.commit()
            connection.close()

            store = OrderIndexStore(config.workflow_database)
            exists = store.connection.execute(
                "select 1 from sqlite_master where type = 'table' and name = 'ignored_server_folders'"
            ).fetchone()
            store.close()
            self.assertIsNone(exists)

    def test_reportless_mixed_folder_requires_review(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            folder = config.source_root / "CS003 PP0047"
            folder.mkdir(parents=True)

            first = scan_server_changes(config)
            self.assertTrue(any(
                item["change_type"] == "missing_report" and item["path"] == str(folder)
                for item in first["server"]["changes"]
            ))
            issue = next(item for item in first["current_issues"] if item["kind"] == "server_missing_report")
            self.assertIn("缺少可识别的报表", issue["message"])

    def test_processed_temporary_folder_uses_three_day_xml_watch_then_is_permanent(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            folder = config.source_root / "temporary-production"
            xml = folder / "New Nesting" / "Optimize file" / "nesting_result.xml"
            xml.parent.mkdir(parents=True)
            xml.write_text('<Nesting><BoardControl OrderID="F100" /></Nesting>', encoding="utf-8")
            (folder / "material.xlsx").write_bytes(b"placeholder")
            processed_at = datetime.now().isoformat(timespec="seconds")

            store = OrderIndexStore(config.workflow_database)
            store.upsert_temporary_order(
                temporary_id="TMP:watch",
                folder_name=folder.name,
                source_folder=str(folder),
                folder_created_at=0,
                content_fingerprint="workbooks",
                processing_status="Traveler 已生成",
                outbound_status="已出库",
                processed_at=processed_at,
            )
            store.save_server_scan_xml_baseline(
                [folder],
                [{
                    "path": str(xml),
                    "source_folder": str(folder),
                    "kind": "optimization_result",
                    "order_id": "",
                    "modified_at": int(xml.stat().st_mtime_ns // 1_000_000),
                }],
                observed_at=processed_at,
            )
            store.commit()
            store.close()

            watched = scan_server_changes(config)["server"]
            self.assertFalse(watched["changed"])
            store = OrderIndexStore(config.workflow_database)
            self.assertEqual(
                store.temporary_order(str(folder))["server_scan_policy"],
                "watching",
            )
            store.connection.execute(
                "update temporary_orders set server_scan_watch_until = ? where source_folder = ?",
                ("2000-01-01T00:00:00", str(folder)),
            )
            store.commit()
            store.close()

            expired = scan_server_changes(config)["server"]
            self.assertFalse(expired["changed"])
            store = OrderIndexStore(config.workflow_database)
            self.assertEqual(
                store.temporary_order(str(folder))["server_scan_policy"],
                "permanent",
            )
            store.close()

    def test_failed_temporary_processing_remains_in_pending_server_changes(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            folder = config.source_root / "temporary-production"
            folder.mkdir(parents=True)
            (folder / "material.xlsx").write_bytes(b"not-a-workbook")

            result = process_server_changes(config)

            self.assertTrue(result["temporary_processing"]["failed"])
            pending_paths = {item["path"] for item in result["pending_server_changes"]}
            self.assertIn(str(folder), pending_paths)
            self.assertIn(str(folder / "material.xlsx"), pending_paths)
            store = OrderIndexStore(config.workflow_database)
            self.assertTrue(any(item["kind"] == "temporary_processing" for item in store.active_issues()))
            self.assertEqual(
                store.connection.execute("select count(*) from source_files").fetchone()[0],
                2,
            )
            store.close()

            repeated = scan_server_changes(config)["server"]
            self.assertTrue(repeated["changed"])
            self.assertTrue(all(
                item["change_type"] == "processing_failed"
                for item in repeated["changes"]
            ))
            self.assertFalse(any(
                "首次发现" in item["message"] or "新增" in item["message"]
                for item in repeated["changes"]
            ))

    def test_temporary_fittings_report_is_deferred_until_user_approves_processing(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            folder = config.source_root / "temporary-production"
            fittings = folder / "Report" / "Fittingslist.xlsx"
            fittings.parent.mkdir(parents=True)
            fittings.write_bytes(b"not-a-workbook")

            result = sync_order_index(config)

            self.assertFalse(any(
                issue["kind"] == "report_error" and issue["path"] == str(fittings)
                for issue in result["current_issues"]
            ))

    def test_manual_temporary_outbound_records_server_baseline_case_insensitively(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            folder = config.source_root / "inserthood cabinet OLD CNC"
            report = folder / "Report"
            report.mkdir(parents=True)
            (folder / "INSERTHOOD CABINET OLD CNC materials.xlsx").write_bytes(b"material")
            (report / "Fittingslist.xlsx").write_bytes(b"fittings")
            (report / "pp-板材清单-new.xlsx").write_bytes(b"board")
            traveler = config.order_root / "INSERTHOOD CABINET OLD CNC" / (
                "Work Order Traveler(INSERTHOOD CABINET OLD CNC).xlsx"
            )

            record_temporary_outbound(
                config,
                traveler,
                {"saved": True, "results": [{"documentNumber": "QTCK-001"}]},
            )

            store = OrderIndexStore(config.workflow_database)
            rows = store.connection.execute(
                "select source_folder, outbound_status, outbound_document from temporary_orders"
            ).fetchall()
            source_paths = {
                row[0]
                for row in store.connection.execute("select path from source_files").fetchall()
            }
            store.close()
            self.assertEqual(rows, [(str(folder), "已出库", "QTCK-001")])
            self.assertEqual(len(source_paths), 4)
            self.assertFalse(scan_server_changes(config)["server"]["changed"])

    def test_shipped_temporary_folder_is_skipped_without_report_rescan(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            folder = config.source_root / "temporary-production"
            report = folder / "Report" / "pp-板材清单.xlsx"
            report.parent.mkdir(parents=True)
            report.write_bytes(b"board")
            xml = folder / "New Nesting" / "Optimize file" / "nesting_result.xml"
            xml.parent.mkdir(parents=True)
            xml.write_text('<Nesting><BoardControl OrderID="F100" /></Nesting>', encoding="utf-8")
            traveler = config.order_root / "TEMPORARY-PRODUCTION" / "Work Order Traveler(TEMPORARY-PRODUCTION).xlsx"
            record_temporary_outbound(
                config,
                traveler,
                {"saved": True, "results": [{"documentNumber": "OUT-001"}]},
            )

            with patch("traveler_assistant.order_index._report_files", wraps=_report_files) as report_files:
                result = scan_server_changes(config)

            self.assertFalse(result["server"]["changed"])
            store = OrderIndexStore(config.workflow_database)
            self.assertEqual(
                store.temporary_order(str(folder))["server_scan_policy"],
                "watching",
            )
            store.close()
            self.assertFalse(any(
                call.args and str(folder) == str(call.args[0])
                for call in report_files.call_args_list
            ))

    def test_old_temporary_folder_is_filtered_before_report_rescan(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                server_scan_baseline_at="2099-01-01T00:00:00-08:00",
            )
            folder = config.source_root / "old-temporary"
            report = folder / "Report" / "pp-板材清单.xlsx"
            report.parent.mkdir(parents=True)
            report.write_bytes(b"board")

            with patch("traveler_assistant.order_index._report_files", wraps=_report_files) as report_files:
                result = scan_server_changes(config)

            self.assertFalse(result["server"]["changed"])
            self.assertFalse(any(
                call.args and str(folder) == str(call.args[0])
                for call in report_files.call_args_list
            ))

    def test_temporary_processing_generates_material_traveler_and_outbounds(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            folder = config.source_root / "temporary-production"
            report = folder / "Report" / "pp-板材清单.xlsx"
            fittings = folder / "Report" / "Fittingslist.xlsx"
            from tests.test_order_workflow import make_board_material_report, make_fittings
            make_board_material_report(report, factory="F100", name="PP9999-KITCHEN")
            make_fittings(fittings, [("F100", 2)])

            with patch(
                "traveler_assistant.inventory.run_jdy",
                return_value={"saved": True, "results": [{"documentNumber": "OUT-1"}]},
            ) as outbound:
                result = process_server_changes(config)

            self.assertEqual(result["temporary_processing"]["failed"], [])
            self.assertEqual(result["temporary_processing"]["succeeded"][0]["order_ids"], ["PP9999"])
            self.assertTrue((folder / "PP9999 materials.xlsx").is_file())
            self.assertTrue((root / "orders" / "PP9999" / "Work Order Traveler(PP9999).xlsx").is_file())
            outbound.assert_called_once()
            self.assertEqual(result["pending_server_changes"], [])

    def test_temporary_processing_can_skip_hardware_in_traveler_and_outbound(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            folder = config.source_root / "temporary-production"
            report = folder / "Report" / "pp-板材清单.xlsx"
            fittings = folder / "Report" / "Fittingslist.xlsx"
            from tests.test_order_workflow import make_board_material_report, make_fittings
            make_board_material_report(report, factory="F100", name="PP9999-KITCHEN")
            make_fittings(fittings, [("F100", 2)])

            with patch(
                "traveler_assistant.inventory.run_jdy",
                return_value={"saved": True, "results": [{"documentNumber": "OUT-2"}]},
            ) as outbound:
                result = process_server_changes(config, include_hardware=False)

            self.assertEqual(result["temporary_processing"]["failed"], [])
            traveler = root / "orders" / "PP9999" / "Work Order Traveler(PP9999).xlsx"
            from openpyxl import load_workbook
            picking = load_workbook(traveler, data_only=False)["Picking List"]
            self.assertFalse(any(
                picking.cell(row, 3).value == "Hinge"
                for row in range(1, picking.max_row + 1)
            ))
            outbound.assert_called_once()

    def test_temporary_folder_without_aimes_identity_uses_folder_name_everywhere(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            folder = config.source_root / "B12"
            report = folder / "Report" / "pp-板材清单.xlsx"
            fittings = folder / "Report" / "Fittingslist.xlsx"
            from tests.test_order_workflow import make_board_material_report, make_fittings
            make_board_material_report(report, factory="", name="")
            make_fittings(fittings, [("", 2)])

            with patch(
                "traveler_assistant.inventory.run_jdy",
                return_value={"saved": True, "results": [{"documentNumber": "OUT-B12"}]},
            ) as outbound:
                result = process_server_changes(config)

            self.assertEqual(result["temporary_processing"]["failed"], [])
            self.assertEqual(result["temporary_processing"]["succeeded"][0]["order_ids"], ["B12"])
            traveler = root / "orders" / "B12" / "Work Order Traveler(B12).xlsx"
            self.assertTrue(traveler.is_file())
            workbook = __import__("openpyxl").load_workbook(traveler, data_only=False)
            self.assertEqual(workbook["WorkOrderTraveler"]["B5"].value, "B12")
            self.assertEqual(workbook["Usage List"]["B1"].value, "B12")
            picking = workbook["Picking List"]
            self.assertTrue(any(
                picking.cell(row, 4).value == "B12"
                for row in range(1, picking.max_row + 1)
            ))
            outbound.assert_called_once()
            store = OrderIndexStore(config.workflow_database)
            record = store.temporary_order(str(folder))
            self.assertEqual(record["folder_name"], "B12")
            self.assertEqual(record["processing_status"], "Traveler 已生成")
            self.assertEqual(record["outbound_status"], "已出库")
            self.assertEqual(record["outbound_document"], "OUT-B12")
            store.close()

    def test_temporary_outbound_is_not_repeated_when_folder_content_is_unchanged(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            folder = config.source_root / "B12"
            report = folder / "Report" / "pp-板材清单.xlsx"
            from tests.test_order_workflow import make_board_material_report
            make_board_material_report(report, factory="", name="")

            with patch(
                "traveler_assistant.inventory.run_jdy",
                return_value={"saved": True, "results": [{"documentNumber": "OUT-B12"}]},
            ) as outbound:
                first = process_server_changes(config)
                second = process_server_changes(config)

            self.assertEqual(first["temporary_processing"]["failed"], [])
            self.assertEqual(second["temporary_processing"]["failed"], [])
            self.assertEqual(second["temporary_processing"]["succeeded"], [])
            outbound.assert_called_once()

    def test_failed_outbound_reuses_unchanged_generated_traveler_on_retry(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            folder = config.source_root / "B12"
            report = folder / "Report" / "pp-板材清单.xlsx"
            from tests.test_order_workflow import make_board_material_report
            make_board_material_report(report, factory="", name="")

            from traveler_assistant.order_workflow import update_order_traveler as real_update
            with patch(
                "traveler_assistant.inventory.run_jdy",
                side_effect=[
                    {"saved": False, "results": []},
                    {"saved": True, "results": [{"documentNumber": "OUT-B12-RETRY"}]},
                ],
            ) as outbound, patch(
                "traveler_assistant.order_workflow.update_order_traveler",
                wraps=real_update,
            ) as updater:
                first = process_server_changes(config)
                second = process_server_changes(config)

            self.assertTrue(first["temporary_processing"]["failed"])
            self.assertEqual(second["temporary_processing"]["failed"], [])
            processed = second["temporary_processing"]["succeeded"][0]["processed"]
            self.assertEqual(processed[0]["traveler_action"], "reused")
            updater.assert_not_called()
            self.assertEqual(outbound.call_count, 2)

    def test_temporary_folder_uses_unique_aimes_review_match_when_available(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "source",
                order_root=root / "orders",
            )
            folder = config.source_root / "B12"
            report = folder / "Report" / "pp-板材清单.xlsx"
            fittings = folder / "Report" / "Fittingslist.xlsx"
            from tests.test_order_workflow import make_board_material_report, make_fittings
            make_board_material_report(report, factory="F100", name="B12")
            make_fittings(fittings, [("F100", 2)])
            store = OrderIndexStore(config.workflow_database)
            store.replace_aimes_review_rows([{
                "ignore_key": "factory:F100",
                "factory_order": "F100",
                "factory_name": "B12",
                "sales_order_name": "临时 B12",
                "reason": "销售单名称不符合标准规则",
            }])
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.inventory.run_jdy",
                return_value={"saved": True, "results": [{"documentNumber": "OUT-B12-AIMES"}]},
            ) as outbound:
                result = process_server_changes(config)

            self.assertEqual(result["temporary_processing"]["failed"], [])
            self.assertEqual(result["temporary_processing"]["succeeded"][0]["order_ids"], ["B12"])
            workbook = __import__("openpyxl").load_workbook(
                root / "orders" / "B12" / "Work Order Traveler(B12).xlsx",
                data_only=False,
            )
            self.assertTrue(any(
                workbook["Picking List"].cell(row, 4).value == "B12"
                for row in range(1, workbook["Picking List"].max_row + 1)
            ))
            outbound.assert_called_once()
    def test_factory_order_initial_date_cutoff_uses_embedded_date(self):
        self.assertTrue(_factory_order_before_initial_date("F2605260119", "", "2026-07-22"))
        self.assertFalse(_factory_order_before_initial_date("F2608010001", "", "2026-07-22"))
        self.assertTrue(_factory_order_before_initial_date("F100", "2026-07-21 10:00:00", "2026-07-22"))
        self.assertFalse(_factory_order_before_initial_date("F100", "", "2026-07-22"))

    def test_initial_date_removes_stale_ownership_issue_and_does_not_recreate_it(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "missing-source",
                initial_date="2026-07-22",
            )
            store = OrderIndexStore(config.workflow_database)
            store.upsert_factory("F2605260119", ownership_status="待确认")
            store.upsert_active_issue(
                issue_key="factory_ownership:F2605260119",
                kind="factory_ownership",
                factory_order="F2605260119",
                message="历史工厂单不应继续提醒",
            )
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.order_index.load_aimes_order_cache",
                return_value=[{
                    "factory_order": "F2605260119",
                    "factory_name": "PP0037 KITCHEN",
                    "sales_order_name": "PP0037",
                    "split_time": "",
                }],
            ):
                result = sync_order_index(config)

            reopened = OrderIndexStore(config.workflow_database)
            self.assertFalse(any(
                item["issue_key"] == "factory_ownership:F2605260119"
                for item in result["current_issues"]
            ))
            self.assertFalse(any(
                item["issue_key"] == "factory_ownership:F2605260119"
                for item in reopened.active_issues()
            ))
            reopened.close()

    def test_server_change_message_identifies_order_factory_and_data(self):
        self.assertEqual(
            _server_data_change_message("modified", ["PP0035-2"], "F20050502", "五金信息"),
            "修改订单 PP0035-2（工厂单 F20050502）的五金信息",
        )

    def test_server_change_message_explains_action_and_path(self):
        item = {"kind": "folder", "order_id": "PP0035-2", "manual_only": False}
        self.assertEqual(
            _server_change_message("modified", item, "/Volumes/server/Optimized Orders/pp0035-2"),
            "Server 订单文件夹修改：pp0035-2（订单 PP0035-2）（路径：/Volumes/server/Optimized Orders/pp0035-2）",
        )
        self.assertEqual(
            _server_data_change_message(
                "added", ["PP0035-2"], "F20050502", "五金信息", "/Volumes/server/Optimized Orders/pp0035-2/Fittingslist.xlsx"
            ),
            "新增订单 PP0035-2（工厂单 F20050502）的五金信息（来源：/Volumes/server/Optimized Orders/pp0035-2/Fittingslist.xlsx）",
        )

    def test_invalid_and_test_aimes_rows_are_warnings_only(self):
        invalid = _aimes_row_issue({
            "factory_order": "F200",
            "factory_name": "KITCHEN",
            "sales_order_name": "ORDER-200",
            "split_time": "2026-08-10 08:30:00",
        })
        test_row = _aimes_row_issue({
            "factory_order": "F201",
            "factory_name": "Kitchen Test",
            "sales_order_name": "PP0035",
            "split_time": "2026-08-10 08:30:00",
        })

        self.assertIn("销售单名称不是", invalid["reason"])
        self.assertIn("包含 test", test_row["reason"])
        self.assertEqual(test_row["ignore_key"], "factory:F201")
        visible, warnings = _partition_aimes_rows(
            [
                {
                    "factory_order": "F2606300182",
                    "factory_name": "CS001-Unnamed",
                    "sales_order_name": "SHERRY 001",
                    "split_time": "2026-06-30 10:00:00",
                }
            ],
            set(),
            {"factory:F2606300182": "CS001"},
        )
        self.assertEqual(visible[0]["sales_order_name"], "CS001")
        self.assertEqual(visible[0]["factory_order"], "F2606300182")
        self.assertEqual(warnings, [])

        unassigned_visible, unassigned_warnings = _partition_aimes_rows(
            [{
                "factory_order": "F2606300182",
                "factory_name": "CS001-Unnamed",
                "sales_order_name": "SHERRY 001",
                "split_time": "2026-06-30 10:00:00",
            }],
            set(),
            {},
        )
        self.assertEqual(unassigned_visible, [])
        self.assertEqual(unassigned_warnings[0]["suggested_order_id"], "CS001")

    def test_aimes_factory_name_order_prefix_mismatch_is_a_warning(self):
        issue = _aimes_row_issue({
            "factory_order": "F2608190230",
            "factory_name": "P0072-BED 2",
            "sales_order_name": "PP0072",
            "split_time": "2026-08-19 13:06:45",
        })
        self.assertIsNotNone(issue)
        self.assertIn("订单前缀 P0072 与销售单名称 PP0072 不一致", issue["reason"])
        visible, warnings = _partition_aimes_rows(
            [{
                "factory_order": "F2608190230",
                "factory_name": "P0072-BED 2",
                "sales_order_name": "PP0072",
                "split_time": "2026-08-19 13:06:45",
            }],
            set(),
        )
        self.assertEqual(visible, [])
        self.assertEqual(warnings[0]["factory_order"], "F2608190230")

    def test_fittings_factory_order_uses_order_folder_hint(self):
        candidates = {}
        _merge_candidate(
            candidates,
            "F2606060141",
            source="server",
            order_id="PP0037",
            folder="/Volumes/server/Optimized Orders/pp0037",
        )

        result = _effective_factory_candidate("F2606060141", candidates["F2606060141"])

        self.assertEqual(result["order_id"], "PP0037")
        self.assertEqual(result["ownership_status"], "已确认")

    def test_unowned_factory_uses_exact_aimes_name_to_derive_order(self):
        candidates = {}
        _merge_candidate(
            candidates,
            "F2606060141",
            source="server",
            folder="/tmp/pp0037",
        )
        config = Config()

        with patch(
            "traveler_assistant.core.lookup_aimes_names",
            return_value={"F2606060141": "PP0037"},
        ):
            error = _exact_resolve_unowned_factories(config, candidates)

        result = _effective_factory_candidate("F2606060141", candidates["F2606060141"])
        self.assertEqual(error, "")
        self.assertEqual(result["order_id"], "PP0037")
        self.assertEqual(result["name_source"], "AIMES精确查询")

    def test_existing_database_factory_skips_exact_aimes_lookup(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(state_dir=Path(temp) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            store.upsert_factory(
                "F2606170156",
                order_id="PP0035",
                factory_name="PP0035-ROOM 5",
                name_source="server_report",
                ownership_status="已确认",
            )
            store.commit()

            candidates = {}
            _merge_candidate(
                candidates,
                "F2606170156",
                source="server",
                folder=str(Path(temp) / "PP0035"),
            )
            _merge_database_factory_candidates(store, candidates)
            with patch(
                "traveler_assistant.core.lookup_aimes_names",
                side_effect=AssertionError("已有本地工厂单身份时不应精确查询 AIMES"),
            ):
                error = _exact_resolve_unowned_factories(config, candidates)

            result = _effective_factory_candidate(
                "F2606170156", candidates["F2606170156"]
            )
            store.close()

        self.assertEqual(error, "")
        self.assertEqual(result["order_id"], "PP0035")
        self.assertEqual(result["factory_name"], "PP0035-ROOM 5")
        self.assertEqual(result["name_source"], "server_report")

    def test_active_issue_is_persisted_and_resolved(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "order-index.sqlite3"
            store = OrderIndexStore(path)
            store.upsert_active_issue(
                issue_key="factory_ownership:F100",
                kind="factory_ownership",
                factory_order="F100",
                path="/tmp/pp0037",
                message="工厂单 F100 的订单归属无法唯一确认",
            )
            store.commit()
            self.assertEqual(store.active_issues()[0]["issue_key"], "factory_ownership:F100")
            store.resolve_active_issue("factory_ownership:F100")
            store.commit()
            self.assertEqual(store.active_issues(), [])
            store.close()

    def test_deleted_aimes_factory_is_audit_only_and_not_in_summaries(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "order-index.sqlite3"
            store = OrderIndexStore(path)
            store.upsert_order("PP9999", order_type="owned")
            store.upsert_aimes_factory(
                "F100", order_id="PP9999", factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999", split_time="2026-08-16T10:00:00",
                seen_at="2026-08-16T10:00:00",
            )
            store.upsert_aimes_factory(
                "F101", order_id="PP9999", factory_name="PP9999-BATH",
                sales_order_name="PP9999", split_time="2026-08-16T10:01:00",
                seen_at="2026-08-16T10:01:00",
            )
            store.mark_aimes_deleted(["F100"], verified_at="2026-08-16T11:00:00")
            store.commit()
            summary = next(item for item in store.summaries() if item["order_id"] == "PP9999")
            status = store.connection.execute(
                "select aimes_status, aimes_deleted_at from factory_orders where factory_order='F100'"
            ).fetchone()
            store.close()
        self.assertEqual([item["factory_order"] for item in summary["factories"]], ["F101"])
        self.assertEqual(tuple(status), ("deleted", "2026-08-16T11:00:00"))

    def test_invalid_aimes_rows_persist_for_reopen_without_entering_business_tables(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config()
            config.state_dir = Path(temp) / "state"
            source_rows = [{
                "factory_order": "F200",
                "factory_name": "TEST ROOM",
                "sales_order_name": "BAD-ORDER",
                "split_time": "2026-08-10 09:30:00",
            }]
            with patch(
                "traveler_assistant.core.refresh_aimes_recent_orders",
                return_value=source_rows,
            ), patch(
                "traveler_assistant.order_index.load_aimes_order_cache",
                return_value=[],
            ):
                result = sync_aimes_index(config, force=True)

            self.assertEqual(result["aimes"]["issue_count"], 0)
            self.assertEqual(result["aimes"]["warning_count"], 1)
            self.assertEqual(result["aimes_issues"], [])
            self.assertEqual(result["aimes_warnings"][0]["factory_order"], "F200")
            store = OrderIndexStore(config.workflow_database)
            self.assertEqual(store.connection.execute("select count(*) from aimes_review_rows").fetchone()[0], 1)
            self.assertEqual(store.connection.execute("select count(*) from factory_orders").fetchone()[0], 0)
            store.close()

            reopened = list_order_index(config)
            self.assertEqual(reopened["aimes_warnings"][0]["factory_order"], "F200")

    def test_skipped_aimes_refresh_keeps_persisted_warning_visible_after_reopen(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(state_dir=Path(temp) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            store.replace_aimes_review_rows([{
                "ignore_key": "factory:F2608190230",
                "factory_order": "F2608190230",
                "factory_name": "P0072-BED 2",
                "sales_order_name": "PP0072",
                "reason": "工厂单名称订单前缀 P0072 与销售单名称 PP0072 不一致",
                "suggested_order_id": "",
                "split_time": "2026-08-19T13:06:45",
            }])
            today = datetime.now().date().isoformat()
            store.record_run(
                f"{today}T08:00:00",
                f"{today}T08:01:00",
                aimes_attempted=True,
                aimes_succeeded=True,
                aimes_count=1,
                server_folder_count=0,
            )
            store.commit()
            store.close()

            result = sync_aimes_index(config, if_needed=True)

            self.assertFalse(result["aimes"]["attempted"])
            self.assertEqual(result["aimes"]["warning_count"], 1)
            self.assertEqual(result["aimes_warnings"][0]["factory_order"], "F2608190230")

    def test_persisted_aimes_warning_can_be_assigned_when_valid_cache_excludes_it(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(state_dir=Path(temp) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP0072", validation_status="正常", stage="已出货")
            store.upsert_factory(
                "F2608190230",
                order_id="PP0072",
                factory_name="P0072-BED 2",
                sales_order_name="PP0072",
                name_source="AIMES",
                ownership_status="已确认",
                outbound_status="已出库",
            )
            store.replace_aimes_review_rows([{
                "ignore_key": "factory:F2608190230",
                "factory_order": "F2608190230",
                "factory_name": "P0072-BED 2",
                "sales_order_name": "PP0072",
                "reason": "工厂单名称订单前缀 P0072 与销售单名称 PP0072 不一致",
                "suggested_order_id": "",
                "split_time": "2026-08-19T13:06:45",
            }])
            store.commit()
            store.close()

            with patch("traveler_assistant.order_index.load_aimes_order_cache", return_value=[]):
                result = assign_aimes_factory_order(
                    config,
                    "factory:F2608190230",
                    "PP0072",
                )

            self.assertEqual(result["aimes_warnings"], [])
            reopened = OrderIndexStore(config.workflow_database)
            assignment = reopened.connection.execute(
                "select assigned_order_id from aimes_order_assignments where ignore_key = ?",
                ("factory:F2608190230",),
            ).fetchone()
            outbound = reopened.connection.execute(
                "select outbound_status from factory_orders where factory_order = ?",
                ("F2608190230",),
            ).fetchone()
            review_count = reopened.connection.execute(
                "select count(*) from aimes_review_rows where ignore_key = ?",
                ("factory:F2608190230",),
            ).fetchone()[0]
            reopened.close()
            self.assertEqual(tuple(assignment), ("PP0072",))
            self.assertEqual(tuple(outbound), ("已出库",))
            self.assertEqual(review_count, 0)

    def test_exactly_verified_aimes_factory_is_persisted_as_aimes_identity(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config()
            config.state_dir = Path(temp) / "state"
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP0035", validation_status="正常")
            store.upsert_factory(
                "F099",
                order_id="PP0035",
                factory_name="PP0035-ROOM 5",
                split_time="2026-08-19T10:00:00",
                name_source="server_report",
                ownership_status="已确认",
            )
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.core.refresh_aimes_recent_orders_and_verify",
                return_value=(
                    [{
                        "factory_order": "F100",
                        "factory_name": "PP0035-KITCHEN",
                        "sales_order_name": "PP0035",
                        "split_time": "2026-08-19 10:00:00",
                    }],
                    {
                        "rows": [{
                            "factory_order": "F099",
                            "factory_name": "PP0035-ROOM 5",
                            "sales_order_name": "PP0035",
                            "split_time": "",
                        }],
                        "missing": [],
                    },
                ),
            ):
                result = sync_aimes_index(config, force=True)

            self.assertTrue(result["aimes"]["succeeded"])
            store = OrderIndexStore(config.workflow_database)
            row = store.connection.execute(
                "select order_id, factory_name, sales_order_name, split_time, name_source, last_aimes_seen from factory_orders where factory_order='F099'"
            ).fetchone()
            store.close()

        self.assertEqual(tuple(row[:5]), ("PP0035", "PP0035-ROOM 5", "PP0035", "2026-08-19T10:00:00", "AIMES"))
        self.assertTrue(row[5])

    def test_business_errors_are_actionable_and_hide_technical_details(self):
        validation = _business_validation_message(
            RuntimeError('Traceback: sqlite3.OperationalError: database is locked')
        )
        aimes = _business_aimes_message(RuntimeError("HTTP status code 500"))
        report = _business_report_message("fittings", Path("FittingslistPC123.xlsx"))

        self.assertIn("请检查", validation)
        self.assertNotIn("Traceback", validation)
        self.assertNotIn("sqlite3", validation)
        self.assertIn("重新获取", aimes)
        self.assertNotIn("500", aimes)
        self.assertIn("FittingslistPC123.xlsx", report)
        self.assertIn("重新扫描 Server", report)

    def test_business_aimes_message_uses_explicit_error_type_before_message_words(self):
        table_error = RuleError(
            "aimes_table_schema",
            "AIMES 工厂订单表缺少必要列；上一阶段登录 AIMES 完成",
        )
        credentials_error = RuleError("aimes_credentials", "账号或密码错误")

        self.assertIn("缺少必要列", _business_aimes_message(table_error))
        self.assertNotIn("用户名和密码", _business_aimes_message(table_error))
        self.assertIn("用户名和密码", _business_aimes_message(credentials_error))

        not_ready_error = RuleError("aimes_table_not_ready", "AIMES_TABLE_NOT_READY：表头")
        self.assertIn("尚未加载完成", _business_aimes_message(not_ready_error))

    def test_old_status_is_migrated_and_validation_reason_is_persisted(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "order-index.sqlite3"
            store = OrderIndexStore(path)
            store.upsert_order("PP9999", validation_status="数据异常", stage="待拆单")
            store.upsert_factory(
                "F100",
                order_id="PP9999",
                factory_name="KITCHEN",
                sales_order_name="PP9999",
                split_time="2026-08-10T08:30:00",
                name_source="AIMES",
                ownership_status="已确认",
            )
            store.connection.execute(
                "update orders set validation_message = ? where order_id = ?",
                ("未找到 material 文件。请补充后重新扫描 Server。", "PP9999"),
            )
            store.commit()
            store.close()

            reopened = OrderIndexStore(path)
            row = reopened.summaries()[0]
            stored_stage = reopened.connection.execute(
                "select stage from orders where order_id = ?", ("PP9999",)
            ).fetchone()[0]
            version = reopened.connection.execute("pragma user_version").fetchone()[0]
            reopened.close()

        self.assertEqual(row["stage"], "数据异常")
        self.assertEqual(stored_stage, "已设计")
        self.assertEqual(row["validation_message"], "未找到 material 文件。请补充后重新扫描 Server。")
        self.assertEqual(version, 11)

    def test_schema_migration_resolves_legacy_warning_from_unique_order_folder(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source_folder = root / "pp0037"
            source_folder.mkdir()
            path = root / "order-index.sqlite3"
            store = OrderIndexStore(path)
            store.upsert_factory(
                "F2606060141",
                source_folder=str(source_folder),
                ownership_status="待确认",
            )
            store.connection.execute("pragma user_version = 6")
            store.commit()
            store.close()

            reopened = OrderIndexStore(path)
            factory = reopened.connection.execute(
                "select order_id, ownership_status, name_source from factory_orders where factory_order = ?",
                ("F2606060141",),
            ).fetchone()
            self.assertEqual(factory, ("PP0037", "已确认", "server_folder"))
            self.assertEqual(reopened.active_issues(), [])
            self.assertIn(
                "自动确认工厂单 F2606060141 归属订单 PP0037",
                reopened.latest_changes()[0]["message"],
            )
            reopened.close()

    def test_aimes_owner_wins_over_stale_server_owner(self):
        candidate = {
            "names": {
                "server": {"PP0035-OFFICE"},
                "aimes": {"OFFICE"},
            },
            "orders": {"PP0035", "PP0035-2"},
            "sales_orders": {"PP0035-2"},
            "split_times": {"2026-08-10T08:30:00"},
            "folders": {"/Volumes/server/Optimized Orders/pp0035-2"},
            "has_hardware": False,
            "optimized": False,
        }

        result = _effective_factory_candidate("F2608050220", candidate)

        self.assertEqual(result["order_id"], "PP0035-2")
        self.assertEqual(result["factory_name"], "OFFICE")
        self.assertEqual(result["sales_order_name"], "PP0035-2")
        self.assertEqual(result["split_time"], "2026-08-10T08:30:00")
        self.assertEqual(result["name_source"], "AIMES")
        self.assertEqual(result["ownership_status"], "已确认")

    def test_aimes_order_validation_and_test_filter(self):
        self.assertEqual(_valid_aimes_order_id("PP0035"), "PP0035")
        self.assertEqual(_valid_aimes_order_id("PP0035-2"), "PP0035-2")
        self.assertEqual(_valid_aimes_order_id("PP0011"), "PP0011")
        self.assertEqual(_valid_aimes_order_id("CS123"), "CS123")
        self.assertEqual(_valid_aimes_order_id("PP0034-2"), "PP0034-2")
        self.assertEqual(_valid_aimes_order_id("PP035"), "")
        self.assertEqual(_valid_aimes_order_id("PP0035-A"), "")
        self.assertIsNone(_visible_aimes_row({
            "factory_order": "F100",
            "factory_name": "Kitchen TEST",
            "sales_order_name": "PP0035",
            "split_time": "2026-08-10 08:30:00",
        }))
        self.assertIsNone(_visible_aimes_row({
            "factory_order": "F101",
            "factory_name": "Kitchen",
            "sales_order_name": "PP0035-test",
            "split_time": "2026-08-10 08:30:00",
        }))

    def test_historical_pp_server_paths_are_in_dashboard_scope(self):
        root = Path("/Volumes/server/Optimized Orders")

        self.assertTrue(_source_path_in_dashboard_scope(root, str(root / "PP0034")))
        self.assertTrue(_source_path_in_dashboard_scope(root, str(root / "PP0034-2" / "report.xlsx")))
        self.assertTrue(_source_path_in_dashboard_scope(root, str(root / "PP0035")))
        self.assertTrue(_source_path_in_dashboard_scope(root, str(root / "PP0035-2" / "report.xlsx")))
        self.assertTrue(_source_path_in_dashboard_scope(root, str(root / "CS123")))

    def test_summary_aggregates_factory_status(self):
        with tempfile.TemporaryDirectory() as temp:
            store = OrderIndexStore(Path(temp) / "order-index.sqlite3")
            store.upsert_order("PP9999", validation_status="正常")
            store.upsert_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                split_time="2026-08-01T10:00:00",
                name_source="AIMES",
                ownership_status="已确认",
                optimized=True,
                outbound_status="已出库",
            )
            store.upsert_factory(
                "F200",
                order_id="PP9999",
                factory_name="PP9999-LAUNDRY",
                sales_order_name="PP9999",
                split_time="2026-08-02T10:00:00",
                name_source="AIMES",
                ownership_status="已确认",
                optimized=True,
                outbound_status="未出库",
            )
            store.commit()

            row = store.summaries()[0]
            store.close()

        self.assertEqual(row["factory_count"], 2)
        self.assertEqual(row["optimized_count"], 2)
        self.assertEqual(row["shipped_count"], 1)
        self.assertEqual(row["stage"], "部分生产，部分出货")
        self.assertEqual(row["optimization_progress"], "2 / 2")
        self.assertEqual(row["outbound_progress"], "1 / 2")
        self.assertEqual(row["latest_split_time"], "2026-08-02T10:00:00")

    def test_standard_outbound_status_reconciles_and_survives_reopen(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(state_dir=Path(temp) / "state")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("CS005", validation_status="正常")
            store.upsert_factory(
                "F2608120222",
                order_id="CS005",
                factory_name="CS005-KITCHEN",
                sales_order_name="CS005",
                name_source="AIMES",
                ownership_status="已确认",
                optimized=True,
                outbound_status="未出库",
            )
            store.commit()
            store.close()

            records = [{
                "order_id": "CS005",
                "remark": "CS005",
                "status": "已出库",
                "document_number": "QTCK20260815001",
                "synced_at": "2026-08-15T16:20:00",
            }]
            with patch("traveler_assistant.order_index._load_outbound_records", return_value=records):
                store = OrderIndexStore(config.workflow_database)
                self.assertEqual(reconcile_outbound_statuses(config, store), 1)
                store.close()
                listed = list_order_index(config)

            self.assertEqual(listed["orders"][0]["stage"], "已出货")
            self.assertEqual(listed["orders"][0]["completed_at"], "2026-08-15T16:20:00")
            reopened = OrderIndexStore(config.workflow_database)
            row = reopened.connection.execute(
                "select outbound_status, outbound_document, outbound_completed_at from factory_orders where factory_order = ?",
                ("F2608120222",),
            ).fetchone()
            reopened.close()
            self.assertEqual(row, ("已出库", "QTCK20260815001", "2026-08-15T16:20:00"))

    def test_fully_shipped_order_is_completed_even_if_optimization_evidence_is_missing(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(state_dir=Path(temp) / "state")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP9998", validation_status="正常")
            for factory_order, factory_name in (("F100", "PP9998-KITCHEN"), ("F101", "PP9998-LAUNDRY")):
                store.upsert_factory(
                    factory_order,
                    order_id="PP9998",
                    factory_name=factory_name,
                    sales_order_name="PP9998",
                    name_source="AIMES",
                    ownership_status="已确认",
                    optimized=False,
                    outbound_status="已出库",
                )
            store.commit()

            row = store.summaries()[0]
            store.close()

        self.assertEqual(row["optimized_count"], 0)
        self.assertEqual(row["shipped_count"], 2)
        self.assertEqual(row["stage"], "已出货")

    def test_grouped_outbound_document_reconciles_all_factory_orders_after_reindex(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(state_dir=Path(temp) / "state")
            config.prepare_storage()
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP9999", validation_status="正常")
            factories = [
                ("F100", "PP9999-MASTER"),
                ("F101", "PP9999-1ST"),
                ("F102", "PP9999-2ND"),
            ]
            for factory_order, factory_name in factories:
                store.upsert_factory(
                    factory_order,
                    order_id="PP9999",
                    factory_name=factory_name,
                    sales_order_name="PP9999",
                    name_source="AIMES",
                    ownership_status="已确认",
                    optimized=True,
                    outbound_status="未出库",
                )
            store.commit()
            store.close()

            connection = sqlite3.connect(config.workflow_database)
            connection.execute(
                """
                insert into outbound_documents(
                    document_number, document_type, order_id, factory_order,
                    status, source, issued_at, updated_at
                ) values(?,?,?,?,?,?,?,?)
                """,
                (
                    "QTCK-GROUPED",
                    "库存出库单",
                    "PP9999",
                    "F100,F101,F102",
                    "已出库",
                    "user-confirmed-grouped",
                    "2026-08-19",
                    "2026-08-19T12:00:00-07:00",
                ),
            )
            connection.commit()
            connection.close()
            ensure_schema(config.workflow_database)

            store = OrderIndexStore(config.workflow_database)
            records = _load_outbound_records(config)
            self.assertEqual(
                sorted(record["factory_order"] for record in records),
                ["F100", "F101", "F102"],
            )
            self.assertEqual(reconcile_outbound_statuses(config, store), 3)
            store.close()

            listed = list_order_index(config)

        order = listed["orders"][0]
        self.assertEqual(order["stage"], "已出货")
        self.assertEqual(order["outbound_progress"], "3 / 3")
        self.assertEqual(
            {
                (item["factory_order"], item["outbound_status"], item["outbound_document"])
                for item in order["factories"]
            },
            {
                ("F100", "已出库", "QTCK-GROUPED"),
                ("F101", "已出库", "QTCK-GROUPED"),
                ("F102", "已出库", "QTCK-GROUPED"),
            },
        )

    def test_order_level_outbound_record_is_not_broadcast_to_split_factories(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(state_dir=Path(temp) / "state")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("CS005", validation_status="正常")
            for factory_order, factory_name in (("F100", "CS005-KITCHEN"), ("F101", "CS005-LAUNDRY")):
                store.upsert_factory(
                    factory_order,
                    order_id="CS005",
                    factory_name=factory_name,
                    sales_order_name="CS005",
                    name_source="AIMES",
                    ownership_status="已确认",
                    optimized=True,
                    outbound_status="未出库",
                )
            store.commit()
            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[{
                    "order_id": "CS005",
                    "remark": "CS005",
                    "status": "已出库",
                    "document_number": "QTCK-ONE",
                }],
            ):
                self.assertEqual(reconcile_outbound_statuses(config, store), 0)
            statuses = store.connection.execute(
                "select outbound_status from factory_orders order by factory_order"
            ).fetchall()
            store.close()
            self.assertEqual(statuses, [("未出库",), ("未出库",)])

    def test_partial_factory_upsert_preserves_persisted_business_statuses(self):
        with tempfile.TemporaryDirectory() as temp:
            store = OrderIndexStore(Path(temp) / "order-index.sqlite3")
            store.upsert_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                report_state="已发现",
                ownership_status="已确认",
                has_hardware=True,
                optimized=True,
                outbound_status="已出库",
                outbound_document="QTCK-ONE",
            )
            store.upsert_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                name_source="AIMES",
            )
            row = store.connection.execute(
                """
                select report_state, ownership_status, has_hardware, optimized,
                       outbound_status, outbound_document
                from factory_orders where factory_order = 'F100'
                """
            ).fetchone()
            store.close()
            self.assertEqual(row, ("已发现", "已确认", 1, 1, "已出库", "QTCK-ONE"))

    def test_fully_shipped_aimes_order_is_not_a_server_scan_candidate(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(state_dir=Path(temp) / "state")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999 KITCHEN",
                sales_order_name="PP9999",
                split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[{
                    "order_id": "PP9999",
                    "remark": "PP9999 KITCHEN",
                    "status": "已出库",
                    "document_number": "QTCK001",
                }],
            ):
                candidates = _orders_requiring_server_scan(config, store)
            store.close()

        self.assertNotIn("PP9999", candidates)

    def test_scan_does_not_parse_material_source_as_traveler_during_outbound_reconcile(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            folder.mkdir(parents=True)
            material_path = folder / "PP0035-2 materials.xlsx"
            from openpyxl import Workbook

            workbook = Workbook()
            workbook.active.title = "Sheet1"
            workbook.save(material_path)

            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999-KITCHEN",
                sales_order_name="PP9999",
                split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            self._set_permanent_server_policy(store, "PP9999", folder)
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[{
                    "order_id": "PP9999",
                    "remark": "PP9999-KITCHEN",
                    "status": "已出库",
                    "document_number": "QTCK001",
                    "raw_fingerprint": "stored-material-fingerprint",
                    "traveler_path": str(material_path),
                }],
            ):
                store = OrderIndexStore(config.workflow_database)
                candidates = _orders_requiring_server_scan(config, store)
                store.close()
                result = scan_server_changes(config)

            self.assertNotIn("PP9999", candidates)
            self.assertFalse(result["server"]["changed"])
            self.assertEqual(result["server"]["changes"], [])

    def test_fully_shipped_folder_resolves_stale_material_validation_issue(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            folder.mkdir(parents=True)
            material_path = folder / "PP9999 materials.xlsx"
            material_path.write_bytes(b"invalid source retained for history")

            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999 KITCHEN",
                sales_order_name="PP9999",
                split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            self._set_permanent_server_policy(store, "PP9999", folder)
            store.upsert_active_issue(
                issue_key=f"material_validation:PP9999:{material_path}",
                kind="material_validation",
                order_id="PP9999",
                path=str(material_path),
                message="历史材料校验失败",
                seen_at="2026-08-10T09:00:00",
            )
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[{
                    "order_id": "PP9999",
                    "remark": "PP9999 KITCHEN",
                    "status": "已出库",
                    "document_number": "QTCK001",
                }],
            ):
                result = scan_server_changes(config)

            self.assertFalse(any(
                issue["kind"] == "material_validation"
                for issue in result["current_issues"]
            ))
            reopened = OrderIndexStore(config.workflow_database)
            issue = reopened.connection.execute(
                "select status, resolved_at from active_issues where issue_key = ?",
                (f"material_validation:PP9999:{material_path}",),
            ).fetchone()
            reopened.close()

        self.assertEqual(issue[0], "resolved")
        self.assertTrue(issue[1])

    def test_completed_production_resolves_post_production_material_issue(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP0063-2"
            folder.mkdir(parents=True)
            material_path = folder / "PP0063-2 materials.xlsx"
            material_path.write_bytes(b"source retained; preview owns validation")

            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP0063-2", source_folder=str(folder))
            store.upsert_aimes_factory(
                "F2606170170",
                order_id="PP0063-2",
                factory_name="PP0063-2 KITCHEN",
                sales_order_name="PP0063-2",
                split_time="2026-06-17T05:43:46",
                seen_at="2026-06-17T05:43:46",
            )
            store.connection.execute(
                """
                insert into material_items(
                    order_id, material_type, color, thickness, quantity, unit,
                    edge, source_type, source_path, source_fingerprint, updated_at
                ) values(?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    "PP0063-2", "panel", "Walnut", "19.1", 2, "张", "",
                    "server", str(material_path), "fingerprint", "2026-08-19T09:38:40",
                ),
            )
            store.connection.execute(
                """
                insert into manual_production_batches(
                    batch_id, batch_number, order_id, production_time, source,
                    status, created_at, updated_at
                ) values(?,?,?,?,?,?,?,?)
                """,
                (
                    1, "MP-20260821-225155-1D99ED", "PP0063-2",
                    "2026-08-21T22:53:19-07:00", "manual", "completed",
                    "2026-08-21T22:51:55-07:00", "2026-08-21T22:53:19-07:00",
                ),
            )
            store.connection.execute(
                "insert into manual_production_batch_factories(batch_id, order_id, factory_order) values(?,?,?)",
                (1, "PP0063-2", "F2606170170"),
            )
            issue_key = f"material_validation:PP0063-2:{material_path}"
            store.upsert_active_issue(
                issue_key=issue_key,
                kind="material_validation",
                order_id="PP0063-2",
                path=str(material_path),
                message="旧解析器错误",
                seen_at="2026-08-31T12:59:10",
            )
            store.commit()
            store.close()

            result = scan_server_changes(config)

            self.assertFalse(any(
                issue["kind"] == "material_validation"
                for issue in result["current_issues"]
            ))
            reopened = OrderIndexStore(config.workflow_database)
            issue = reopened.connection.execute(
                "select status, resolved_at from active_issues where issue_key = ?",
                (issue_key,),
            ).fetchone()
            order = reopened.connection.execute(
                "select material_status from orders where order_id = 'PP0063-2'"
            ).fetchone()
            change = reopened.connection.execute(
                "select kind from sync_changes where kind = 'material_validation_reconciled'"
            ).fetchone()
            reopened.close()

        self.assertEqual(issue[0], "resolved")
        self.assertTrue(issue[1])
        self.assertEqual(order[0], "板材 · 封边")
        self.assertEqual(change[0], "material_validation_reconciled")

    def test_fully_shipped_folder_resolves_stale_hardware_selection_issue(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            folder.mkdir(parents=True)

            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999 KITCHEN",
                sales_order_name="PP9999",
                split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            self._set_permanent_server_policy(store, "PP9999", folder)
            issue_key = f"hardware_selection:PP9999:{folder}"
            store.upsert_active_issue(
                issue_key=issue_key,
                kind="hardware_selection",
                order_id="PP9999",
                path=str(folder),
                message="五金清单缺少 Order No. 区块",
                seen_at="2026-08-10T09:00:00",
            )
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[{
                    "order_id": "PP9999",
                    "remark": "PP9999 KITCHEN",
                    "status": "已出库",
                    "document_number": "QTCK001",
                }],
            ):
                result = scan_server_changes(config)

            self.assertFalse(any(
                issue["kind"] == "hardware_selection"
                for issue in result["current_issues"]
            ))
            reopened = OrderIndexStore(config.workflow_database)
            issue = reopened.connection.execute(
                "select status, resolved_at from active_issues where issue_key = ?",
                (issue_key,),
            ).fetchone()
            reopened.close()

        self.assertEqual(issue[0], "resolved")
        self.assertTrue(issue[1])

    def test_fully_shipped_folder_resolves_stale_order_validation_issue(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            config.prepare_storage()
            folder = config.source_root / "PP9999"
            folder.mkdir(parents=True)

            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999 KITCHEN",
                sales_order_name="PP9999",
                split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            self._set_permanent_server_policy(store, "PP9999", folder)
            issue_key = "order_validation:PP9999"
            store.upsert_active_issue(
                issue_key=issue_key,
                kind="order_validation",
                order_id="PP9999",
                path=str(folder),
                message="历史订单校验失败",
                seen_at="2026-08-10T09:00:00",
            )
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[{
                    "order_id": "PP9999",
                    "remark": "PP9999 KITCHEN",
                    "status": "已出库",
                    "document_number": "QTCK001",
                }],
            ):
                result = scan_server_changes(config)

            self.assertFalse(any(
                issue["kind"] == "order_validation"
                for issue in result["current_issues"]
            ))
            reopened = OrderIndexStore(config.workflow_database)
            issue = reopened.connection.execute(
                "select status, resolved_at from active_issues where issue_key = ?",
                (issue_key,),
            ).fetchone()
            reopened.close()

        self.assertEqual(issue[0], "resolved")
        self.assertTrue(issue[1])

    def test_automatic_server_snapshot_skips_shipped_order_until_aimes_adds_factory(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            shipped_folder = config.source_root / "PP9999"
            active_folder = config.source_root / "PP8888"
            unindexed_folder = config.source_root / "PP7777"
            shipped_folder.mkdir(parents=True)
            active_folder.mkdir(parents=True)
            unindexed_folder.mkdir(parents=True)
            (shipped_folder / "PP9999 materials.xlsx").write_bytes(b"shipped")
            (active_folder / "PP8888 materials.xlsx").write_bytes(b"active")
            (unindexed_folder / "PP7777 materials.xlsx").write_bytes(b"unindexed")
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="PP9999",
                factory_name="PP9999 KITCHEN",
                sales_order_name="PP9999",
                split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            store.upsert_aimes_factory(
                "F200",
                order_id="PP8888",
                factory_name="PP8888 KITCHEN",
                sales_order_name="PP8888",
                split_time="2026-08-10T08:31:00",
                seen_at="2026-08-10T08:31:00",
            )
            store.commit()
            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[{
                    "order_id": "PP9999",
                    "remark": "PP9999 KITCHEN",
                    "status": "已出库",
                    "document_number": "QTCK001",
                }],
            ):
                _, snapshot = _server_snapshot(config, store)
                # A newly confirmed shipped order is watched for seven days;
                # it is not skipped immediately.
                self.assertIn(str(shipped_folder), snapshot)
                self.assertIn(str(active_folder), snapshot)
                self.assertIn(str(unindexed_folder), snapshot)

                # A newly persisted AIMES factory order reopens the order for
                # the next automatic scan.
                store.upsert_aimes_factory(
                    "F101",
                    order_id="PP9999",
                    factory_name="PP9999 CLOSET",
                    sales_order_name="PP9999",
                    split_time="2026-08-17T08:30:00",
                    seen_at="2026-08-17T08:30:00",
                )
                store.commit()
                _, reopened_snapshot = _server_snapshot(config, store)
            store.close()

        self.assertIn(str(shipped_folder), reopened_snapshot)

    def test_new_current_aimes_factory_reopens_server_scan_candidate(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(
                state_dir=Path(temp) / "state",
                source_root=Path(temp) / "source",
            )
            config.source_root.mkdir()
            (config.source_root / "PP9999").mkdir()
            row = {
                "factory_order": "F100",
                "factory_name": "PP9999 KITCHEN",
                "sales_order_name": "PP9999",
                "split_time": "2026-08-10 08:30:00",
            }
            with patch(
                "traveler_assistant.core.refresh_aimes_recent_orders",
                return_value=[row],
            ), patch(
                "traveler_assistant.order_index.load_aimes_order_cache",
                return_value=[],
            ):
                result = sync_order_index(config, aimes_if_needed=True)

        self.assertEqual(result["sync"]["server_folder_count"], 1)
        self.assertEqual(result["orders"][0]["order_id"], "PP9999")

    def test_skipped_standard_order_does_not_resolve_its_old_issue(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config(
                state_dir=Path(temp) / "state",
                source_root=Path(temp) / "source",
            )
            config.source_root.mkdir()
            folder = config.source_root / "PP9999"
            folder.mkdir()
            store = OrderIndexStore(config.workflow_database)
            self._set_permanent_server_policy(store, "PP9999", folder)
            store.upsert_active_issue(
                issue_key="factory_ownership:F100",
                kind="factory_ownership",
                order_id="PP9999",
                factory_order="F100",
                path=str(folder),
                message="需要人工确认",
            )
            store.commit()
            store.close()

            result = sync_order_index(config)

            self.assertTrue(any(item["issue_key"] == "factory_ownership:F100" for item in result["current_issues"]))

    def test_orders_sort_by_latest_factory_split_time(self):
        with tempfile.TemporaryDirectory() as temp:
            store = OrderIndexStore(Path(temp) / "order-index.sqlite3")
            for order_id in ("PP0034", "PP0035", "PP0036"):
                store.upsert_order(order_id, validation_status="正常")
            store.upsert_factory("F100", order_id="PP0035", factory_name="Kitchen", sales_order_name="PP0035", split_time="2026-08-01T10:00:00", name_source="AIMES", ownership_status="已确认")
            store.upsert_factory("F101", order_id="PP0035", factory_name="Master", sales_order_name="PP0035", split_time="2026-08-03T10:00:00", name_source="AIMES", ownership_status="已确认")
            store.upsert_factory("F200", order_id="PP0036", factory_name="Office", sales_order_name="PP0036", split_time="2026-08-02T10:00:00", name_source="AIMES", ownership_status="已确认")
            store.upsert_factory("F099", order_id="PP0034", factory_name="Old", sales_order_name="PP0034", split_time="2026-08-04T10:00:00", name_source="AIMES", ownership_status="已确认")
            store.commit()
            rows = store.summaries()
            store.close()

        self.assertEqual([row["order_id"] for row in rows], ["PP0034", "PP0035", "PP0036"])
        self.assertEqual([row["factory_order"] for row in rows[1]["factories"]], ["F101", "F100"])

    def test_summary_includes_confirmed_server_report_factory_assigned_to_normal_order(self):
        with tempfile.TemporaryDirectory() as temp:
            store = OrderIndexStore(Path(temp) / "order-index.sqlite3")
            store.upsert_order("PP0035", validation_status="正常")
            store.upsert_factory(
                "F100",
                order_id="PP0035",
                factory_name="PP0035-ROOM 5",
                name_source="server_report",
                ownership_status="已确认",
            )
            store.commit()
            rows = store.summaries()
            store.close()

        summary = next(row for row in rows if row["order_id"] == "PP0035")
        self.assertEqual([item["factory_order"] for item in summary["factories"]], ["F100"])
        self.assertEqual(summary["factory_count"], 1)

    def test_aimes_if_needed_runs_once_per_day_after_success(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config()
            config.state_dir = Path(temp) / "state"
            config.source_root = Path(temp) / "source"
            config.source_root.mkdir()

            cached_row = {
                "factory_order": "F100",
                "factory_name": "KITCHEN",
                "sales_order_name": "PP9999",
                "split_time": "2026-08-10 08:30:00",
            }
            with patch(
                "traveler_assistant.core.refresh_aimes_recent_names",
                return_value={"F100": "KITCHEN"},
            ), patch(
                "traveler_assistant.core.refresh_aimes_recent_orders",
                return_value=[cached_row],
            ) as refresh, patch(
                "traveler_assistant.order_index.load_aimes_order_cache",
                side_effect=[[], [cached_row]],
            ):
                first = sync_order_index(config, aimes_if_needed=True)
                second = sync_order_index(config, aimes_if_needed=True)

            self.assertEqual(first["sync"]["aimes_count"], 1)
            self.assertEqual(second["sync"]["aimes_count"], 1)
            self.assertFalse(second["sync"]["aimes_attempted"])
            self.assertEqual(refresh.call_count, 1)
            self.assertEqual(refresh.call_args.args[1], 50)
            self.assertEqual(second["orders"][0]["order_id"], "PP9999")

    def test_aimes_only_sync_reports_change_then_skips_after_daily_success(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config()
            config.state_dir = Path(temp) / "state"
            row = {
                "factory_order": "F100",
                "factory_name": "KITCHEN",
                "sales_order_name": "PP0035-2",
                "split_time": "2026-08-10 08:30:00",
            }
            with patch(
                "traveler_assistant.core.refresh_aimes_recent_orders",
                return_value=[row],
            ) as refresh, patch(
                "traveler_assistant.order_index.load_aimes_order_cache",
                side_effect=[[], [row]],
            ):
                first = sync_aimes_index(config, if_needed=True)
                second = sync_aimes_index(config, if_needed=True)

            self.assertTrue(first["aimes"]["attempted"])
            self.assertTrue(first["aimes"]["succeeded"])
            self.assertTrue(first["aimes"]["changed"])
            self.assertEqual(first["orders"][0]["order_id"], "PP0035-2")
            self.assertFalse(second["aimes"]["attempted"])
            self.assertTrue(second["aimes"]["skipped_today"])
            self.assertEqual(refresh.call_count, 1)
            self.assertEqual(refresh.call_args.args[1], 50)

    def test_server_scan_is_non_mutating_until_full_processing(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config()
            config.state_dir = Path(temp) / "state"
            config.source_root = Path(temp) / "source"
            folder = config.source_root / "PP0035-2"
            folder.mkdir(parents=True)
            nesting = folder / "New Nesting" / "Optimize file" / "layout file" / "nesting_result.xml"
            nesting.parent.mkdir(parents=True)
            nesting.write_text(
                '<Nesting><BoardControl OrderID="F100" /></Nesting>',
                encoding="utf-8",
            )
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="PP0035-2",
                factory_name="PP0035-2 KITCHEN",
                sales_order_name="PP0035-2",
                split_time="2026-08-25T08:30:00",
                seen_at="2026-08-25T08:30:00",
            )
            store.commit()
            store.close()

            first = scan_server_changes(config)["server"]
            repeated = scan_server_changes(config)["server"]
            self.assertTrue(first["changed"])
            self.assertEqual(first["changes"], repeated["changes"])
            self.assertEqual(first["scan_stats"]["order_folder_count"], 1)
            self.assertEqual(first["scan_stats"]["related_xml_count"], 1)
            self.assertEqual(first["scan_stats"]["related_excel_count"], 0)
            self.assertEqual(first["scan_stats"]["quick_checked_file_count"], 1)
            self.assertEqual(first["scan_stats"]["reused_folder_count"], 0)
            self.assertEqual(first["scan_stats"]["deep_scanned_folder_count"], 1)
            self.assertEqual(first["scan_stats"]["added_count"], 1)
            self.assertEqual(first["scan_stats"]["modified_count"], 0)
            self.assertEqual(first["scan_stats"]["deleted_count"], 0)
            self.assertTrue(first["changes"][0]["event_time"])
            trace = scan_server_changes(config)["operation_trace"]["server"]
            self.assertIn(
                "快速检查 1 个相关 XML 文件，复用 0 个订单文件夹，深度扫描 1 个订单文件夹，总用时",
                trace[0],
            )
            self.assertIn("扫描范围：订单文件夹 1 个，相关 XML 文件 1 个", trace[1])

            sync_order_index(config)
            self.assertFalse(scan_server_changes(config)["server"]["changed"])

            xml_stat = nesting.stat()
            os.utime(
                nesting,
                ns=(xml_stat.st_atime_ns, xml_stat.st_mtime_ns + 1_000_000_000),
            )
            xml_changed = scan_server_changes(config)["server"]
            self.assertTrue(xml_changed["changed"])
            self.assertTrue(any(
                item["path"] == str(nesting) and item["change_type"] == "modified"
                for item in xml_changed["changes"]
            ))
            sync_order_index(config)
            self.assertFalse(scan_server_changes(config)["server"]["changed"])

            folder_stat = folder.stat()
            os.utime(
                folder,
                ns=(folder_stat.st_atime_ns, folder_stat.st_mtime_ns + 1_000_000_000),
            )
            folder_only_change = scan_server_changes(config)["server"]
            self.assertFalse(folder_only_change["changed"])

            store = OrderIndexStore(config.workflow_database)
            folder_changes_before = store.connection.execute(
                "select count(*) from sync_changes where kind = 'folder_changed'"
            ).fetchone()[0]
            store.close()
            sync_order_index(config)
            store = OrderIndexStore(config.workflow_database)
            folder_changes_after = store.connection.execute(
                "select count(*) from sync_changes where kind = 'folder_changed'"
            ).fetchone()[0]
            store.close()
            self.assertEqual(folder_changes_after, folder_changes_before)

            report = folder / "material.xlsx"
            report.write_bytes(b"placeholder")
            changed = scan_server_changes(config)["server"]
            self.assertFalse(changed["changed"])
            self.assertEqual(changed["scan_stats"]["related_xml_count"], 1)
            self.assertEqual(changed["scan_stats"]["related_excel_count"], 0)
            self.assertEqual(changed["scan_stats"]["quick_checked_file_count"], 1)
            self.assertEqual(changed["scan_stats"]["reused_folder_count"], 0)
            self.assertEqual(changed["scan_stats"]["deep_scanned_folder_count"], 1)
            self.assertEqual(changed["scan_stats"]["added_count"], 0)
            self.assertEqual(changed["scan_stats"]["modified_count"], 0)
            self.assertEqual(changed["scan_stats"]["deleted_count"], 0)
            self.assertFalse(any(item["path"] == str(report) for item in changed["changes"]))

            sync_order_index(config)
            self.assertFalse(scan_server_changes(config)["server"]["changed"])

            report.unlink()
            removed = scan_server_changes(config)["server"]
            self.assertEqual(removed["scan_stats"]["related_xml_count"], 1)
            self.assertEqual(removed["scan_stats"]["related_excel_count"], 0)
            self.assertEqual(removed["scan_stats"]["added_count"], 0)
            self.assertEqual(removed["scan_stats"]["modified_count"], 0)
            self.assertEqual(removed["scan_stats"]["deleted_count"], 0)
            sync_order_index(config)
            self.assertFalse(scan_server_changes(config)["server"]["changed"])

    def test_confirm_without_business_changes_only_updates_xml_scan_baseline(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state", source_root=root / "source")
            folder = config.source_root / "PP9999"
            xml = folder / "New Nesting" / "Optimize file" / "nesting_result.xml"
            xml.parent.mkdir(parents=True)
            xml.write_text('<Nesting><BoardControl OrderID="F100" /></Nesting>', encoding="utf-8")

            payload = {
                "source_folders": [str(folder)],
                "orders": [{
                    "order_id": "PP9999",
                    "validation_status": "正常",
                    "validation_message": "",
                    "material_changes": [],
                    "factories": [],
                    "hardware_changes": [],
                }],
                "has_business_changes": False,
                "hardware_mapping_requirements": [],
                "write_records": {},
            }
            result = confirm_server_material_preview_memory(
                config,
                payload,
                confirm_write=True,
            )

            self.assertTrue(result["server_write_confirmed"])
            self.assertTrue(result["server_write_skipped"])
            self.assertIn("仅更新 Server XML 扫描基线", result["server_write_skip_reason"])
            store = OrderIndexStore(config.workflow_database)
            try:
                self.assertEqual(
                    store.connection.execute(
                        "select path, kind, modified_at from server_scan_xml_state"
                    ).fetchall(),
                    [(str(xml), "optimization_result", int(xml.stat().st_mtime_ns // 1_000_000))],
                )
                self.assertEqual(store.connection.execute("select count(*) from orders").fetchone()[0], 0)
                self.assertEqual(store.connection.execute("select count(*) from material_items").fetchone()[0], 0)
                self.assertEqual(store.connection.execute("select count(*) from hardware_items").fetchone()[0], 0)
            finally:
                store.close()

    def test_preview_scopes_optimization_artifacts_to_selected_order(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            folder = root / "server" / "PP0062"
            other_folder = root / "server" / "OTHER"
            folder.mkdir(parents=True)
            other_folder.mkdir(parents=True)
            preview_path = root / "preview.sqlite3"
            preview = OrderIndexStore(preview_path)
            preview.upsert_order("PP0062", source_folder=str(folder), validation_status="正常")
            preview.upsert_factory("F2609060245", order_id="PP0062", source_folder=str(folder))
            preview.upsert_order("OTHER", source_folder=str(other_folder), validation_status="正常")
            preview.upsert_factory("FOTHER", order_id="OTHER", source_folder=str(other_folder))
            artifact_columns = "source_path, order_id, factory_order, file_modified_at, file_created_at, completed_at, copied_at, first_seen_at, last_seen_at, size"
            preview.connection.execute(
                f"insert into optimization_artifacts({artifact_columns}) values(?,?,?,?,?,?,?,?,?,?)",
                (str(folder / "nesting_result.xml"), "PP0062", "F2609060245", 10.0, 9.0,
                 "2026-09-06T10:00:00", "", "2026-09-06T11:00:00", "2026-09-06T11:00:00", 100),
            )
            preview.connection.execute(
                f"insert into optimization_artifacts({artifact_columns}) values(?,?,?,?,?,?,?,?,?,?)",
                (str(other_folder / "nesting_result.xml"), "OTHER", "FOTHER", 10.0, 9.0,
                 "2026-09-06T10:00:00", "", "2026-09-06T11:00:00", "2026-09-06T11:00:00", 100),
            )
            preview.commit()
            preview.close()

            payload = _server_preview_payload(
                config, preview_path, "", [folder], include_hardware=False
            )
            self.assertEqual(
                [(row["order_id"], row["factory_order"])
                 for row in payload["write_records"]["optimization_artifacts"]],
                [("PP0062", "F2609060245")],
            )

    def test_business_confirmation_persists_optimization_evidence_for_list_index(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            folder = root / "server" / "PP0062"
            folder.mkdir(parents=True)
            preview_path = root / "preview.sqlite3"
            preview = OrderIndexStore(preview_path)
            preview.upsert_order("PP0062", source_folder=str(folder), validation_status="正常")
            preview.upsert_factory(
                "F2609060245", order_id="PP0062", source_folder=str(folder),
                ownership_status="已确认", optimized=True,
            )
            preview.connection.execute(
                """insert into optimization_artifacts(
                    source_path, order_id, factory_order, file_modified_at,
                    file_created_at, completed_at, copied_at, first_seen_at,
                    last_seen_at, size
                ) values(?,?,?,?,?,?,?,?,?,?)""",
                (str(folder / "nesting_result.xml"), "PP0062", "F2609060245", 10.0, 9.0,
                 "2026-09-06T10:00:00", "", "2026-09-06T11:00:00", "2026-09-06T11:00:00", 100),
            )
            preview.commit()
            preview.close()

            payload = _server_preview_payload(
                config, preview_path, "", [folder], include_hardware=False
            )
            self.assertTrue(payload["has_business_changes"])
            confirmed = confirm_server_material_preview_memory(
                config, payload, confirm_write=True
            )
            self.assertEqual(confirmed["optimization_artifact_count"], 1)
            summary = list_order_index(config)["orders"]
            self.assertEqual(summary[0]["stage"], "已优化")
            store = OrderIndexStore(config.workflow_database)
            try:
                self.assertEqual(
                    store.connection.execute(
                        "select order_id, factory_order, completed_at from optimization_artifacts"
                    ).fetchall(),
                    [("PP0062", "F2609060245", "2026-09-06T10:00:00")],
                )
            finally:
                store.close()

    def test_memory_confirmation_upserts_selected_optimization_evidence_idempotently(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            folder = root / "server" / "PP0062"
            folder.mkdir(parents=True)
            source_path = str(folder / "nesting_result.xml")
            artifact_columns = "source_path, order_id, factory_order, file_modified_at, file_created_at, completed_at, copied_at, first_seen_at, last_seen_at, size"

            current = OrderIndexStore(config.workflow_database)
            current.upsert_order("PP0062", source_folder=str(folder), validation_status="正常")
            current.upsert_factory(
                "F2609060245", order_id="PP0062", source_folder=str(folder),
                ownership_status="已确认", optimized=True,
            )
            current.upsert_order("OTHER", validation_status="正常")
            current.upsert_factory("FOTHER", order_id="OTHER", ownership_status="已确认", optimized=True)
            current.connection.execute(
                f"insert into optimization_artifacts(id,{artifact_columns}) values(?,?,?,?,?,?,?,?,?,?,?)",
                (11, source_path, "PP0062", "F2609060245", 10.0, 9.0,
                 "2026-09-06T10:00:00", "", "2026-09-06T11:00:00", "2026-09-06T11:00:00", 100),
            )
            current.connection.execute(
                f"insert into optimization_artifacts(id,{artifact_columns}) values(?,?,?,?,?,?,?,?,?,?,?)",
                (12, "/other/nesting_result.xml", "OTHER", "FOTHER", 10.0, 9.0,
                 "2026-09-06T10:00:00", "", "2026-09-06T11:00:00", "2026-09-06T11:00:00", 100),
            )
            current.commit()
            current.close()

            payload = {
                "source_folders": [str(folder)],
                "orders": [{
                    "order_id": "PP0062",
                    "validation_status": "正常",
                    "validation_message": "",
                    "material_changes": [],
                    "factories": [],
                    "hardware_changes": [],
                }],
                "has_business_changes": False,
                "write_records": {
                    "optimization_artifacts": [
                        {"id": 12, "source_path": source_path, "order_id": "PP0062", "factory_order": "F2609060245",
                         "file_modified_at": 10.0, "file_created_at": 9.0, "completed_at": "2099-01-01T00:00:00",
                         "copied_at": "", "first_seen_at": "2099-01-01T00:00:00", "last_seen_at": "2026-09-06T10:00:00", "size": 100},
                        {"id": 13, "source_path": "/other/nesting_result.xml", "order_id": "OTHER", "factory_order": "FOTHER",
                         "file_modified_at": 10.0, "file_created_at": 9.0, "completed_at": "2099-01-01T00:00:00",
                         "copied_at": "", "first_seen_at": "2099-01-01T00:00:00", "last_seen_at": "2099-01-01T00:00:00", "size": 100},
                    ],
                },
            }
            self.assertEqual(_memory_factory_selection(payload), [])
            first = confirm_server_material_preview_memory(config, payload, confirm_write=True)
            second = confirm_server_material_preview_memory(config, payload, confirm_write=True)
            self.assertTrue(first["server_write_skipped"])
            self.assertEqual(first["optimization_artifact_count"], 1)
            self.assertEqual(second["optimization_artifact_count"], 1)

            store = OrderIndexStore(config.workflow_database)
            try:
                rows = store.connection.execute(
                    "select id, order_id, factory_order, completed_at, first_seen_at, last_seen_at from optimization_artifacts order by id"
                ).fetchall()
                self.assertEqual(len(rows), 2)
                self.assertEqual(rows[0][0], 11)
                self.assertEqual(rows[0][1:5], (
                    "PP0062", "F2609060245", "2026-09-06T10:00:00", "2026-09-06T11:00:00"
                ))
                self.assertEqual(rows[0][5], "2026-09-06T11:00:00")
                self.assertEqual(rows[1][1], "OTHER")
                self.assertEqual(list_order_index(config)["orders"][0]["stage"], "已优化")
            finally:
                store.close()

    def test_memory_evidence_confirmation_rolls_back_on_baseline_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(state_dir=root / "state")
            config.prepare_storage()
            folder = root / "server" / "PP0062"
            folder.mkdir(parents=True)
            payload = {
                "source_folders": [str(folder)],
                "orders": [{"order_id": "PP0062", "validation_status": "正常", "factories": [], "material_changes": [], "hardware_changes": []}],
                "has_business_changes": False,
                "write_records": {"optimization_artifacts": [{
                    "source_path": str(folder / "nesting_result.xml"), "order_id": "PP0062", "factory_order": "F1",
                    "file_modified_at": 1.0, "file_created_at": 1.0, "completed_at": "2026-01-01T00:00:00",
                    "copied_at": "", "first_seen_at": "2026-01-01T00:00:00", "last_seen_at": "2026-01-01T00:00:00", "size": 1,
                }]},
            }
            try:
                shared = sqlite3.connect(config.workflow_database)
                bootstrap = OrderIndexStore(config.workflow_database, connection=shared)
                install_shared_workflow_connection(config.workflow_database, shared)
                with patch.object(OrderIndexStore, "save_server_scan_xml_baseline", side_effect=RuntimeError("baseline failed")):
                    with self.assertRaises(RuntimeError):
                        confirm_server_material_preview_memory(config, payload, confirm_write=True)
                self.assertFalse(shared.in_transaction)
                self.assertEqual(shared.execute("select count(*) from optimization_artifacts").fetchone()[0], 0)
            finally:
                clear_shared_workflow_connection()
                bootstrap.close()
                shared.close()

    def test_server_scan_baseline_covers_both_server_roots(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "Optimized Orders",
            )
            folder = root / "CUT TO SIZE" / "CS003"
            folder.mkdir(parents=True)
            nesting = folder / "New Nesting" / "Optimize file" / "nesting_result.xml"
            nesting.parent.mkdir(parents=True)
            nesting.write_text(
                '<Nesting><BoardControl OrderID="F100" /></Nesting>',
                encoding="utf-8",
            )
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="CS003",
                factory_name="CS003 KITCHEN",
                sales_order_name="CS003",
                split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            store.commit()
            store.close()

            with patch(
                "traveler_assistant.order_index._load_outbound_records",
                return_value=[],
            ):
                first = scan_server_changes(config)["server"]
                self.assertTrue(any(item["path"] == str(nesting) for item in first["changes"]))

                sync_order_index(config)
                repeated = scan_server_changes(config)["server"]

        self.assertFalse(repeated["changed"])
        self.assertEqual(repeated["changes"], [])

    def test_report_edits_are_ignored_by_xml_only_server_scan(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = Config(
                state_dir=root / "state",
                source_root=root / "server" / "Optimized Orders",
            )
            folder = config.source_root.parent / "CUT TO SIZE" / "CS005"
            folder.mkdir(parents=True)
            store = OrderIndexStore(config.workflow_database)
            store.upsert_aimes_factory(
                "F100",
                order_id="CS005",
                factory_name="CS005-KITCHEN",
                sales_order_name="CS005",
                split_time="2026-08-10T08:30:00",
                seen_at="2026-08-10T08:30:00",
            )
            store.upsert_source_file(
                folder,
                source_folder=folder,
                kind="folder",
                order_id="CS005",
                changed_at="2026-08-14T10:00:00",
            )
            store.commit()
            store.close()

            materials = folder / "CS005 materials.xlsx"
            materials.write_bytes(b"generated material")
            before_registration = scan_server_changes(config)["server"]
            self.assertFalse(before_registration["changed"])
            self.assertFalse(any(item["path"] == str(materials) for item in before_registration["changes"]))

            store = OrderIndexStore(config.workflow_database)
            _record_generated_material_baseline(store, folder, materials, order_id="CS005")
            store.commit()
            store.close()
            self.assertFalse(scan_server_changes(config)["server"]["changed"])

            materials.write_bytes(b"external edit")
            after_edit = scan_server_changes(config)["server"]
            self.assertFalse(after_edit["changed"])
            self.assertFalse(any(item["path"] == str(materials) for item in after_edit["changes"]))

    def test_selected_server_folder_reuses_index_processing_for_one_folder(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config()
            config.state_dir = Path(temp) / "state"
            config.source_root = (Path(temp) / "source").resolve()
            folder = config.source_root / "PP0035-2"
            folder.mkdir(parents=True)

            store = OrderIndexStore(config.workflow_database)
            store.upsert_order("PP0035-2", order_type="temporary")
            store.commit()
            store.close()

            result = process_server_folder(config, folder)

            self.assertEqual(result["sync"]["server_folder_count"], 1)
            store = OrderIndexStore(config.workflow_database)
            order_type = store.connection.execute(
                "select order_type from orders where order_id = ?", ("PP0035-2",)
            ).fetchone()[0]
            source_paths = {
                row[0] for row in store.connection.execute("select path from source_files").fetchall()
            }
            store.close()
            self.assertEqual(order_type, "owned")
            self.assertIn(str(folder.resolve()), source_paths)
            self.assertFalse(scan_server_changes(config)["server"]["changed"])

    def test_selected_non_order_folder_is_rejected_without_processing_children(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config()
            config.state_dir = Path(temp) / "state"
            config.source_root = (Path(temp) / "source").resolve()
            selected = config.source_root / "temporary-production"
            child_order = selected / "PP0035-2"
            child_order.mkdir(parents=True)
            (child_order / "material.xlsx").write_bytes(b"placeholder")

            with self.assertRaisesRegex(RuleError, "选错了文件夹"):
                process_server_folder(config, selected)

            self.assertFalse((config.workflow_database).exists())

    def test_selected_non_order_folder_with_recognized_report_is_temporary_order(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config()
            config.state_dir = Path(temp) / "state"
            config.source_root = (Path(temp) / "source").resolve()
            selected = config.source_root / "temporary-production"
            selected.mkdir(parents=True)
            (selected / "material.xlsx").write_bytes(b"placeholder")

            result = process_server_folder(config, selected)

            self.assertEqual(result["sync"]["server_folder_count"], 1)
            store = OrderIndexStore(config.workflow_database)
            source_paths = {
                row[0] for row in store.connection.execute("select path from source_files").fetchall()
            }
            store.close()
            self.assertTrue(
                str(selected / "material.xlsx") in source_paths
            )
            self.assertFalse(scan_server_changes(config)["server"]["changed"])

    def test_scan_reports_unprocessed_non_order_folder_as_manual_only(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Config()
            config.state_dir = Path(temp) / "state"
            config.source_root = (Path(temp) / "source").resolve()
            selected = config.source_root / "temporary-production"
            selected.mkdir(parents=True)
            (selected / "material.xlsx").write_bytes(b"placeholder")

            server = scan_server_changes(config)["server"]
            manual_changes = [item for item in server["changes"] if item.get("manual_only")]

            self.assertTrue(manual_changes)
            self.assertTrue(any("临时订单文件夹" in item["message"] for item in manual_changes))


if __name__ == "__main__":
    unittest.main()

"""Traveler 解析、商品映射、库存检查及出库流程。

本模块将本地需求计算与浏览器自动化分开：解析和映射可离线执行，只有最终
确认的出库操作会写入库存系统；中央 SQLite 保存商品事实、映射、操作状态
及已核验结果，用于进程或浏览器故障后的恢复。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import threading
import time
from urllib.parse import urlsplit
from urllib.request import urlopen
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field, replace
from datetime import datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .core import (
    Config,
    RuleError,
    _normalize_name,
    _text,
    factory_name_order_mismatch,
    progress,
)
from .operation_log import configure_operation_log, log_database_statement, log_progress_payload, safe_exception_details, safe_rule_error_text
from .database import (
    catalog_material_attributes,
    connect_database,
    enable_foreign_keys,
    ensure_outbound_document_factory_links,
    ensure_schema,
)


TRAVELER_RE = re.compile(r"^Work Order Traveler\(.+\)\.xlsx$", re.IGNORECASE)
PP_FOLDER_RE = re.compile(r"^(?:PP\d{4}(?:-\d+)?|CS\d{3})$", re.IGNORECASE)
TB_RE = re.compile(r"^TB(\d+)$", re.IGNORECASE)
PANEL_RE = re.compile(r"^(\d+(?:\.\d+)?)mm--(.+)$", re.IGNORECASE)
EDGE_RE = re.compile(r"^Edge banding-+(.+)$", re.IGNORECASE)
REQUIRED_PRODUCT_HEADERS = {"商品类别", "*商品编号", "商品名称", "规格型号", "状态"}
INVENTORY_DATABASE_VERSION = 1
INVENTORY_CDP_DEFAULT_ENDPOINT = "http://127.0.0.1:9222"
INVENTORY_LOGIN_URL = "https://www.jdy.com/login/"
# 登录后会根据租户、区域或产品版本跳转到不同的 jdy.com 子域名，
# 例如 vip2-hz.jdy.com/default-new.jsp，而不是固定的 /workbench/ 路径。
INVENTORY_WORKBENCH_PREFIXES = (
    "https://service.jdy.com/workbench/",
    "http://service.jdy.com/workbench/",
    "https://www.jdy.com/workbench/",
    "http://www.jdy.com/workbench/",
)
INVENTORY_DOMAIN_SUFFIX = ".jdy.com"
# 这是每张单据的浏览器子进程超时保护；App 另有更长的无活动监控时限，避免多单操作仅因首单耗时就被取消。
INVENTORY_DOCUMENT_TIMEOUT_SECONDS = 90

# 以下是跨报表来源代码稳定的五金默认显示名，以规范库存 SKU 为键，使 Hinge/TestFullHinge 及左右导轨别名共用名称；不参与材料匹配或出库载荷构造。
HARDWARE_PRODUCT_DISPLAY_NAMES = {
    "M1001": "Hinge",
    "M1002": "H-Rail",
    "M1003": "L-Rail",
    "M1013": "Shelf Holder",
}


@dataclass
class TravelerItem:
    row: int
    section: str
    name: str
    quantity: float
    document_remark: str = ""
    product_code: str = ""

    def source_snapshot(self) -> dict:
        """序列化用于原始指纹的五个历史来源字段。规范 SKU 独立保存，不改变旧 Traveler 快照、出库原始指纹或待处理操作载荷。

        参数：self：当前实例。
        """
        return {
            "row": self.row,
            "section": self.section,
            "name": self.name,
            "quantity": self.quantity,
            "document_remark": self.document_remark,
        }


def _with_product_code(item: TravelerItem, product_code: str) -> TravelerItem:
    """给来源条目绑定规范 SKU 并返回该条目，保持历史序列化来源身份不变。显式字段使 SKU 在 dataclasses.replace 等正常复制中仍得以保留。

    参数：item：当前来源材料或五金条目；product_code：已确认的库存商品 SKU。
    """
    item.product_code = str(product_code or "").strip().upper()
    return item


@dataclass
class TravelerData:
    path: Path
    pp_folder: str
    order_id: str
    order_name: str
    items: list[TravelerItem]
    zero_items: list[TravelerItem]
    documents: dict[str, list[TravelerItem]]
    modified_at: str
    fingerprint: str

    def content_snapshot(self) -> dict:
        """按单据备注排序生成来源快照，包含正数量与零数量条目。

        参数：self：当前实例。
        """
        return {
            "order_id": self.order_id,
            "documents": {
                remark: [item.source_snapshot() for item in items]
                for remark, items in sorted(self.documents.items())
            },
            "items": [item.source_snapshot() for item in self.items],
            "zero_items": [item.source_snapshot() for item in self.zero_items],
        }


@dataclass
class Product:
    category: str
    code: str
    name: str
    spec: str
    status: str
    remark: str = ""
    unit: str = ""
    cost_price: float | None = None
    brand: str = ""
    material_kind: str = ""
    material_color: str = ""
    material_thickness: str = ""
    catalog_present: bool = True


@dataclass
class OutboundItem:
    traveler_name: str
    product_code: str
    product_name: str
    quantity: float
    section: str
    match_source: str
    document_remark: str = ""
    unit: str = ""


@dataclass
class InventoryPreview:
    traveler: TravelerData
    outbound_items: list[OutboundItem]
    ignored_items: list[dict] = field(default_factory=list)
    missing_items: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    selected_document_remarks: tuple[str, ...] = ()
    # 订单级材料单可覆盖一个或多个明确选择的工厂单；预览保留该选择，成功持久化时仅关联对应身份，不能把订单级记录隐式扩散到拆分订单。
    selected_factory_orders: tuple[str, ...] = ()
    source_type: str = "traveler"
    scope_decisions: list[dict] = field(default_factory=list)
    excluded_items: list[dict] = field(default_factory=list)
    no_outbound_required: bool = False
    # 房间预览只覆盖订单的一个工厂房间片段；保存局部范围时，不得将已有订单级出库记录视为已消失单据。
    partial_scope: bool = False
    # 部分流程只处理混合订单的一种单据；直接工厂出货只处理五金，不应把预览与早先生产的订单级材料记录比较。
    document_kinds: tuple[str, ...] = ()

    @property
    def ready(self) -> bool:
        """判断预览是否已解决全部缺失映射。

        参数：self：当前实例。
        """
        return not self.missing_items

    def payload(self) -> dict:
        """组装可序列化的出库预览，包含来源、映射问题、范围决定及单据明细。

        参数：self：当前实例。
        """
        documents = self.document_payloads()
        selected = self._selected_document_set()
        return {
            "source_type": self.source_type,
            "traveler": {
                "path": str(self.traveler.path),
                "pp_folder": self.traveler.pp_folder,
                "order_id": self.traveler.order_id,
                "order_name": self.traveler.order_name,
                "modified_at": self.traveler.modified_at,
                "fingerprint": self.traveler.fingerprint,
            },
            "outbound_items": [asdict(item) for item in self.outbound_items],
            # 结构化结果保留零数量以便诊断，但 Swift 预览有意不将其渲染为出库行。
            "zero_items": [
                item.source_snapshot() for item in self.traveler.zero_items
                if not selected or _normalize_name(item.document_remark) in selected
            ],
            "ignored_items": self.ignored_items,
            "missing_items": self.missing_items,
            "scope_decisions": self.scope_decisions,
            "excluded_items": self.excluded_items,
            "no_outbound_required": self.no_outbound_required,
            "warnings": self.warnings,
            "documents": documents,
            "ready": self.ready,
        }

    def _selected_document_set(self) -> set[str]:
        """规范化所选单据备注；选择工厂单时同时纳入订单级材料单。

        参数：self：当前实例。
        """
        selected = {
            _normalize_name(remark)
            for remark in self.selected_document_remarks
            if _normalize_name(remark)
        }
        if selected:
            selected.add(_normalize_name(self.traveler.order_id))
        return selected

    def _document_is_selected(self, remark: str) -> bool:
        """检查单据备注是否在当前选择范围，未指定范围时全部纳入。

        参数：self：当前实例；remark：用于单据识别的备注。
        """
        selected = self._selected_document_set()
        return not selected or _normalize_name(remark) in selected

    def document_payloads(self) -> list[dict]:
        """按来源单据备注分组已映射和忽略条目，并标明单据类型。

        参数：self：当前实例。
        """
        mapped: dict[str, list[OutboundItem]] = {}
        for item in self.outbound_items:
            mapped.setdefault(item.document_remark, []).append(item)
        ignored_by_remark: dict[str, list[dict]] = {}
        for item in self.ignored_items:
            ignored_by_remark.setdefault(str(item.get("document_remark", "")), []).append(item)
        return [
            {
                "remark": remark,
                "kind": (
                    "materials" if all(item.section == "板材与封边" for item in mapped.get(remark, []))
                    else "hardware" if all(item.section == "五金" for item in mapped.get(remark, []))
                    else "materials_hardware"
                ),
                "items": [asdict(item) for item in mapped.get(remark, [])],
                "ignored_items": ignored_by_remark.get(remark, []),
            }
            for remark in self.traveler.documents
            if self._document_is_selected(remark)
        ]


def stock_requirements(preview: InventoryPreview, include_hardware: bool = False) -> list[dict]:
    """按 SKU 汇总实时库存需求；默认只覆盖订单级材料。显式包含五金时跨工厂单合并，因为库存余额按 SKU 返回；本步骤不写出库。

    参数：preview：已解析 SKU 及出库范围的预览；include_hardware：是否同时统计五金；默认仅板材和封边。
    """
    if not preview.ready:
        raise RuleError(
            "inventory_precheck",
            "Traveler 存在未映射材料，不能查询完整库存",
            missing_items=preview.missing_items,
        )
    selected = [
        item for item in preview.outbound_items
        if include_hardware or item.section == "板材与封边"
    ]
    return _group_stock_requirements(selected)


def _group_stock_requirements(items: Iterable[OutboundItem]) -> list[dict]:
    """按 SKU 汇总需求数量并保留来源名称及单位。

    参数：items：待汇总、解析或计算指纹的材料/五金条目。
    """
    grouped: dict[str, dict] = {}
    for item in items:
        row = grouped.setdefault(item.product_code, {
            "productCode": item.product_code,
            "productName": item.product_name,
            "unit": item.unit,
            "requiredQuantity": 0.0,
            "travelerNames": [],
        })
        row["requiredQuantity"] += float(item.quantity)
        if item.traveler_name not in row["travelerNames"]:
            row["travelerNames"].append(item.traveler_name)
    return [grouped[code] for code in sorted(grouped)]


def database_stock_requirements(config: Config, order_id: str) -> tuple[str, list[dict]]:
    """仅根据 SQLite 材料事实生成订单级库存需求，未映射时阻止查询。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号。
    """
    normalized_order_id = order_id.strip().upper()
    detail, _ = _database_factory_rows(config, normalized_order_id)
    if not detail.get("materials"):
        raise RuleError(
            "inventory_database_missing",
            f"数据库中没有订单 {normalized_order_id} 的材料明细",
        )
    _, documents, _, _ = database_document_items(config, normalized_order_id)
    mappings = InventoryMappings(config.workflow_database)
    outbound: list[OutboundItem] = []
    missing: list[dict] = []
    with ProductDatabase(bootstrap_product_database(config)) as catalog:
        for item in documents.get(normalized_order_id, []):
            try:
                outbound.extend(match_item(catalog, mappings, item))
            except RuleError as exc:
                missing.append({"name": item.name, "message": str(exc), **exc.context})
    if missing:
        raise RuleError(
            "inventory_precheck",
            "当前订单存在无法映射到库存商品的板材或封边，不能查询完整库存",
            missing_items=missing,
        )
    return normalized_order_id, _group_stock_requirements(outbound)


def order_stock_requirements(config: Config, order_folder: Path) -> tuple[str, list[dict]]:
    """读取当前订单来源材料预览并解析为库存 SKU 需求。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_folder：订单来源文件夹。
    """
    database_order_id = order_folder.name.strip().upper()
    if database_order_id:
        try:
            return database_stock_requirements(config, database_order_id)
        except RuleError as exc:
            if exc.code != "inventory_database_missing":
                raise
        except (OSError, sqlite3.Error):
            pass

    from .order_workflow import preview_order

    preview = preview_order(config, order_folder)
    mappings = InventoryMappings(config.workflow_database)
    outbound: list[OutboundItem] = []
    missing: list[dict] = []
    plywood_names = {18.0: "18mm--Plywood", 14.5: "14.5mm--Plywood", 5.4: "5.4mm--Plywood"}
    source_items: list[TravelerItem] = []
    for index, material in enumerate(preview.materials, start=1):
        if material.quantity <= 0:
            continue
        if material.kind == "plywood":
            name = next(
                (label for thickness, label in plywood_names.items() if abs(material.thickness - thickness) < 0.01),
                f"{material.thickness:g}mm--Plywood",
            )
        else:
            name = f"{material.thickness:g}mm--{material.color}"
        source_items.append(TravelerItem(index, "板材与封边", name, material.quantity, preview.order_id))
    for color, quantity in preview.edge_banding.items():
        if quantity > 0:
            source_items.append(TravelerItem(
                len(source_items) + 1,
                "板材与封边",
                f"Edge banding--{color}",
                quantity,
                preview.order_id,
            ))
    with ProductDatabase(bootstrap_product_database(config)) as catalog:
        for item in source_items:
            try:
                outbound.extend(match_item(catalog, mappings, item))
            except RuleError as exc:
                missing.append({"name": item.name, "message": str(exc), **exc.context})
    if missing:
        raise RuleError(
            "inventory_precheck",
            "当前订单存在无法映射到库存商品的板材或封边，不能查询完整库存",
            missing_items=missing,
        )
    return preview.order_id, _group_stock_requirements(outbound)


def _database_factory_rows(config: Config, order_id: str) -> tuple[dict, dict[str, dict]]:
    """读取订单详情并按非空编号建立工厂单索引，空订单号时报错。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号。
    """
    from .order_details import order_detail

    normalized_order_id = order_id.strip().upper()
    if not normalized_order_id:
        raise RuleError("inventory_argument", "数据库出库需要订单号")
    detail = order_detail(config, normalized_order_id)
    factory_rows = {
        str(row.get("factory_order", "")).strip().upper(): row
        for row in detail.get("factory_orders", [])
        if str(row.get("factory_order", "")).strip()
    }
    return detail, factory_rows


def _usable_hardware_factory_orders(detail: dict) -> set[str]:
    """从订单详情提取含正数量五金事实的工厂单集合。

    参数：detail：订单详情及材料五金事实。
    """
    result: set[str] = set()
    for row in detail.get("hardware", []):
        factory_order = str(row.get("factory_order", "")).strip().upper()
        if not factory_order:
            continue
        try:
            quantity = float(row.get("quantity", 0) or 0)
        except (TypeError, ValueError):
            quantity = 0
        if quantity > 0:
            result.add(factory_order)
    return result


def database_outbound_fingerprint(config: Config, order_id: str, factory_order: str) -> str:
    """计算指定工厂单对应的持久出库来源事实指纹。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号；factory_order：工厂单编号。
    """
    normalized_order_id = order_id.strip().upper()
    normalized_factory_order = factory_order.strip().upper()
    detail, factory_rows = _database_factory_rows(config, normalized_order_id)
    if normalized_factory_order not in factory_rows:
        raise RuleError(
            "inventory_factory_unknown",
            f"数据库中找不到订单 {normalized_order_id} 的工厂单：{normalized_factory_order or '空白'}",
        )
    scope = outbound_scope_decisions(config, normalized_order_id, [normalized_factory_order])
    hardware = [
        row for row in detail.get("hardware", [])
        if str(row.get("factory_order", "")).strip().upper() == normalized_factory_order
    ]
    return _fingerprint({
        "order_id": normalized_order_id,
        "factory_order": normalized_factory_order,
        "material_scope": scope["material"],
        "materials": detail.get("materials", []),
        "hardware_scope": scope["hardware"].get(normalized_factory_order, {}),
        "hardware": hardware,
    })


def mark_customer_supplied_outbound(
    config: Config,
    order_id: str,
    selected_factory_orders: Iterable[str] | None = None,
) -> dict:
    """为客户提供材料的订单登记所选工厂单完成状态，不创建库存出库单。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号；selected_factory_orders：明确选择的工厂单编号集合；为空时按订单范围处理。
    """
    normalized_order_id = order_id.strip().upper()
    detail, factory_rows = _database_factory_rows(config, normalized_order_id)
    scope = outbound_scope_decisions(config, normalized_order_id, selected_factory_orders)
    if scope["material"]["requirement"] != "customer_supplied":
        raise RuleError("inventory_scope", "当前订单没有确认“客户提供材料（不入库存）”")
    requested = {
        str(value).strip().upper()
        for value in (selected_factory_orders or [])
        if str(value).strip()
    }
    factory_ids = requested or set(factory_rows)
    unknown = sorted(factory_ids - set(factory_rows))
    if unknown:
        raise RuleError(
            "inventory_factory_unknown",
            "数据库中找不到所选工厂单，已停止出库：" + "、".join(unknown),
            factory_orders=unknown,
        )
    if not factory_ids:
        raise RuleError("inventory_factory_unknown", f"订单 {normalized_order_id} 没有可确认出库的工厂单")
    available_hardware = _usable_hardware_factory_orders(detail)
    hardware_in_selection = sorted(available_hardware.intersection(factory_ids))
    if hardware_in_selection:
        raise RuleError(
            "inventory_scope",
            "所选工厂单存在五金数据，不能只写数据库：" + "、".join(hardware_in_selection),
            factory_orders=hardware_in_selection,
        )

    fingerprints = {
        factory_order: database_outbound_fingerprint(config, normalized_order_id, factory_order)
        for factory_order in sorted(factory_ids)
    }
    connection = connect_database(config.workflow_database)
    try:
        placeholders = ",".join("?" for _ in factory_ids)
        rows = connection.execute(
            f"select factory_order, case when stage='已出货' then '已出库' else '未出库' end as outbound_status from factory_orders "
            f"where order_id=? and factory_order in ({placeholders}) and aimes_status='active'",
            [normalized_order_id, *sorted(factory_ids)],
        ).fetchall()
        by_factory = {str(row[0]).upper(): str(row[1] or "") for row in rows}
        missing = sorted(factory_ids - set(by_factory))
        if missing:
            raise RuleError(
                "inventory_factory_unknown",
                "数据库中找不到所选工厂单，已停止出库：" + "、".join(missing),
                factory_orders=missing,
            )
        already_outbound = sorted(
            factory_order for factory_order, status in by_factory.items()
            if status == "已出库"
        )
        if already_outbound:
            raise RuleError(
                "inventory_already_outbound",
                "工厂单 " + "、".join(already_outbound) + " 已出库，不能重复出库",
                factory_orders=already_outbound,
            )
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        for factory_order, fingerprint in fingerprints.items():
            connection.execute(
                """
                update factory_orders
                set stage='已出货', outbound_document='',
                    outbound_mode='customer_supplied', outbound_fingerprint=?, updated_at=?
                where order_id=? and factory_order=? and aimes_status='active'
                """,
                (fingerprint, now, normalized_order_id, factory_order),
            )
        connection.commit()
    finally:
        connection.close()
    return {
        "order_id": normalized_order_id,
        "factory_orders": sorted(factory_ids),
        "outbound_status": "已出库",
        "outbound_mode": "customer_supplied",
        "inventory_document": False,
    }


def mark_no_hardware_outbound(
    config: Config,
    order_id: str,
    selected_factory_orders: Iterable[str] | None = None,
) -> dict:
    """为没有正数量有效五金的所选工厂单登记出货状态。出货是工作流状态，五金出库单是可选副作用；存在未映射五金时仍须预检失败，不能直接标为已出货。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号；selected_factory_orders：明确选择的工厂单编号集合；为空时按订单范围处理。
    """
    normalized_order_id = order_id.strip().upper()
    _, factory_rows = _database_factory_rows(config, normalized_order_id)
    requested = {
        str(value).strip().upper()
        for value in (selected_factory_orders or [])
        if str(value).strip()
    }
    factory_ids = requested or set(factory_rows)
    unknown = sorted(factory_ids - set(factory_rows))
    if unknown:
        raise RuleError(
            "inventory_factory_unknown",
            "数据库中找不到所选工厂单，已停止出货：" + "、".join(unknown),
            factory_orders=unknown,
        )
    if not factory_ids:
        raise RuleError("inventory_factory_unknown", f"订单 {normalized_order_id} 没有可确认出货的工厂单")
    detail, _ = _database_factory_rows(config, normalized_order_id)
    hardware_in_selection = _usable_hardware_factory_orders(detail).intersection(factory_ids)
    if hardware_in_selection:
        raise RuleError(
            "inventory_scope",
            "所选工厂单存在五金数据，不能按无五金只更新状态：" + "、".join(sorted(hardware_in_selection)),
            factory_orders=sorted(hardware_in_selection),
        )

    fingerprints = {
        factory_order: database_outbound_fingerprint(config, normalized_order_id, factory_order)
        for factory_order in sorted(factory_ids)
    }
    connection = connect_database(config.workflow_database)
    try:
        placeholders = ",".join("?" for _ in factory_ids)
        rows = connection.execute(
            f"select factory_order, case when stage='已出货' then '已出库' else '未出库' end as outbound_status from factory_orders "
            f"where order_id=? and factory_order in ({placeholders}) and aimes_status='active'",
            [normalized_order_id, *sorted(factory_ids)],
        ).fetchall()
        by_factory = {str(row[0]).upper(): str(row[1] or "") for row in rows}
        missing = sorted(factory_ids - set(by_factory))
        if missing:
            raise RuleError(
                "inventory_factory_unknown",
                "数据库中找不到所选工厂单，已停止出货：" + "、".join(missing),
                factory_orders=missing,
            )
        already_outbound = sorted(
            factory_order for factory_order, status in by_factory.items()
            if status == "已出库"
        )
        if already_outbound:
            raise RuleError(
                "inventory_already_outbound",
                "工厂单 " + "、".join(already_outbound) + " 已出库，不能重复出货",
                factory_orders=already_outbound,
            )
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        for factory_order, fingerprint in fingerprints.items():
            connection.execute(
                """
                update factory_orders
                set stage='已出货', outbound_document='',
                    outbound_mode='no_hardware', outbound_fingerprint=?, updated_at=?
                where order_id=? and factory_order=? and aimes_status='active'
                """,
                (fingerprint, now, normalized_order_id, factory_order),
            )
        connection.commit()
    finally:
        connection.close()
    return {
        "order_id": normalized_order_id,
        "factory_orders": sorted(factory_ids),
        "outbound_status": "已出库",
        "outbound_mode": "no_hardware",
        "inventory_document": False,
    }


def database_document_items(
    config: Config,
    order_id: str,
    selected_factory_orders: Iterable[str] | None = None,
    *,
    production_request_id: str = "",
    production_materials: Iterable[dict] | None = None,
    shipment_only: bool = False,
) -> tuple[str, dict[str, list[TravelerItem]], list[TravelerItem], dict[str, dict]]:
    """直接从持久事实组织出库来源条目，不读写 Traveler。材料按订单号分单，五金按工厂名称分单，并沿用 Traveler 的规范名称以定位历史出库记录。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号；selected_factory_orders：明确选择的工厂单编号集合；为空时按订单范围处理；production_request_id：本次生产请求的幂等标识；production_materials：本次明确指定的生产材料；为空时使用订单材料；shipment_only：是否只组织发货五金而排除订单材料。
    """
    detail, factory_rows = _database_factory_rows(config, order_id)
    normalized_order_id = order_id.strip().upper()
    requested = {
        str(value).strip().upper()
        for value in (selected_factory_orders or [])
        if str(value).strip()
    }
    unknown = sorted(requested - set(factory_rows))
    if unknown:
        raise RuleError(
            "inventory_factory_unknown",
            "数据库中找不到所选工厂单，已停止出库：" + "、".join(unknown),
            factory_orders=unknown,
        )
    invalid_names = []
    for factory_order, factory in factory_rows.items():
        if requested and factory_order not in requested:
            continue
        mismatch = factory_name_order_mismatch(
            str(factory.get("factory_name", "")),
            str(factory.get("sales_order_name", "")),
        )
        if mismatch:
            prefix, order = mismatch
            invalid_names.append(
                f"{factory_order}（名称前缀 {prefix}，销售单 {order}）"
            )
    if invalid_names:
        raise RuleError(
            "factory_name_order_mismatch",
            "工厂单名称与销售单号不一致，已停止生成出库备注，请先修正 AIMES 后重新获取："
            + "、".join(sorted(invalid_names)),
            factory_orders=[item.split("（", 1)[0] for item in sorted(invalid_names)],
        )

    order_type = str(detail.get("order", {}).get("order_type", "")).strip()
    scope = outbound_scope_decisions(config, normalized_order_id, requested)
    material_requirement = scope["material"]["requirement"]
    if production_request_id and production_materials is None:
        raise RuleError("production_materials", "生产请求必须包含本次确认的材料明细，请重新预览")
    production_mode = production_materials is not None or bool(production_request_id)
    materials = [] if shipment_only else list(detail.get("materials", []))
    if production_materials is not None:
        materials = [dict(row) for row in production_materials]
    documents: dict[str, list[TravelerItem]] = {normalized_order_id: []}
    zero_items: list[TravelerItem] = []
    row_number = 1
    if material_requirement != "customer_supplied" and material_requirement not in {"remainder", "not_required"}:
        material_rows = materials
    else:
        material_rows = []
    for row in material_rows:
        product_code = str(row.get("product_code", "") or "").strip().upper()
        if not product_code:
            raise RuleError(
                "inventory_material_sku",
                "数据库材料缺少已确认商品 SKU，不能按名称重新匹配",
            )
        kind = str(row.get("material_type", "")).strip().casefold()
        color = str(row.get("color", "")).strip()
        thickness = float(str(row.get("thickness", "0") or "0"))
        if kind == "plywood":
            name = f"{thickness:g}mm--Plywood"
        elif kind in {"panel", "back"}:
            name = f"{thickness:g}mm--{color}"
        elif kind == "edge":
            name = f"Edge banding--{color}"
        else:
            raise RuleError(
                "inventory_material",
                f"数据库中的材料类型暂不支持：{kind or '空白'}",
            )
        quantity = float(row.get("quantity", 0) or 0)
        item = _with_product_code(
            TravelerItem(
                row_number, "板材与封边", name, quantity, normalized_order_id
            ),
            product_code,
        )
        row_number += 1
        if quantity <= 0:
            zero_items.append(item)
        else:
            documents[normalized_order_id].append(item)

    for row in ([] if production_mode else detail.get("hardware", [])):
        factory_order = str(row.get("factory_order", "")).strip().upper()
        if requested and factory_order not in requested:
            continue
        hardware_requirement = scope["hardware"].get(factory_order, {"requirement": "required"})["requirement"]
        if hardware_requirement in {"remainder", "not_required"}:
            continue
        factory = factory_rows.get(factory_order, {})
        remark = str(factory.get("factory_name", "")).strip() or factory_order
        documents.setdefault(remark, [])
        product_code = str(row.get("product_code", "")).strip()
        name = product_code
        item = _with_product_code(
            TravelerItem(
                row_number, "五金", name, float(row.get("quantity", 0) or 0), remark
            ),
            product_code,
        )
        quantity = item.quantity
        row_number += 1
        if quantity <= 0:
            zero_items.append(item)
        else:
            documents[remark].append(item)

    if not any(documents.values()):
        explicit_no_outbound = material_requirement in {"customer_supplied", "remainder", "not_required"}
        if requested:
            explicit_no_outbound = explicit_no_outbound or all(
                scope["hardware"].get(factory_order, {"requirement": "required"})["requirement"]
                in {"remainder", "not_required"}
                for factory_order in requested
            )
        if shipment_only:
            explicit_no_outbound = True
        if not explicit_no_outbound:
            raise RuleError("inventory_empty", f"数据库中没有订单 {normalized_order_id} 可出库的材料或五金")
    return normalized_order_id, documents, zero_items, factory_rows


def outbound_scope_decisions(
    config: Config,
    order_id: str,
    selected_factory_orders: Iterable[str] | None = None,
) -> dict:
    """读取显式出库范围决定而不修改来源事实。缺少决定时保守处理：已有事实视为需要出库，无事实的订单保持未决，不能自动归为余料生产。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号；selected_factory_orders：明确选择的工厂单编号集合；为空时按订单范围处理。
    """
    detail, factory_rows = _database_factory_rows(config, order_id)
    normalized = order_id.strip().upper()
    requested = {
        str(value).strip().upper()
        for value in (selected_factory_orders or [])
        if str(value).strip()
    }
    rows = connect_database(config.workflow_database)
    rows.row_factory = sqlite3.Row
    try:
        decisions = rows.execute(
            "select id, scope_type, factory_order, requirement, reason, source_fingerprint, updated_at "
            "from outbound_scope_decisions where order_id=?",
            (normalized,),
        ).fetchall()
    finally:
        rows.close()
    by_key = {(row["scope_type"], row["factory_order"]): dict(row) for row in decisions}
    material_exists = bool(detail.get("materials"))
    material = by_key.get(("material", "")) or {
        "scope_type": "material",
        "factory_order": "",
        "requirement": "required" if material_exists else "pending",
        "reason": "",
        "source_fingerprint": "",
        "updated_at": "",
    }
    available_hardware = _usable_hardware_factory_orders(detail)
    factory_ids = requested or set(factory_rows)
    hardware = {}
    for factory_order in sorted(factory_ids):
        if factory_order not in factory_rows or factory_order not in available_hardware:
            continue
        hardware[factory_order] = by_key.get(("hardware", factory_order)) or {
            "scope_type": "hardware",
            "factory_order": factory_order,
            "requirement": "required",
            "reason": "",
            "source_fingerprint": "",
            "updated_at": "",
        }
    persisted_decisions = [
        dict(row)
        for row in [material, *hardware.values()]
        if str(row.get("updated_at") or "").strip()
    ]
    last_decision = max(
        persisted_decisions,
        key=lambda row: (str(row.get("updated_at", "")), int(row.get("id", 0) or 0)),
        default=None,
    )
    return {
        "order_id": normalized,
        "order_type": str(detail.get("order", {}).get("order_type", "")),
        "material": material,
        "hardware": hardware,
        "decisions": [material, *hardware.values()],
        "last_decision": last_decision,
    }


def set_outbound_scope(
    config: Config,
    order_id: str,
    scope_type: str,
    requirement: str,
    *,
    factory_order: str = "",
    reason: str = "",
) -> dict:
    """校验订单类型、材料或五金范围及理由，保存明确的出库范围决定。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号；scope_type：material 材料范围或 hardware 五金范围；requirement：需要出库、客户提供、余料或无需出库的明确决定；factory_order：工厂单编号；reason：忽略规则或不出库决定的原因。
    """
    normalized = order_id.strip().upper()
    scope_type = scope_type.strip().lower()
    requirement = requirement.strip().lower()
    factory_order = factory_order.strip().upper()
    allowed = {"required", "customer_supplied", "remainder", "not_required"}
    if scope_type not in {"material", "hardware"} or requirement not in allowed:
        raise RuleError("inventory_scope", "出库范围类型或决定无效")
    detail, factory_rows = _database_factory_rows(config, normalized)
    order_type = str(detail.get("order", {}).get("order_type", "")).strip()
    if order_type == "owned":
        raise RuleError("inventory_scope", "自有订单不支持设置出库范围")
    if scope_type == "material" and factory_order:
        raise RuleError("inventory_scope", "订单材料出库范围不能填写工厂单")
    if scope_type == "hardware":
        if not factory_order or factory_order not in factory_rows:
            raise RuleError("inventory_scope", f"数据库中找不到订单 {normalized} 的工厂单：{factory_order or '空白'}")
        if factory_order not in _usable_hardware_factory_orders(detail):
            raise RuleError(
                "inventory_scope",
                f"数据库中没有订单 {normalized} 的工厂单 {factory_order} 的可出库五金数据",
            )
        if requirement == "customer_supplied":
            raise RuleError("inventory_scope", "客户提供材料只适用于来料加工订单的板材和封边，不适用于五金")
    if scope_type == "material" and requirement == "customer_supplied" and order_type != "cutToSize":
        raise RuleError("inventory_scope", "只有来料加工订单支持“客户提供，不出库”")
    if requirement in {"customer_supplied", "remainder", "not_required"} and not reason.strip():
        raise RuleError("inventory_scope_reason", "不出库决定必须填写原因")
    now = datetime.now().astimezone().isoformat(timespec="seconds")
    connection = connect_database(config.workflow_database)
    try:
        connection.execute(
            """insert into outbound_scope_decisions(
                order_id, scope_type, factory_order, requirement, reason,
                source_fingerprint, created_at, updated_at
            ) values(?,?,?,?,?,?,?,?)
            on conflict(order_id, scope_type, factory_order) do update set
                requirement=excluded.requirement, reason=excluded.reason,
                source_fingerprint=excluded.source_fingerprint, updated_at=excluded.updated_at""",
            (normalized, scope_type, factory_order, requirement, reason.strip(), "", now, now),
        )
        connection.commit()
    finally:
        connection.close()
    return outbound_scope_decisions(config, normalized, [factory_order] if factory_order else None)


def build_database_preview(
    config: Config,
    order_id: str,
    selected_factory_orders: Iterable[str] | None = None,
    *,
    production_request_id: str = "",
    production_materials: Iterable[dict] | None = None,
    shipment_only: bool = False,
) -> InventoryPreview:
    """根据 SQLite 事实生成 SKU 出库预览并收集映射冲突，不生成 Traveler 文件。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号；selected_factory_orders：明确选择的工厂单编号集合；为空时按订单范围处理；production_request_id：本次生产请求的幂等标识；production_materials：本次明确指定的生产材料；为空时使用订单材料；shipment_only：是否只组织发货五金而排除订单材料。
    """
    _assert_single_server_material_source(config, order_id)
    normalized_order_id, documents, zero_items, _ = database_document_items(
        config, order_id, selected_factory_orders,
        production_request_id=production_request_id,
        production_materials=production_materials,
        shipment_only=shipment_only,
    )
    selected_factory_ids = tuple(sorted({
        str(value).strip().upper()
        for value in (selected_factory_orders or [])
        if str(value).strip()
    }))
    scope = outbound_scope_decisions(config, normalized_order_id, selected_factory_orders)
    detail, factory_rows = _database_factory_rows(config, normalized_order_id)
    excluded_items: list[dict] = []
    material_requirement = scope["material"]["requirement"]
    if shipment_only:
        material_requirement = "required"
    if material_requirement in {"customer_supplied", "remainder", "not_required"}:
        label = {
            "customer_supplied": "客户提供",
            "remainder": "余料生产",
            "not_required": "不需要出库",
        }[material_requirement]
        for row in detail.get("materials", []):
            kind = str(row.get("material_type", "")).strip().casefold()
            color = str(row.get("color", "")).strip()
            thickness = float(str(row.get("thickness", "0") or "0"))
            name = f"{thickness:g}mm--Plywood" if kind == "plywood" else (
                f"{thickness:g}mm--{color}" if kind in {"panel", "back"}
                else f"Edge banding--{color}" if kind == "edge" else kind
            )
            excluded_items.append({
                "name": name,
                "quantity": float(row.get("quantity", 0) or 0),
                "section": "板材与封边",
                "document_remark": normalized_order_id,
                "status": label,
                "reason": scope["material"].get("reason", ""),
            })
    for row in detail.get("hardware", []):
        factory_order = str(row.get("factory_order", "")).strip().upper()
        if selected_factory_orders and factory_order not in {str(value).strip().upper() for value in selected_factory_orders}:
            continue
        decision = scope["hardware"].get(factory_order, {})
        requirement = decision.get("requirement")
        if requirement not in {"remainder", "not_required"}:
            continue
        excluded_items.append({
            "name": str(row.get("name", "")).strip() or str(row.get("product_code", "")).strip(),
            "quantity": float(row.get("quantity", 0) or 0),
            "section": "五金",
            "document_remark": str(factory_rows.get(factory_order, {}).get("factory_name", "")).strip() or factory_order,
            "status": "余料生产" if requirement == "remainder" else "不需要出库",
            "reason": decision.get("reason", ""),
        })
    mappings = InventoryMappings(config.workflow_database)
    selected_remarks = tuple(
        remark for remark in documents
        if remark != normalized_order_id
    )
    outbound: list[OutboundItem] = []
    ignored: list[dict] = []
    missing: list[dict] = []
    catalog_path = bootstrap_product_database(config)
    catalog_context = (
        ProductDatabase(catalog_path)
        if catalog_path.suffix.lower() in {".sqlite", ".sqlite3", ".db"}
        else None
    )
    try:
        catalog = catalog_context or ProductCatalog(catalog_path)
        for remark, items in documents.items():
            for item in items:
                reason = None if item.product_code else mappings.ignored_reason(item.name)
                if reason is not None:
                    ignored.append({**item.source_snapshot(), "reason": reason})
                    continue
                try:
                    outbound.extend(match_item(catalog, mappings, item))
                except RuleError as exc:
                    missing.append({**item.source_snapshot(), "code": exc.code, "message": str(exc), **exc.context})
    finally:
        if catalog_context is not None:
            catalog_context.close()

    duplicate_keys = [
        key for key, count in Counter(
            (item.document_remark, item.product_code) for item in outbound
        ).items()
        if count > 1 and _normalize_name(key[1]) not in {"M1068", "M1069"}
    ]
    if duplicate_keys:
        missing.append({
            "code": "duplicate_product_code",
            "message": "多个数据库订单项目映射到同一商品编号",
            "items": [
                asdict(item) for item in outbound
                if (item.document_remark, item.product_code) in duplicate_keys
            ],
        })

    snapshot = {
        "order_id": normalized_order_id,
        "documents": {
            remark: [item.source_snapshot() for item in items]
            for remark, items in sorted(documents.items())
        },
        "zero_items": [item.source_snapshot() for item in zero_items],
    }
    traveler = TravelerData(
        path=config.workflow_database.resolve(),
        pp_folder=normalized_order_id,
        order_id=normalized_order_id,
        order_name=normalized_order_id,
        items=[item for items in documents.values() for item in items],
        zero_items=zero_items,
        documents=documents,
        modified_at=datetime.now().isoformat(timespec="seconds"),
        fingerprint=_fingerprint(snapshot),
    )
    return InventoryPreview(
        traveler,
        outbound,
        ignored,
        missing,
        selected_document_remarks=selected_remarks,
        selected_factory_orders=selected_factory_ids,
        source_type="database",
        scope_decisions=scope["decisions"],
        excluded_items=excluded_items,
        no_outbound_required=not outbound and not missing and bool(
            excluded_items
            or scope["material"]["requirement"] in {"customer_supplied", "remainder", "not_required"}
            or shipment_only
        ),
        document_kinds=("hardware",) if shipment_only else (),
    )


def _assert_single_server_material_source(config: Config, order_id: str) -> None:
    """阻止同一订单同时使用多个基础 Server 材料根目录。补单按设计可叠加，但两份非补单工作簿可能是迁移残留，不能合计造成重复出库。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号。
    """
    normalized_order_id = str(order_id or "").strip().upper()
    if not normalized_order_id:
        return
    connection = connect_database(config.workflow_database)
    try:
        paths = [
            str(row[0] or "")
            for row in connection.execute(
                "select distinct source_path from material_items "
                "where order_id=? and source_type='aihouse' and source_path <> '' "
                "order by source_path",
                (normalized_order_id,),
            ).fetchall()
        ]
    finally:
        connection.close()
    base_paths = [
        path for path in paths
        if not any(
            part.casefold() == "recut" or part.casefold().endswith("-recut")
            for part in Path(path).parts[:-1]
        )
    ]
    if len(base_paths) > 1:
        raise RuleError(
            "duplicate_material_source",
            f"订单 {normalized_order_id} 存在多个 Server 材料来源，已停止出库以防重复扣库存："
            + "；".join(base_paths),
            order_id=normalized_order_id,
            source_paths=base_paths,
        )


def build_factory_room_preview(
    config: Config,
    order_id: str,
    factory_order: str,
) -> InventoryPreview:
    """为明确指定的 Server 房间生成局部出库预览。订单材料仍是 SQLite 订单级事实；仅读取用户已确认且工作簿明确归属的房间行，并加入所选工厂单五金，不从其他房间推算数量。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号；factory_order：工厂单编号。
    """
    from .order_workflow import parse_material_room_rows

    normalized_order = order_id.strip().upper()
    normalized_factory = factory_order.strip().upper()
    detail, factory_rows = _database_factory_rows(config, normalized_order)
    factory = factory_rows.get(normalized_factory)
    if factory is None:
        raise RuleError(
            "inventory_factory_unknown",
            f"数据库中找不到订单 {normalized_order} 的工厂单：{normalized_factory or '空白'}",
        )
    factory_name = str(factory.get("factory_name", "")).strip() or normalized_factory
    material_paths = sorted({
        Path(str(row.get("source_path", "")).strip())
        for row in detail.get("materials", [])
        if str(row.get("source_path", "")).strip()
    })
    material_path = next((path for path in material_paths if path.is_file()), None)
    if material_path is None:
        raise RuleError(
            "inventory_material_source",
            f"订单 {normalized_order} 找不到可读取的材料源文件，无法按房间出库",
            factory_order=normalized_factory,
        )
    target_key = _normalize_name(factory_name)
    room_rows = parse_material_room_rows(material_path)
    selected_rows = [
        row for row in room_rows
        if _normalize_name(str(row[0]).strip()) == target_key
    ]
    if not selected_rows:
        raise RuleError(
            "inventory_room_unknown",
            f"材料文件中找不到与工厂名称完全对应的房间：{factory_name}",
            factory_order=normalized_factory,
            source_path=str(material_path),
        )

    def material_identity(kind: str, color: str, thickness: object) -> tuple[str, str, float]:
        """规范化材料种类、颜色及厚度，生成房间材料与已确认事实的比较键。

        参数：kind：材料种类；color：材料颜色；thickness：材料标称厚度。
        """
        try:
            normalized_thickness = round(float(str(thickness or "0")), 4)
        except ValueError:
            normalized_thickness = 0.0
        return (
            str(kind or "").strip().casefold(),
            _normalize_name(str(color or "")),
            normalized_thickness,
        )

    material_codes: dict[tuple[str, str, float], set[str]] = defaultdict(set)
    for row in detail.get("materials", []):
        if Path(str(row.get("source_path", "")).strip()) != material_path:
            continue
        code = str(row.get("product_code", "") or "").strip().upper()
        if code:
            material_codes[material_identity(
                str(row.get("material_type", "")),
                str(row.get("color", "")),
                row.get("thickness", ""),
            )].add(code)

    def confirmed_material_code(kind: str, color: str, thickness: object) -> str:
        """按材料身份提取唯一已确认 SKU，不允许通过名称重新匹配。

        参数：kind：房间材料种类；color：材料颜色；thickness：材料标称厚度。
        """
        identity = material_identity(kind, color, thickness)
        codes = material_codes.get(identity, set())
        if len(codes) != 1:
            raise RuleError(
                "inventory_material_sku",
                "房间材料无法唯一对应已确认商品 SKU，不能按名称重新匹配",
                material_kind=identity[0],
                material_color=color,
                material_thickness=identity[2],
                product_codes=sorted(codes),
            )
        return next(iter(codes))

    documents: dict[str, list[TravelerItem]] = {factory_name: []}
    row_number = 1
    for _, items, edges in selected_rows:
        for material in items:
            if material.quantity <= 0:
                continue
            if material.kind == "plywood":
                name = f"{material.thickness:g}mm--Plywood"
            elif material.kind in {"panel", "back"}:
                name = f"{material.thickness:g}mm--{material.color}"
            elif material.kind == "edge":
                name = f"Edge banding--{material.color}"
            else:
                raise RuleError("inventory_material", f"数据库中的材料类型暂不支持：{material.kind or '空白'}")
            documents[factory_name].append(_with_product_code(
                TravelerItem(row_number, "板材与封边", name, float(material.quantity), factory_name),
                confirmed_material_code(material.kind, material.color, material.thickness),
            ))
            row_number += 1
        for color, quantity in edges.items():
            if quantity > 0:
                documents[factory_name].append(_with_product_code(
                    TravelerItem(row_number, "板材与封边", f"Edge banding--{color}", float(quantity), factory_name),
                    confirmed_material_code("edge", color, 0),
                ))
                row_number += 1

    for row in detail.get("hardware", []):
        if str(row.get("factory_order", "")).strip().upper() != normalized_factory:
            continue
        quantity = float(row.get("quantity", 0) or 0)
        if quantity <= 0:
            continue
        product_code = str(row.get("product_code", "")).strip()
        name = product_code
        documents[factory_name].append(_with_product_code(
            TravelerItem(row_number, "五金", name, quantity, factory_name),
            product_code,
        ))
        row_number += 1
    if not documents[factory_name]:
        raise RuleError(
            "inventory_empty",
            f"工厂单 {normalized_factory} 对应房间没有可出库的材料或五金",
        )

    mappings = InventoryMappings(config.workflow_database)
    outbound: list[OutboundItem] = []
    ignored: list[dict] = []
    missing: list[dict] = []
    catalog_path = bootstrap_product_database(config)
    catalog_context = (
        ProductDatabase(catalog_path)
        if catalog_path.suffix.lower() in {".sqlite", ".sqlite3", ".db"}
        else None
    )
    try:
        catalog = catalog_context or ProductCatalog(catalog_path)
        for item in documents[factory_name]:
            reason = None if item.product_code else mappings.ignored_reason(item.name)
            if reason is not None:
                ignored.append({**item.source_snapshot(), "reason": reason})
                continue
            try:
                outbound.extend(match_item(catalog, mappings, item))
            except RuleError as exc:
                missing.append({**item.source_snapshot(), "code": exc.code, "message": str(exc), **exc.context})
    finally:
        if catalog_context is not None:
            catalog_context.close()
    traveler = TravelerData(
        path=material_path.resolve(),
        pp_folder=Path(str(detail.get("order", {}).get("source_folder", ""))).name or normalized_order,
        order_id=normalized_order,
        order_name=normalized_order,
        items=list(documents[factory_name]),
        zero_items=[],
        documents=documents,
        modified_at=datetime.fromtimestamp(material_path.stat().st_mtime).isoformat(timespec="seconds"),
        fingerprint=_fingerprint({"factory_order": normalized_factory, "documents": {factory_name: [item.source_snapshot() for item in documents[factory_name]]}}),
    )
    return InventoryPreview(
        traveler=traveler,
        outbound_items=outbound,
        ignored_items=ignored,
        missing_items=missing,
        selected_document_remarks=(factory_name,),
        selected_factory_orders=(normalized_factory,),
        source_type="database_room",
        partial_scope=True,
    )


def changed_factory_orders_for_documents(
    config: Config,
    order_id: str,
    selected_factory_orders: Iterable[str],
    documents: Iterable[dict],
) -> set[str]:
    """将已变化单据的备注映射回所选工厂单编号。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号；selected_factory_orders：明确选择的工厂单编号集合；为空时按订单范围处理；documents：待反查工厂单的出库单据。
    """
    _, factory_rows = _database_factory_rows(config, order_id)
    changed_remarks = {
        _normalize_name(document.get("remark", ""))
        for document in documents
        if document.get("changed")
    }
    order_key = _normalize_name(order_id)
    material_changed = order_key in changed_remarks
    changed: set[str] = set()
    for value in selected_factory_orders:
        factory_order = str(value).strip().upper()
        factory_name = str(factory_rows.get(factory_order, {}).get("factory_name", "")).strip()
        if material_changed or _normalize_name(factory_order) in changed_remarks or _normalize_name(factory_name) in changed_remarks:
            changed.add(factory_order)
    return changed


def _fingerprint(payload: dict) -> str:
    """对排序且紧凑编码的 JSON 计算稳定 SHA-256 摘要。

    参数：payload：待序列化或登记的结构化载荷。
    """
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(encoded).hexdigest()


def _find_header(row: Iterable, label: str) -> int | None:
    """查找行中唯一匹配表头并返回一基列号，重复表头时报错。

    参数：row：工作表一行的单元格值；label：目标表头或错误提示中的字段名。
    """
    matches = [index for index, value in enumerate(row, 1) if _text(value) == label]
    if len(matches) > 1:
        raise RuleError("traveler_schema", f"Picking List 同一行出现多个“{label}”表头")
    return matches[0] if matches else None


def _order_name_candidates(ws) -> list[str]:
    """扫描工厂单名称标签右侧的首个非空值，收集名称候选。

    参数：ws：待读取的工作表。
    """
    candidates = []
    for row in ws.iter_rows():
        for index, cell in enumerate(row):
            if _text(cell.value) == "Name/工厂单名称":
                for neighbor in row[index + 1:]:
                    value = _text(neighbor.value)
                    if value:
                        candidates.append(value)
                        break
    return candidates


def _normalized_label(value) -> str:
    """去除表头空白并转小写，便于统一匹配。

    参数：value：待规范化或校验的原始值。
    """
    return re.sub(r"\s+", "", _text(value)).lower()


def _find_label_row(ws, label: str) -> int:
    """在工作表中查找规范化标签所在行，缺失时报错。

    参数：ws：待读取的工作表；label：目标表头或错误提示中的字段名。
    """
    wanted = _normalized_label(label)
    for row in range(1, ws.max_row + 1):
        if any(_normalized_label(ws.cell(row, col).value) == wanted for col in range(1, ws.max_column + 1)):
            return row
    raise RuleError("traveler_usage_schema", f"Usage List 缺少“{label}”")


def _displayed_integer(cell, label: str) -> float:
    """读取非负板材数量，按整数显示格式四舍五入；无法解释为整数时报错。

    参数：cell：含数值及显示格式的单元格；label：目标表头或错误提示中的字段名。
    """
    value = cell.value
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        raise RuleError("traveler_quantity", f"{label} 不是有效数字：{value}")
    number = float(value)
    if number < 0:
        raise RuleError("traveler_quantity", f"{label} 不能为负数：{number:g}")
    if abs(number - round(number)) <= 1e-9:
        return float(round(number))
    first_format = _text(cell.number_format).split(";", 1)[0]
    if first_format.lower() != "general" and "." not in first_format and re.search(r"[0#?]", first_format):
        return float(Decimal(str(number)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    raise RuleError("traveler_quantity", f"{label} 数量为 {number:g}，板材数量必须为整数")


def _nonnegative_number(cell, label: str) -> float:
    """读取有限非负数量，空单元格按零处理，非法值时报错。

    参数：cell：含数值及显示格式的单元格；label：目标表头或错误提示中的字段名。
    """
    value = cell.value
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        raise RuleError("traveler_quantity", f"{label} 不是有效数字：{value}")
    number = float(value)
    if number < 0:
        raise RuleError("traveler_quantity", f"{label} 不能为负数：{number:g}")
    return number


def _canonical_usage_color(value: str) -> str:
    """将 Usage List 的 Khaki 别名归一为 Penelope FA44，其他颜色原样返回。

    参数：value：待规范化或校验的原始值。
    """
    color = _text(value)
    if re.fullmatch(r"khaki(?:\s*\(7x9\))?", color, re.IGNORECASE):
        return "Penelope FA44"
    return color


def _usage_list_items(ws, order_id: str) -> tuple[list[TravelerItem], list[TravelerItem]]:
    """校验 Usage List 表头及颜色，汇总板材和封边并分别返回正数量与零数量条目。

    参数：ws：待读取的工作表；order_id：销售订单号。
    """
    total_row = _find_label_row(ws, "Total Qty:")
    header_map = {
        _normalized_label(ws.cell(2, col).value): col
        for col in range(1, ws.max_column + 1)
        if _text(ws.cell(2, col).value)
    }
    required_headers = (
        "3/4 Plywood",
        "5/8 Plywood",
        "1/4 Plywood",
        "3/4 Finish Panel",
        "1/4 Finish Panel",
        "Edge Banding (m)",
        "Color",
    )
    missing = [label for label in required_headers if _normalized_label(label) not in header_map]
    if missing:
        raise RuleError("traveler_usage_schema", f"Usage List 缺少列：{', '.join(missing)}")

    totals: dict[str, tuple[int, float, bool]] = {}

    def accumulate(row: int, name: str, value, integer: bool) -> None:
        """校验非负数量并按材料名称累计，保留首次行号及整数数量标志。

        参数：row：待读取的行或来源行号；name：来源材料名称或查询名称；value：待规范化或校验的原始值；integer：是否在最终汇总时按整数处理数量。
        """
        if value in (None, ""):
            number = 0.0
        elif isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
            raise RuleError("traveler_quantity", f"Usage List 第 {row} 行 {name} 不是有效数字：{value}")
        else:
            number = float(value)
        if number < 0:
            raise RuleError("traveler_quantity", f"Usage List 第 {row} 行 {name} 不能为负数：{number:g}")
        first_row, current, current_integer = totals.get(name, (row, 0.0, integer))
        totals[name] = (first_row, current + number, current_integer)

    plywood_columns = (
        ("3/4 Plywood", "18mm--Plywood"),
        ("5/8 Plywood", "14.5mm--Plywood"),
        ("1/4 Plywood", "5.4mm--Plywood"),
    )
    panel_columns = (
        ("3/4 Finish Panel", "19.1mm"),
        ("1/4 Finish Panel", "8mm"),
    )
    color_col = header_map[_normalized_label("Color")]
    edge_col = header_map[_normalized_label("Edge Banding (m)")]
    for row in range(3, total_row):
        for label, name in plywood_columns:
            accumulate(row, name, ws.cell(row, header_map[_normalized_label(label)]).value, True)
        color = _canonical_usage_color(ws.cell(row, color_col).value)
        has_colored_material = any(
            ws.cell(row, header_map[_normalized_label(label)]).value not in (None, "", 0, 0.0)
            for label, _ in panel_columns
        ) or ws.cell(row, edge_col).value not in (None, "", 0, 0.0)
        if has_colored_material and not color:
            raise RuleError("traveler_usage_schema", f"Usage List 第 {row} 行有板材或封边数量但缺少颜色")
        if not color:
            continue
        for label, thickness in panel_columns:
            accumulate(
                row,
                f"{thickness}--{color}",
                ws.cell(row, header_map[_normalized_label(label)]).value,
                True,
            )
        accumulate(row, f"Edge banding--{color}", ws.cell(row, edge_col).value, False)

    positive: list[TravelerItem] = []
    zero: list[TravelerItem] = []
    for name, (row, quantity, integer) in totals.items():
        if integer:
            quantity = float(Decimal(str(quantity)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        item = TravelerItem(row, "板材与封边", name, quantity, order_id)
        (positive if quantity > 0 else zero).append(item)
    return positive, zero


def parse_traveler(path: Path) -> TravelerData:
    """校验 Traveler 文件及表格结构，提取材料、五金、单据和来源指纹。

    参数：path：待解析的 Traveler 文件。
    """
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise RuleError("traveler_missing", f"Traveler 文件不存在或格式不是 xlsx：{path}")
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise RuleError("traveler_open", f"无法打开 Traveler：{path.name}") from exc
    required = {"WorkOrderTraveler", "Usage List", "Picking List"}
    if not required.issubset(wb.sheetnames):
        raise RuleError("traveler_schema", f"Traveler 缺少必要工作表：{sorted(required - set(wb.sheetnames))}")
    work_names = _order_name_candidates(wb["WorkOrderTraveler"])
    work_order = work_names[0].upper() if work_names else ""
    usage = wb["Usage List"]
    usage_order = ""
    for row in usage.iter_rows():
        for index, cell in enumerate(row):
            if _normalized_label(cell.value) == "job:":
                usage_order = next((_text(item.value).upper() for item in row[index + 1:] if _text(item.value)), "")
                break
        if usage_order:
            break
    folder_order = next((parent.name.upper() for parent in path.parents if PP_FOLDER_RE.fullmatch(parent.name)), "")
    temporary_folder_name = path.parent.name.strip()
    temporary_order_id = temporary_folder_name.upper() if temporary_folder_name and not folder_order and not PP_FOLDER_RE.fullmatch(temporary_folder_name) else ""
    filename_match = re.fullmatch(r"Work Order Traveler\(([^)]+)\)\.xlsx", path.name, re.IGNORECASE)
    filename_order = filename_match.group(1).strip().upper() if filename_match else ""
    candidates = [value for value in (folder_order, filename_order, usage_order, work_order) if PP_FOLDER_RE.fullmatch(value)]
    temporary_identity = not candidates and bool(temporary_order_id)
    if not candidates:
        if not temporary_identity:
            raise RuleError("traveler_identity", "无法从订单文件夹、文件名、Usage List 或 WorkOrderTraveler 取得订单号")
        order_id = temporary_order_id
    else:
        order_id = candidates[0]
    if any(value != order_id for value in candidates[1:]):
        raise RuleError("traveler_identity", f"Traveler 订单号不一致：{candidates}")

    ws = wb["Picking List"]
    section = ""
    current_factory = ""
    name_col = qty_col = None
    fitting_totals: dict[tuple[str, str], TravelerItem] = {}
    fitting_zero: list[TravelerItem] = []
    documents: dict[str, list[TravelerItem]] = {}
    for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
        for index, value in enumerate(row):
            if _text(value) == "Name/工厂单名称":
                current_factory = next((_text(item) for item in row[index + 1:] if _text(item)), "")
                if current_factory:
                    documents.setdefault(order_id if temporary_identity else current_factory, [])
                section = ""
                name_col = qty_col = None
                break
        first = _text(row[0]) if row else ""
        if first and not isinstance(row[0], (int, float)) and first not in {"No.", "Picking List领料单", "Name/工厂单名称"}:
            if _find_header(row, "Name名字") is None and _find_header(row, "QTY数量") is None:
                section = first
                name_col = qty_col = None
                continue
        found_name = _find_header(row, "Name名字")
        found_qty = _find_header(row, "QTY数量")
        if found_name or found_qty:
            if not found_name or not found_qty:
                raise RuleError("traveler_schema", f"Picking List 第 {row_number} 行名称与数量表头不完整")
            name_col, qty_col = found_name, found_qty
            continue
        if not row or not isinstance(row[0], (int, float)):
            continue
        if not name_col or not qty_col:
            raise RuleError("traveler_schema", f"Picking List 第 {row_number} 行数据没有对应表头")
        name = _text(row[name_col - 1] if name_col <= len(row) else None)
        value = row[qty_col - 1] if qty_col <= len(row) else None
        if not name:
            if value not in (None, "", 0, 0.0):
                raise RuleError("traveler_schema", f"Picking List 第 {row_number} 行有数量但没有名称")
            continue
        if not current_factory:
            raise RuleError("traveler_identity", f"Picking List 第 {row_number} 行五金没有对应工厂单名称")
        if value in (None, "", 0, 0.0):
            fitting_zero.append(TravelerItem(row_number, section, name, 0.0, order_id if temporary_identity else current_factory))
            continue
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise RuleError("traveler_quantity", f"{name} 的数量不是有效数字：{value}", row=row_number)
        if float(value) < 0:
            raise RuleError("traveler_quantity", f"{name} 的数量不能为负数：{value}", row=row_number)
        document_remark = order_id if temporary_identity else current_factory
        key = (document_remark, _normalize_name(name))
        if key in fitting_totals:
            fitting_totals[key].quantity += float(value)
        else:
            fitting_totals[key] = TravelerItem(row_number, section, name, float(value), document_remark)
    material_items, material_zero = _usage_list_items(usage, order_id)
    documents[order_id] = material_items
    for item in fitting_totals.values():
        documents.setdefault(item.document_remark, []).append(item)
    items = material_items + list(fitting_totals.values())
    zero_items = material_zero + fitting_zero
    pp_folder = folder_order or path.parent.name
    snapshot = {
        "order_id": order_id,
        "documents": {
            remark: [item.source_snapshot() for item in values]
            for remark, values in sorted(documents.items())
        },
        "zero_items": [item.source_snapshot() for item in zero_items],
    }
    return TravelerData(
        path=path.resolve(),
        pp_folder=pp_folder,
        order_id=order_id,
        order_name=order_id,
        items=items,
        zero_items=zero_items,
        documents=documents,
        modified_at=datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
        fingerprint=_fingerprint(snapshot),
    )


def _catalog_cost_price(value, label: str) -> float | None:
    """解析非负有限的预计采购价，空值保留未知，非法数字时报错。

    参数：value：待规范化或校验的原始值；label：目标表头或错误提示中的字段名。
    """
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        raise RuleError("product_catalog_data", f"商品 {label} 的预计采购价不是有效数字：{value}")
    try:
        price = float(value)
    except (TypeError, ValueError) as exc:
        raise RuleError("product_catalog_data", f"商品 {label} 的预计采购价不是有效数字：{value}") from exc
    if not math.isfinite(price) or price < 0:
        raise RuleError("product_catalog_data", f"商品 {label} 的预计采购价不能为负数或非有限数字：{value}")
    return price


class ProductCatalog:
    def __init__(self, path: Path):
        """读取商品工作簿并初始化商品及代码索引。

        参数：self：当前实例；path：待读取或初始化的文件路径。
        """
        self.path = path
        self.products: list[Product] = []
        self.by_code: dict[str, list[Product]] = {}
        self._load()

    def _load(self) -> None:
        """读取商品工作簿的必要列和可选属性，建立商品列表及代码索引。

        参数：self：当前实例。
        """
        if not self.path.is_file():
            raise RuleError("product_catalog_missing", f"尚未导入库存商品资料：{self.path}")
        try:
            wb = load_workbook(self.path, data_only=True, read_only=True)
        except Exception as exc:
            raise RuleError("product_catalog_open", f"无法打开库存商品资料：{self.path.name}") from exc
        ws = wb[wb.sheetnames[0]]
        header_row = None
        header_map = {}
        for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
            values = {_text(value): index for index, value in enumerate(row)}
            if REQUIRED_PRODUCT_HEADERS.issubset(values):
                header_row = row_number
                header_map = values
                break
        if not header_row:
            raise RuleError("product_catalog_schema", f"商品资料缺少必要表头：{sorted(REQUIRED_PRODUCT_HEADERS)}")
        optional = {"备注": None, "计量单位": None, "预计采购价": None, "品牌": None}
        for label in optional:
            optional[label] = header_map.get(label)
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            code = _text(row[header_map["*商品编号"]] if header_map["*商品编号"] < len(row) else None)
            name = _text(row[header_map["商品名称"]] if header_map["商品名称"] < len(row) else None)
            if not code and not name:
                continue
            product = Product(
                category=_text(row[header_map["商品类别"]] if header_map["商品类别"] < len(row) else None),
                code=code,
                name=name,
                spec=_text(row[header_map["规格型号"]] if header_map["规格型号"] < len(row) else None),
                status=_text(row[header_map["状态"]] if header_map["状态"] < len(row) else None),
                brand=_text(row[optional["品牌"]] if optional["品牌"] is not None and optional["品牌"] < len(row) else None),
                remark=_text(row[optional["备注"]] if optional["备注"] is not None and optional["备注"] < len(row) else None),
                unit=_text(row[optional["计量单位"]] if optional["计量单位"] is not None and optional["计量单位"] < len(row) else None),
                cost_price=_catalog_cost_price(
                    row[optional["预计采购价"]]
                    if optional["预计采购价"] is not None and optional["预计采购价"] < len(row)
                    else None,
                    code or name,
                ),
            )
            self.products.append(product)
            self.by_code.setdefault(code.upper(), []).append(product)

    def require_code(self, code: str) -> Product:
        """要求商品代码唯一、名称完整且启用，否则抛出可诊断错误。

        参数：self：当前实例；code：库存商品代码。
        """
        matches = self.by_code.get(code.upper(), [])
        if len(matches) != 1:
            raise RuleError("product_conflict", f"商品编号 {code} 匹配到 {len(matches)} 条记录", product_code=code)
        product = matches[0]
        if not product.code or not product.name:
            raise RuleError("product_invalid", f"当前需要的商品 {code} 缺少编号或名称")
        if product.status and product.status != "启用":
            raise RuleError("product_disabled", f"商品已停用：{code} {product.name}")
        return product

    def find(self, *, category: str | None = None, name: str | None = None, contains: str | None = None,
             spec_thickness: float | None = None) -> list[Product]:
        """按可选分类、名称、关键词及规格厚度筛选商品，厚度使用已有别名容差规则。

        参数：self：当前实例；category：可选商品分类条件；name：来源材料名称或查询名称；contains：可选模糊查询关键词；spec_thickness：可选规格厚度，使用现行别名及容差规则。
        """
        results = self.products
        if category:
            results = [item for item in results if _normalize_name(item.category) == _normalize_name(category)]
        if name:
            results = [item for item in results if _normalize_name(item.name) == _normalize_name(name)]
        if contains:
            token = _normalize_name(contains)
            results = [item for item in results if token in _normalize_name(item.name) or token in _normalize_name(item.remark)]
        if spec_thickness is not None:
            aliases = {14.5: 15.0, 5.4: 5.2, 8.0: 9.0}
            requested = float(spec_thickness)
            expected_values = (8.0, 9.0) if requested in {8.0, 9.0} else (
                aliases.get(requested, requested),
            )
            results = [
                item for item in results
                if (numbers := re.findall(r"\d+(?:\.\d+)?", item.spec))
                and any(abs(float(numbers[0]) - expected) < 0.6 for expected in expected_values)
            ]
        return results


def _ensure_product_cost_column(connection: sqlite3.Connection) -> None:
    """检查商品表并在缺失时新增采购价字段后提交。

    参数：connection：现有 SQLite 连接。
    """
    columns = {
        str(row[1])
        for row in connection.execute("pragma table_info(products)").fetchall()
    }
    if "cost_price" not in columns:
        connection.execute("alter table products add column cost_price real")
        connection.commit()


def _ensure_product_brand_column(connection: sqlite3.Connection) -> None:
    """检查商品表并在缺失时新增品牌字段后提交。

    参数：connection：现有 SQLite 连接。
    """
    columns = {
        str(row[1])
        for row in connection.execute("pragma table_info(products)").fetchall()
    }
    if "brand" not in columns:
        connection.execute("alter table products add column brand text not null default ''")
        connection.commit()


class ProductDatabase:
    """供运行时查询使用的 SQLite 商品目录。"""

    def __init__(self, path: Path):
        """连接并校验商品数据库，按现有流程补齐可选字段。

        参数：self：当前实例；path：待读取或初始化的文件路径。
        """
        self.path = path
        if not path.is_file():
            raise RuleError("product_database_missing", f"尚未导入库存商品数据库：{path}")
        try:
            self.connection = connect_database(path)
            self.connection.set_trace_callback(lambda statement: log_database_statement(path, statement))
            version = int(self.connection.execute("pragma user_version").fetchone()[0])
            # 商品目录位于共享业务数据库，user_version 由业务数据库结构统一管理。
            if version not in (0, INVENTORY_DATABASE_VERSION) and version < 1:
                raise RuleError("product_database_schema", f"库存商品数据库版本不受支持：{version}")
            if self.connection.execute(
                "select 1 from sqlite_master where type = 'table' and name = 'products'"
            ).fetchone() is None:
                raise RuleError("product_database_schema", "库存商品数据库缺少 products 表")
            _ensure_product_cost_column(self.connection)
            _ensure_product_brand_column(self.connection)
        except RuleError:
            self.connection.close()
            raise
        except sqlite3.Error as exc:
            if hasattr(self, "connection"):
                self.connection.close()
            raise RuleError("product_database_open", f"无法打开库存商品数据库：{path}") from exc

    def close(self) -> None:
        """关闭当前商品数据库连接。

        参数：self：当前实例。
        """
        self.connection.close()

    def __enter__(self) -> "ProductDatabase":
        """进入上下文管理时返回当前商品数据库实例。

        参数：self：当前实例。
        """
        return self

    def __exit__(self, exc_type, exc_value, traceback) -> None:
        """离开上下文时关闭数据库连接，不吞掉原异常。

        参数：self：当前实例；exc_type：上下文退出时的异常类型；exc_value：上下文退出时的异常对象；traceback：上下文退出时的异常调用栈。
        """
        self.close()

    @staticmethod
    def _from_row(row: sqlite3.Row | tuple) -> Product:
        """按查询列顺序将数据库行转换为商品对象。

        参数：row：按固定商品查询列顺序返回的数据库行。
        """
        return Product(
            category=str(row[0] or ""),
            code=str(row[1] or ""),
            name=str(row[2] or ""),
            spec=str(row[3] or ""),
            status=str(row[4] or ""),
            brand=str(row[5] or ""),
            remark=str(row[6] or ""),
            unit=str(row[7] or ""),
            cost_price=None if row[8] is None else float(row[8]),
            material_kind=str(row[9] or ""),
            material_color=str(row[10] or ""),
            material_thickness=str(row[11] or ""),
            catalog_present=bool(row[12]),
        )

    @property
    def products(self) -> list[Product]:
        """按商品代码读取完整商品目录，包括最新目录已移除但仍保留的记录。

        参数：self：当前实例。
        """
        rows = self.connection.execute(
            """
            select category, code, name, spec, status, brand, remark, unit, cost_price,
                   material_kind, material_color, material_thickness, catalog_present
            from products order by code
            """
        ).fetchall()
        return [self._from_row(row) for row in rows]

    def count(self) -> int:
        """返回商品表记录总数。

        参数：self：当前实例。
        """
        return int(self.connection.execute("select count(*) from products").fetchone()[0])

    def require_code(self, code: str) -> Product:
        """要求商品代码唯一、名称完整且启用，否则抛出可诊断错误。最新目录已移除的商品也不能使用。

        参数：self：当前实例；code：库存商品代码。
        """
        rows = self.connection.execute(
            """
            select category, code, name, spec, status, brand, remark, unit, cost_price,
                   material_kind, material_color, material_thickness, catalog_present
            from products where normalized_code = ?
            """,
            (_normalize_name(code),),
        ).fetchall()
        if len(rows) != 1:
            raise RuleError("product_conflict", f"商品编号 {code} 匹配到 {len(rows)} 条记录", product_code=code)
        product = self._from_row(rows[0])
        if not product.code or not product.name:
            raise RuleError("product_invalid", f"当前需要的商品 {code} 缺少编号或名称")
        if not product.catalog_present:
            raise RuleError("product_catalog_missing", f"商品已不在最新商品目录中：{code} {product.name}")
        if product.status and product.status != "启用":
            raise RuleError("product_disabled", f"商品已停用：{code} {product.name}")
        return product

    def find(
        self,
        *,
        category: str | None = None,
        name: str | None = None,
        contains: str | None = None,
        spec_thickness: float | None = None,
    ) -> list[Product]:
        """按可选分类、名称、关键词及规格厚度筛选商品，厚度使用已有别名容差规则。

        参数：self：当前实例；category：可选商品分类条件；name：来源材料名称或查询名称；contains：可选模糊查询关键词；spec_thickness：可选规格厚度，使用现行别名及容差规则。
        """
        clauses = []
        parameters: list[str] = []
        if category:
            clauses.append("normalized_category = ?")
            parameters.append(_normalize_name(category))
        if name:
            clauses.append("normalized_name = ?")
            parameters.append(_normalize_name(name))
        if contains:
            token = _normalize_name(contains)
            clauses.append(
                "(normalized_code like ? or normalized_name like ? or normalized_spec like ? "
                "or normalized_category like ? or normalized_remark like ?)"
            )
            parameters.extend([f"%{token}%"] * 5)
        query = """
            select category, code, name, spec, status, brand, remark, unit, cost_price,
                   material_kind, material_color, material_thickness, catalog_present
            from products
        """
        if clauses:
            query += " where " + " and ".join(clauses)
        query += " order by code"
        results = [
            self._from_row(row)
            for row in self.connection.execute(query, parameters).fetchall()
        ]
        results = [product for product in results if product.catalog_present]
        if spec_thickness is None:
            return results
        aliases = {14.5: 15.0, 5.4: 5.2, 8.0: 9.0}
        requested = float(spec_thickness)
        expected_values = (8.0, 9.0) if requested in {8.0, 9.0} else (
            aliases.get(requested, requested),
        )
        return [
            product
            for product in results
            if (numbers := re.findall(r"\d+(?:\.\d+)?", product.spec))
            and any(abs(float(numbers[0]) - expected) < 0.6 for expected in expected_values)
        ]


def _create_product_database(path: Path) -> None:
    """创建商品表及查询索引，补齐采购价字段并关闭连接。

    参数：path：待读取或初始化的文件路径。
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = connect_database(path)
    connection.set_trace_callback(lambda statement: log_database_statement(path, statement))
    try:
        connection.executescript(
            f"""
            create table if not exists products(
                category text not null default '',
                code text primary key,
                name text not null default '',
                spec text not null default '',
                status text not null default '',
                brand text not null default '',
                remark text not null default '',
                unit text not null default '',
                cost_price real,
                normalized_code text not null,
                normalized_name text not null,
                normalized_spec text not null,
                normalized_category text not null,
                normalized_remark text not null,
                material_kind text not null default '',
                material_color text not null default '',
                material_thickness text not null default '',
                catalog_present integer not null default 1
            );
            create index if not exists idx_products_name on products(normalized_name);
            create index if not exists idx_products_category on products(normalized_category);
            """
        )
        _ensure_product_cost_column(connection)
    finally:
        connection.close()


def _replace_product_database(path: Path, products: list[Product]) -> None:
    """校验目录后事务更新商品，保留已引用材料属性和历史商品，禁止改变已引用五金单位。

    参数：path：待读取或初始化的文件路径；products：已从新目录解析的商品列表。
    """
    missing = [product.code or product.name for product in products if not product.code or not product.name]
    duplicate_codes = [
        code for code, count in Counter(_normalize_name(product.code) for product in products).items()
        if code and count > 1
    ]
    if missing:
        raise RuleError("product_catalog_data", "商品资料存在缺少商品编号或商品名称的记录")
    if duplicate_codes:
        raise RuleError(
            "product_catalog_data",
            f"商品资料存在重复商品编号：{', '.join(duplicate_codes[:10])}",
            product_codes=duplicate_codes,
        )
    _create_product_database(path)
    connection = connect_database(path)
    connection.set_trace_callback(lambda statement: log_database_statement(path, statement))
    try:
        connection.execute("begin immediate")
        reference_tables = [
            table for table in (
                "material_items",
                "production_materials",
                "server_material_allocations",
            )
            if connection.execute(
                "select 1 from sqlite_master where type='table' and name=?",
                (table,),
            ).fetchone()
        ]
        referenced_codes: set[str] = set()
        for table in reference_tables:
            referenced_codes.update(
                str(row[0] or "").strip().upper()
                for row in connection.execute(
                    f"select distinct product_code from {table}"
                ).fetchall()
                if str(row[0] or "").strip()
            )
        bound_attributes = {
            str(row[0] or "").strip().upper(): (
                str(row[1] or ""), str(row[2] or ""), str(row[3] or "")
            )
            for row in connection.execute(
                """select code,material_kind,material_color,material_thickness
                   from products"""
            ).fetchall()
            if str(row[0] or "").strip().upper() in referenced_codes
        }
        if connection.execute("select 1 from sqlite_master where name='hardware_items'").fetchone():
            hardware_units = dict(connection.execute(
                "select p.code,p.unit from products p where exists "
                "(select 1 from hardware_items h where h.product_code=p.code)"))
            for product in products:
                if product.code in hardware_units and hardware_units[product.code] != product.unit:
                    raise RuleError('hardware_product_unit_locked',
                        f"商品 {product.code} 已被五金记录引用，不能将单位从 {hardware_units[product.code]} 改为 {product.unit}；请先核对数量换算")
        connection.execute("update products set catalog_present=0")
        connection.executemany(
            """
            insert into products(
                category, code, name, spec, status, brand, remark, unit, cost_price,
                normalized_code, normalized_name, normalized_spec,
                normalized_category, normalized_remark,
                material_kind, material_color, material_thickness, catalog_present
            ) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)
            on conflict(code) do update set
                category=excluded.category,
                name=excluded.name,
                spec=excluded.spec,
                status=excluded.status,
                brand=excluded.brand,
                remark=excluded.remark,
                unit=excluded.unit,
                cost_price=excluded.cost_price,
                normalized_code=excluded.normalized_code,
                normalized_name=excluded.normalized_name,
                normalized_spec=excluded.normalized_spec,
                normalized_category=excluded.normalized_category,
                normalized_remark=excluded.normalized_remark,
                material_kind=excluded.material_kind,
                material_color=excluded.material_color,
                material_thickness=excluded.material_thickness,
                catalog_present=1
            """,
            [
                (
                    product.category,
                    product.code,
                    product.name,
                    product.spec,
                    product.status,
                    product.brand,
                    product.remark,
                    product.unit,
                    product.cost_price,
                    _normalize_name(product.code),
                    _normalize_name(product.name),
                    _normalize_name(product.spec),
                    _normalize_name(product.category),
                    _normalize_name(product.remark),
                    *bound_attributes.get(
                        product.code.strip().upper(),
                        catalog_material_attributes(
                            product.category, product.name, product.spec, product.code
                        ),
                    ),
                )
                for product in products
            ],
        )
        connection.execute(f"pragma user_version = {INVENTORY_DATABASE_VERSION}")
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


class InventoryMappings:
    """由数据库保存的材料规则。JSON 兼容模式仅保留给隔离的旧测试夹具；应用调用方传集中业务 SQLite，因此运行时只读写 SQLite。"""
    def __init__(self, path: Path, *, connection: sqlite3.Connection | None = None):
        """初始化材料映射存储，支持应用数据库及隔离旧 JSON 夹具。

        参数：self：当前实例；path：待读取或初始化的文件路径；connection：可选外部管理的 SQLite 连接。
        """
        self.path = path
        self._connection = connection
        self._legacy_path = path if path.suffix.lower() == ".json" else None
        self.database_path = (
            path.parent.parent / "workflow.sqlite3"
            if self._legacy_path is not None
            else path.parent / "workflow.sqlite3"
            if path.name == "order-index.sqlite3"
            else path
        )
        self.manual: dict[str, str] = {}
        self.manual_display_names: dict[str, str] = {}
        self.ignored: dict[str, str] = {}
        if self._legacy_path is not None:
            if path.is_file():
                try:
                    data = json.loads(path.read_text(encoding="utf-8"))
                    self.manual = {_normalize_name(k): str(v) for k, v in data.get("manual", {}).items()}
                    self.manual_display_names = {
                        _normalize_name(k): str(v)
                        for k, v in data.get("manual_display_names", {}).items()
                    }
                    self.ignored = {_normalize_name(k): str(v) for k, v in data.get("ignored", {}).items()}
                except (OSError, json.JSONDecodeError) as exc:
                    raise RuleError("inventory_mapping", f"库存商品映射无法读取：{path}") from exc
            return
        if self._connection is None:
            ensure_schema(self.database_path)

    def _rows(self, rule_type: str | None = None) -> list[sqlite3.Row]:
        """按可选规则类型读取材料解析规则，外部传入的连接由调用方管理。

        参数：self：当前实例；rule_type：mapping 映射或 ignore 忽略规则类型。
        """
        connection = self._connection or connect_database(self.database_path)
        connection.row_factory = sqlite3.Row
        try:
            if rule_type:
                return connection.execute(
                    "select source_name, normalized_name, product_code, display_name, reason from inventory_resolution_rules "
                    "where rule_type=? order by source_name",
                    (rule_type,),
                ).fetchall()
            return connection.execute(
                "select source_name, normalized_name, product_code, display_name, reason from inventory_resolution_rules "
                "order by source_name"
            ).fetchall()
        finally:
            if self._connection is None:
                connection.close()

    def entries(self) -> tuple[list[dict], list[dict]]:
        """返回供界面展示的人工映射和忽略规则列表。

        参数：self：当前实例。
        """
        if self._legacy_path is not None:
            return (
                [
                    {
                        "name": name,
                        "product_code": code,
                        "display_name": self.manual_display_names.get(name, "").strip()
                        or HARDWARE_PRODUCT_DISPLAY_NAMES.get(code.strip().upper(), "")
                        or name,
                    }
                    for name, code in sorted(self.manual.items())
                ],
                [{"name": name, "reason": reason} for name, reason in sorted(self.ignored.items())],
            )
        rows = self._rows()
        manual = [
            {
                "name": str(row["source_name"]),
                "product_code": str(row["product_code"] or ""),
                "display_name": self._effective_display_name(
                    str(row["product_code"] or ""),
                    str(row["source_name"] or ""),
                    str(row["display_name"] or ""),
                ),
            }
            for row in rows if row["product_code"]
        ]
        ignored = [
            {"name": str(row["source_name"]), "reason": str(row["reason"] or "")}
            for row in rows if not row["product_code"]
        ]
        return manual, ignored

    def _effective_display_name(
        self, product_code: str, source_name: str = "", stored_name: str = ""
    ) -> str:
        """优先选择显式名称或同 SKU 的已存别名，最后回退默认名称，不修改来源事实。

        参数：self：当前实例；product_code：已确认的库存商品 SKU；source_name：原始来源名称；stored_name：已存的显式显示名称。
        """
        explicit = stored_name.strip()
        if explicit:
            return explicit
        code = product_code.strip().upper()
        if self._legacy_path is not None:
            legacy_name = self.manual_display_names.get(_normalize_name(source_name), "").strip()
            return legacy_name or HARDWARE_PRODUCT_DISPLAY_NAMES.get(code, "") or source_name.strip()
        connection = self._connection or connect_database(self.database_path)
        try:
            row = connection.execute(
                """select display_name from inventory_resolution_rules
                   where rule_type='mapping' and product_code=? and trim(display_name)<>''
                   order by updated_at desc, id desc limit 1""",
                (code,),
            ).fetchone()
            if row and str(row[0] or "").strip():
                return str(row[0]).strip()
        finally:
            if self._connection is None:
                connection.close()
        return HARDWARE_PRODUCT_DISPLAY_NAMES.get(code, "") or source_name.strip()

    def display_name_for_product(self, product_code: str, source_name: str = "") -> str:
        """读取商品 SKU 对应的有效显示名称。

        参数：self：当前实例；product_code：已确认的库存商品 SKU；source_name：原始来源名称。
        """
        return self._effective_display_name(product_code, source_name)

    def display_name_for_hardware(
        self, product_code: str, source_name: str = "", source_code: str = ""
    ) -> str:
        """先按 SKU，再按稳定来源别名解析五金显示名称。

        参数：self：当前实例；product_code：已确认的库存商品 SKU；source_name：原始来源名称；source_code：仅用于来源识别的原始报表代码。
        """
        code = product_code.strip().upper()
        name = self._effective_display_name(code, source_name)
        if name and name != source_name.strip():
            return name
        source_display = (
            HARDWARE_NAME_DISPLAY_NAMES.get(_normalize_name(source_name))
            or HARDWARE_DISPLAY_NAMES.get(_normalize_name(source_code))
        )
        return source_display or name or source_code.strip() or source_name.strip()

    def ignored_reason(self, name: str) -> str | None:
        """按规范化来源名称查找忽略原因，无规则返回空值。

        参数：self：当前实例；name：来源材料名称或查询名称。
        """
        normalized = _normalize_name(name)
        if self._legacy_path is not None:
            return self.ignored.get(normalized)
        connection = self._connection or connect_database(self.database_path)
        try:
            row = connection.execute(
                "select reason from inventory_resolution_rules where normalized_name=? and rule_type='ignore'",
                (normalized,),
            ).fetchone()
            return str(row[0]) if row else None
        finally:
            if self._connection is None:
                connection.close()

    def manual_code(self, name: str) -> str | None:
        """按规范化来源名称查找人工确认的 SKU，无规则返回空值。

        参数：self：当前实例；name：来源材料名称或查询名称。
        """
        normalized = _normalize_name(name)
        if self._legacy_path is not None:
            return self.manual.get(normalized)
        connection = self._connection or connect_database(self.database_path)
        try:
            row = connection.execute(
                "select product_code from inventory_resolution_rules where normalized_name=? and rule_type='mapping'",
                (normalized,),
            ).fetchone()
            return str(row[0]) if row else None
        finally:
            if self._connection is None:
                connection.close()

    def save_ignored(self, name: str, reason: str) -> None:
        """校验名称并保存忽略规则，未填写理由时使用默认说明。

        参数：self：当前实例；name：来源材料名称或查询名称；reason：忽略规则或不出库决定的原因。
        """
        normalized = _normalize_name(name)
        if not normalized:
            raise RuleError("inventory_mapping", "忽略材料名称不能为空")
        value = reason.strip() or "用户在出库预览中选择忽略"
        if self._legacy_path is not None:
            self.ignored[normalized] = value
            self._save()
            return
        self._upsert("ignore", name.strip(), normalized, None, value)

    def save_manual(self, name: str, product_code: str, display_name: str = "") -> None:
        """校验名称与 SKU，保存人工映射并统一同商品的显示别名。

        参数：self：当前实例；name：来源材料名称或查询名称；product_code：已确认的库存商品 SKU；display_name：用于界面显示的商品别名。
        """
        normalized = _normalize_name(name)
        if not normalized:
            raise RuleError("inventory_mapping", "映射材料名称不能为空")
        code = product_code.strip().upper()
        if not code:
            raise RuleError("inventory_mapping", "映射商品 SKU 不能为空")
        if self._legacy_path is not None:
            self.manual[normalized] = code
            self.manual_display_names[normalized] = (
                display_name.strip()
                or self.display_name_for_hardware(code, name, "")
                or name.strip()
            )
            self.ignored.pop(normalized, None)
            self._save()
            return
        effective_display_name = (
            display_name.strip()
            or self.display_name_for_product(code, name)
            or name.strip()
        )
        self._upsert("mapping", name.strip(), normalized, code, "", effective_display_name)
        self._set_product_display_name(code, effective_display_name)

    def remove_ignored(self, name: str) -> None:
        """删除指定来源名称的忽略规则。

        参数：self：当前实例；name：来源材料名称或查询名称。
        """
        normalized = _normalize_name(name)
        if self._legacy_path is not None:
            self.ignored.pop(normalized, None)
            self._save()
            return
        self._delete(normalized, "ignore")

    def remove_manual(self, name: str) -> None:
        """删除指定来源名称的人工商品映射。

        参数：self：当前实例；name：来源材料名称或查询名称。
        """
        normalized = _normalize_name(name)
        if self._legacy_path is not None:
            self.manual.pop(normalized, None)
            self._save()
            return
        self._delete(normalized, "mapping")

    def _upsert(
        self,
        rule_type: str,
        source_name: str,
        normalized: str,
        product_code: str | None,
        reason: str,
        display_name: str = "",
    ) -> None:
        """按规范化来源名新增或覆盖规则，数据库失败时回滚。

        参数：self：当前实例；rule_type：mapping 映射或 ignore 忽略规则类型；source_name：原始来源名称；normalized：已规范化的来源名称键；product_code：已确认的库存商品 SKU；reason：忽略规则或不出库决定的原因；display_name：用于界面显示的商品别名。
        """
        connection = connect_database(self.database_path)
        try:
            connection.execute(
                """insert into inventory_resolution_rules(
                    rule_type, source_name, normalized_name, product_code, display_name, reason, created_at, updated_at
                ) values(?,?,?,?,?,?,?,?)
                on conflict(normalized_name) do update set
                    rule_type=excluded.rule_type,
                    source_name=excluded.source_name,
                    product_code=excluded.product_code,
                    display_name=case when excluded.rule_type='mapping' then excluded.display_name else '' end,
                    reason=excluded.reason,
                    updated_at=excluded.updated_at""",
                (rule_type, source_name, normalized, product_code, display_name, reason, datetime.now().astimezone().isoformat(timespec="seconds"), datetime.now().astimezone().isoformat(timespec="seconds")),
            )
            connection.commit()
        except sqlite3.Error as exc:
            connection.rollback()
            raise RuleError("inventory_mapping", f"库存规则无法保存：{source_name}") from exc
        finally:
            connection.close()

    def _set_product_display_name(self, product_code: str, display_name: str) -> None:
        """将同一 SKU 的所有来源别名更新为统一显示名称。

        参数：self：当前实例；product_code：已确认的库存商品 SKU；display_name：用于界面显示的商品别名。
        """
        if self._legacy_path is not None:
            for normalized, code in self.manual.items():
                if code.strip().upper() == product_code.strip().upper():
                    self.manual_display_names[normalized] = display_name
            self._save()
            return
        connection = connect_database(self.database_path)
        try:
            connection.execute(
                """update inventory_resolution_rules
                   set display_name=?, updated_at=?
                   where rule_type='mapping' and product_code=?""",
                (display_name, datetime.now().astimezone().isoformat(timespec="seconds"), product_code.strip().upper()),
            )
            connection.commit()
        finally:
            connection.close()

    def _delete(self, normalized: str, rule_type: str) -> None:
        """按规范化来源名和规则类型删除数据库规则并提交。

        参数：self：当前实例；normalized：已规范化的来源名称键；rule_type：mapping 映射或 ignore 忽略规则类型。
        """
        connection = connect_database(self.database_path)
        try:
            connection.execute(
                "delete from inventory_resolution_rules where normalized_name=? and rule_type=?",
                (normalized, rule_type),
            )
            connection.commit()
        finally:
            connection.close()

    def _save(self) -> None:
        """通过临时文件原子替换旧版隔离测试用 JSON 映射规则。

        参数：self：当前实例。
        """
        self.path.parent.mkdir(parents=True, exist_ok=True)
        payload = {
            "version": 1,
            "manual": self.manual,
            "manual_display_names": self.manual_display_names,
            "ignored": self.ignored,
        }
        temporary = self.path.with_suffix(self.path.suffix + ".tmp")
        temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(self.path)


HARDWARE_DISPLAY_NAMES = {
    "WJCBT": "Shelf Holder",
    "71T950A": "Hinge",
    "HRAIL": "H-Rail",
    "LRAIL": "L-Rail",
}

# 部分 AICNC 报表以左右侧来源名称标识低导轨；应按原始来源名解析，避免误映射为 H-Rail。
HARDWARE_NAME_DISPLAY_NAMES = {
    "LOWERLEFTRAIL": "L-Rail",
    "LOWERRIGHTRAIL": "L-Rail",
}


def ignored_hardware_reason(
    mappings: InventoryMappings,
    name: str = "",
    code: str = "",
    source_code: str = "",
) -> str | None:
    """依次按原始五金名称、代码及显示别名查找忽略原因。

    参数：mappings：人工映射及忽略规则集合；name：来源材料名称或查询名称；code：库存商品代码；source_code：仅用于来源识别的原始报表代码。
    """
    candidates = [name, code, source_code]
    name_display_name = HARDWARE_NAME_DISPLAY_NAMES.get(_normalize_name(name))
    if name_display_name:
        candidates.append(name_display_name)
    display_name = HARDWARE_DISPLAY_NAMES.get(_normalize_name(source_code or code))
    if display_name:
            candidates.append(display_name)
    for candidate in candidates:
        reason = mappings.ignored_reason(candidate)
        if reason is not None:
            return reason
    return None


FIXED_CODES = {
    "18MMPLYWOOD": "M0004",
    "14.5MMPLYWOOD": "M0003",
    "5.4MMPLYWOOD": "M0002",
    "ADJUSTABLESHELFHOLDER": "M1013",
    "SHELHOLDER": "M1013",
    "HINGE": "M1001",
    "TESTFULLHINGE": "M1001",
    "HRAIL": "M1002",
    "LRAIL": "M1003",
    "LOWERLEFTRAIL": "M1003",
    "LOWERRIGHTRAIL": "M1003",
    "M.C(L)": "M1089",
}


def _single(catalog: ProductCatalog, matches: list[Product], traveler_name: str, source: str) -> tuple[Product, str]:
    """要求候选商品唯一并再次校验代码及启用状态，返回商品和匹配来源。

    参数：catalog：用于校验唯一商品的目录；matches：候选商品列表；traveler_name：Traveler 中的原始名称；source：匹配规则来源说明。
    """
    if len(matches) != 1:
        raise RuleError(
            "product_match",
            f"{traveler_name} 匹配到 {len(matches)} 个库存商品，需要人工指定",
            traveler_name=traveler_name,
            candidates=[asdict(item) for item in matches],
        )
    return catalog.require_code(matches[0].code), source


def match_item(catalog: ProductCatalog, mappings: InventoryMappings, item: TravelerItem) -> list[OutboundItem]:
    """优先使用已确认 SKU，否则按人工与固定映射规则解析材料或五金，拒绝歧义。

    参数：catalog：用于校验唯一商品的目录；mappings：人工映射及忽略规则集合；item：当前来源材料或五金条目。
    """
    def outbound(product: Product, quantity: float, source: str) -> OutboundItem:
        """用已确认商品和数量创建当前来源条目的出库明细。

        参数：product：已匹配并确认的商品；quantity：待出库数量；source：匹配规则来源说明。
        """
        return OutboundItem(
            item.name, product.code, product.name, quantity, item.section, source,
            item.document_remark, product.unit,
        )

    def edge_outbound(product: Product, source: str) -> OutboundItem:
        """将封边数量四舍五入为整数后创建出库明细。

        参数：product：已匹配并确认的商品；source：商品匹配规则来源说明。
        """
        rounded = float(Decimal(str(item.quantity)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        return outbound(
            product,
            rounded,
            f"{source}（{item.quantity:g}m→{rounded:g}m，四舍五入取整）",
        )

    canonical_code = str(getattr(item, "product_code", "") or "").strip().upper()
    if canonical_code:
        product = catalog.require_code(canonical_code)
        if _normalize_name(product.category) == _normalize_name("Edge band"):
            return [edge_outbound(product, "数据库 canonical SKU")]
        return [outbound(product, item.quantity, "数据库 canonical SKU")]
    ignored = mappings.ignored_reason(item.name)
    if ignored is not None:
        return []
    manual = mappings.manual_code(item.name)
    if manual:
        product = catalog.require_code(manual)
        if _normalize_name(product.category) == _normalize_name("Edge band"):
            return [edge_outbound(product, "人工指定")]
        return [outbound(product, item.quantity, "人工指定")]
    # 数据库五金用 product_code 保存规范库存 SKU；直接接受该身份，避免出库时对 Left Rail 等 Server 原始显示名重新匹配。
    try:
        product = catalog.require_code(item.name)
    except RuleError:
        product = None
    if product is not None:
        if _normalize_name(product.category) == _normalize_name("Edge band"):
            return [edge_outbound(product, "数据库 canonical SKU")]
        return [outbound(product, item.quantity, "数据库 canonical SKU")]
    normalized = _normalize_name(item.name)
    if normalized == "PUSHOPEN":
        results = []
        for code in ("M1068", "M1069"):
            product = catalog.require_code(code)
            results.append(outbound(product, item.quantity, "Push Open 1:1拆分"))
        return results
    fixed = FIXED_CODES.get(normalized)
    if fixed:
        product = catalog.require_code(fixed)
        return [outbound(product, item.quantity, "已确认规则")]
    if match := TB_RE.fullmatch(item.name.strip()):
        token = f"TB{match.group(1)}用"
        product, source = _single(catalog, catalog.find(category="Hardware/Trash Can", contains=token), item.name, "TB型号规则")
        return [outbound(product, item.quantity, source)]
    if match := EDGE_RE.fullmatch(item.name.strip()):
        color = match.group(1).strip()
        matches = catalog.find(category="Edge band", name=f"{color} Edge Banding")
        source = "封边颜色精确名称"
        if not matches:
            color_token = _normalize_name(color)
            matches = [
                product for product in catalog.find(category="Edge band")
                if (normalized := _normalize_name(product.name)).startswith(color_token)
                and "ABSBANDING" in normalized
                and normalized.endswith("24MM")
            ]
            source = "封边颜色 + ABS banding + 24mm后缀"
        product, source = _single(catalog, matches, item.name, source)
        return [edge_outbound(product, source)]
    if match := PANEL_RE.fullmatch(item.name.strip()):
        thickness = float(match.group(1))
        color = match.group(2).strip()
        if color.lower() != "plywood":
            product, source = _single(
                catalog,
                catalog.find(category="Panel", name=color, spec_thickness=thickness),
                item.name,
                "Panel颜色和规格",
            )
            return [outbound(product, item.quantity, source)]
    exact_matches = catalog.find(name=item.name)
    if exact_matches:
        product, source = _single(catalog, exact_matches, item.name, "库存商品名称精确匹配")
        return [outbound(product, item.quantity, source)]
    raise RuleError(
        "product_match",
        f"找不到与 {item.name} 名称完全一致的库存商品，需要人工指定",
        traveler_name=item.name,
        candidates=[],
    )


def resolve_inventory_items(
    config: Config,
    items: Iterable[tuple[TravelerItem, str]],
) -> dict:
    """在材料或五金成为有效事实前解析 SKU。来源代码只临时参与匹配，不能把 AICNC 的 WJ-CBD 等代码直接当库存 SKU；仅接受明确忽略或唯一启用商品。

    参数：config：包含状态库、订单来源及备份路径的运行配置；items：由来源条目和原始来源代码组成的二元组集合。
    """
    items = list(items)
    if not items:
        return {"outbound": [], "ignored": [], "missing": [], "accepted": []}
    mappings = InventoryMappings(config.workflow_database)
    outbound: list[OutboundItem] = []
    ignored: list[dict] = []
    missing: list[dict] = []
    accepted: list[dict] = []
    with ProductDatabase(bootstrap_product_database(config)) as catalog:
        for item, source_code in items:
            source_code = _text(source_code)
            # 已知来源代码在流程中有稳定显示名；该名称仅用于目录匹配，原始代码仅保留在临时解析记录中，不写成 SKU 事实。
            match_item_value = item
            display_name = (
                HARDWARE_NAME_DISPLAY_NAMES.get(_normalize_name(item.name))
                or HARDWARE_DISPLAY_NAMES.get(_normalize_name(source_code))
            )
            if item.section == "五金" and display_name and not _normalize_name(item.name) == _normalize_name(display_name):
                match_item_value = replace(item, name=display_name)
            reason = (
                ignored_hardware_reason(mappings, item.name, source_code)
                if item.section == "五金"
                else mappings.ignored_reason(item.name)
            )
            if reason is not None:
                accepted.append({
                    "name": item.name,
                    "source_code": source_code,
                    "product_codes": [],
                    "ignored": True,
                })
                ignored.append({
                    **item.source_snapshot(),
                    "source_code": source_code,
                    "reason": reason,
                })
                continue
            try:
                matched = match_item(catalog, mappings, match_item_value)
                accepted.append({
                    "name": item.name,
                    "source_code": source_code,
                    "product_codes": [entry.product_code for entry in matched],
                    "ignored": False,
                })
                outbound.extend(matched)
            except RuleError as exc:
                accepted.append({
                    "name": item.name,
                    "source_code": source_code,
                    "product_codes": [],
                    "ignored": False,
                })
                missing.append({
                    **item.source_snapshot(),
                    "source_code": source_code,
                    "code": exc.code,
                    "message": str(exc),
                    **exc.context,
                })
    return {"outbound": outbound, "ignored": ignored, "missing": missing, "accepted": accepted}


def resolved_product_code(resolution: dict, index: int) -> str:
    """要求解析结果含一个已确认目录 SKU，绝不回退使用报表代码。

    参数：resolution：材料解析返回的结构化结果；index：待提取条目在解析结果中的下标。
    """
    accepted = resolution.get("accepted", []) if isinstance(resolution, dict) else []
    if 0 <= index < len(accepted):
        codes = accepted[index].get("product_codes", []) or []
        if len(codes) == 1 and str(codes[0]).strip():
            return str(codes[0]).strip()
    raise RuleError("hardware_mapping_required", "写入前必须明确映射到唯一商品 SKU", item_index=index)


def confirm_product_material_attributes(
    connection: sqlite3.Connection,
    product_code: str,
    material_kind: str,
    material_color: str = "",
    material_thickness: object = "",
) -> None:
    """首次确认使用 SKU 时绑定商品拥有的材料属性，保留原始目录名称及规格。SKU 已被订单材料引用后，后续映射不能暗改其颜色或标称厚度。

    参数：connection：现有 SQLite 连接；product_code：已确认的库存商品 SKU；material_kind：待绑定的材料种类；material_color：待绑定的材料颜色；material_thickness：待绑定的材料标称厚度。
    """
    code = str(product_code or "").strip().upper()
    kind = str(material_kind or "").strip().casefold()
    color = str(material_color or "").strip() if kind in {"panel", "edge", "back"} else ""
    thickness_text = str(material_thickness or "").strip()
    try:
        thickness = f"{float(thickness_text):g}" if thickness_text else ""
    except ValueError as exc:
        raise RuleError("product_material_attributes", f"商品 {code} 的材料厚度无效：{thickness_text}") from exc
    row = connection.execute(
        """select material_kind,material_color,material_thickness,catalog_present,status
           from products where code=?""",
        (code,),
    ).fetchone()
    if row is None:
        raise RuleError("product_conflict", f"商品编号 {code} 匹配到 0 条记录", product_code=code)
    if not bool(row[3]) or (row[4] and str(row[4]) != "启用"):
        raise RuleError("product_disabled", f"商品不可用于新的材料确认：{code}")
    has_confirmed_fact = any(
        connection.execute(
            f"select 1 from {table} where product_code=? limit 1", (code,)
        ).fetchone() is not None
        for table in (
            "material_items",
            "production_materials",
            "server_material_allocations",
        )
        if connection.execute(
            "select 1 from sqlite_master where type='table' and name=?", (table,)
        ).fetchone() is not None
    )
    existing = (str(row[0] or ""), str(row[1] or ""), str(row[2] or ""))
    desired = (kind, color, thickness)
    if has_confirmed_fact:
        same_color = _normalize_name(existing[1]) == _normalize_name(desired[1])
        same_thickness = (
            (not existing[2] and not desired[2])
            or (
                existing[2] and desired[2]
                and abs(float(existing[2]) - float(desired[2])) < 0.01
            )
        )
        if existing[0] != desired[0] or not same_color or not same_thickness:
            raise RuleError(
                "product_material_attribute_conflict",
                f"商品 {code} 已关联其他材料属性，不能静默重绑",
                product_code=code,
                existing={"material_kind": existing[0], "material_color": existing[1], "material_thickness": existing[2]},
                requested={"material_kind": kind, "material_color": color, "material_thickness": thickness},
            )
        return
    connection.execute(
        """update products set material_kind=?,material_color=?,material_thickness=?
           where code=?""",
        (kind, color, thickness, code),
    )


def build_preview(
    path: Path,
    catalog_path: Path,
    mapping_path: Path,
    selected_document_remarks: Iterable[str] | None = None,
) -> InventoryPreview:
    """只读解析 Traveler 和本地商品映射，返回出库预览及冲突，不提交出库。缺失、忽略和重复匹配均显式报告，交由调用方阻止确认而非猜测商品。

    参数：path：待解析的 Traveler 文件；catalog_path：商品目录文件或数据库路径；mapping_path：映射规则数据库或隔离测试 JSON 路径；selected_document_remarks：明确选择的单据备注；为空时包含全部。
    """
    traveler = parse_traveler(path)
    mappings = InventoryMappings(mapping_path)
    requested_remarks = tuple(
        dict.fromkeys(
            str(remark).strip()
            for remark in (selected_document_remarks or [])
            if str(remark).strip()
        )
    )
    selected = {
        _normalize_name(remark)
        for remark in requested_remarks
    }
    # 选择工厂单出库时始终纳入订单级材料单；五金仅保留所选工厂名称。允许缺少名称，因为 CS 订单的 Traveler 可以没有五金区域。
    if selected:
        selected.add(_normalize_name(traveler.order_id))
    outbound = []
    ignored = []
    missing = []
    catalog_context = (
        ProductDatabase(catalog_path)
        if catalog_path.suffix.lower() in {".sqlite", ".sqlite3", ".db"}
        else None
    )
    try:
        catalog = catalog_context or ProductCatalog(catalog_path)
        for item in traveler.items:
            if selected and _normalize_name(item.document_remark) not in selected:
                continue
            reason = mappings.ignored_reason(item.name)
            if reason is not None:
                ignored.append({**item.source_snapshot(), "reason": reason})
                continue
            try:
                outbound.extend(match_item(catalog, mappings, item))
            except RuleError as exc:
                missing.append({
                    **item.source_snapshot(),
                    "code": exc.code,
                    "message": str(exc),
                    **exc.context,
                })
    finally:
        if catalog_context is not None:
            catalog_context.close()
    duplicate_keys = [
        key for key, count in Counter(
            (item.document_remark, item.product_code) for item in outbound
        ).items()
        if count > 1 and _normalize_name(key[1]) not in {"M1068", "M1069"}
    ]
    if duplicate_keys:
        conflicts = [
            asdict(item) for item in outbound
            if (item.document_remark, item.product_code) in duplicate_keys
        ]
        missing.append({"code": "duplicate_product_code", "message": "多个 Traveler 项目映射到同一商品编号", "items": conflicts})
    return InventoryPreview(
        traveler,
        outbound,
        ignored,
        missing,
        selected_document_remarks=requested_remarks,
    )


def set_ignored_mapping(config: Config, name: str, ignored: bool, reason: str = "") -> dict:
    """新增或移除未来导入使用的忽略规则，并返回规则摘要。

    参数：config：包含状态库、订单来源及备份路径的运行配置；name：来源材料名称或查询名称；ignored：是否启用忽略规则；reason：忽略规则或不出库决定的原因。
    """
    mappings = InventoryMappings(config.workflow_database)
    if ignored:
        mappings.save_ignored(name, reason)
        removed_database_rows = 0  # 忽略规则仅影响未来导入，不修改已确认 SKU 事实。
    else:
        mappings.remove_ignored(name)
        removed_database_rows = 0
    return {
        "ok": True,
        "traveler_name": name,
        "ignored": ignored,
        "removed_database_rows": removed_database_rows,
    }


def update_ignored_mapping(config: Config, old_name: str, name: str, reason: str = "") -> dict:
    """重命名或修改导入忽略规则，不改已确认 SKU 事实。

    参数：config：包含状态库、订单来源及备份路径的运行配置；old_name：更新前的来源名称；name：来源材料名称或查询名称；reason：忽略规则或不出库决定的原因。
    """
    mappings = InventoryMappings(config.workflow_database)
    old_normalized = _normalize_name(old_name)
    new_normalized = _normalize_name(name)
    if not old_normalized or not new_normalized:
        raise RuleError("inventory_mapping", "忽略项目名称不能为空")
    if old_normalized != new_normalized:
        mappings.remove_ignored(old_name)
    mappings.save_ignored(name, reason)
    removed_database_rows = 0  # 忽略规则仅影响未来导入，不修改已确认 SKU 事实。
    return {
        "ok": True,
        "old_name": old_name,
        "traveler_name": name,
        "ignored": True,
        "removed_database_rows": removed_database_rows,
    }


def list_inventory_mappings(config: Config) -> dict:
    """读取全部人工映射、显示别名和忽略规则。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    mappings = InventoryMappings(config.workflow_database)
    manual, ignored = mappings.entries()
    return {
        "ok": True,
        "manual": {item["name"]: item["product_code"] for item in manual},
        "manual_display_names": {item["name"]: item["display_name"] for item in manual},
        "ignored": {item["name"]: item["reason"] for item in ignored},
    }


def search_inventory_products(config: Config, query: str) -> dict:
    """按查询词搜索商品，并返回可供人工映射选择的记录。

    参数：config：包含状态库、订单来源及备份路径的运行配置；query：商品搜索关键词。
    """
    token = _normalize_name(query)
    if not token:
        raise RuleError("inventory_argument", "请输入商品编号、名称或规格")
    with ProductDatabase(bootstrap_product_database(config)) as catalog:
        matches = [
            product for product in catalog.find(contains=token)
            if not product.status or product.status == "启用"
        ]
    matches.sort(key=lambda product: (
        0 if token in {_normalize_name(product.code), _normalize_name(product.name)} else 1,
        product.code,
    ))
    return {"query": query, "products": [asdict(product) for product in matches[:50]]}


def save_manual_mapping(
    config: Config, name: str, product_code: str, display_name: str = ""
) -> dict:
    """校验目标商品并保存来源名称到 SKU 的人工映射。

    参数：config：包含状态库、订单来源及备份路径的运行配置；name：来源材料名称或查询名称；product_code：已确认的库存商品 SKU；display_name：用于界面显示的商品别名。
    """
    with ProductDatabase(bootstrap_product_database(config)) as catalog:
        product = catalog.require_code(product_code)
    mappings = InventoryMappings(config.workflow_database)
    mappings.save_manual(name, product.code, display_name)
    return {
        "ok": True,
        "traveler_name": name,
        "display_name": mappings.display_name_for_hardware(product.code, name),
        "product": asdict(product),
    }


def remove_manual_mapping(config: Config, name: str) -> dict:
    """移除来源名称的人工映射并返回删除结果。

    参数：config：包含状态库、订单来源及备份路径的运行配置；name：来源材料名称或查询名称。
    """
    mappings = InventoryMappings(config.workflow_database)
    mappings.remove_manual(name)
    return {"ok": True, "traveler_name": name, "removed": True}


def update_manual_mapping(
    config: Config, old_name: str, name: str, product_code: str, display_name: str = ""
) -> dict:
    """校验新商品并更新映射名称、SKU 和显示别名。

    参数：config：包含状态库、订单来源及备份路径的运行配置；old_name：更新前的来源名称；name：来源材料名称或查询名称；product_code：已确认的库存商品 SKU；display_name：用于界面显示的商品别名。
    """
    mappings = InventoryMappings(config.workflow_database)
    old_normalized = _normalize_name(old_name)
    new_normalized = _normalize_name(name)
    if not old_normalized or not new_normalized:
        raise RuleError("inventory_mapping", "映射材料名称不能为空")
    with ProductDatabase(bootstrap_product_database(config)) as catalog:
        product = catalog.require_code(product_code)
    if old_normalized != new_normalized:
        mappings.remove_manual(old_name)
    mappings.save_manual(name, product.code, display_name)
    return {
        "ok": True,
        "old_name": old_name,
        "traveler_name": name,
        "display_name": mappings.display_name_for_hardware(product.code, name),
        "product": asdict(product),
    }


class InventoryOperationJournal:
    """保存浏览器与数据库联合操作的意图及结果。日志是操作元数据，不是生产或出库事实；在浏览器动作前写入，使重试能区分外部已保存与尚未到达金蝶的操作。"""

    def __init__(self, database: Path):
        """确保库存操作日志结构可用并保存数据库路径。

        参数：self：当前实例；database：集中业务数据库路径。
        """
        ensure_schema(database)
        self.database = database

    @staticmethod
    def _canonical(value: object) -> str:
        """将对象序列化为排序且紧凑的 JSON，供幂等身份比较。

        参数：value：待规范化或校验的原始值。
        """
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))

    def prepare(
        self,
        operation_kind: str,
        order_id: str,
        factory_orders: Iterable[str],
        payload: dict,
    ) -> dict:
        """按动作、订单、工厂单和业务载荷生成幂等操作记录，保留已有外部确认的恢复信息。

        参数：self：当前实例；operation_kind：需要登记的库存操作类别；order_id：销售订单号；factory_orders：本次操作关联的工厂单集合；payload：待序列化或登记的结构化载荷。
        """
        normalized_order = str(order_id or "").strip().upper()
        normalized_factories = sorted({
            str(value or "").strip().upper()
            for value in factory_orders
            if str(value or "").strip()
        })
        payload_json = self._canonical(payload)
        identity_payload = json.loads(payload_json)
        identity_payload.pop("recovery_context", None)
        identity_draft = identity_payload.get("production_draft")
        if isinstance(identity_draft, dict):
            # 每次安全重试都会生成新的内存批次号；不能因此掩盖相同订单、工厂单及材料数量已有的外部保存确认。
            identity_draft.pop("batch_number", None)  # 兼容先前保存的恢复载荷
            identity_draft.pop("request_id", None)
            identity_draft.pop("production_time", None)
        identity_payload_json = self._canonical(identity_payload)
        payload_fingerprint = hashlib.sha256(identity_payload_json.encode("utf-8")).hexdigest()
        identity = self._canonical({
            "kind": operation_kind,
            "order_id": normalized_order,
            "factory_orders": normalized_factories,
            "payload_fingerprint": payload_fingerprint,
        })
        operation_id = hashlib.sha256(identity.encode("utf-8")).hexdigest()
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        connection = connect_database(self.database)
        connection.row_factory = sqlite3.Row
        try:
            connection.execute(
                """insert into inventory_operations(
                       operation_id, operation_kind, order_id,
                       factory_orders_json, payload_json, payload_fingerprint,
                       status, document_results_json, attempt_count,
                       last_error, created_at, updated_at
                   ) values(?,?,?,?,?,?,'prepared','[]',0,'',?,?)
                   on conflict(operation_id) do update set
                       payload_json=case
                           when inventory_operations.status in ('external_confirmed','local_committed','verification_required','partial_external_confirmed')
                           then inventory_operations.payload_json
                           else excluded.payload_json
                       end,
                       updated_at=excluded.updated_at""",
                (
                    operation_id,
                    operation_kind,
                    normalized_order,
                    self._canonical(normalized_factories),
                    payload_json,
                    payload_fingerprint,
                    now,
                    now,
                ),
            )
            connection.commit()
            row = connection.execute(
                "select * from inventory_operations where operation_id=?",
                (operation_id,),
            ).fetchone()
            return dict(row) if row is not None else {}
        finally:
            connection.close()

    def update(
        self,
        operation_id: str,
        status: str,
        *,
        results: Iterable[dict] | None = None,
        error: str = "",
        increment_attempt: bool = False,
    ) -> None:
        """更新库存操作状态、错误和可选单据结果，按需增加尝试次数。

        参数：self：当前实例；operation_id：持久库存操作标识；status：新的操作状态；results：外部返回的单据结果列表；error：本次失败的诊断说明；increment_attempt：是否将尝试次数加一。
        """
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        connection = connect_database(self.database)
        try:
            assignments = ["status=?", "last_error=?", "updated_at=?"]
            values: list[object] = [status, str(error or ""), now]
            if results is not None:
                assignments.append("document_results_json=?")
                values.append(self._canonical(list(results)))
            if increment_attempt:
                assignments.append("attempt_count=attempt_count+1")
            values.append(operation_id)
            connection.execute(
                f"update inventory_operations set {', '.join(assignments)} where operation_id=?",
                values,
            )
            connection.commit()
        finally:
            connection.close()

    @staticmethod
    def decoded_payload(row: dict) -> dict:
        """安全解析操作记录的请求载荷，格式不符返回空字典。

        参数：row：库存操作日志数据库记录。
        """
        try:
            value = json.loads(str(row.get("payload_json", "{}")))
        except json.JSONDecodeError:
            return {}
        return value if isinstance(value, dict) else {}

    @staticmethod
    def decoded_results(row: dict) -> list[dict]:
        """安全解析操作记录的单据结果，仅保留字典元素。

        参数：row：库存操作日志数据库记录。
        """
        try:
            value = json.loads(str(row.get("document_results_json", "[]")))
        except json.JSONDecodeError:
            return []
        return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []


RECOVERABLE_INVENTORY_STATES = ("submitting", "verification_required", "external_confirmed", "partial_external_confirmed")


def _completed_shipment_documents(connection, operation: dict) -> list[str]:
    """只读识别被后续出货完成的旧操作；单号、原计划和本地工厂关联必须完整一致。"""
    payload = InventoryOperationJournal.decoded_payload(operation)
    documents = payload.get("documents")
    # 生产草稿和状态恢复失败还需要业务核对，不能仅凭出库单关闭。
    if (operation.get("operation_kind") != "shipment" or payload.get("production_draft")
            or str(operation.get("last_error", "")).startswith("本地单据已恢复")
            or not isinstance(documents, list) or not documents):
        return []

    def quantities(items):
        """保留唯一 SKU 和有效数量；坏数据或重复 SKU 不参与自动消除。"""
        if not isinstance(items, list) or not items:
            return None
        result = {}
        try:
            for item in items:
                code = str(item["productCode"]).strip().upper()
                quantity = Decimal(str(item["quantity"]))
                if not code or code in result or not quantity.is_finite() or quantity <= 0:
                    return None
                result[code] = quantity
        except (KeyError, TypeError, ValueError, ArithmeticError):
            return None
        return result

    try:
        started = datetime.fromisoformat(operation["created_at"])
        scope = json.loads(operation["factory_orders_json"])
        if not isinstance(scope, list) or not scope or not all(isinstance(value, str) for value in scope):
            return []
    except (KeyError, TypeError, ValueError):
        return []
    later_cursor = connection.execute("""select * from inventory_operations
        where order_id=? and operation_kind='shipment' and status='local_committed'""", (operation["order_id"],))
    later = [dict(zip([column[0] for column in later_cursor.description], row)) for row in later_cursor]
    original_results = InventoryOperationJournal.decoded_results(operation)
    numbers, remarks, covered_factories = [], set(), set()
    used_later = False
    for document in documents:
        if not isinstance(document, dict) or document.get("kind") != "hardware":
            return []
        remark = document.get("remark")
        expected = quantities(document.get("items"))
        if not isinstance(remark, str) or not remark or remark in remarks or expected is None:
            return []
        remarks.add(remark)
        candidates = set()
        original = [r for r in original_results if r.get("remark") == remark]
        if original:
            if len(original) != 1 or not (original[0].get("saved") or original[0].get("unchanged")):
                return []
            original_number = original[0].get("documentNumber", "")
            if not isinstance(original_number, str):
                return []
            candidates.add(original_number)
        else:
            for completed in later:
                try:
                    if datetime.fromisoformat(completed["created_at"]) <= started:
                        continue
                except (TypeError, ValueError):
                    continue
                completed_payload = InventoryOperationJournal.decoded_payload(completed)
                if completed_payload.get("production_draft"):
                    continue
                plans = completed_payload.get("documents", [])
                if not isinstance(plans, list):
                    continue
                matches = [d for d in plans if isinstance(d, dict) and d.get("remark") == remark]
                if (len(matches) != 1 or matches[0].get("kind") != "hardware"
                        or quantities(matches[0].get("items")) != expected):
                    continue
                for result in InventoryOperationJournal.decoded_results(completed):
                    if result.get("remark") == remark and (result.get("saved") or result.get("unchanged")):
                        candidate_number = result.get("documentNumber", "")
                        if not isinstance(candidate_number, str):
                            return []
                        candidates.add(candidate_number)
            used_later = True
        if len(candidates) != 1:
            return []
        number = candidates.pop()
        known = document.get("knownDocumentNumber")
        if not isinstance(number, str) or not re.fullmatch(r"QTCK\d+", number) or (known and known != number):
            return []
        local = connection.execute("""select items_json from outbound_documents
            where document_number=? and order_id=? and document_type='hardware' and status='已出库'""",
            (number, operation["order_id"])).fetchone()
        if not local:
            return []
        try:
            if quantities(json.loads(local[0])) != expected:
                return []
        except (TypeError, ValueError):
            return []
        factories = connection.execute("""select factory_order from factory_orders
            where order_id=? and (factory_order=? or factory_name=?)""",
            (operation["order_id"], remark, remark)).fetchall()
        if len(factories) != 1 or factories[0][0] not in scope:
            return []
        covered_factories.add(factories[0][0])
        links = connection.execute("""select order_id, factory_order from outbound_document_factories
            where document_number=?""", (number,)).fetchall()
        if [tuple(link) for link in links] != [(operation["order_id"], factories[0][0])]:
            return []
        numbers.append(number)
    if any(not isinstance(result.get("remark"), str) or result["remark"] not in remarks for result in original_results):
        return []
    return numbers if used_later and len(set(numbers)) == len(documents) and covered_factories == set(scope) else []


def pending_inventory_operations(database: Path, *, connection=None) -> list[dict]:
    """将已提交但尚未核对完成的操作投影为待处理项，不持久化提示文案。"""
    owned = connection is None
    connection = connection or connect_database(database)
    try:
        cursor = connection.execute("""select * from inventory_operations
            where status in ('submitting','verification_required','external_confirmed','partial_external_confirmed')
            order by updated_at desc""")
        rows = [dict(zip([column[0] for column in cursor.description], row)) for row in cursor]
        result = []
        for row in rows:
            if _completed_shipment_documents(connection, row):
                continue
            key, order, error = row["operation_id"], row["order_id"], row["last_error"]
            created, updated, results = row["created_at"], row["updated_at"], row["document_results_json"]
            numbers = [str(row.get("documentNumber", "")) for row in InventoryOperationJournal.decoded_results(
                {"document_results_json": results}) if row.get("documentNumber")]
            message = f"{created[:10]} 发起的出库操作还需要核对。系统可能已出库，本地记录尚未确认完整。点击下方按钮重新检查；核对一致后补齐本地记录，不会再次扣库存。"
            if numbers:
                message += " 已返回单据：" + "、".join(numbers)
            if error:
                message += "。上次未能完成检查：" + str(error).replace("；库存未确认成功，本次未写入本地业务数据库", "")
            result.append(dict(issue_key="inventory_recovery:" + key, kind="inventory_recovery",
                               order_id=order, factory_order="", path="", message=message,
                               status="open", first_seen=created, last_seen=updated, resolved_at=""))
        return result
    finally:
        if owned:
            connection.close()


def recover_inventory_operation(config: Config, operation_id: str) -> dict:
    """逐张只读核对单号、备注、SKU 和数量，全部一致后原子补齐本地事实；绝不提交外部出库。"""
    journal = InventoryOperationJournal(config.workflow_database)
    connection = connect_database(config.workflow_database)
    connection.row_factory = sqlite3.Row
    try:
        row = connection.execute("select * from inventory_operations where operation_id=?", (operation_id,)).fetchone()
        if row is None:
            raise RuleError("inventory_operation_missing", "找不到该操作，请刷新待处理中心")
        row = dict(row)
        if row["status"] == "local_committed":
            return {"ok": True, "syncRecorded": True, "alreadyRecovered": True}
        if row["status"] not in RECOVERABLE_INVENTORY_STATES:
            raise RuleError("inventory_operation_state", "该操作没有需要恢复的外部结果")
        completed = _completed_shipment_documents(connection, row)
        if completed:
            return {"ok": True, "syncRecorded": True, "alreadyRecovered": True,
                    "documentNumbers": completed,
                    "message": "后续操作已补齐本地出库记录，旧提示已移除；本次未修改库存或出库记录"}
    finally:
        connection.close()
    payload = journal.decoded_payload(row)
    if row.get('operation_kind') == 'aicnc_rework':
        from .aicnc_import import recover
        return recover(config, payload['optimization_id'])
    documents = payload.get("documents", [])
    if not documents or not all(isinstance(doc, dict) and doc.get("remark") and doc.get("items") for doc in documents):
        raise RuleError("inventory_operation_corrupt", "恢复记录缺少原始单据明细，无法安全恢复；请核对操作日志")
    if len({doc["remark"] for doc in documents}) != len(documents):
        raise RuleError("inventory_operation_corrupt", "恢复记录的单据备注重复，已停止")
    original_results = {item.get("remark"): item for item in journal.decoded_results(row)}
    results = []
    try:
        for document in documents:
            known = str(original_results.get(document["remark"], {}).get("documentNumber") or document.get("knownDocumentNumber") or "")
            verified = run_jdy(config, "verifyOutbound", order_name=document["remark"],
                               verification_document={"orderName": document["remark"], "items": document["items"],
                                                      "knownDocumentNumber": known})
            number = str(verified.get("documentNumber", ""))
            if (not verified.get("verified") or not re.fullmatch(r"QTCK\d+", number)
                    or verified.get("remark") != document["remark"] or (known and known != number)):
                raise RuleError("inventory_verification_required", "外部单据未完整匹配，未修改本地；请核对单号、备注和商品数量后重新检查")
            results.append({**verified, "saved": True, "kind": document["kind"]})
        if len({item["documentNumber"] for item in results}) != len(results):
            raise RuleError("inventory_verification_required", "多个备注指向同一单据，已停止恢复")
        context = payload.get("recovery_context", {})
        # 旧记录仍有完整计划及生产草稿；按已保存的操作身份恢复，绝不重新解析变动后的 Server。
        source_type = context.get("source_type", "database" if payload.get("production_draft") else "traveler")
        traveler = TravelerData(Path(context.get("source_path") or "."), row["order_id"], row["order_id"],
                                row["order_id"], [], [], {}, "", "")
        preview = InventoryPreview(traveler, [], selected_factory_orders=tuple(json.loads(row["factory_orders_json"])),
                                   source_type=source_type)
        InventorySyncStore(config.workflow_database, config.backup_root).save_success(
            preview, results, production_draft=payload.get("production_draft"), operation_id=operation_id,
            prepared_documents=documents,
        )
    except Exception as exc:
        journal.update(operation_id, "verification_required", error=str(exc))
        raise
    from .order_index import reconcile_outbound_statuses, record_temporary_outbound
    try:
        reconcile_outbound_statuses(config, order_ids=[row["order_id"]], factory_orders=preview.selected_factory_orders or None)
        if context.get("source_path"):
            record_temporary_outbound(config, preview.traveler.path, {"saved": True, "results": results})
    except Exception as exc:
        journal.update(operation_id, "verification_required", results=results,
                       error="本地单据已恢复，订单状态核对未完成：" + str(exc))
        raise
    return {"ok": True, "syncRecorded": True, "results": results,
            "message": "已核对外部单据并恢复本地记录，未再次扣库存"}


class InventorySyncStore:
    def __init__(self, path: Path, backup_root: Path):
        # 出库事实已迁入 workflow.sqlite3；path 仅保留调用方传入的旧文件名，不再读取或写入该文件，防止归档旧 JSON 意外恢复为事实来源。
        """初始化集中出库事实存储与备份位置，不读取历史 JSON。

        参数：self：当前实例；path：用于推导同目录 workflow.sqlite3 的历史同步文件路径；不读写旧 JSON；backup_root：配置的备份根目录。
        """
        self.path = path
        self.database = path.parent / "workflow.sqlite3"
        self.backup_root = backup_root / "Inventory Sync Records"
        ensure_schema(self.database)

    @staticmethod
    def key(order_id: str, remark: str) -> str:
        """对规范化订单号及单据备注计算稳定身份键。

        参数：order_id：销售订单号；remark：用于单据识别的备注。
        """
        return hashlib.sha256(
            f"{_normalize_name(order_id)}\n{_normalize_name(remark)}".encode()
        ).hexdigest()

    @staticmethod
    def raw_document_fingerprint(items: list[TravelerItem]) -> str:
        """按来源名称和数量排序生成原始单据指纹。

        参数：items：待汇总、解析或计算指纹的材料/五金条目。
        """
        return _fingerprint({
            "items": sorted(
                (
                    _normalize_name(item.name),
                    float(item.quantity),
                )
                for item in items
            )
        })

    @staticmethod
    def mapped_document_fingerprint(items: list[OutboundItem]) -> str:
        """按大写 SKU 和数量排序生成实际出库明细指纹。

        参数：items：待汇总、解析或计算指纹的材料/五金条目。
        """
        return _fingerprint({
            "items": sorted(
                (item.product_code.upper(), float(item.quantity))
                for item in items
            )
        })

    @staticmethod
    def canonical_document_fingerprint(
        items: list[TravelerItem],
    ) -> str | None:
        """在每个来源条目都有 SKU 时计算规范单据指纹。直接读 Traveler 的条目缺少 SKU，仍依赖历史原始指纹；数据库投影可据此发现名称数量不变但 SKU 改变的情况。

        参数：items：待汇总、解析或计算指纹的材料/五金条目。
        """
        values: list[tuple[str, float]] = []
        for item in items:
            product_code = str(item.product_code or "").strip().upper()
            if not product_code:
                return None
            quantity = float(item.quantity)
            if EDGE_RE.fullmatch(str(item.name or "").strip()):
                quantity = float(
                    Decimal(str(quantity)).quantize(
                        Decimal("1"), rounding=ROUND_HALF_UP
                    )
                )
            values.append((product_code, quantity))
        return _fingerprint({"items": sorted(values)})

    def _records(self) -> list[dict]:
        """读取持久出库单及工厂单关联，展开为按单据备注可检索的历史记录。

        参数：self：当前实例。
        """
        connection = connect_database(self.database)
        try:
            rows = connection.execute(
                """select document_number, document_type, order_id, factory_order,
                          status, issued_at, source_path, document_url, items_json,
                          raw_fingerprint, mapped_fingerprint
                   from outbound_documents order by document_number"""
            ).fetchall()
            records: list[dict] = []
            for row in rows:
                links = [
                    str(item[0]).strip()
                    for item in connection.execute(
                        "select factory_order from outbound_document_factories where document_number=? order by factory_order",
                        (row[0],),
                    ).fetchall()
                    if str(item[0]).strip()
                ]
                remarks = links or [str(row[3] or row[2] or "").strip()]
                if row[1] == 'rework_materials':
                    from .aicnc_import import optimization_id
                    remarks = [f"{row[2]} 返工 {optimization_id(Path(row[6]))}"]
                try:
                    items = json.loads(row[8] or "[]")
                except (TypeError, json.JSONDecodeError):
                    items = []
                for remark in remarks:
                    records.append({
                        "document_number": row[0],
                        "kind": row[1],
                        "order_id": row[2],
                        "remark": remark,
                        "status": row[4],
                        "synced_at": row[5],
                        "traveler_path": row[6],
                        "document_url": row[7],
                        "items": items,
                        "raw_fingerprint": row[9],
                        "mapped_fingerprint": row[10],
                    })
            return records
        finally:
            connection.close()

    def record_for_document(self, order_id: str, remark: str) -> dict | None:
        """按订单号和备注查找首个持久出库记录。

        参数：self：当前实例；order_id：销售订单号；remark：用于单据识别的备注。
        """
        order_key = _normalize_name(order_id)
        remark_key = _normalize_name(remark)
        return next(
            (
                record for record in self._records()
                if _normalize_name(record.get("order_id", "")) == order_key
                and _normalize_name(record.get("remark", "")) == remark_key
            ),
            None,
        )

    def records_for_order(self, order_id: str) -> list[dict]:
        """读取指定订单的全部持久出库记录。

        参数：self：当前实例；order_id：销售订单号。
        """
        normalized = _normalize_name(order_id)
        return [
            record for record in self._records()
            if _normalize_name(str(record.get("order_id", ""))) == normalized
        ]

    def status_for(self, traveler: TravelerData) -> tuple[str, str]:
        """比较当前来源单据与历史指纹，返回出库状态及已知单号。

        参数：self：当前实例；traveler：已解析的 Traveler 来源数据。
        """
        records = self.records_for_order(traveler.order_id)
        if not records:
            return "未出库", ""
        current_remarks = set(traveler.documents)
        if any(
            record.get("remark") not in current_remarks
            or not traveler.documents.get(str(record.get("remark", "")), [])
            for record in records
        ):
            return "需要人工处理", ""
        document_numbers = []
        missing = False
        changed = False
        for remark, items in traveler.documents.items():
            if not items:
                continue
            record = self.record_for_document(traveler.order_id, remark)
            if not record:
                missing = True
                continue
            document_numbers.append(str(record.get("document_number", "")))
            if record.get("raw_fingerprint") != self.raw_document_fingerprint(items):
                changed = True
                continue
            canonical_fingerprint = self.canonical_document_fingerprint(items)
            if (
                canonical_fingerprint is not None
                and record.get("mapped_fingerprint")
                and record.get("mapped_fingerprint") != canonical_fingerprint
            ):
                changed = True
        if changed:
            return "需要更新", "、".join(filter(None, document_numbers))
        if missing:
            return "未出库", "、".join(filter(None, document_numbers))
        return "已出库", "、".join(filter(None, document_numbers))

    def prepare_documents(self, preview: InventoryPreview) -> list[dict]:
        """比较预览与历史单据生成新建或更新载荷；完整范围内单据消失时要求人工处理。

        参数：self：当前实例；preview：已解析 SKU 及出库范围的预览。
        """
        selected = preview._selected_document_set()
        document_kinds = {
            str(kind).strip().casefold()
            for kind in preview.document_kinds
            if str(kind).strip()
        }
        previous = {
            str(record.get("remark", "")): record
            for record in self.records_for_order(preview.traveler.order_id)
            if record.get("kind") != "rework_materials"
            and (not document_kinds or str(record.get("kind", "")).strip().casefold() in document_kinds)
            and (not selected or _normalize_name(str(record.get("remark", ""))) in selected)
        }
        current = {
            remark for remark in preview.traveler.documents
            if not selected or _normalize_name(remark) in selected
        }
        disappeared = sorted(
            remark for remark in previous
            if remark and (remark not in current or not preview.traveler.documents.get(remark))
        )
        if disappeared and not preview.partial_scope:
            raise RuleError(
                "inventory_manual_void",
                "以下已出库工厂单在当前 Traveler 中消失或五金为空，请人工删除或作废旧出库单后再处理："
                + "、".join(disappeared),
                remarks=disappeared,
            )
        payloads = []
        for document in preview.document_payloads():
            remark = document["remark"]
            items = [
                OutboundItem(**item) for item in document["items"]
            ]
            if not items:
                continue
            record = self.record_for_document(preview.traveler.order_id, remark)
            mapped_fingerprint = self.mapped_document_fingerprint(items)
            payloads.append({
                "remark": remark,
                "kind": document["kind"],
                "items": [
                    {"productCode": item.product_code, "quantity": item.quantity}
                    for item in items
                ],
                # Traveler 来源变化时，即使映射后商品代码相同也应更新原出库单；原始指纹反映 Server/报表变化，映射指纹保护实际明细。
                "changed": (
                    not record
                    or record.get("mapped_fingerprint") != mapped_fingerprint
                    or record.get("raw_fingerprint") != self.raw_document_fingerprint(
                        preview.traveler.documents.get(remark, [])
                    )
                ),
                "knownDocumentNumber": str(record.get("document_number", "")) if record else "",
                "mappedFingerprint": mapped_fingerprint,
                "rawFingerprint": self.raw_document_fingerprint(
                    preview.traveler.documents.get(remark, [])
                ),
            })
        return payloads

    def save_success(
        self,
        preview: InventoryPreview,
        results: list[dict],
        production_draft: dict | None = None,
        operation_id: str = "",
        commit_operation: bool = True,
        commit_production: bool = True,
        prepared_documents: list[dict] | None = None,
    ) -> None:
        """备份后保存外部已成功单据及工厂关联，按选项提交生产事实和操作恢复状态。

        参数：self：当前实例；preview：已解析 SKU 及出库范围的预览；results：外部返回的单据结果列表；production_draft：待提交的生产事实草稿；为空时不登记生产；operation_id：持久库存操作标识；commit_operation：是否将库存操作状态标记为本地提交完成；commit_production：是否同时提交生产草稿。
        """
        self._backup_current()
        prepared = {item["remark"]: item for item in (prepared_documents if prepared_documents is not None else self.prepare_documents(preview))}
        connection = connect_database(self.database)
        try:
            for result in results:
                if not result.get("saved") and not result.get("unchanged"):
                    continue
                remark = str(result.get("remark", ""))
                plan = prepared.get(remark)
                document_number = str(result.get("documentNumber", "")).strip()
                if not plan or not document_number:
                    continue
                if prepared_documents is not None:
                    existing = connection.execute("select order_id from outbound_documents where document_number=?", (document_number,)).fetchone()
                    if existing and existing[0] != preview.traveler.order_id:
                        raise RuleError("inventory_document_conflict", "该单据在本地属于其他订单，已停止恢复，请核对单号")
                has_factory_orders = connection.execute(
                    "select 1 from sqlite_master where type='table' and name='factory_orders'"
                ).fetchone() is not None
                factory = connection.execute(
                    "select factory_order from factory_orders where order_id=? and (factory_name=? or factory_order=?) limit 1",
                    (preview.traveler.order_id, remark, remark),
                ).fetchone() if has_factory_orders else None
                is_production_material_commit = (
                    production_draft is not None
                    and preview.source_type == "database"
                    and _normalize_name(remark) == _normalize_name(preview.traveler.order_id)
                )
                linked_factory_values = (
                    () if is_production_material_commit else preview.selected_factory_orders
                    if (
                        preview.source_type == "database"
                        and _normalize_name(remark) == _normalize_name(preview.traveler.order_id)
                        and preview.selected_factory_orders
                    ) else (factory[0],) if factory else (remark,)
                )
                now = datetime.now().astimezone().isoformat(timespec="seconds")
                connection.execute(
                    """insert into outbound_documents(
                        document_number,document_type,order_id,factory_order,status,source,issued_at,source_path,
                        document_url,items_json,raw_fingerprint,mapped_fingerprint,updated_at
                    ) values(?,?,?,?,?,?,?,?,?,?,?,?,?)
                    on conflict(document_number) do update set
                        document_type=excluded.document_type, order_id=excluded.order_id,
                        factory_order=excluded.factory_order, status=excluded.status,
                        source=excluded.source, issued_at=excluded.issued_at,
                        source_path=excluded.source_path, document_url=excluded.document_url,
                        items_json=excluded.items_json, raw_fingerprint=excluded.raw_fingerprint,
                        mapped_fingerprint=excluded.mapped_fingerprint, updated_at=excluded.updated_at""",
                    (
                        document_number,
                        "production_materials" if is_production_material_commit else str(result.get("kind", "")),
                        preview.traveler.order_id,
                        factory[0] if factory else remark,
                        "已出库", "金蝶", str(result.get("syncedAt", "")), str(preview.traveler.path),
                        str(result.get("url", "")), json.dumps(plan["items"], ensure_ascii=False, separators=(",", ":")),
                        plan["rawFingerprint"], plan["mappedFingerprint"], now,
                    ),
                )
                if is_production_material_commit:
                    connection.execute(
                        "delete from outbound_document_factories where document_number=?",
                        (document_number,),
                    )
                for linked_factory_value in linked_factory_values:
                    ensure_outbound_document_factory_links(
                        connection, document_number, preview.traveler.order_id, linked_factory_value
                    )
            if production_draft is not None and commit_production:
                from .production import record_completed_production
                record_completed_production(connection, production_draft)
            if operation_id and commit_operation:
                updated = connection.execute(
                    """update inventory_operations
                       set status='local_committed', last_error='', updated_at=?
                       where operation_id=?""",
                    (datetime.now().astimezone().isoformat(timespec="seconds"), operation_id),
                ).rowcount
                if updated != 1:
                    raise RuleError("inventory_operation_missing", "库存操作恢复记录不存在，已停止本地业务提交")
            connection.commit()
        finally:
            connection.close()

    def _backup_current(self) -> None:
        """备份当前业务库并保留最近 50 份；远端路径或元数据复制失败时回退本地备份。

        参数：self：当前实例。
        """
        if not self.database.is_file():
            return
        backup_root = self.backup_root
        try:
            backup_root.mkdir(parents=True, exist_ok=True)
        except OSError:
            # 正式备份常位于 Server 挂载目录；挂载缺失不能丢弃已确认的库存同步，应在本地业务库旁保留可恢复副本。
            backup_root = self.path.parent / "database-backups" / "Inventory Sync Records"
            backup_root.mkdir(parents=True, exist_ok=True)
        destination = backup_root / f"workflow-before-outbound {datetime.now():%Y-%m-%d %H%M%S.%f}.sqlite3"
        try:
            shutil.copy2(self.database, destination)
        except OSError:
            # 部分 SMB/macOS 挂载允许写内容却拒绝 copy2() 的元数据复制；备份失败不能阻止已确认金蝶单据落库，需在本地状态卷重试。
            destination.unlink(missing_ok=True)
            backup_root = self.path.parent / "database-backups" / "Inventory Sync Records"
            backup_root.mkdir(parents=True, exist_ok=True)
            destination = backup_root / f"workflow-before-outbound {datetime.now():%Y-%m-%d %H%M%S.%f}.sqlite3"
            shutil.copyfile(self.database, destination)
        backups = sorted(backup_root.glob("workflow-before-outbound *.sqlite3"), key=lambda item: item.stat().st_mtime, reverse=True)
        for old in backups[50:]:
            old.unlink()


def _catalog_path(config: Config) -> Path:
    """返回当前商品工作簿备份路径。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    return config.state_dir / "inventory" / "current-products.xlsx"


def _database_path(config: Config) -> Path:
    """返回唯一支持的集中商品目录数据库路径。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    return config.state_dir / "workflow.sqlite3"


def _catalog_info(config: Config) -> dict:
    """读取商品目录位置、数量、更新时间及是否超过 30 天的提示。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    database = _database_path(config)
    if not database.is_file():
        return {}
    loaded = ProductDatabase(database)
    try:
        backup = _catalog_path(config)
        modified_path = backup if backup.is_file() else database
        modified = datetime.fromtimestamp(modified_path.stat().st_mtime)
        return {
            "path": str(database),
            "backup_path": str(backup),
            "count": loaded.count(),
            "modified_at": modified.isoformat(timespec="seconds"),
            "stale": datetime.now() - modified > timedelta(days=30),
        }
    finally:
        loaded.close()


def _sync_path(config: Config) -> Path:
    """返回集中出库同步事实数据库路径。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    return config.workflow_database


def _persist_completed_outbound_results(
    config: Config,
    preview: InventoryPreview | None,
    confirm_save: bool,
    responses: list[dict],
) -> list[str]:
    """后续单据失败前先持久化已成功单据。多单逐个执行，早先成功必须在重试前登记；此步骤不把整单标为出货，仍需全部事实对账。

    参数：config：包含状态库、订单来源及备份路径的运行配置；preview：已解析 SKU 及出库范围的预览；confirm_save：是否明确允许外部保存出库单；responses：此前各单据的浏览器结果。
    """
    if not (confirm_save and preview is not None):
        return []
    completed = [
        item for item in responses
        if item.get("saved") or item.get("unchanged")
    ]
    if not completed:
        return []
    InventorySyncStore(_sync_path(config), config.backup_root).save_success(
        preview, completed
    )
    return [
        str(item.get("documentNumber", "")).strip()
        for item in completed
        if str(item.get("documentNumber", "")).strip()
    ]


def _persist_single_outbound_result(
    config: Config,
    preview: InventoryPreview | None,
    result: dict,
    production_draft: dict | None = None,
) -> str:
    """在尝试下一单前记录一张外部已确认单据，缺少完整单号时阻止继续。

    参数：config：包含状态库、订单来源及备份路径的运行配置；preview：已解析 SKU 及出库范围的预览；result：当前单据的浏览器结果；production_draft：待提交的生产事实草稿；为空时不登记生产。
    """
    if preview is None or not (result.get("saved") or result.get("unchanged")):
        return ""
    document_number = str(result.get("documentNumber", "")).strip()
    if not document_number:
        raise RuleError(
            "jdy_save_result",
            "库存系统返回成功但缺少完整单据编号；请先查询库存历史核实，不要直接重复出库",
        )
    InventorySyncStore(_sync_path(config), config.backup_root).save_success(
        preview,
        [result],
        # 全部请求单据确认前，操作仍处于部分完成状态；当前已确认单据立即写入持久台账。
        operation_id="",
        production_draft=production_draft,
        commit_production=False,
    )
    return document_number


def bootstrap_catalog(config: Config) -> Path:
    """当前商品工作簿缺失时，校验并复制随包提供的最新目录。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    destination = _catalog_path(config)
    if destination.is_file():
        return destination
    resource_dir = Path(__file__).resolve().parent.parent / "resources" / "inventory"
    sources = sorted(resource_dir.glob("*.xlsx"), key=lambda item: item.stat().st_mtime, reverse=True)
    if not sources:
        return destination
    ProductCatalog(sources[0])
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(sources[0], destination)
    return destination


def bootstrap_product_database(config: Config) -> Path:
    """返回集中运行时商品库，必要时从商品工作簿初始化目录。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    if not config.storage_prepared:
        config.prepare_storage()
    destination = _database_path(config)
    source = bootstrap_catalog(config)
    if destination.is_file():
        try:
            connection = sqlite3.connect(destination)
            has_products = connection.execute(
                "select 1 from sqlite_master where type='table' and name='products'"
            ).fetchone() is not None
            product_count = (
                int(connection.execute("select count(*) from products").fetchone()[0])
                if has_products else 0
            )
            connection.close()
        except sqlite3.Error:
            has_products = False
        if has_products and product_count:
            return destination
        if source.is_file():
            import_catalog(config, source)
        return destination
    if not source.is_file():
        return destination
    import_catalog(config, source)
    return destination


def reconcile_folder_status(config: Config, folder: str) -> dict:
    """查询库存历史，将文件夹下 Traveler 与单号匹配并列出缺失及歧义，不写对账结果。

    参数：config：包含状态库、订单来源及备份路径的运行配置；folder：待对账的订单文件夹编号。
    """
    folder = folder.strip().upper()
    if not re.fullmatch(r"(?:PP\d{4}|CS\d{3})", folder):
        raise RuleError("inventory_argument", f"文件夹编号格式无效：{folder}")
    paths = sorted((config.order_root / folder).glob("Work Order Traveler(*).xlsx"))
    travelers = [parse_traveler(path) for path in paths]
    response = run_jdy(config, "findOutbound", order_name=folder)
    rows = [str(row) for row in response.get("matches", [])]
    store = InventorySyncStore(_sync_path(config), config.backup_root)
    found = []
    not_found = []
    needs_review = []
    for traveler in travelers:
        token = _normalize_name(traveler.order_name)
        candidates = [row for row in rows if token and token in _normalize_name(row)]
        document_numbers = sorted(set(
            number
            for row in candidates
            for number in re.findall(r"QTCK\d+", row, flags=re.IGNORECASE)
        ))
        if len(document_numbers) == 1:
            number = document_numbers[0].upper()
            found.append({"order_name": traveler.order_name, "document_number": number})
        elif not document_numbers:
            not_found.append({
                "order_name": traveler.order_name,
                "reason": "库存系统中没有查询到对应出库记录",
            })
        else:
            needs_review.append({
                "order_name": traveler.order_name,
                "reason": "找到多个出库单号，需要人工选择",
                "candidates": document_numbers,
            })
    return {
        "ok": True,
        "folder": folder,
        "found": found,
        "not_found": not_found,
        "needs_review": needs_review,
    }


def list_travelers(config: Config, include_history: bool = False) -> dict:
    """解析订单目录下 Traveler 并结合持久记录生成状态，收集单文件错误。

    参数：config：包含状态库、订单来源及备份路径的运行配置；include_history：是否包含起始日期之前已出库的历史文件。
    """
    database = bootstrap_product_database(config)
    initial = datetime.strptime(config.initial_date, "%Y-%m-%d")
    store = InventorySyncStore(_sync_path(config), config.backup_root)
    entries = []
    errors = []
    for path in sorted(config.order_root.rglob("Work Order Traveler(*).xlsx")):
        if not TRAVELER_RE.fullmatch(path.name):
            continue
        try:
            traveler = parse_traveler(path)
            status, document = store.status_for(traveler)
            modified = datetime.fromisoformat(traveler.modified_at)
            if not include_history and modified < initial and status == "已出库":
                continue
            entries.append({
                "path": str(path),
                "pp_folder": traveler.pp_folder,
                "file_name": path.name,
                "order_name": traveler.order_name,
                "modified_at": traveler.modified_at,
                "fingerprint": traveler.fingerprint,
                "status": status,
                "document_number": document,
            })
        except RuleError as exc:
            errors.append({"path": str(path), "code": exc.code, "message": str(exc), **exc.context})
    return {"travelers": entries, "errors": errors, "catalog": _catalog_info(config)}


def list_traveler_names(config: Config) -> dict:
    """仅枚举 Traveler 文件元信息与已知出库记录，延迟工作簿内容解析。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    store = InventorySyncStore(_sync_path(config), config.backup_root)
    records_by_path = {
        str(Path(record.get("traveler_path", "")).resolve()): record
        for record in store._records()
        if record.get("traveler_path")
    }
    entries = []
    for path in sorted(config.order_root.rglob("Work Order Traveler(*).xlsx")):
        if not TRAVELER_RE.fullmatch(path.name):
            continue
        resolved = str(path.resolve())
        record = records_by_path.get(resolved, {})
        entries.append({
            "path": resolved,
            "pp_folder": path.parent.name,
            "file_name": path.name,
            "order_name": "",
            "modified_at": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
            "fingerprint": "",
            "status": record.get("status", "未出库"),
            "document_number": record.get("document_number", ""),
        })
    return {
        "travelers": entries,
        "errors": [],
        "catalog": _catalog_info(config),
        "lazy": True,
    }


def import_catalog(config: Config, source: Path) -> dict:
    """校验商品工作簿后更新集中目录，再原子替换工作簿备份并清理旧副本。

    参数：config：包含状态库、订单来源及备份路径的运行配置；source：待导入的商品工作簿路径。
    """
    if not config.storage_prepared:
        config.prepare_storage()
    catalog = ProductCatalog(source)
    destination = _catalog_path(config)
    destination.parent.mkdir(parents=True, exist_ok=True)
    draft = destination.with_suffix(".tmp.xlsx")
    shutil.copy2(source, draft)
    ProductCatalog(draft)
    _replace_product_database(_database_path(config), catalog.products)
    os.replace(draft, destination)
    for old in destination.parent.glob("current-products*.xlsx"):
        if old != destination and old.is_file():
            old.unlink()
    return {"path": str(destination), "source": str(source), "count": len(catalog.products)}


def _catalog_change_summary(previous: list[Product], current: list[Product]) -> dict[str, int]:
    """按规范商品代码比较两版目录，统计新增、字段变化及移除数量。

    参数：previous：更新前商品列表；current：更新后商品列表。
    """
    def signature(product: Product) -> tuple[str, ...]:
        """提取商品业务字段元组并规范价格文本，供目录变更比较。

        参数：product：已匹配并确认的商品。
        """
        return (
            product.category,
            product.name,
            product.spec,
            product.status,
            product.brand,
            product.remark,
            product.unit,
            "" if product.cost_price is None else f"{product.cost_price:.12g}",
        )

    previous_by_code = {
        _normalize_name(product.code): product
        for product in previous if _normalize_name(product.code)
    }
    current_by_code = {
        _normalize_name(product.code): product
        for product in current if _normalize_name(product.code)
    }
    previous_codes = set(previous_by_code)
    current_codes = set(current_by_code)
    return {
        "added_count": len(current_codes - previous_codes),
        "updated_count": sum(
            signature(previous_by_code[code]) != signature(current_by_code[code])
            for code in previous_codes & current_codes
        ),
        "removed_count": len(previous_codes - current_codes),
    }


def _existing_catalog_products(config: Config) -> list[Product]:
    """优先读取商品工作簿，缺失时读取数据库目录，均无则返回空列表。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    catalog_path = _catalog_path(config)
    if catalog_path.is_file():
        return ProductCatalog(catalog_path).products
    database_path = _database_path(config)
    if database_path.is_file():
        with ProductDatabase(database_path) as database:
            return database.products
    return []


def update_catalog_online(config: Config) -> dict:
    """从金蝶导出最新商品目录，校验后更新本地目录并返回变化统计。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    inventory_dir = _catalog_path(config).parent
    inventory_dir.mkdir(parents=True, exist_ok=True)
    download = inventory_dir / f".products-download-{os.getpid()}.xlsx"
    try:
        progress("正在从库存系统导出商品资料")
        previous_products = _existing_catalog_products(config)
        run_jdy(config, "exportProducts", download_path=download)
        if not download.is_file():
            raise RuleError("product_catalog_download", "库存系统导出完成，但没有找到下载的商品资料文件")
        current_products = ProductCatalog(download).products
        summary = _catalog_change_summary(previous_products, current_products)
        result = import_catalog(config, download)
        progress(
            f"商品资料已更新，共 {result['count']} 个商品；"
            f"新增 {summary['added_count']} 个，更新 {summary['updated_count']} 个，"
            f"删除 {summary['removed_count']} 个"
        )
        return {"ok": True, **result, **summary}
    finally:
        download.unlink(missing_ok=True)


def _local_setting(config: Config, name: str) -> str:
    """读取本地配置项文本，文件缺失或不可解析时返回空串。

    参数：config：包含状态库、订单来源及备份路径的运行配置；name：settings.json 中的配置键名。
    """
    settings = config.state_dir / "settings.json"
    if not settings.is_file():
        return ""
    try:
        return str(json.loads(settings.read_text(encoding="utf-8")).get(name, "")).strip()
    except (OSError, json.JSONDecodeError):
        return ""


def _inventory_cdp_endpoint() -> str:
    """读取库存专用 Chrome 调试端点环境配置，空值使用默认地址。

    参数：无。
    """
    return os.environ.get(
        "TRAVELER_CHROME_CDP_ENDPOINT", INVENTORY_CDP_DEFAULT_ENDPOINT
    ).strip() or INVENTORY_CDP_DEFAULT_ENDPOINT


def _inventory_cdp_pages(endpoint: str) -> list[dict]:
    """读取本地浏览器调试标签页列表，不接触 Cookie 或凭据。

    参数：endpoint：库存专用 Chrome 本地调试地址。
    """
    try:
        with urlopen(endpoint.rstrip("/") + "/json/list", timeout=1.0) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except (OSError, ValueError, UnicodeDecodeError):
        return []
    return [item for item in payload if isinstance(item, dict)] if isinstance(payload, list) else []


def _find_existing_inventory_page(endpoint: str) -> dict | None:
    """在调试标签页中优先选择库存业务页，其次服务工作台。

    参数：endpoint：库存专用 Chrome 本地调试地址。
    """
    pages = [
        page for page in _inventory_cdp_pages(endpoint)
        if page.get("type") == "page"
        and _is_inventory_authenticated_url(str(page.get("url", "")))
    ]
    return next(
        (page for page in pages if not _is_inventory_service_workbench_url(str(page.get("url", "")))),
        None,
    ) or (pages[0] if pages else None)


def _is_inventory_domain_url(url: str) -> bool:
    """判断网址是否使用 HTTP 协议且属于金蝶主域或子域。

    参数：url：待分类的网址。
    """
    try:
        parsed = urlsplit(url)
    except ValueError:
        return False
    hostname = (parsed.hostname or "").lower().rstrip(".")
    return (
        parsed.scheme in {"http", "https"}
        and (hostname == "jdy.com" or hostname.endswith(INVENTORY_DOMAIN_SUFFIX))
    )


def _is_inventory_login_or_global_url(url: str) -> bool:
    """识别登录、退出或全球站网址，解析失败按入口处理。

    参数：url：待分类的网址。
    """
    try:
        parsed = urlsplit(url)
    except ValueError:
        return True
    path = (parsed.path or "/").lower()
    return bool(
        re.search(r"/(?:login|logout)(?:/|$)", path)
        or path.rstrip("/") == "/global"
        or parsed.query.lower().find("logout=true") >= 0
    )


def _is_inventory_authenticated_url(url: str) -> bool:
    """按网址判断候选页属于库存业务范围且非登录入口，不实际验证会话。

    参数：url：待分类的网址。
    """
    return _is_inventory_domain_url(url) and not _is_inventory_login_or_global_url(url)


def _is_inventory_service_workbench_url(url: str) -> bool:
    """识别金蝶服务工作台入口域名和路径。

    参数：url：待分类的网址。
    """
    try:
        parsed = urlsplit(url)
    except ValueError:
        return False
    return (
        (parsed.hostname or "").lower().rstrip(".") in {"service.jdy.com", "www.jdy.com"}
        and (parsed.path or "/").lower().startswith("/workbench")
    )


def _inventory_chrome_profile(config: Config) -> Path:
    """返回库存专用浏览器用户资料目录。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    return config.state_dir / "inventory" / "browser-profile-zh"


def _inventory_chrome_executable() -> Path | None:
    """按环境指定路径、系统安装及可执行搜索路径查找 Chrome。

    参数：无。
    """
    configured = os.environ.get("TRAVELER_BROWSER_EXECUTABLE", "").strip()
    candidates = [Path(configured)] if configured else []
    candidates.extend([
        Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"),
        Path("/Applications/Google Chrome Beta.app/Contents/MacOS/Google Chrome Beta"),
    ])
    if executable := shutil.which("google-chrome"):
        candidates.append(Path(executable))
    return next(
        (candidate for candidate in candidates if candidate.is_file() and os.access(candidate, os.X_OK)),
        None,
    )


def open_inventory_chrome(config: Config) -> dict:
    """复用或启动 App 专用可交互 Chrome，让用户完成库存登录。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    endpoint = _inventory_cdp_endpoint()
    pages = _inventory_cdp_pages(endpoint)
    existing_candidates = [
        page for page in pages
        if page.get("type") == "page"
        and _is_inventory_authenticated_url(str(page.get("url", "")))
    ]
    existing = next(
        (
            page for page in existing_candidates
            if not _is_inventory_service_workbench_url(str(page.get("url", "")))
        ),
        None,
    ) or (existing_candidates[0] if existing_candidates else None)
    if existing:
        return {
            "ok": True,
            "reused": True,
            "url": str(existing.get("url", "")),
        }
    if any(
        page.get("type") == "page"
        and _is_inventory_domain_url(str(page.get("url", "")))
        for page in pages
    ):
        return {
            "ok": True,
            "waitingForLogin": True,
        }

    executable = _inventory_chrome_executable()
    if executable is None:
        raise RuleError(
            "inventory_runtime",
            "未找到 Google Chrome；请安装 Chrome，或用 TRAVELER_BROWSER_EXECUTABLE 指定路径",
        )
    parsed_endpoint = urlsplit(endpoint)
    port = parsed_endpoint.port or 9222
    profile = _inventory_chrome_profile(config)
    profile.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.Popen(
            [
                str(executable),
                f"--remote-debugging-port={port}",
                f"--user-data-dir={profile}",
                "--no-first-run",
                "--no-default-browser-check",
                INVENTORY_LOGIN_URL,
            ],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            start_new_session=True,
        )
    except OSError as exc:
        raise RuleError("inventory_runtime", f"无法打开库存专用 Chrome：{exc}") from exc
    return {
        "ok": True,
        "launched": True,
        "cdpEndpoint": endpoint,
        "profileDir": str(profile),
    }


def close_inventory_chrome(config: Config) -> dict:
    """仅关闭库存调试端点对应的专用 Chrome。

    参数：config：包含状态库、订单来源及备份路径的运行配置。
    """
    root = Path(__file__).resolve().parent.parent
    node, node_modules = _resolve_jdy_runtime(root)
    helper = root / "tools" / "jdy_inventory.mjs"
    env = os.environ.copy()
    env["NODE_PATH"] = str(node_modules)
    try:
        result = subprocess.run(
            [str(node), str(helper)],
            input=json.dumps({
                "action": "closeChrome",
                "cdpEndpoint": _inventory_cdp_endpoint(),
            }),
            text=True,
            capture_output=True,
            env=env,
            check=False,
            timeout=5,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuleError("inventory_runtime", "关闭库存专用 Chrome 超过 5 秒未完成") from exc
    if result.returncode != 0:
        raise RuleError("inventory_runtime", "关闭库存专用 Chrome 失败")
    try:
        payload = json.loads(result.stdout)
    except json.JSONDecodeError as exc:
        raise RuleError("inventory_runtime", "关闭库存专用 Chrome 返回结果无法解析") from exc
    return payload if isinstance(payload, dict) else {"ok": True, "closed": False}


def _keychain_password(account: str) -> str:
    """通过本地辅助程序读取指定账号的钥匙串密码，缺失或超时报错。

    参数：account：库存登录账号。
    """
    if not account:
        raise RuleError("jdy_credentials", "请先在配置中心填写库存系统用户名")
    helper = Path(__file__).resolve().parent.parent / "bin" / "keychain-read"
    if not helper.is_file():
        raise RuleError("jdy_credentials", "库存系统钥匙串读取器尚未构建")
    try:
        result = subprocess.run(
            [str(helper), account],
            capture_output=True, text=True, check=False,
            timeout=10,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuleError("jdy_credentials", "读取库存系统密码超过 10 秒未完成，请检查钥匙串访问权限") from exc
    if result.returncode != 0 or not result.stdout.strip():
        raise RuleError("jdy_credentials", "库存系统密码尚未保存到 macOS 钥匙串")
    return result.stdout.strip()


def _jdy_error_detail(stderr: str) -> str:
    """从浏览器标准错误中提取有用失败原因，过滤进度事件和调用栈噪声。

    参数：stderr：浏览器子进程的标准错误文本。
    """
    def concise(detail: str) -> str:
        """将已知浏览器失败转换为简短操作提示，并去除冗长页面后缀。

        参数：detail：待简化的浏览器错误文本。
        """
        if "Failed to create a ProcessSingleton" in detail or "profile is already in use" in detail:
            return "上一次库存查询浏览器尚未完全退出，请稍候后再次点击查询"
        detail = detail.split("；当前页面：", 1)[0].strip()
        if "库存系统登录未成功" in detail:
            return "库存系统登录未成功；请检查密码是否已更新，或网站是否要求验证码/扫码登录，然后再次查询"
        if "验证码或安全验证" in detail or "安全验证未完成" in detail:
            return "库存系统要求完成验证码或安全验证；请先使用可见浏览器完成登录验证，再重试"
        if "左侧菜单中等待可见“仓库”" in detail:
            return "库存业务工作台已打开，但左侧“仓库”菜单未在等待时间内加载完成；请保持库存系统页面可见后重试"
        if "左侧菜单中等待可见“商品”" in detail:
            return "库存业务工作台已打开，但左侧“商品”菜单未在等待时间内加载完成；请保持库存系统页面可见后重试"
        if "otheroutbound_menu" in detail.lower() or "仓库菜单已打开，但找不到" in detail:
            return "已找到左侧“仓库”菜单，但“其他出库单”入口未出现；请保持库存系统页面可见后重试"
        if "找不到其他出库单记录列表" in detail or "记录列表控件" in detail:
            return "已进入“其他出库单”，但记录列表控件未加载完成；请保持库存系统页面可见后重试"
        if "内嵌表单" in detail or "编辑表单未加载完成" in detail:
            return "已进入“其他出库单”，但出库表单未加载完成；请保持库存系统页面可见后重试"
        if "保存前发现" in detail or "保存前商品" in detail:
            return detail
        return detail

    lines = [line.strip() for line in stderr.splitlines() if line.strip()]
    for index, text in enumerate(lines):
        prefix = "库存系统自动操作失败："
        if text.startswith(prefix):
            detail = text[len(prefix):].strip()
            if index + 1 < len(lines) and lines[index + 1].startswith("；当前页面"):
                detail += lines[index + 1]
            return concise(detail)
    for text in reversed(lines):
        if (
            text.startswith("Node.js v")
            or text.startswith("at ")
            or text.startswith("triggerUncaughtException")
            or text.startswith("file://")
        ):
            continue
        try:
            event = json.loads(text)
            if isinstance(event, dict) and event.get("event") == "progress":
                continue
        except json.JSONDecodeError:
            pass
        return concise(text)
    return "库存系统浏览器操作失败，未返回具体原因"


def _resolve_jdy_runtime(root: Path) -> tuple[Path, Path]:
    """定位可执行 Node.js 和完整 Playwright 依赖，缺失时给出诊断。

    参数：root：应用资源或项目根目录。
    """
    configured_node = os.environ.get("TRAVELER_NODE", "").strip()
    if configured_node:
        node_candidates = [Path(configured_node)]
    else:
        node_candidates = [root / "bin" / "node"]
        if path_node := shutil.which("node"):
            node_candidates.append(Path(path_node))
    node = next(
        (candidate for candidate in node_candidates if candidate.is_file() and os.access(candidate, os.X_OK)),
        None,
    )
    if node is None:
        raise RuleError(
            "inventory_runtime",
            "未找到 Node.js；请重新构建 App，或用 TRAVELER_NODE 指定可执行文件",
        )

    configured_modules = os.environ.get("TRAVELER_NODE_MODULES", "").strip()
    node_modules = Path(configured_modules) if configured_modules else root / "node_modules"
    if not all((node_modules / package).is_dir() for package in ("playwright", "playwright-core")):
        raise RuleError(
            "inventory_runtime",
            "未找到 Playwright；请重新构建 App，或用 TRAVELER_NODE_MODULES 指定依赖目录",
        )
    return node, node_modules


def _jdy_failure_boundary(action: str) -> str:
    """区分只读核对失败与出库未确认，避免将读取错误写成库存失败。"""
    if action == "verifyOutbound":
        return "本次自动核对未完成，未修改库存或本地出库记录；请在库存系统核对单据后重试"
    return "库存未确认成功，本次未写入本地业务数据库"


def run_jdy(config: Config, action: str, traveler_path: Path | None = None, confirm_save: bool = False,
            download_path: Path | None = None, order_name: str = "",
            stock_items: list[dict] | None = None,
            selected_document_remarks: Iterable[str] | None = None,
            selected_factory_orders: Iterable[str] | None = None,
            order_id: str = "",
            room_material: bool = False,
            production_request_id: str = "",
            production_materials: Iterable[dict] | None = None,
            shipment_only: bool = False,
            verification_document: dict | None = None) -> dict:
    """组织库存浏览器动作、逐单确认与本地恢复记录，实时转发进度并防止重复外部出库。

    参数：config：包含状态库、订单来源及备份路径的运行配置；action：库存浏览器动作名称；traveler_path：可选 Traveler 来源文件；confirm_save：是否明确允许外部保存出库单；download_path：商品目录导出的目标文件；order_name：历史出库查询的订单名称或备注；stock_items：按 SKU 汇总的库存查询需求；selected_document_remarks：明确选择的单据备注；为空时包含全部；selected_factory_orders：明确选择的工厂单编号集合；为空时按订单范围处理；order_id：销售订单号；room_material：是否仅处理明确选择的单个房间材料；production_request_id：本次生产请求的幂等标识；production_materials：本次明确指定的生产材料；为空时使用订单材料；shipment_only：是否只组织发货五金而排除订单材料。
    """
    if action == "outbound" and order_id.strip():
        from .production import assert_order_active
        connection = connect_database(config.workflow_database)
        try:
            assert_order_active(connection, order_id)
        finally:
            connection.close()
    operation_started = time.perf_counter()
    progress(f"库存系统：开始准备 {action} 操作")
    current_production_materials = (
        [dict(item) for item in production_materials]
        if production_materials is not None
        else None
    )
    inventory_production_materials = current_production_materials
    if current_production_materials is not None:
        from .production import cumulative_production_materials

        inventory_production_materials = cumulative_production_materials(
            config,
            order_id,
            current_production_materials,
        )
    username = _local_setting(config, "jdy_username")
    cdp_endpoint = _inventory_cdp_endpoint()
    existing_inventory_page = _find_existing_inventory_page(cdp_endpoint)
    password = ""
    if existing_inventory_page is not None:
        existing_url = str(existing_inventory_page.get("url", ""))
        if _is_inventory_service_workbench_url(existing_url):
            progress(
                "库存系统：发现库存服务工作台，自动化将先点击“进入使用”进入业务系统"
                f"（{existing_url}）"
            )
        else:
            progress(
                "库存系统：发现已登录的库存专用 Chrome 页面，将直接复用，不重新登录"
                f"（{existing_url}）"
            )
    elif not (action in {"outbound", "optimizationOutbound"} and confirm_save):
        credentials_started = time.perf_counter()
        password = _keychain_password(username)
        progress(f"库存系统：未发现已登录页面，账号与钥匙串密码读取完成（用时 {time.perf_counter() - credentials_started:.2f} 秒）")
    preview = None
    store = None
    documents: list[dict] = []
    production_draft = None
    operation_journal = None
    operation_id = ""
    recovered_responses: list[dict] = []
    partial_recovery = False
    if action == "outbound":
        if order_id.strip():
            if room_material:
                selected = list(selected_factory_orders or [])
                if len(selected) != 1:
                    raise RuleError("inventory_argument", "按房间出库必须只选择一个工厂单")
                preview = build_factory_room_preview(config, order_id, selected[0])
            else:
                preview = build_database_preview(
                    config, order_id, selected_factory_orders,
                    production_request_id=production_request_id,
                    production_materials=inventory_production_materials,
                    shipment_only=shipment_only,
                )
        else:
            if not traveler_path:
                raise RuleError("inventory_argument", "出库必须提供订单号或 Traveler")
            preview = build_preview(
                traveler_path,
                bootstrap_product_database(config),
                config.workflow_database,
                selected_document_remarks=selected_document_remarks,
            )
        from .production import assert_order_active
        connection = connect_database(config.workflow_database)
        try:
            assert_order_active(connection, preview.traveler.order_id)
        finally:
            connection.close()
        if not preview.ready:
            names = sorted({
                str(item.get("name", "")).strip()
                for item in preview.missing_items
                if str(item.get("name", "")).strip()
            })
            missing = "、".join(names) or "材料"
            raise RuleError(
                "inventory_mapping_required",
                f"当前出库数据存在未映射材料：{missing}。请先在订单材料确认阶段完成材料映射后再出库",
                missing_items=preview.missing_items,
                traveler_path=str(traveler_path) if traveler_path else "",
            )
        store = InventorySyncStore(_sync_path(config), config.backup_root)
        documents = store.prepare_documents(preview)
        if not documents:
            if preview.no_outbound_required:
                if shipment_only and order_id.strip():
                    completion = mark_no_hardware_outbound(
                        config,
                        order_id,
                        selected_factory_orders,
                    )
                    return {
                        "ok": True,
                        "saved": True,
                        "no_outbound_required": True,
                        "status_only": True,
                        "database_updated": True,
                        **completion,
                        "scope_decisions": preview.scope_decisions,
                        "message": "所选工厂单没有可出库五金，已确认出货并仅更新本地状态，未创建库存出库单",
                    }
                material_scope = next(
                    (
                        decision for decision in preview.scope_decisions
                        if decision.get("scope_type") == "material"
                    ),
                    {},
                )
                if (
                    order_id.strip()
                    and material_scope.get("requirement") == "customer_supplied"
                ):
                    completion = mark_customer_supplied_outbound(
                        config,
                        order_id,
                        selected_factory_orders,
                    )
                    return {
                        "ok": True,
                        "saved": True,
                        "no_outbound_required": True,
                        "customer_supplied": True,
                        "database_updated": True,
                        **completion,
                        "scope_decisions": preview.scope_decisions,
                        "message": "客户提供材料且没有需要出库的五金，已确认出库并仅更新数据库，未创建库存出库单",
                    }
                return {
                    "ok": True,
                    "saved": False,
                    "no_outbound_required": True,
                    "scope_decisions": preview.scope_decisions,
                    "message": "本订单已明确标记为无需出库，未打开库存系统，也未创建出库单",
                }
            raise RuleError("inventory_empty", "当前订单没有需要出库的板材、封边或五金")
        if selected_factory_orders and shipment_only:
            from .production import assert_shipment_allowed
            assert_shipment_allowed(config, preview.traveler.order_id, selected_factory_orders)
        elif selected_factory_orders:
            from .order_index import assert_factory_orders_outbound_allowed

            changed_factory_orders = changed_factory_orders_for_documents(
                config,
                preview.traveler.order_id,
                selected_factory_orders,
                documents,
            )
            assert_factory_orders_outbound_allowed(
                config,
                preview.traveler.order_id,
                selected_factory_orders,
                changed_factory_orders=changed_factory_orders,
            )

        if current_production_materials is not None:
            production_draft = {
                "request_id": production_request_id,
                "order_id": preview.traveler.order_id,
                "selected_factory_orders": list(preview.selected_factory_orders),
                "materials": current_production_materials,
            }
        if confirm_save:
            pending = pending_inventory_operations(config.workflow_database)
            if any(item["order_id"] == preview.traveler.order_id for item in pending):
                raise RuleError("inventory_verification_required", "该订单仍有未核对的库存操作，请先在待处理中心核对并恢复，不要重复出库")
            operation_kind = (
                "production" if production_draft is not None
                else "shipment" if shipment_only
                else "outbound"
            )
            operation_journal = InventoryOperationJournal(config.workflow_database)
            operation_row = operation_journal.prepare(
                operation_kind,
                preview.traveler.order_id,
                preview.selected_factory_orders,
                {
                    "documents": documents,
                    "production_draft": production_draft,
                    "recovery_context": {
                        "source_path": str(preview.traveler.path),
                        "source_type": preview.source_type,
                    },
                },
            )
            operation_id = str(operation_row.get("operation_id", ""))
            operation_status = str(operation_row.get("status", ""))
            if operation_status in {"verification_required", "submitting", "partial_external_confirmed", "external_confirmed"}:
                raise RuleError(
                    "inventory_verification_required",
                    "上次库存操作的保存结果尚未确认；为避免重复出库，本次未再次提交。请先在库存历史中按订单备注核对后再处理",
                )
            if operation_status in {"external_confirmed", "local_committed", "partial_external_confirmed"}:
                recovered_responses = operation_journal.decoded_results(operation_row)
                saved_payload = operation_journal.decoded_payload(operation_row)
                saved_draft = saved_payload.get("production_draft")
                if isinstance(saved_draft, dict):
                    production_draft = saved_draft
                if not recovered_responses:
                    raise RuleError(
                        "inventory_operation_corrupt",
                        "库存操作恢复记录缺少已确认的单据结果，请人工核对库存历史",
                    )
                if operation_status == "local_committed":
                    return {
                        "ok": True,
                        "saved": True,
                        "results": recovered_responses,
                        "recovered": True,
                        "syncRecorded": True,
                        "productionCompleted": production_draft is not None,
                    }
                partial_recovery = operation_status == "partial_external_confirmed"
                progress(
                    "库存系统：发现上次已确认的外部保存结果，将跳过已完成单据，"
                    + ("继续处理未完成单据" if partial_recovery else "只补做本地事务，不重复操作库存系统")
                )
            else:
                operation_journal.update(
                    operation_id,
                    "submitting",
                    increment_attempt=True,
                )

        if confirm_save and existing_inventory_page is None and not recovered_responses:
            if operation_journal is not None and operation_id:
                operation_journal.update(
                    operation_id,
                    "failed",
                    error="没有发现已登录的库存专用 Chrome",
                )
            raise RuleError(
                "inventory_session_required",
                "没有发现已登录的库存专用 Chrome；请先在设置中打开库存系统并完成登录，再重试。本次未操作库存系统，也未写入本地业务数据",
            )

    root = Path(__file__).resolve().parent.parent
    helper = root / "tools" / "jdy_inventory.mjs"
    node = Path("")
    node_modules = Path("")
    if not recovered_responses or partial_recovery:
        runtime_started = time.perf_counter()
        node, node_modules = _resolve_jdy_runtime(root)
        progress(f"库存系统：自动化运行环境准备完成（用时 {time.perf_counter() - runtime_started:.2f} 秒）")
    request = {
        "action": action,
        "username": username,
        "profileDir": str(config.state_dir / "inventory" / "browser-profile-zh"),
        # 浏览器自动化默认在后台运行，避免抢走当前应用焦点；可见模式覆盖仅对当前进程生效，用于登录恢复和诊断。
        "headless": os.environ.get("TRAVELER_BROWSER_VISIBLE", "").strip().lower()
        not in {"1", "true", "yes", "on"},
        # 只连接通过此本地调试端点启动的 Chrome，不触碰用户普通启动的浏览器。
        "cdpEndpoint": cdp_endpoint,
        # 单次辅助进程返回 JSON 后必须退出；已有受控浏览器可连接后断开而不关闭，新启动浏览器则在任务结束后关闭，避免 Swift 无限等待子进程。
        "keepBrowserOpen": False,
        "diagnosticsDir": str(config.state_dir / "inventory" / "diagnostics"),
    }
    if existing_inventory_page is None:
        request["password"] = password
    if order_name:
        request["orderName"] = order_name
        today = datetime.now().date()
        request["queryDateFrom"] = (today - timedelta(days=550)).isoformat()
        request["queryDateTo"] = (today + timedelta(days=45)).isoformat()
    if action == "outbound":
        request.update({
            "orderName": preview.traveler.order_id,
            "documents": documents,
            "confirmSave": confirm_save,
            "queryDateFrom": (datetime.now().date() - timedelta(days=550)).isoformat(),
            "queryDateTo": (datetime.now().date() + timedelta(days=45)).isoformat(),
        })
    elif action == "optimizationOutbound":
        if not confirm_save or not verification_document:
            raise RuleError('write_confirmation_required', '返工材料出库需要确认及完整单据')
        if existing_inventory_page is None:
            raise RuleError('inventory_session_required', '请先打开库存系统并登录')
        request.update(verification_document)
        request.update(action='outbound', confirmSave=True, orderName=verification_document['remark'])
    elif action == "verifyOutbound":
        if not verification_document:
            raise RuleError("inventory_argument", "核对缺少原始单据明细")
        request.update(verification_document)
        request["action"] = "verifyOutbound"
        request["confirmSave"] = False
    elif action == "exportProducts":
        if not download_path:
            raise RuleError("inventory_argument", "更新商品资料缺少下载目标")
        request["downloadPath"] = str(download_path)
    elif action == "stockBalance":
        if not stock_items:
            raise RuleError("inventory_empty", "没有需要查询库存的板材或封边")
        request["items"] = stock_items
    env = os.environ.copy()
    env["NODE_PATH"] = str(node_modules)
    requests = [request]
    if action == "outbound":
        requests = [
            {
                **request,
                "orderName": document["remark"],
                "remark": document["remark"],
                "kind": document["kind"],
                "items": document["items"],
                "changed": document["changed"],
                "knownDocumentNumber": document["knownDocumentNumber"],
            }
            for document in request["documents"]
        ]
    responses = list(recovered_responses)
    completed_remarks = {
        str(item.get("remark", "")).strip()
        for item in recovered_responses
        if str(item.get("remark", "")).strip()
    }
    browser_requests = [
        item for item in requests
        if str(item.get("remark", "")).strip() not in completed_remarks
    ]
    for browser_request in browser_requests:
        browser_started = time.perf_counter()
        progress(f"库存系统：开始浏览器操作（第 {len(responses) + 1}/{len(requests)} 个单据）")
        stderr_lines: list[str] = []
        process = subprocess.Popen(
            [str(node), str(helper)],
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            env=env,
            bufsize=1,
        )

        def forward_browser_progress() -> None:
            """持续读取子进程标准错误，记录结构化进度并立即转发给界面。

            参数：无；使用当前浏览器子进程与日志缓冲。
            """
            assert process.stderr is not None
            for line in process.stderr:
                stderr_lines.append(line)
                try:
                    event = json.loads(line)
                except json.JSONDecodeError:
                    event = None
                if isinstance(event, dict) and event.get("event") == "progress":
                    log_progress_payload(event)
                # 立即转发：SwiftUI 在浏览器仍工作时持续读取管道，避免等 Node 辅助进程结束后才回放全部页面进度。
                sys.stderr.write(line)
                sys.stderr.flush()

        stderr_thread = threading.Thread(target=forward_browser_progress, daemon=True)
        stderr_thread.start()
        try:
            assert process.stdin is not None
            process.stdin.write(json.dumps(browser_request, ensure_ascii=False))
            process.stdin.close()
            # Node 辅助程序只向标准输出写最终结果；上方独立线程持续读取标准错误以实时显示页面进度。
            process.wait(timeout=INVENTORY_DOCUMENT_TIMEOUT_SECONDS)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait()
            stderr_thread.join(timeout=1)
            stdout_text = process.stdout.read() if process.stdout is not None else ""
            stderr_text = "".join(stderr_lines)
            detail = _jdy_error_detail(
                "\n".join(part for part in (stderr_text, stdout_text) if part)
            )
            try:
                partial_numbers = _persist_completed_outbound_results(
                    config, preview, confirm_save, responses
                )
            except Exception as sync_error:  # 保留最初的超时诊断
                detail = f"{detail}；已完成单据本地同步失败：{sync_error}"
            else:
                if partial_numbers:
                    detail = (
                        f"{detail}；已完成单据已同步：{'、'.join(partial_numbers)}；"
                        "后续单据结果需要先查询库存历史再重试"
                    )
                else:
                    detail = f"{detail}；{_jdy_failure_boundary(action)}"
            if operation_journal is not None and operation_id:
                operation_journal.update(
                    operation_id,
                    "verification_required",
                    results=responses,
                    error=detail,
                )
            remark = str(browser_request.get("remark", "")).strip()
            document_label = "单据核对 " if action == "verifyOutbound" else (f"出库单 {remark} " if remark else "当前出库单 ")
            raise RuleError(
                "jdy_timeout",
                f"{document_label}浏览器操作超过 {INVENTORY_DOCUMENT_TIMEOUT_SECONDS} 秒未完成：{detail}",
            )
        stderr_thread.join(timeout=1)
        stdout_text = process.stdout.read() if process.stdout is not None else ""
        stderr_text = "".join(stderr_lines)
        progress(
            f"库存系统：浏览器操作进程结束（用时 {time.perf_counter() - browser_started:.2f} 秒）"
        )
        if process.returncode != 0:
            partial_numbers = _persist_completed_outbound_results(
                config, preview, confirm_save, responses
            )
            detail = _jdy_error_detail(stderr_text)
            if partial_numbers:
                detail = (
                    f"{detail}；已完成单据已同步：{'、'.join(partial_numbers)}；"
                    "后续单据结果需要先查询库存历史再重试"
                )
            else:
                detail = f"{detail}；{_jdy_failure_boundary(action)}"
            if operation_journal is not None and operation_id:
                ambiguous = any(
                    marker in detail
                    for marker in ("可能已经保存", "请先查询库存历史", "请人工核实", "保存后")
                )
                operation_journal.update(
                    operation_id,
                    "verification_required" if ambiguous or responses else "failed",
                    results=responses,
                    error=detail,
                )
            raise RuleError("jdy_browser", detail)
        try:
            parsed_result = json.loads(stdout_text)
            if not isinstance(parsed_result, dict):
                raise ValueError("库存系统返回结果不是对象")
            responses.append(parsed_result)
        except (json.JSONDecodeError, ValueError) as exc:
            if operation_journal is not None and operation_id:
                operation_journal.update(
                    operation_id,
                    "verification_required",
                    results=responses,
                    error="库存系统返回结果无法解析，保存状态待核对",
                )
            raise RuleError("jdy_browser", "库存系统返回结果无法解析") from exc
        if action == "outbound" and confirm_save and operation_journal is not None and operation_id:
            operation_journal.update(
                operation_id,
                "external_confirmed" if len(responses) == len(requests) else "partial_external_confirmed",
                results=responses,
            )
            try:
                confirmed_number = _persist_single_outbound_result(
                    config, preview, parsed_result, production_draft=production_draft
                )
            except Exception as exc:
                operation_journal.update(operation_id, "verification_required", results=responses,
                                         error="外部返回成功，本地保存未完成，请在待处理中心核对并恢复")
                raise RuleError("local_sync_failed", "外部返回成功，本地保存未完成；请在待处理中心核对并恢复，不要重复出库") from exc
            if confirmed_number:
                if len(responses) < len(requests):
                    operation_journal.update(
                        operation_id,
                        "partial_external_confirmed",
                        results=responses,
                    )
                    progress(
                        f"库存系统：已完成第 {len(responses)}/{len(requests)} 个单据，"
                        f"{confirmed_number} 已同步本地，后续重试将跳过该单据"
                    )
                else:
                    operation_journal.update(
                        operation_id,
                        "external_confirmed",
                        results=responses,
                    )
    try:
        response = responses[0] if action != "outbound" else {
            "ok": True,
            "saved": confirm_save and all(
                item.get("saved") or item.get("unchanged") for item in responses
            ),
            "results": responses,
            "simulated": not confirm_save,
        }
        if action == "outbound" and response.get("saved"):
            results = response["results"]
            if preview is None or any(
                not str(item.get("documentNumber", "")).strip()
                for item in results
            ):
                if operation_journal is not None and operation_id:
                    operation_journal.update(
                        operation_id,
                        "verification_required",
                        results=results,
                        error="库存系统返回成功但缺少完整单据编号",
                    )
                raise RuleError(
                    "jdy_save_result",
                    "库存系统可能已经保存，但没有返回完整单据编号；请先查询库存历史核实，不要直接重复出库",
                )
            if operation_journal is not None and operation_id and not recovered_responses:
                operation_journal.update(
                    operation_id,
                    "external_confirmed",
                    results=results,
                )
            store = InventorySyncStore(_sync_path(config), config.backup_root)
            try:
                store.save_success(
                    preview,
                    results,
                    production_draft=production_draft,
                    operation_id=operation_id,
                )
            except Exception as exc:
                raise RuleError(
                    "local_sync_failed",
                    "库存系统已成功返回出库单，但本地数据库同步失败；请先按单据号核对库存系统，再执行本地同步，不要重复出库",
                    document_numbers=[
                        str(item.get("documentNumber", "")).strip()
                        for item in results
                        if str(item.get("documentNumber", "")).strip()
                    ],
                ) from exc
            if production_draft is not None:
                response["production"] = {
                    "request_id": production_draft.get("request_id", ""),
                    "order_id": production_draft["order_id"],
                    "factory_orders": sorted(set(production_draft["selected_factory_orders"])),
                    "status": "completed",
                }
                response["productionCompleted"] = True
            response["syncRecorded"] = True
            try:
                from .order_index import (
                    record_standard_outbound_baseline,
                    record_temporary_outbound,
                    reconcile_outbound_statuses,
                )
                reconcile_outbound_statuses(config, order_ids=[preview.traveler.order_id],
                                            factory_orders=preview.selected_factory_orders or None)
                record_temporary_outbound(config, preview.traveler.path, response)
                response["serverBaselineRecorded"] = record_standard_outbound_baseline(
                    config, preview.traveler.order_id
                )
            except (OSError, sqlite3.Error):
                # 库存系统已返回保存成功的单据，需保留该事实供界面使用；即使辅助索引暂不可用，正常出库同步记录仍是权威来源。
                response["orderIndexReconciled"] = False
                response["temporaryLedgerRecorded"] = False
                response["serverBaselineRecorded"] = False
            else:
                response["orderIndexReconciled"] = True
                response["temporaryLedgerRecorded"] = True
        progress(f"库存系统：{action} 操作结束（总用时 {time.perf_counter() - operation_started:.2f} 秒）")
        return response
    except (TypeError, KeyError) as exc:
        raise RuleError("jdy_browser", "库存系统返回结果结构不完整") from exc


def _check_requirements_stock(config: Config, requirements: list[dict]) -> list[dict]:
    """实时查询 SKU 库存，核对商品身份并计算短缺数量。

    参数：config：包含状态库、订单来源及备份路径的运行配置；requirements：含 SKU、商品名称及需求数量的记录。
    """
    response = run_jdy(config, "stockBalance", stock_items=requirements)
    raw_results = response.get("results")
    if not isinstance(raw_results, list):
        raise RuleError("jdy_browser", "库存系统没有返回库存查询明细")
    by_code = {
        str(item.get("productCode", "")).upper(): item
        for item in raw_results if isinstance(item, dict)
    }
    rows = []
    for requirement in requirements:
        code = requirement["productCode"].upper()
        actual = by_code.get(code)
        if actual is None:
            raise RuleError("jdy_browser", f"库存系统没有返回 {code} 的查询结果")
        actual_name = str(actual.get("productName", "")).strip()
        if actual_name and _normalize_name(actual_name) != _normalize_name(requirement["productName"]):
            raise RuleError(
                "inventory_product_mismatch",
                f"库存商品身份不一致：{code} 本地为 {requirement['productName']}，网页为 {actual_name}",
            )
        try:
            available = float(actual.get("availableQuantity", 0))
        except (TypeError, ValueError) as exc:
            raise RuleError("jdy_browser", f"库存系统返回的 {code} 数量不是数字") from exc
        required = float(requirement["requiredQuantity"])
        rows.append({
            **requirement,
            "availableQuantity": available,
            "shortageQuantity": max(0.0, required - available),
            "sufficient": available >= required,
        })
    return rows


def check_stock(config: Config, traveler_path: Path, include_hardware: bool = False) -> dict:
    """比较 Traveler 映射需求与实时库存，不写出库且不要求 Agent。真正库存写入须在后续单独授权的操作中执行。

    参数：config：包含状态库、订单来源及备份路径的运行配置；traveler_path：可选 Traveler 来源文件；include_hardware：是否同时统计五金；默认仅板材和封边。
    """
    preview = build_preview(traveler_path, bootstrap_product_database(config), config.workflow_database)
    requirements = stock_requirements(preview, include_hardware)
    rows = _check_requirements_stock(config, requirements)
    return {
        "ok": True,
        "traveler": str(traveler_path.resolve()),
        "includeHardware": include_hardware,
        "rows": rows,
        "hasShortage": any(not row["sufficient"] for row in rows),
    }


def check_order_stock(config: Config, order_folder: Path) -> dict:
    """根据订单来源目录材料需求查询实时库存并汇总短缺。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_folder：订单来源文件夹。
    """
    order_id, requirements = order_stock_requirements(config, order_folder)
    rows = _check_requirements_stock(config, requirements)
    return {
        "ok": True,
        "order_id": order_id,
        "folder": str(order_folder.resolve()),
        "rows": rows,
        "hasShortage": any(not row["sufficient"] for row in rows),
    }


def check_database_stock(config: Config, order_id: str) -> dict:
    """根据持久订单材料事实检查实时库存，无需 Traveler 文件。

    参数：config：包含状态库、订单来源及备份路径的运行配置；order_id：销售订单号。
    """
    normalized_order_id, requirements = database_stock_requirements(config, order_id)
    rows = _check_requirements_stock(config, requirements)
    return {
        "ok": True,
        "order_id": normalized_order_id,
        "source_type": "database",
        "rows": rows,
        "hasShortage": any(not row["sufficient"] for row in rows),
    }


def inventory_main(argv: list[str] | None = None) -> int:
    """解析库存命令行选项，分派查询、预览、规则维护和明确确认的出库操作。

    参数：argv：命令行参数；为空时读取进程参数。
    """
    parser = argparse.ArgumentParser(prog="pp-flowhub inventory")
    parser.add_argument("action", choices=(
        "list", "list-names", "preview", "order-preview", "get-outbound-scope", "import-products", "preflight", "outbound",
        "recover-operation", "find-outbound", "reconcile-folder", "ignore-item", "unignore-item",
        "search-products", "set-mapping", "update-mapping", "remove-mapping", "list-mappings", "update-ignore", "set-outbound-scope", "update-products", "stock-check", "open-chrome", "close-chrome",
    ))
    parser.add_argument("--traveler", type=Path)
    parser.add_argument("--order-id", default="")
    parser.add_argument("--source", type=Path)
    parser.add_argument("--include-history", action="store_true")
    parser.add_argument("--order-root", type=Path)
    parser.add_argument("--state-dir", type=Path)
    parser.add_argument("--confirm-save", action="store_true")
    parser.add_argument("--operation-id", default="")
    parser.add_argument("--order-name", default="")
    parser.add_argument("--name", default="")
    parser.add_argument("--ignored", choices=("true", "false"))
    parser.add_argument("--traveler-name", action="append", default=[])
    parser.add_argument("--item-name", action="append", default=[])
    parser.add_argument("--old-name", default="")
    parser.add_argument("--document-remark", action="append", default=[])
    parser.add_argument("--factory-order", action="append", default=[])
    parser.add_argument("--scope-type", choices=("material", "hardware"))
    parser.add_argument("--requirement", choices=("required", "customer_supplied", "remainder", "not_required"))
    parser.add_argument("--reason", default="")
    parser.add_argument("--folder", default="")
    parser.add_argument("--query", default="")
    parser.add_argument("--product-code", default="")
    parser.add_argument("--display-name", default="")
    parser.add_argument("--include-hardware", action="store_true")
    parser.add_argument("--room-material", action="store_true")
    parser.add_argument("--production-request", default="")
    parser.add_argument("--production-materials-json", default="")
    parser.add_argument("--shipment-only", action="store_true")
    args = parser.parse_args(argv)
    config = Config()
    config.load_settings()
    if args.order_root:
        config.order_root = args.order_root
    if args.state_dir:
        config.state_dir = args.state_dir
    config.prepare_storage()
    logger = configure_operation_log(config)
    command_started = time.perf_counter()
    logger.event("backend.command.started", "开始库存系统操作", details={"action": args.action})
    try:
        if args.action == "list":
            result = list_travelers(config, args.include_history)
        elif args.action == "list-names":
            result = list_traveler_names(config)
        elif args.action == "preview":
            if not args.traveler:
                raise RuleError("inventory_argument", "preview 必须提供 --traveler")
            result = build_preview(args.traveler, bootstrap_product_database(config), config.workflow_database).payload()
        elif args.action == "order-preview":
            if not args.order_id:
                raise RuleError("inventory_argument", "order-preview 必须提供 --order-id")
            result = build_database_preview(
                config, args.order_id, args.factory_order,
                production_request_id=args.production_request,
                shipment_only=args.shipment_only,
            ).payload()
        elif args.action == "get-outbound-scope":
            if not args.order_id:
                raise RuleError("inventory_argument", "读取出库范围需要订单号")
            result = outbound_scope_decisions(config, args.order_id, args.factory_order)
        elif args.action == "set-outbound-scope":
            if not args.order_id or not args.scope_type or not args.requirement:
                raise RuleError("inventory_argument", "保存出库范围需要订单号、范围类型和决定")
            result = set_outbound_scope(
                config,
                args.order_id,
                args.scope_type,
                args.requirement,
                factory_order=args.factory_order[0] if args.factory_order else "",
                reason=args.reason,
            )
        elif args.action == "import-products":
            if not args.source:
                raise RuleError("inventory_argument", "import-products 必须提供 --source")
            result = import_catalog(config, args.source)
        elif args.action == "update-products":
            result = update_catalog_online(config)
        elif args.action == "open-chrome":
            result = open_inventory_chrome(config)
        elif args.action == "close-chrome":
            result = close_inventory_chrome(config)
        elif args.action == "preflight":
            result = run_jdy(config, "preflight")
        elif args.action == "stock-check":
            if not args.traveler:
                raise RuleError("inventory_argument", "stock-check 必须提供 --traveler")
            result = check_stock(config, args.traveler, args.include_hardware)
        elif args.action == "recover-operation":
            result = recover_inventory_operation(config, args.operation_id)
        elif args.action == "find-outbound":
            if not args.order_name:
                raise RuleError("inventory_argument", "查询出库单必须提供 --order-name")
            result = run_jdy(config, "findOutbound", order_name=args.order_name)
        elif args.action == "reconcile-folder":
            if not args.folder:
                raise RuleError("inventory_argument", "更新文件夹状态必须提供 --folder")
            result = reconcile_folder_status(config, args.folder)
        elif args.action == "search-products":
            result = search_inventory_products(config, args.query)
        elif args.action == "set-mapping":
            item_names = args.item_name or args.traveler_name
            if not item_names or not args.product_code:
                raise RuleError("inventory_argument", "保存映射必须提供材料名称和商品编号")
            result = save_manual_mapping(config, item_names[0], args.product_code, args.display_name)
        elif args.action == "update-mapping":
            item_names = args.item_name or args.traveler_name
            if not args.old_name or not item_names or not args.product_code:
                raise RuleError("inventory_argument", "update-mapping 需要 --old-name、材料名称和商品编号")
            result = update_manual_mapping(config, args.old_name, item_names[0], args.product_code, args.display_name)
        elif args.action == "remove-mapping":
            item_names = args.item_name or args.traveler_name
            if not item_names:
                raise RuleError("inventory_argument", "删除映射必须提供材料名称")
            result = remove_manual_mapping(config, item_names[0])
        elif args.action == "list-mappings":
            result = list_inventory_mappings(config)
        elif args.action == "update-ignore":
            if not args.old_name or not args.name or args.ignored is None:
                raise RuleError("inventory_argument", "update-ignore 需要 --old-name、--name 和 --ignored")
            if args.ignored != "true":
                raise RuleError("inventory_argument", "update-ignore 只能用于保存忽略项目")
            result = update_ignored_mapping(config, args.old_name, args.name, args.reason)
        elif args.action in {"ignore-item", "unignore-item"}:
            item_names = args.item_name or args.traveler_name
            if not item_names:
                raise RuleError("inventory_argument", "修改忽略状态必须提供材料名称")
            results = [
                set_ignored_mapping(
                    config, name, args.action == "ignore-item", args.reason
                )
                for name in item_names
            ]
            result = {"ok": True, "items": results}
        else:
            production_materials = None
            if args.production_materials_json:
                try:
                    production_materials = json.loads(args.production_materials_json)
                except json.JSONDecodeError as exc:
                    raise RuleError("inventory_argument", "生产材料不是有效 JSON") from exc
                if not isinstance(production_materials, list):
                    raise RuleError("inventory_argument", "生产材料必须是数组")
            result = run_jdy(
                config,
                "outbound",
                args.traveler,
                args.confirm_save,
                selected_document_remarks=args.document_remark,
                selected_factory_orders=args.factory_order,
                order_id=args.order_id,
                room_material=args.room_material,
                production_request_id=args.production_request,
                production_materials=production_materials,
                shipment_only=args.shipment_only,
            )
        logger.event(
            "backend.command.completed",
            "库存系统操作完成",
            details={"action": args.action, "duration_seconds": round(time.perf_counter() - command_started, 6)},
        )
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except RuleError as exc:
        logger.event(
            "backend.command.failed",
            "库存系统操作失败",
            details={"action": args.action, "code": exc.code, "error": safe_rule_error_text(str(exc)), "order_id": args.order_id, "duration_seconds": round(time.perf_counter() - command_started, 6)},
        )
        print(json.dumps({"fatal": {"code": exc.code, "message": str(exc), **exc.context}}, ensure_ascii=False, indent=2))
        return 2
    except Exception as exc:
        logger.event(
            "backend.command.failed", "库存系统操作发生未预期错误",
            details={**safe_exception_details(exc, action=args.action, order_id=args.order_id),
                     "duration_seconds": round(time.perf_counter() - command_started, 6)},
        )
        print(json.dumps({"fatal": {"code": "inventory_processing_error", "message": "库存操作发生未预期错误，请查看操作日志。"}}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    raise SystemExit(inventory_main())

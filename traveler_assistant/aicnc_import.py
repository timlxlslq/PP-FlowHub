"""新版 AICNC：目录发现、只读预览、确认入账与返工库存恢复。

优化文件夹末尾的 AICNC 编号是业务身份。处理后不检查文件版本；预览快照只用于本次确认与恢复。
结构升级由显式迁移执行，普通读取不会创建新版表。
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import sqlite3
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

from .core import Config, RuleError, parse_fittings_groups
from .database import connect_database

ORDER = r"(?:PP\d{4}(?:-\d+)?|CS\d{3})"
OPTIMIZATION_RE = re.compile(rf"^({ORDER})_(?:.+-)?(\d{{14}})$", re.I)
ORDER_RE = re.compile(rf"^{ORDER}$", re.I)


def now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def encode(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def digest(value) -> str:
    return hashlib.sha256(encode(value).encode()).hexdigest()


def enabled(connection) -> bool:
    return connection.execute("select 1 from sqlite_master where name='aicnc_import_settings' and type='table'").fetchone() is not None


def require_enabled(connection) -> None:
    if not enabled(connection):
        raise RuleError("aicnc_migration_required", "新版优化台账尚未启用，需要先完成已批准的数据库迁移")


def optimization_id(folder: Path) -> str:
    match = OPTIMIZATION_RE.fullmatch(folder.name)
    if not match:
        raise RuleError("aicnc_folder", "优化文件夹必须使用 AICNC 生成的“订单号_可选说明-年月日时分秒”名称")
    try:
        datetime.strptime(match[2], "%Y%m%d%H%M%S")
    except ValueError as exc:
        raise RuleError("aicnc_folder", "优化文件夹中的日期或时间无效") from exc
    return match[2]


def migrate(connection, legacy: list[dict]) -> None:
    """只供明确批准的迁移及隔离测试调用；不从 ensure_schema 自动调用。"""
    if enabled(connection):
        return
    statements = [
        """create table aicnc_import_settings(key text primary key,value text not null)""",
        """create table aicnc_optimizations(
            optimization_id text primary key, source_folder text not null,
            status text not null check(status in ('processing','completed','ignored')),
            plan_json text not null default '{}', created_at text not null, completed_at text not null default '')""",
        """create table aicnc_material_allocations(
            optimization_id text not null references aicnc_optimizations(optimization_id),
            order_id text not null, purpose text not null check(purpose in ('normal','rework')),
            product_code text not null references products(code) on update cascade on delete restrict,
            quantity real not null check(quantity>=0),
            primary key(optimization_id,order_id,purpose,product_code))""",
        """create table aicnc_legacy_watch(
            path text primary key, files_json text not null, retired_at text not null default '')""",
    ]
    connection.execute('savepoint aicnc_migration')
    try:
        for sql in statements:
            connection.execute(sql)
        connection.execute("insert into aicnc_import_settings values('enabled_at',?)", (now(),))
        connection.execute("insert into aicnc_import_settings values('cleanup_dismissed','0')")
        connection.executemany("insert into aicnc_legacy_watch(path,files_json) values(?,?)",
                               [(r['path'], encode(r['files'])) for r in legacy])
        connection.execute('release aicnc_migration')
    except Exception:
        connection.execute('rollback to aicnc_migration')
        connection.execute('release aicnc_migration')
        raise


def legacy_snapshot(config, store, timing_sink=None):
    """仅检查切换时已登记的旧路径，永不深入新增优化文件夹。"""
    from .order_index import (_available_server_roots, _server_folder_scan_allowed,
                              _server_snapshot_folder, _server_folder_handling_mode,
                              _folder_name_order_ids, _is_standard_order_folder)
    roots = _available_server_roots(config)
    if not roots:
        raise RuleError('server_unavailable', 'Server 不可用，旧监控状态尚未确认')
    snapshot = {}
    for (path,) in store.connection.execute("select path from aicnc_legacy_watch where retired_at='' ").fetchall():
        folder = Path(path)
        if not any(folder.parent == root for root in roots):
            continue  # 不可用的来源不能被当作已退出。
        if not _server_folder_scan_allowed(config, store, folder):
            store.connection.execute("update aicnc_legacy_watch set retired_at=? where path=?", (now(), path))
            continue
        if not folder.is_dir():
            continue
        previous = store.temporary_order(path) or {}
        xml_only = not _is_standard_order_folder(folder.name) and previous.get('server_scan_policy') in {'watching','manual_pending'}
        entries, timing = _server_snapshot_folder((folder, xml_only))
        mode = _server_folder_handling_mode(store, folder)
        references = previous.get('reference_order_ids', _folder_name_order_ids(folder))
        for item in entries.values():
            item['handling_mode'] = mode
            item['reference_order_ids'] = references
            if mode in {'supplemental','external_manual'}:
                item['manual_only'] = True
                item['order_id'] = '、'.join(references)
        snapshot.update(entries)
        if timing_sink is not None:
            timing_sink.append(timing)
    return roots[0], snapshot


def discover(config, connection) -> list[dict]:
    """仅列举两层目录；完成/忽略后不进入目录，不 stat 内部文件。"""
    from .order_index import _available_server_roots, _folder_created_at
    require_enabled(connection)
    handled = {r[0] for r in connection.execute("select optimization_id from aicnc_optimizations where status in ('completed','ignored')")}
    legacy = {r[0] for r in connection.execute('select path from aicnc_legacy_watch')}
    historical_roots = legacy | {r[0] for r in connection.execute('select source_folder from temporary_orders')}
    historical_roots.update(r[0] for r in connection.execute('select path from server_folder_ignores'))
    historical_roots.update(r[0] for r in connection.execute('select source_folder from orders'))
    enabled_at = connection.execute("select value from aicnc_import_settings where key='enabled_at'").fetchone()
    cutover = datetime.fromisoformat(enabled_at[0]).timestamp() if enabled_at else 0

    def predates_cutover(folder):
        # 仅给不规范目录划定新旧边界，不比较修改时间，也不限制有效新版编号。
        try:
            return _folder_created_at(folder) < cutover
        except OSError:
            return False

    archived_orders = {r[0] for r in connection.execute("select source_folder from orders where server_scan_policy in ('legacy','permanent')")}
    historical_children = set()
    for path, source in connection.execute('select path,source_folder from source_files union select path,source_folder from server_scan_xml_state'):
        historical_roots.add(source)
        try:
            relative = Path(path).relative_to(source)
        except ValueError:
            continue
        if relative.parts:
            historical_children.add(str(Path(source) / relative.parts[0]))
    found = {}
    changes = []
    for root in _available_server_roots(config):
        for order in sorted(root.iterdir()):
            if order.is_symlink() or not order.is_dir() or order.name.startswith('.'):
                continue
            if not ORDER_RE.fullmatch(order.name):
                if str(order) not in historical_roots and (
                    OPTIMIZATION_RE.fullmatch(order.name) or not predates_cutover(order)
                ):
                    changes.append(_change(order, '请手工将此文件夹移入以订单号命名的目录，或整理目录名称', review=True))
                continue
            for folder in sorted(order.iterdir()):
                if folder.is_symlink() or not folder.is_dir() or folder.name.startswith('.'):
                    continue
                if not OPTIMIZATION_RE.fullmatch(folder.name):
                    if (str(order) not in archived_orders and str(folder) not in historical_children
                            and not predates_cutover(folder)):
                        changes.append(_change(folder, '内层目录不是 AICNC 优化编号，请手工整理', review=True))
                    continue
                try:
                    key = optimization_id(folder)
                except RuleError:
                    changes.append(_change(folder, '优化编号中的日期或时间无效，请手工核对', review=True))
                    continue
                if key in handled:
                    continue
                found.setdefault(key, []).append(folder)
    for key, paths in found.items():
        if len(paths) > 1:
            for path in paths:
                changes.append(_change(path, f'优化编号 {key} 出现在多个位置，请保留一个来源后重新扫描', review=True))
        else:
            changes.append(_change(paths[0], '新优化待确认；请选择用途并核对材料分配'))
    return changes


def _change(path, message, review=False):
    return dict(id=f'aicnc:{path}', change_type='added', kind='folder', order_id=path.parent.name,
                source_folder=str(path), path=str(path), message=message, manual_only=False,
                mixed_order=False, handling_mode='layout_review' if review else 'aicnc', event_time='')


def quantity(value, *, integer=False) -> float:
    try:
        result = float(value)
    except (ValueError, TypeError) as exc:
        raise RuleError('aicnc_quantity', f'数量无效：{value}') from exc
    if not math.isfinite(result) or result < 0 or (integer and not result.is_integer()):
        raise RuleError('aicnc_quantity', f'数量必须是有限非负{"整数" if integer else "数"}：{value}')
    return result


def parse_board(path: Path) -> dict:
    """只读大板统计和封边汇总；同色封边以十进制累计，随后沿用整米规则。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb[wb.sheetnames[0]].values)
    finally:
        wb.close()
    clean = lambda v: re.sub(r'\s+', '', str(v or ''))
    identities = ''
    for row in rows[:5]:
        for i, value in enumerate(row):
            if clean(value) == '订单号':
                identities = next((str(v).strip() for v in row[i+1:] if v is not None), '')
    factories = []
    for entry in re.split('[,，;；\n]+', identities):
        factory, sep, name = entry.strip().partition('-')
        if not sep or not re.fullmatch(r'F\d+', factory, re.I) or not name.strip():
            raise RuleError('aicnc_identity', f'{path.name} 工厂单身份无效：{entry}')
        if factory.upper() in {f['factory_order'] for f in factories}:
            raise RuleError('aicnc_identity', '板材清单中工厂单重复')
        factories.append(dict(factory_order=factory.upper(), factory_name=name.strip()))
    def section(label):
        indices = [i for i, row in enumerate(rows) if label in {clean(v) for v in row}]
        if len(indices) != 1:
            raise RuleError('aicnc_schema', f'{path.name} 必须有一个“{label}”区域')
        return indices[0]
    large, edge = section('大板统计'), section('封边汇总')
    headers = {clean(v): i for i, v in enumerate(rows[large+1]) if v is not None}
    if not {'颜色','规格','数量'} <= headers.keys():
        raise RuleError('aicnc_schema', '大板统计缺少颜色、规格或数量列')
    materials = []
    for row in rows[large+2:edge]:
        raw = str(row[headers['颜色']] or '').strip()
        match = re.fullmatch(r'([\d.]+)mm/([^/]+)/([^/]+)', raw, re.I)
        if not match:
            # 大板区在“小板小计”等区块标题处结束；小板内容完全不参与读取。
            if any(clean(v) in {'小板小计','小板统计','封边小计'} for v in row):
                break
            if any(v is not None for v in row):
                raise RuleError('aicnc_schema', f'无法识别大板材料：{raw}')
            continue
        thickness, color, base = float(match[1]), match[2].strip(), match[3].strip()
        kind = 'plywood' if base.casefold() == 'plywood' and color.casefold() == 'finished' else 'panel' if base.casefold() == 'mdf' else ''
        if not kind:
            raise RuleError('aicnc_material', f'板材尚未支持安全归类：{raw}')
        nominal = 5.4 if kind == 'plywood' and thickness == 5 else 8 if kind == 'panel' and thickness == 9 else thickness
        q = quantity(row[headers['数量']], integer=True)
        materials.append(dict(material_type=kind,color='' if kind=='plywood' else color,
                              thickness=f'{nominal:g}',quantity=q,raw_name=raw,
                              raw_spec=str(row[headers['规格']] or ''),raw_quantity=q,unit='张'))
    headers = {clean(v): i for i, v in enumerate(rows[edge+1]) if v is not None}
    if not {'颜色','封边条/米'} <= headers.keys():
        raise RuleError('aicnc_schema', '封边汇总缺少颜色或封边条/米')
    edges = defaultdict(Decimal)
    for row in rows[edge+2:]:
        color = str(row[headers['颜色']] or '').strip()
        if not color:
            continue
        q = quantity(row[headers['封边条/米']])
        edges[color] += Decimal(str(q))
    for color, q in edges.items():
        materials.append(dict(material_type='edge',color=color,thickness='',
            quantity=float(q.quantize(Decimal('1'), rounding=ROUND_HALF_UP)),raw_name=color,
            raw_spec='',raw_quantity=float(q),unit='米'))
    if not materials or not any(m['material_type'] != 'edge' and m['quantity'] for m in materials):
        raise RuleError('aicnc_material', '大板统计没有有效板材用量')
    return dict(factories=factories, materials=materials, report_path=str(path))


def _rows(connection, sql, args=()):
    cursor = connection.execute(sql, args)
    names = [v[0] for v in cursor.description]
    return [dict(zip(names, row)) for row in cursor.fetchall()]


def business_revision(connection, factories):
    # 本地相关事实改变后拒绝旧确认；不检查 Server 文件版本。
    result = []
    for f in factories:
        key = f['factory_order']
        result.append(_rows(connection, 'select factory_order,order_id,stage,ownership_status,aimes_status from factory_orders where factory_order=?', (key,)))
        result.append(_rows(connection, "select product_code,quantity,source_type from hardware_items where factory_order=? order by product_code,quantity,source_type", (key,)))
    return digest(result)


def preview(config: Config, folder: Path) -> dict:
    from .inventory import TravelerItem, resolve_inventory_items, resolved_product_code
    from .order_index import _server_root_candidates, _materialize_memory_hardware
    from .order_workflow import _material_inventory_name
    connection = connect_database(config.workflow_database)
    try:
        require_enabled(connection)
        folder = folder.expanduser().resolve()
        key = optimization_id(folder)
        if not ORDER_RE.fullmatch(folder.parent.name) or folder.parent.parent not in [p.resolve() for p in _server_root_candidates(config)]:
            raise RuleError('aicnc_folder', '请先将优化文件夹放到 Server 的订单号目录下')
        previous = connection.execute('select status,plan_json from aicnc_optimizations where optimization_id=?', (key,)).fetchone()
        if previous:
            if previous[0] in ('completed','ignored'):
                raise RuleError('aicnc_handled', '此优化已处理，不再读取内部报表')
            return dict(aicnc_preview=dict(json.loads(previous[1]), resuming=True))
        reports = [p for p in (folder/'Report').glob('*.xlsx') if not p.name.startswith('~$')]
        boards = [p for p in reports if '板材清单' in p.name]
        if len(boards) != 1:
            raise RuleError('aicnc_schema', 'Report 中必须有一份新版板材清单')
        data = parse_board(boards[0])
        for f in data['factories']:
            rows = _rows(connection, 'select * from factory_orders where factory_order=?', (f['factory_order'],))
            known = rows[0] if rows else {}
            match = re.match(rf'^({ORDER})(?:$|[-\s])', f['factory_name'], re.I)
            suggested = match[1].upper() if match else ''
            f.update(order_id=known.get('order_id') or suggested, known_order_id=known.get('order_id',''),
                     stage=known.get('stage',''), purpose='rework' if known.get('stage') in ('已生产','已出货') else 'normal' if known.get('stage') in ('已拆单','已优化') else '',
                     hardware_choice='same', hardware=[], existing_hardware=[])
            if known and known.get('aimes_status') == 'deleted':
                raise RuleError('aicnc_identity', f"工厂单 {f['factory_order']} 已失效，请核对归属")
        data['report_materials'] = [dict(m) for m in data['materials']]
        materials = [m for m in data['materials'] if m['quantity']>0]
        resolution = resolve_inventory_items(config, [(TravelerItem(i,'板材与封边',
            (f"Edge banding--{m['color']}" if m['material_type']=='edge' else _material_inventory_name(m['material_type'],m['thickness'],m['color'])),m['quantity'],''),'') for i,m in enumerate(materials)])
        if resolution.get('missing'):
            raise RuleError('order_inventory_mapping_required','材料未完成商品 SKU 处理：'+ '、'.join(str(v.get('name','')) for v in resolution['missing']))
        mapped = []
        for i,m in enumerate(materials):
            if resolution['accepted'][i].get('ignored'):
                continue
            code = resolved_product_code(resolution,i)
            if code:
                mapped.append(dict(m,product_code=code))
        # 相同 SKU 多行先合并，分配和扣库存均使用累计数量。
        by_code = {}
        for m in mapped:
            code = m['product_code']
            if code in by_code:
                old = by_code[code]
                if any(old[k] != m[k] for k in ('material_type','color','thickness')):
                    raise RuleError('aicnc_product', f'不同材料不能映射到同一商品：{code}')
                old['quantity'] += m['quantity']
                old['raw_quantity'] += m['raw_quantity']
            else:
                by_code[code] = dict(m)
        mapped = list(by_code.values())
        data['materials'] = mapped
        unshipped = tuple(f['factory_order'] for f in data['factories'] if f['stage']!='已出货')
        fittings = [p for p in reports if 'fittingslist' in p.name.casefold()]
        if unshipped and len(fittings)>1:
            raise RuleError('aicnc_schema','新版优化只能包含一份五金清单')
        source_groups = []
        if unshipped and fittings:
            for factory, items in parse_fittings_groups(fittings[0],included_factories=unshipped):
                f = next(v for v in data['factories'] if v['factory_order']==factory)
                source_groups.append(dict(factory_order=factory,order_id=f['order_id'],source_path=str(fittings[0]),
                                          items=[vars(i) for i in items]))
        for f in data['factories']:
            if f['stage']=='已出货':
                f['hardware_choice']='shipped'
                continue
            f['existing_hardware']=_rows(connection,"select h.product_code,h.quantity,p.name from hardware_items h join products p on p.code=h.product_code where h.factory_order=? and h.source_type='aicnc' order by h.product_code",(f['factory_order'],))
            group=next((g for g in source_groups if g['factory_order']==f['factory_order']),None)
            if group is None:
                f['hardware_choice']='absent'
                continue
            records={}
            try:
                _materialize_memory_hardware(config,dict(hardware_source_items=[group]),records,{f['factory_order']},set())
                f['hardware']=records['hardware_items']
                for item in f['hardware']:
                    product = connection.execute('select name from products where code=?',(item['product_code'],)).fetchone()
                    item['name'] = product[0] if product else item['product_code']
            except RuleError as exc:
                f['hardware_choice']=''
                f['hardware_error']=str(exc)
                f['hardware_source']=group
                continue
            def totals(rows):
                out=defaultdict(float)
                for r in rows: out[r['product_code']]+=r['quantity']
                return dict(out)
            f['hardware_choice']='same' if totals(f['hardware'])==totals(f['existing_hardware']) else '' if f['existing_hardware'] else 'replace'
        data.update(optimization_id=key,source_folder=str(folder),revision=business_revision(connection,data['factories']))
        groups={(f['order_id'],f['purpose']) for f in data['factories']}
        data['allocations']=[dict(order_id=o,purpose=p,product_code=m['product_code'],quantity=m['quantity'])
                             for o,p in groups for m in mapped] if len(groups)==1 and all(next(iter(groups))) else []
        data['materials']=list(mapped)
        return dict(aicnc_preview=data)
    finally:
        connection.close()


def validate(connection, data):
    """按最终人工分类重新校验；报表原值不随选择而改写。"""
    require_enabled(connection)
    if not re.fullmatch(r'\d{14}', str(data.get('optimization_id',''))):
        raise RuleError('aicnc_identity','优化编号必须是 AICNC 生成的 14 位时间编号')
    if not data.get('factories'):
        raise RuleError('aicnc_identity','缺少工厂单')
    if business_revision(connection,data['factories']) != data['revision']:
        raise RuleError('aicnc_stale','工厂单或五金事实已变化，请重新预览')
    groups=set()
    for f in data['factories']:
        if f['purpose'] not in ('normal','rework') or not ORDER_RE.fullmatch(f['order_id']):
            raise RuleError('aicnc_classification','请确认每个工厂单的订单归属及正常/返工用途')
        if f.get('known_order_id') and f['known_order_id']!=f['order_id']:
            raise RuleError('aicnc_identity','已有工厂单归属与选择不一致，请先修正工厂单身份')
        order=connection.execute('select stage from orders where order_id=?',(f['order_id'],)).fetchone()
        if order and order[0]=='已中止':
            raise RuleError('order_aborted','已中止订单不能追加材料或补出库')
        if f['hardware_choice'] not in ('same','absent','shipped','keep','replace'):
            raise RuleError('aicnc_hardware','请确认保留原五金或使用新五金')
        if f.get('hardware_error') and f['hardware_choice']!='keep':
            raise RuleError('hardware_mapping_required',f['hardware_error'])
        groups.add((f['order_id'],f['purpose']))
    if not data.get('materials'):
        raise RuleError('aicnc_material','没有可入账的板材或封边，请核对商品映射，或忽略整个优化')
    total=defaultdict(float)
    for m in data['materials']:
        total[m['product_code']]+=quantity(m['quantity'],integer=True)
    allocated=defaultdict(float)
    seen=set()
    for row in data['allocations']:
        if (row['order_id'],row['purpose']) not in groups or row['product_code'] not in total:
            raise RuleError('aicnc_allocation','分配包含本次优化之外的订单、用途或商品')
        identity=(row['order_id'],row['purpose'],row['product_code'])
        if identity in seen: raise RuleError('aicnc_allocation','材料分配行重复')
        seen.add(identity)
        allocated[row['product_code']]+=quantity(row['quantity'],integer=True)
    if any(abs(allocated[k]-v)>1e-6 for k,v in total.items()):
        raise RuleError('aicnc_allocation','各订单及用途的分配合计必须等于本次材料总量')
    # 库存动作前先在保存点中验证商品材料属性，验证不会留下写入。
    from .inventory import confirm_product_material_attributes
    connection.execute('savepoint aicnc_attributes')
    try:
        for m in data['materials']:
            confirm_product_material_attributes(connection,m['product_code'],m['material_type'],m['color'],m['thickness'])
    finally:
        connection.execute('rollback to aicnc_attributes')
        connection.execute('release aicnc_attributes')
    for code in total:
        product=connection.execute("select status,catalog_present from products where code=?",(code,)).fetchone()
        if not product or product[0] not in ('','启用') or not product[1]:
            raise RuleError('aicnc_product',f'商品已停用或不存在：{code}')


def ignore(config, folder):
    connection=connect_database(config.workflow_database)
    try:
        require_enabled(connection)
        folder = Path(folder).expanduser().resolve()
        from .order_index import _server_root_candidates
        if not ORDER_RE.fullmatch(folder.parent.name) or folder.parent.parent not in [p.resolve() for p in _server_root_candidates(config)]:
            raise RuleError('aicnc_folder', '只能忽略订单号目录内的单个优化文件夹')
        key=optimization_id(folder)
        previous=connection.execute('select status from aicnc_optimizations where optimization_id=?',(key,)).fetchone()
        if previous and previous[0]=='processing':
            raise RuleError('aicnc_processing','返工库存操作尚未完成核对，不能忽略')
        with connection:
            connection.execute("insert or ignore into aicnc_optimizations values(?,?,'ignored','{}',?,?)",(key,str(folder),now(),now()))
        return dict(ok=True,source_folder=str(folder))
    finally: connection.close()


def documents_for(data):
    groups=defaultdict(list)
    for row in data['allocations']:
        if row['purpose']=='rework' and row['quantity']>0:
            groups[row['order_id']].append(dict(productCode=row['product_code'],quantity=row['quantity']))
    return [dict(order_id=o,remark=f"{o} 返工 {data['optimization_id']}",kind='rework_materials',
                 items=items,rawFingerprint=digest(items),mappedFingerprint=digest(items),
                 knownDocumentNumber='',changed=True) for o,items in sorted(groups.items())]


def confirm(config, data, *, confirmed=False):
    """本地事实一次事务提交；返工外部操作先留可核对记录，不确定结果禁止盲重试。"""
    if not confirmed: raise RuleError('write_confirmation_required','需要明确确认本次材料写入及返工出库')
    connection=connect_database(config.workflow_database)
    try:
        require_enabled(connection)
        key=data['optimization_id']
        previous=connection.execute('select status,plan_json from aicnc_optimizations where optimization_id=?',(key,)).fetchone()
        if previous and previous[0] in ('completed','ignored'):
            return dict(ok=True,already_processed=True,source_folder=data['source_folder'])
        if previous:
            data=json.loads(previous[1])
        else:
            validate(connection,data)
            with connection:
                connection.execute("insert into aicnc_optimizations values(?,?,'processing',?,?,'')",(key,data['source_folder'],encode(data),now()))
    finally: connection.close()
    results=execute_rework(config,data)
    commit(config,data,results)
    return dict(ok=True,source_folder=data['source_folder'],optimization_id=key,documents=results)


def execute_rework(config,data,*,verify_only=False):
    from .inventory import InventoryOperationJournal, run_jdy
    journal=InventoryOperationJournal(config.workflow_database)
    results=[]
    for doc in documents_for(data):
        row=journal.prepare('aicnc_rework',doc['order_id'],[],dict(documents=[doc],optimization_id=data['optimization_id']))
        key=row['operation_id']
        prior=journal.decoded_results(row)
        if row['status']=='local_committed':
            results.extend(dict(r,operation_id=key,order_id=doc['order_id']) for r in prior)
            continue
        if row['status'] in ('submitting','verification_required','external_confirmed','partial_external_confirmed') or verify_only:
            try:
                known=prior[0].get('documentNumber','') if prior else ''
                result=run_jdy(config,'verifyOutbound',order_name=doc['remark'],verification_document={
                    'orderName':doc['remark'],'items':doc['items'],'knownDocumentNumber':known})
                if not result.get('verified') or result.get('remark')!=doc['remark'] or not re.fullmatch(r'QTCK\d+',str(result.get('documentNumber',''))) or (known and known!=result['documentNumber']):
                    raise RuleError('inventory_verification_required','返工出库单尚未核对成功，请核对库存系统单据后恢复')
            except Exception as exc:
                journal.update(key,'verification_required',error=str(exc))
                raise
        else:
            # 会话前置失败未触及外部保存，可以直接重试；提交后任何未知结果必须核对。
            from .inventory import _find_existing_inventory_page, _inventory_cdp_endpoint
            if not _find_existing_inventory_page(_inventory_cdp_endpoint()):
                journal.update(key,'failed',error='请先打开库存系统并登录')
                raise RuleError('inventory_session_required','请先在设置中打开库存系统并登录，再确认返工出库')
            claim = connect_database(config.workflow_database)
            try:
                with claim:
                    changed = claim.execute("update inventory_operations set status='submitting',attempt_count=attempt_count+1,updated_at=? where operation_id=? and status in ('prepared','failed')",(now(),key)).rowcount
                if changed != 1:
                    raise RuleError('inventory_verification_required','此返工出库正在处理或等待核对，请稍后恢复，不能再次提交')
            finally:
                claim.close()
            try:
                result=run_jdy(config,'optimizationOutbound',confirm_save=True,verification_document=doc)
                if not result.get('saved') or not re.fullmatch(r'QTCK\d+',str(result.get('documentNumber',''))):
                    raise RuleError('inventory_verification_required','返工出库结果不明确，请核对单据后恢复')
            except Exception as exc:
                journal.update(key,'verification_required',error=str(exc))
                raise
        result=dict(result,remark=doc['remark'],kind=doc['kind'],saved=True)
        journal.update(key,'external_confirmed',results=[result])
        results.append(dict(result,operation_id=key,order_id=doc['order_id']))
    return results


def commit(config,data,results):
    from .hardware_facts import replace_factory_hardware
    from .inventory import confirm_product_material_attributes
    from .order_index import OrderIndexStore
    store=OrderIndexStore(config.workflow_database)
    c=store.connection
    try:
        c.execute('begin immediate')
        previous=c.execute('select status from aicnc_optimizations where optimization_id=?',(data['optimization_id'],)).fetchone()
        if previous and previous[0]=='completed':
            c.rollback(); return
        # 保留冻结的确认计划；库存往返期间不允许悄悄覆盖发生变化的工厂单/五金。
        validate(c,data)
        timestamp=now()
        for f in data['factories']:
            store.upsert_order(f['order_id'],source_folder=str(Path(data['source_folder']).parent) if Path(data['source_folder']).parent.name.upper()==f['order_id'] else '')
            exists=c.execute('select stage from factory_orders where factory_order=?',(f['factory_order'],)).fetchone()
            if not exists:
                store.upsert_factory(f['factory_order'],order_id=f['order_id'],factory_name=f['factory_name'],ownership_status='已确认',name_source='Server')
            if f['purpose']=='normal':
                completed = datetime.strptime(data['optimization_id'], '%Y%m%d%H%M%S').isoformat(timespec='seconds')
                c.execute('''insert into optimization_artifacts(source_path,order_id,factory_order,file_modified_at,
                    completed_at,first_seen_at,last_seen_at) values(?,?,?,0,?,?,?)''',
                    (data['report_path'],f['order_id'],f['factory_order'],completed,timestamp,timestamp))
                c.execute("""update factory_orders set
                    stage=case when stage='已拆单' then '已优化' else stage end,
                    report_state='已发现',ownership_status='已确认',last_server_seen=?,updated_at=?,source_folder=?,
                    optimization_first_completed_at=case when optimization_first_completed_at='' then ? else min(optimization_first_completed_at,?) end,
                    optimization_latest_completed_at=max(optimization_latest_completed_at,?),
                    optimization_first_seen_at=case when optimization_first_seen_at='' then ? else optimization_first_seen_at end,
                    optimization_latest_seen_at=?,optimization_source_path=?
                    where factory_order=?""",(timestamp,timestamp,data['source_folder'],completed,completed,completed,timestamp,timestamp,data['report_path'],f['factory_order']))
            if f['hardware_choice']=='replace' and f['stage']!='已出货':
                rows=[dict(r,order_id=f['order_id']) for r in f['hardware']]
                replace_factory_hardware(c,f['factory_order'],rows,source_path=data['source_folder'],allow_empty=True)
                c.execute('update factory_orders set has_hardware=? where factory_order=?',(bool(rows),f['factory_order']))
        source=data['source_folder']
        for m in data['materials']:
            confirm_product_material_attributes(c,m['product_code'],m['material_type'],m['color'],m['thickness'])
        aggregate=defaultdict(float)
        for r in data['allocations']:
            aggregate[(r['order_id'],r['product_code'])]+=r['quantity']
            c.execute('insert into aicnc_material_allocations values(?,?,?,?,?)',(data['optimization_id'],r['order_id'],r['purpose'],r['product_code'],r['quantity']))
        for (owner,code),q in aggregate.items():
            if q:
                c.execute("insert into material_items(order_id,product_code,quantity,source_type,source_path,source_fingerprint,updated_at) values(?,?,?,'aicnc_optimization',?,?,?)",(owner,code,q,source,digest([source,code,q]),timestamp))
        rework=[r for r in data['allocations'] if r['purpose']=='rework' and r['quantity']>0]
        if rework:
            record=c.execute("insert into production_records(production_time,source,status,created_at,updated_at) values(?,'aicnc_rework','completed',?,?)",(timestamp,timestamp,timestamp)).lastrowid
            c.executemany('insert into production_materials values(?,?,?,?)',[(record,r['order_id'],r['product_code'],r['quantity']) for r in rework])
        docs={d['order_id']:d for d in documents_for(data)}
        if set(docs)!={r['order_id'] for r in results}:
            raise RuleError('inventory_verification_required','返工单据尚未全部完成，未写入本地材料')
        for r in results:
            d=docs[r['order_id']]
            c.execute("""insert into outbound_documents(document_number,document_type,order_id,factory_order,status,source,issued_at,source_path,document_url,items_json,raw_fingerprint,mapped_fingerprint,updated_at)
                values(?,'rework_materials',?,'','已出库','金蝶',?,?,?,?,?,?,?)""",
                (r['documentNumber'],r['order_id'],r.get('syncedAt') or timestamp,source,r.get('url',''),encode(d['items']),d['rawFingerprint'],d['mappedFingerprint'],timestamp))
            c.execute("update inventory_operations set status='local_committed',last_error='',updated_at=? where operation_id=?",(timestamp,r['operation_id']))
        c.execute("update aicnc_optimizations set status='completed',completed_at=? where optimization_id=?",(timestamp,data['optimization_id']))
        normal_orders = {f['order_id'] for f in data['factories'] if f['purpose']=='normal'}
        for order in store.summaries(persist=False):
            if order['order_id'] in normal_orders:
                c.execute("update orders set stage=?,updated_at=? where order_id=? and stage<>'已中止'",(order['stage'],timestamp,order['order_id']))
        store.commit()
    except Exception:
        c.rollback(); raise
    finally: store.close()


def recover(config, key):
    c=connect_database(config.workflow_database)
    try:
        row=c.execute('select plan_json from aicnc_optimizations where optimization_id=?',(key,)).fetchone()
        if not row: raise RuleError('aicnc_recovery','找不到优化确认计划')
        data=json.loads(row[0])
    finally: c.close()
    results=execute_rework(config,data,verify_only=True)
    commit(config,data,results)
    return dict(ok=True,syncRecorded=True,message='已核对返工出库单并补记材料与消耗，未再次扣库存')

"""共享配置、校验和外部身份查询辅助能力。

本模块为工作流提供配置解析、业务错误、通用标准化及 AIMES 查询边界。
界面动作的选择由 CLI、路由和工具网关负责，不在此处决定。
"""

from __future__ import annotations

import json
import os
import re
import sqlite3
import subprocess
import sys
import time
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .report_read_context import cached_report
from .streaming_process import run_with_progress
from .operation_log import log_aimes_failure, log_progress_payload, redact
from .database import (
    database_path,
    ensure_schema,
    read_cache,
    write_cache,
)


FACTORY_RE = re.compile(r"^F\d+$", re.IGNORECASE)
FACTORY_NAME_ORDER_PREFIX_RE = re.compile(
    r"^([A-Z]{1,3}\d{3,5}(?:-\d+)?)(?=$|[-\s_])",
    re.IGNORECASE,
)
AIMES_BULK_FETCH_LIMIT = 50


class RuleError(RuntimeError):
    def __init__(self, code: str, message: str, **context):
        """建立带错误码和上下文的业务异常；code 为错误码，message 为提示，context 为定位字段。"""
        super().__init__(message)
        self.code = code
        self.context = context


def _extract_aimes_failure(stderr: str) -> str:
    """从标准错误文本 stderr 提取最终 AIMES 错误，跳过 JSON 进度并移除附带页面地址。"""
    plain_lines: list[str] = []
    final_errors: list[str] = []
    for raw_line in str(stderr or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        try:
            event = json.loads(line)
        except json.JSONDecodeError:
            event = None
        if isinstance(event, dict) and event.get("event") == "progress":
            continue
        match = re.search(r"AIMES 自动查询失败\s*[:：]\s*(.*)", line)
        if match:
            final_errors.append(match.group(1).strip())
        else:
            plain_lines.append(line)
    if final_errors:
        error = final_errors[-1]
        # 当前页面地址只是诊断信息；先去掉它，使业务错误及分类不依赖页面细节，后续仍会脱敏。
        error = re.split(r"[;；]\s*当前页面\s*[:：]", error, maxsplit=1)[0]
        return error.strip()
    return plain_lines[-1] if plain_lines else ""


def _aimes_failure_code(error: str) -> str:
    """把最终错误说明 error 分类为稳定的业务错误码，供界面和调用方识别。"""
    message = str(error or "").strip().casefold()
    if any(marker in message for marker in ("账号或密码错误", "用户名或密码", "登录失败", "凭证无效")):
        return "aimes_credentials"
    if "aimes_table_not_ready" in message or "表格尚未加载完成" in message:
        return "aimes_table_not_ready"
    if "缺少必要列" in message or "表头" in message:
        return "aimes_table_schema"
    if "找不到工厂单名称" in message:
        return "aimes_factory_name_missing"
    if any(marker in message for marker in ("超时", "timeout", "timed out")):
        return "aimes_timeout"
    return "aimes_unavailable"


def factory_name_order_prefix(factory_name: str) -> str:
    """从工厂单名称 factory_name 提取明确的订单号前缀，未识别时返回空字符串。"""
    match = FACTORY_NAME_ORDER_PREFIX_RE.match(str(factory_name or "").strip())
    return match.group(1).upper() if match else ""


def factory_name_order_mismatch(factory_name: str, sales_order_name: str) -> tuple[str, str] | None:
    """检查工厂单名称与销售订单归属，不猜测修正值。

    参数：factory_name 为工厂单名称；sales_order_name 为销售订单号。主单允许对应分单后缀；
    没有明确订单前缀的旧房间名称不判为冲突，不一致时返回两个前缀。"""
    prefix = factory_name_order_prefix(factory_name)
    order = str(sales_order_name or "").strip().upper()
    if not prefix or not order or prefix == order:
        return None
    if re.fullmatch(rf"{re.escape(order)}-\d+", prefix):
        return None
    return prefix, order


def progress(message: str, **details) -> None:
    """输出并记录一条进度 JSON；message 为说明，details 为阶段、耗时等附加字段。"""
    payload = {"event": "progress", "message": message, **details}
    log_progress_payload(payload)
    print(json.dumps(payload, ensure_ascii=False), file=sys.stderr, flush=True)


@dataclass
class Config:
    source_root: Path = Path("/Volumes/server/Optimized Orders")
    order_root: Path = Path.home() / "Documents/pp-flowhub/runtime/travelers"
    template: Path = Path(__file__).resolve().parent.parent / "resources/templates/Work Order Traveler.xlsx"
    # Traveler 文件备份路径可配置；数据库备份使用本地状态目录，与此设置独立。
    backup_root: Path = Path("/Volumes/server/g/pp-flowhub/database-backups")
    state_dir: Path = Path.home() / "Documents/pp-flowhub/data"
    initial_date: str = "2026-07-22"
    server_scan_baseline_folder: str = "CS003 PP0047"
    server_scan_baseline_at: str = "2026-07-25T11:00:25-07:00"
    aimes_username: str = ""
    aimes_keychain_service: str = "com.pacificpride.ppflowhub.aimes"
    aimes_retry_delays: tuple[float, float] = (2.0, 4.0)
    operation_log_enabled: bool = True
    storage_prepared: bool = False
    test_source: bool = False
    # Server 预览复用进程内连接，不将连接对象写入设置或命令行状态。
    workflow_connection: sqlite3.Connection | None = None
    # 一次性维护或诊断可请求重算；App 常驻服务关闭此项，普通读取直接使用事务维护的状态。
    reconcile_outbound_on_read: bool = True

    @property
    def operation_log_file(self) -> Path:
        """返回当前状态目录中的操作日志路径；无显式参数。"""
        return self.state_dir / "operation-log.jsonl"

    @property
    def workflow_database(self) -> Path:
        """返回当前状态目录中的中央业务数据库路径；无显式参数。"""
        return database_path(self.state_dir)

    @property
    def database_backup_root(self) -> Path:
        """返回当前状态目录中的数据库备份目录；无显式参数，与 Traveler 备份设置独立。"""
        return self.state_dir / "database-backups"

    def prepare_storage(self) -> None:
        """检查测试来源隔离并确保中央数据库结构可用；无显式参数。

        旧存储已在数据库切换时迁移到归档目录，启动不再读取或修改旧文件；
        workflow.sqlite3 是运行时持久事实的唯一来源。"""
        from .hardware_facts import assert_source_isolation
        assert_source_isolation(self.workflow_database, [self.source_root], test_mode=self.test_source)
        self.storage_prepared = True
        ensure_schema(database_path(self.state_dir))

    @property
    def settings_file(self) -> Path:
        """返回当前状态目录中的设置文件路径；无显式参数。"""
        return self.state_dir / "settings.json"

    def load_settings(self, source_profile: str | None = None) -> None:
        """读取设置并校验日期；source_profile 为 server、local 或默认来源配置选择。"""
        self.test_source = source_profile == "local"
        if not self.settings_file.is_file():
            return
        try:
            values = json.loads(self.settings_file.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise RuleError("settings_invalid", f"设置文件无法读取：{self.settings_file}") from exc
        if source_profile == "server":
            path_keys = {
                "source_root": ("server_source_root", "source_root"),
                "order_root": ("production_order_root", "order_root"),
                "backup_root": ("production_backup_root", "backup_root"),
            }
        elif source_profile == "local":
            local_root = values.get("local_test_root")
            if local_root:
                root = Path(str(local_root)).expanduser()
                self.source_root = root / "Optimized Orders"
                self.order_root = root / "Generated Travelers"
                self.backup_root = root / "Backups"
            path_keys = {}
        else:
            path_keys = {
                "source_root": ("source_root",),
                "order_root": ("order_root",),
                "backup_root": ("backup_root",),
            }
        for attribute, keys in path_keys.items():
            for key in keys:
                if key in values:
                    setattr(self, attribute, Path(str(values[key])).expanduser())
                    break
        # 忽略旧版 iCloud Traveler 备份路径，避免过期设置在数据库切换后重新引入云端持久化。
        if "icloud" in str(self.backup_root).lower() or "mobile documents" in str(self.backup_root).lower():
            self.backup_root = Path("/Volumes/server/g/pp-flowhub/database-backups")
        if "initial_date" in values:
            self.initial_date = str(values["initial_date"])
        if "server_scan_baseline_at" in values:
            self.server_scan_baseline_at = str(values["server_scan_baseline_at"])
        if "aimes_username" in values:
            self.aimes_username = str(values["aimes_username"])
        if isinstance(values.get("operation_log_enabled"), bool):
            self.operation_log_enabled = values["operation_log_enabled"]
        try:
            datetime.strptime(self.initial_date, "%Y-%m-%d")
        except ValueError as exc:
            raise RuleError("settings_invalid", f"初始扫描日期格式错误：{self.initial_date}") from exc
        try:
            datetime.fromisoformat(self.server_scan_baseline_at)
        except ValueError as exc:
            raise RuleError("settings_invalid", f"Server 扫描基准时间格式错误：{self.server_scan_baseline_at}") from exc

    @property
    def factory_names_file(self) -> Path:
        """返回历史工厂单名称文件的路径；无显式参数，仅计算路径，不读取文件。"""
        return self.state_dir / "factory-names.json"

    @property
    def aimes_orders_file(self) -> Path:
        """返回历史 AIMES 订单文件的路径；无显式参数，仅计算路径，不读取文件。"""
        return self.state_dir / "aimes-orders.json"

    @property
    def material_assignments_file(self) -> Path:
        """返回历史材料归属文件的路径；无显式参数，仅计算路径，不读取文件。"""
        return self.state_dir / "material-assignments.json"

    @property
    def node_path(self) -> str:
        """查找 Node 运行程序，依次使用环境配置、随包程序、工作区程序或系统命令；无显式参数。"""
        configured = os.environ.get("TRAVELER_NODE", "").strip()
        if configured:
            return configured
        bundled = Path(__file__).resolve().parent.parent / "bin" / "node"
        if bundled.is_file():
            return str(bundled)
        workspace_runtime = self.playwright_node_modules.resolve().parent / "bin" / "node"
        return str(workspace_runtime) if workspace_runtime.is_file() else "node"

    @property
    def playwright_node_modules(self) -> Path:
        """返回 Playwright 依赖目录，优先读取环境变量配置；无显式参数。"""
        return Path(os.environ.get("TRAVELER_NODE_MODULES", Path(__file__).resolve().parent.parent / "node_modules"))


def load_factory_name_cache(config: Config) -> dict[str, str]:
    """读取并标准化工厂单名称缓存；config 提供中央数据库配置。"""
    values = read_cache(config.workflow_database, "factory_names", None, {})
    if not isinstance(values, dict):
        return {}
    return {str(key).upper(): str(value).strip() for key, value in values.items() if str(value).strip()}


def save_factory_name_cache(config: Config, values: dict[str, str]) -> None:
    """保存工厂单名称缓存；config 为数据库配置，values 为工厂单号到名称的映射。"""
    write_cache(config.workflow_database, "factory_names", dict(sorted(values.items())))


def load_aimes_order_cache(config: Config) -> list[dict[str, str]]:
    """读取并标准化已有 AIMES 订单缓存，过滤无工厂单号的记录；config 为数据库配置。"""
    values = read_cache(config.workflow_database, "aimes_orders", None, [])
    if not isinstance(values, list):
        return []
    return [
        {
            "factory_order": str(row.get("factory_order", "")).upper().strip(),
            "factory_name": str(row.get("factory_name", "")).strip(),
            "sales_order_name": str(row.get("sales_order_name", "")).upper().strip(),
            "split_time": str(row.get("split_time", "")).strip(),
        }
        for row in values
        if isinstance(row, dict) and str(row.get("factory_order", "")).strip()
    ]


def save_aimes_order_cache(config: Config, values: list[dict[str, str]]) -> None:
    """保存 AIMES 订单记录缓存；config 为数据库配置，values 为订单记录列表。"""
    write_cache(config.workflow_database, "aimes_orders", values)


def load_material_assignments(config: Config) -> dict[str, str]:
    """读取材料归属映射并过滤空路径；config 为数据库配置。"""
    values = read_cache(config.workflow_database, "material_assignments", None, {})
    return {str(key): str(value) for key, value in values.items() if str(value).strip()}


def save_material_assignment(config: Config, key: str, path: str) -> None:
    """更新一项材料归属并保存映射；config 为配置，key 为归属键，path 为所选路径。"""
    values = load_material_assignments(config)
    values[str(key)] = str(path)
    write_cache(config.workflow_database, "material_assignments", dict(sorted(values.items())))


def _run_aimes_lookup(
    config: Config,
    factory_orders: list[str],
    *,
    recent_limit: int = 0,
    include_order_metadata: bool = False,
    verify_factory_orders: bool = False,
):
    """读取钥匙串凭据并执行 AIMES 查询，处理进度、超时、重试及业务错误。

    参数：config 为连接和运行配置；factory_orders 为待查工厂单号；recent_limit 为最近记录数；
    include_order_metadata 决定是否返回完整订单字段；verify_factory_orders 决定是否精确核验存在性。"""
    missing = [order.upper() for order in factory_orders if order]
    if not missing and not recent_limit:
        return {"rows": []} if include_order_metadata else {}
    if not config.aimes_username.strip():
        raise RuleError("aimes_credentials", "缺少 AIMES 用户名，无法查询工厂单名称", factory_orders=missing)
    helper = Path(__file__).resolve().parent.parent / "tools/aimes_lookup.mjs"
    if not helper.is_file():
        raise RuleError("aimes_unavailable", "AIMES 查询程序不存在", factory_orders=missing)
    try:
        password = subprocess.check_output([
            "/usr/bin/security", "find-generic-password", "-w", "-a", config.aimes_username,
            "-s", config.aimes_keychain_service,
        ], text=True, stderr=subprocess.DEVNULL).strip()
    except (subprocess.CalledProcessError, FileNotFoundError) as exc:
        raise RuleError("aimes_credentials", "无法从 macOS 钥匙串读取 AIMES 密码", factory_orders=missing) from exc
    payload = json.dumps({
        "username": config.aimes_username,
        "password": password,
        "factoryOrders": missing,
        "recentLimit": recent_limit,
        "includeOrderMetadata": include_order_metadata,
        "verifyFactoryOrders": verify_factory_orders,
    })
    env = os.environ.copy()
    env["NODE_PATH"] = str(config.playwright_node_modules)
    last_error = ""
    operation_timings: list[dict[str, object]] = []
    last_progress_stage = ""
    sensitive_values = tuple(value for value in (config.aimes_username, password) if value)

    def consume_progress(stderr: str, attempt_timings: list[dict[str, object]]) -> None:
        """解析并脱敏转发进度；stderr 为标准错误文本，attempt_timings 为本次查询阶段耗时的收集列表。"""
        nonlocal last_progress_stage
        for line in stderr.splitlines():
            try:
                event = json.loads(line)
            except json.JSONDecodeError:
                continue
            if not isinstance(event, dict) or event.get("event") != "progress":
                continue
            safe_event = redact(event, sensitive_values=sensitive_values)
            log_progress_payload(safe_event, sensitive_values=sensitive_values)
            print(json.dumps(safe_event, ensure_ascii=False), file=sys.stderr, flush=True)
            stage = str(event.get("stage", "")).strip()
            if stage:
                last_progress_stage = stage
            duration = event.get("duration_seconds")
            if stage and isinstance(duration, (int, float)):
                attempt_timings.append({
                    "stage": stage,
                    "label": str(event.get("stage_label", stage)),
                    "duration_seconds": round(float(duration), 2),
                })

    deadline = time.monotonic() + 90
    attempt = 0
    for attempt in range(2):
        attempt_timings: list[dict[str, object]] = []
        progress(f"正在获取 AIMES 数据（第 {attempt + 1}/2 次）", factory_orders=missing)
        try:
            completed = run_with_progress(
                [config.node_path, str(helper)], input_text=payload, env=env,
                timeout=max(0.001, deadline - time.monotonic()),
                on_stderr_line=lambda line: consume_progress(line, attempt_timings),
            )
            result = json.loads(completed.stdout)
            if not isinstance(result, dict):
                raise json.JSONDecodeError("AIMES 返回结果不是对象", completed.stdout, 0)
            node_timings = result.pop("timings", [])
            if not isinstance(node_timings, list):
                node_timings = []
            normalized_timings = [
                {
                    "stage": str(item.get("stage", "")).strip(),
                    "label": str(item.get("label", item.get("stage", ""))).strip(),
                    "duration_seconds": round(float(item.get("duration_seconds", 0)), 2),
                }
                for item in node_timings
                if isinstance(item, dict) and str(item.get("label", item.get("stage", ""))).strip()
            ]
            # 单次尝试耗时覆盖整个辅助进程，不是子阶段；此处只保留不重叠阶段，总耗时由调用方记录。
            result["_aimes_timings"] = operation_timings + (normalized_timings or attempt_timings)
            result["_aimes_retry_count"] = attempt
            return result
        except subprocess.CalledProcessError as exc:
            last_error = _extract_aimes_failure(exc.stderr or "")
            if not last_error:
                last_error = f"AIMES 查询进程退出（状态码 {exc.returncode}）"
            operation_timings.extend(attempt_timings)
            progress(
                f"AIMES 第 {attempt + 1}/2 次尝试失败，准备重试",
                retry_attempt=attempt + 1,
            )
        except subprocess.TimeoutExpired:
            last_error = "AIMES 查询进程超时"
            operation_timings.extend(attempt_timings)
            progress(
                f"AIMES 第 {attempt + 1}/2 次尝试失败，准备重试",
                retry_attempt=attempt + 1,
            )
        except OSError as exc:
            last_error = f"AIMES 查询进程无法启动：{exc}"
            operation_timings.extend(attempt_timings)
            progress(
                f"AIMES 第 {attempt + 1}/2 次尝试失败，准备重试",
                retry_attempt=attempt + 1,
            )
        except json.JSONDecodeError:
            last_error = "AIMES 返回结果不是有效 JSON"
            operation_timings.extend(attempt_timings)
            progress(
                f"AIMES 第 {attempt + 1}/2 次尝试失败，准备重试",
                retry_attempt=attempt + 1,
            )
        # 登录后的页面或查询失败在同一浏览器内重试；只有进程或会话失败才重启，已知数据错误不重启。
        if _aimes_failure_code(last_error) in {
            "aimes_credentials",
            "aimes_table_schema",
            "aimes_factory_name_missing",
        } or "AIMES_STEP_FAILED" in last_error:
            break
        if time.monotonic() >= deadline:
            break
        if attempt < 1:
            retry_started = time.perf_counter()
            progress(f"等待 {config.aimes_retry_delays[attempt]:.1f} 秒后重试 AIMES")
            time.sleep(min(config.aimes_retry_delays[attempt], max(0, deadline - time.monotonic())))
            operation_timings.append({
                "stage": "retry_wait",
                "label": f"重试等待（第 {attempt + 1} 次后）",
                "duration_seconds": round(time.perf_counter() - retry_started, 2),
            })
    failure_code = _aimes_failure_code(last_error)
    safe_error = str(redact(last_error, sensitive_values=sensitive_values)).strip()[-1000:]
    if not safe_error:
        safe_error = "AIMES 查询失败"
    log_aimes_failure(
        error=safe_error,
        code=failure_code,
        stage=last_progress_stage,
        enabled=bool(config.operation_log_enabled),
        sensitive_values=sensitive_values,
    )
    raise RuleError(
        failure_code,
        f"AIMES 查询失败：{safe_error}",
        factory_orders=missing,
        aimes_timings=operation_timings,
        aimes_retry_count=attempt,
        aimes_failure_stage=last_progress_stage,
    )


def lookup_aimes_names(config: Config, factory_orders: list[str], recent_limit: int = 0) -> dict[str, str]:
    """查询并校验工厂单名称；config 为配置，factory_orders 为工厂单号列表，recent_limit 为附加近期读取数。"""
    missing = [order.upper() for order in factory_orders if order]
    result = _run_aimes_lookup(config, missing, recent_limit=recent_limit)
    if not isinstance(result, dict):
        raise RuleError("aimes_unavailable", "AIMES 返回的工厂单名称格式无效", factory_orders=missing)
    values = {
        str(factory).upper(): str(name).strip()
        for factory, name in result.items()
        if str(factory).strip() and str(name).strip()
    }
    if missing and not all(values.get(factory) for factory in missing):
        raise RuleError("aimes_unavailable", "AIMES 返回结果缺少工厂单名称", factory_orders=missing)
    return values


def lookup_aimes_recent_orders(
    config: Config,
    limit: int = AIMES_BULK_FETCH_LIMIT,
    *,
    include_trace: bool = False,
) -> list[dict[str, str]] | tuple[list[dict[str, str]], list[dict[str, object]]]:
    """读取最近 AIMES 订单并标准化字段。

    参数：config 为配置；limit 为读取数量；include_trace 决定是否同时返回各阶段耗时。"""
    result = _run_aimes_lookup(config, [], recent_limit=limit, include_order_metadata=True)
    rows = result.get("rows") if isinstance(result, dict) else None
    if not isinstance(rows, list):
        raise RuleError("aimes_unavailable", "AIMES 返回的订单明细格式无效")
    normalized = [
        {
            "factory_order": str(row.get("factory_order", "")).upper().strip(),
            "factory_name": str(row.get("factory_name", "")).strip(),
            "sales_order_name": str(row.get("sales_order_name", "")).upper().strip(),
            "split_time": str(row.get("split_time", "")).strip(),
        }
        for row in rows
        if isinstance(row, dict)
    ]
    if include_trace:
        timings = result.get("_aimes_timings", [])
        return normalized, _non_aggregate_aimes_timings(timings)
    return normalized


def _non_aggregate_aimes_timings(values: object) -> list[dict[str, object]]:
    """从耗时数据 values 保留独立阶段，过滤旧的尝试和总计项，避免重复累计。"""
    if not isinstance(values, list):
        return []
    return [
        item for item in values
        if isinstance(item, dict)
        and str(item.get("stage", "")).strip() not in {"attempt", "total"}
        and "总计用时" not in str(item.get("label", ""))
    ]


def verify_aimes_factory_orders(config: Config, factory_orders: list[str]) -> dict[str, list[dict[str, str]] | list[str]]:
    """精确查询工厂单存在性，区分记录缺失和传输失败。

    参数：config 为配置；factory_orders 为待核验工厂单号。返回找到的记录及缺失单号。"""
    requested = [str(order).upper().strip() for order in factory_orders if str(order).strip()]
    result = _run_aimes_lookup(config, requested, verify_factory_orders=True)
    rows = result.get("rows") if isinstance(result, dict) else None
    missing = result.get("missing") if isinstance(result, dict) else None
    if not isinstance(rows, list) or not isinstance(missing, list):
        raise RuleError("aimes_unavailable", "AIMES 返回的精确核验格式无效", factory_orders=requested)
    normalized_rows = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        normalized_rows.append({
            "factory_order": str(row.get("factory_order", "")).upper().strip(),
            "factory_name": str(row.get("factory_name", "")).strip(),
            "sales_order_name": str(row.get("sales_order_name", "")).upper().strip(),
            "split_time": str(row.get("split_time", "")).strip(),
        })
    return {
        "rows": normalized_rows,
        "missing": [str(value).upper().strip() for value in missing if str(value).strip()],
    }


def refresh_aimes_recent_orders_and_verify(
    config: Config,
    limit: int,
    factory_orders: list[str],
    *,
    timing_sink: list[dict[str, object]] | None = None,
) -> tuple[list[dict[str, str]], dict[str, list[dict[str, str]] | list[str]]]:
    """在同一浏览器会话读取最近订单并核验指定工厂单。

    参数：config 为配置；limit 为近期读取数；factory_orders 为精确核验列表；timing_sink 为可选耗时收集列表。"""
    result = _run_aimes_lookup(
        config,
        [str(order).upper().strip() for order in factory_orders if str(order).strip()],
        recent_limit=limit,
        include_order_metadata=True,
        verify_factory_orders=True,
    )
    recent_rows = result.get("rows") if isinstance(result, dict) else None
    verified_rows = result.get("verify_rows") if isinstance(result, dict) else None
    missing = result.get("missing") if isinstance(result, dict) else None
    if not isinstance(recent_rows, list) or not isinstance(verified_rows, list) or not isinstance(missing, list):
        raise RuleError("aimes_unavailable", "AIMES 返回的获取与精确核验格式无效")

    def normalize(rows: list[object]) -> list[dict[str, str]]:
        """标准化订单列表 rows 的身份、名称及拆单时间字段，过滤非字典记录。"""
        return [
            {
                "factory_order": str(row.get("factory_order", "")).upper().strip(),
                "factory_name": str(row.get("factory_name", "")).strip(),
                "sales_order_name": str(row.get("sales_order_name", "")).upper().strip(),
                "split_time": str(row.get("split_time", "")).strip(),
            }
            for row in rows
            if isinstance(row, dict)
        ]

    if timing_sink is not None:
        timing_sink.extend(_non_aggregate_aimes_timings(result.get("_aimes_timings", [])))
    return normalize(recent_rows), {
        "rows": normalize(verified_rows),
        "missing": [str(value).upper().strip() for value in missing if str(value).strip()],
    }


def refresh_aimes_recent_orders(
    config: Config,
    limit: int = AIMES_BULK_FETCH_LIMIT,
    *,
    persist: bool = True,
    timing_sink: list[dict[str, object]] | None = None,
) -> list[dict[str, str]]:
    """获取结构化 AIMES 订单，可选择保存原始结果。

    参数：config 为配置；limit 为近期读取数；persist 为缓存写入开关；timing_sink 为可选耗时收集列表。
    订单索引传入 persist=False 以先校验；无效销售订单名只能作为警告，不能进入业务映射缓存。"""
    traced = lookup_aimes_recent_orders(config, limit, include_trace=True)
    rows, timings = traced
    if timing_sink is not None:
        timing_sink.extend(timings)
    if persist and rows:
        save_aimes_order_cache(config, rows)
        names = {
            row["factory_order"]: row["factory_name"]
            for row in rows
            if row["factory_order"] and row["factory_name"]
        }
        cache = load_factory_name_cache(config)
        cache.update(names)
        save_factory_name_cache(config, cache)
    return rows


def refresh_aimes_recent_names(config: Config, limit: int = AIMES_BULK_FETCH_LIMIT) -> dict[str, str]:
    """读取最近 AIMES 工厂单名称但不保存缓存；config 为配置，limit 为近期读取数。"""
    rows = refresh_aimes_recent_orders(config, limit, persist=False)
    return {
        row["factory_order"]: row["factory_name"]
        for row in rows
        if row["factory_order"] and row["factory_name"]
    }


@dataclass
class FittingItem:
    name: str
    code: str
    size: str
    unit: str
    quantity: float


def _text(value) -> str:
    """把 value 转为去除首尾空白的文本；None 转为空字符串。"""
    return "" if value is None else str(value).strip()


def _number(value) -> float:
    """把 value 转为浮点数；空值按零处理，其他无法解析的值抛出业务错误。"""
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise RuleError("invalid_number", f"无法识别数值：{value}") from exc


def _normalize_name(value: str) -> str:
    """去除名称 value 中的空白、下划线和连字符并转为大写，供五金名称比较。"""
    return re.sub(r"[\s_-]+", "", value).upper()


RAIL_PAIR_NAMES = (
    ("LEFTRAIL", "RIGHTRAIL", "Left Rail", "Right Rail"),
    ("LOWERLEFTRAIL", "LOWERRIGHTRAIL", "Lower Left Rail", "Lower Right Rail"),
)


@cached_report
def parse_fittings_groups(
    path: Path,
    *,
    allow_missing_factory: bool = False,
    fallback_factory: str = "",
    included_factories: tuple[str, ...] | None = None,
) -> list[tuple[str, list[FittingItem]]]:
    """解析并校验五金报表，按工厂单返回五金记录，并合并数量一致的左右导轨。

    参数：path 为报表路径；allow_missing_factory 是否允许缺少有效工厂单号；fallback_factory 为替代单号。"""
    invalid_dimension = False
    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    if re.search(br'<dimension[^>]+ref="[^"]*\?[^"]*"', archive.read(name)[:1024]):
                        invalid_dimension = True
                        break
    except (OSError, zipfile.BadZipFile):
        pass

    try:
        workbook = load_workbook(path, data_only=True, read_only=not invalid_dimension)
    except ValueError:
        try:
            workbook = load_workbook(path, data_only=True, read_only=False)
        except Exception as exc:
            raise RuleError("fittings_file_invalid", f"五金清单文件损坏或无法读取：{path.name}") from exc
    if workbook.sheetnames != ["Page1"]:
        raise RuleError("fittings_schema", f"五金清单工作表结构变化：{workbook.sheetnames}", path=str(path))
    sheet = workbook["Page1"]
    starts = [
        row for row in range(1, sheet.max_row + 1)
        if re.sub(r"\s+", "", _text(sheet.cell(row, 1).value)) in {"OrderNo.", "订单号"}
    ]
    if not starts:
        raise RuleError("fittings_schema", "五金清单缺少 Order No. 区块")

    groups = []
    for index, start in enumerate(starts):
        factory = next((_text(sheet.cell(start, col).value).upper()
                        for col in range(2, sheet.max_column + 1)
                        if FACTORY_RE.fullmatch(_text(sheet.cell(start, col).value).upper())), "")
        if not FACTORY_RE.fullmatch(factory):
            if allow_missing_factory and fallback_factory:
                factory = fallback_factory.strip()
            else:
                raise RuleError("fittings_identity", f"五金清单工厂单号异常：{factory}")
        if included_factories is not None and factory not in included_factories:
            continue
        header = start + 5
        # 按标题定位，兼容新版紧凑报表；已出货区块在读取数量前跳过。
        legacy = re.sub(r"\s+", "", _text(sheet.cell(start, 1).value)) == "订单号"
        labels = {_text(sheet.cell(header, col).value): col for col in range(1, sheet.max_column + 1)}
        expected = ("名称", "编号", "尺寸", "数量") if legacy else ("Name", "Code", "Size", "Quantity")
        if any(label not in labels for label in expected):
            raise RuleError("fittings_schema", f"五金清单第 {start} 行开始的区块字段发生变化")
        name_col, code_col, size_col, qty_col = (labels[label] for label in expected)
        unit_col = labels.get("单位" if legacy else "Unit", 9)
        end = starts[index + 1] if index + 1 < len(starts) else sheet.max_row + 1
        items = []
        for row in range(header + 1, end):
            name = _text(sheet.cell(row, name_col).value)
            if not name or name in {"Total", "小计", "合计"}:
                continue
            items.append(FittingItem(
                name=name,
                code=_text(sheet.cell(row, code_col).value),
                size=_text(sheet.cell(row, size_col).value),
                unit=_text(sheet.cell(row, unit_col).value),
                quantity=_number(sheet.cell(row, qty_col).value),
            ))
        # 左右导轨一对表示一套实物，在此校验并合并，使下游每套只收到一条来源记录。
        # 高导轨和低导轨使用相同规则，但五金报表中的名称不同。
        for left_name, right_name, left_label, right_label in RAIL_PAIR_NAMES:
            left = [item for item in items if _normalize_name(item.name) == left_name]
            right = [item for item in items if _normalize_name(item.name) == right_name]
            if (left or right) and (
                bool(left) != bool(right)
                or len(left) != 1
                or len(right) != 1
                or (
                    left
                    and right
                    and left[0].quantity != right[0].quantity
                )
            ):
                left_quantity = left[0].quantity if len(left) == 1 else None
                right_quantity = right[0].quantity if len(right) == 1 else None
                left_display = (
                    f"{left_quantity:g}"
                    if left_quantity is not None
                    else "缺失" if not left else f"{len(left)} 行"
                )
                right_display = (
                    f"{right_quantity:g}"
                    if right_quantity is not None
                    else "缺失" if not right else f"{len(right)} 行"
                )
                raise RuleError(
                    "paired_rail_quantity_mismatch",
                    f"工厂单 {factory} 的 {left_label} / {right_label} 数量不一致，"
                    f"{left_label}={left_display}、{right_label}={right_display}；"
                    "请人工核对后再写入。",
                    factory_order=factory,
                    left_quantity=left_quantity,
                    right_quantity=right_quantity,
                )
            if len(left) == 1 and len(right) == 1 and left[0].quantity == right[0].quantity:
                items = [item for item in items if item is not right[0]]
        groups.append((factory, items))
    return groups
